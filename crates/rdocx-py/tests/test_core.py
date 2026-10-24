import io
import posixpath
import re
import struct
import subprocess
import time
import zipfile
import zlib

import pytest

@pytest.mark.parametrize('write', [False, True])
def test_fx257_frozen_section_margin_names_the_live_owner_route(write):
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held section')
    run, live = paragraph.runs[0], document.section(0)
    snapshot = document.sections[0]
    before = document.to_bytes()
    with pytest.raises(NotImplementedError, match=r'document\.section\(0\)\.left_margin'):
        if write:
            snapshot.left_margin = rdocx.Inches(2)
        else:
            _ = snapshot.left_margin
    assert document.to_bytes() == before
    assert paragraph.text == run.text == 'held section' and live.xml
    assert snapshot.margin_left == rdocx.Inches(1)
    live.left_margin = rdocx.Inches(2)
    assert snapshot.margin_left == rdocx.Inches(1)
    assert rdocx.Document.from_bytes(document.to_bytes()).sections[0].margin_left == rdocx.Inches(2)


def test_fx257_immutable_tab_alignment_names_delete_and_readd():
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held tabs')
    run = paragraph.runs[0]
    tabs = paragraph.paragraph_format.tab_stops
    tab = tabs.add_tab_stop(rdocx.Inches(1), rdocx.WD_TAB_ALIGNMENT.LEFT, rdocx.WD_TAB_LEADER.DOTS)
    before = document.to_bytes()
    for value in [rdocx.WD_TAB_ALIGNMENT.RIGHT, object()]:
        with pytest.raises(NotImplementedError, match='delete.*readd|delete.*re-add'):
            tab.alignment = value
        assert document.to_bytes() == before
        assert paragraph.text == run.text == 'held tabs'
        assert tab.alignment == rdocx.WD_TAB_ALIGNMENT.LEFT
    position, leader = tab.position, tab.leader
    del tabs[0]
    tabs.add_tab_stop(position, rdocx.WD_TAB_ALIGNMENT.RIGHT, leader)
    reopened = rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].paragraph_format.tab_stops[0]
    assert (reopened.position, reopened.alignment, reopened.leader) == (position, rdocx.WD_TAB_ALIGNMENT.RIGHT, leader)
    assert tab.alignment == rdocx.WD_TAB_ALIGNMENT.LEFT


def test_fx257_custom_property_mapping_reuses_checked_publication():
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held mapping')
    run, section = paragraph.runs[0], document.section(0)
    first, second = document.custom_properties, document.custom_properties
    before = document.to_bytes()
    assert 'Client' not in first and '' not in first
    assert document.to_bytes() == before
    with zipfile.ZipFile(io.BytesIO(before)) as package:
        assert 'docProps/custom.xml' not in package.namelist()
    for key in [1, True, None, object()]:
        with pytest.raises(TypeError):
            first[key] = 'bad'
        with pytest.raises(TypeError):
            del first[key]
        with pytest.raises(TypeError):
            _ = key in first
        assert document.to_bytes() == before
    for name in ['', 'missing']:
        with pytest.raises(KeyError):
            del first[name]
        assert document.to_bytes() == before
    with pytest.raises(ValueError):
        first[''] = 'bad'
    for invalid in [object(), float('nan'), float('inf'), 2**40]:
        with pytest.raises((TypeError, ValueError, OverflowError)):
            first['Client'] = invalid
        assert document.to_bytes() == before
    assert paragraph.text == run.text == 'held mapping' and section.xml
    first['Client'] = 'Ada'
    assert 'Client' in second and 'client' not in second
    frozen = first.get('Client')
    assert frozen.value == 'Ada'
    first['Client'] = 'Ada'
    same = document.to_bytes()
    second['Client'] = 'Ada'
    assert document.to_bytes() == same
    second['Client'] = 'Grace'
    assert first.get('Client').value == 'Grace' and frozen.value == 'Ada'
    assert _document_xml(document) == _document_xml(rdocx.Document.from_bytes(before))
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.custom_properties.get('Client').value == 'Grace'
    del first['Client']
    assert 'Client' not in second and first.delete('Client') is False
    assert paragraph.text == run.text == 'held mapping' and section.xml

    seed = rdocx.Document()
    seed.custom_properties.set('seed', 'retained empty name')
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(seed.to_bytes())) as package, zipfile.ZipFile(output, 'w') as target:
        for entry in package.infolist():
            raw = package.read(entry.filename)
            if entry.filename == 'docProps/custom.xml':
                raw = raw.replace(b'name="seed"', b'name=""')
            target.writestr(entry, raw)
    imported = rdocx.Document.from_bytes(output.getvalue())
    assert imported.custom_properties.get('').value == 'retained empty name'
    before = imported.to_bytes()
    assert '' not in imported.custom_properties
    with pytest.raises(KeyError):
        del imported.custom_properties['']
    assert imported.to_bytes() == before
    assert imported.custom_properties.get('').value == 'retained empty name'


@pytest.mark.parametrize('owner_name,attribute', [('paragraph', '_p'), ('run', '_r'), ('run', '_element')])
def test_fx257_private_word_names_point_to_checked_xml(owner_name, attribute):
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held xml')
    run, section = paragraph.runs[0], document.section(0)
    owner = paragraph if owner_name == 'paragraph' else run
    before = document.to_bytes()
    with pytest.raises(NotImplementedError, match=rf'{owner_name}\.xml.*{owner_name}\.replace_xml\(xml\)'):
        getattr(owner, attribute)
    assert document.to_bytes() == before
    assert paragraph.text == run.text == 'held xml' and section.xml
    owner.replace_xml(owner.xml)
    assert document.to_bytes() == before
    owner.replace_xml(owner.xml.replace('held xml', 'changed xml'))
    assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == 'changed xml'
    stale = owner
    document.insert_section(0)
    with pytest.raises(rdocx.StaleElementError):
        getattr(stale, attribute)


def test_fx257_legacy_dropdown_reports_received_type_and_index_route():
    import rdocx
    import xml.etree.ElementTree as ET
    document = _replace_document_body(rdocx.Document(), _fx241_legacy_typed('pick', 'dropdown') + '<w:p><w:r><w:t>held legacy</w:t></w:r></w:p><w:sectPr/>')
    snapshot = document.legacy_form_fields()[0]
    paragraph, run = document.paragraphs[-1], document.paragraphs[-1].runs[0]
    before = document.to_bytes()
    for value in ['two', True, 1.5, object()]:
        with pytest.raises(TypeError) as error:
            document.set_legacy_form_field_value('pick', value)
        message = str(error.value)
        assert type(value).__name__ in message and 'nonnegative' in message
        assert 'legacy_form_fields' in message and 'index' in message
        assert document.to_bytes() == before
    for value in [-1, 2, 2**100]:
        with pytest.raises((OverflowError, rdocx.RdocxError)):
            document.set_legacy_form_field_value('pick', value)
        assert document.to_bytes() == before
    assert paragraph.text == run.text == 'held legacy'
    assert snapshot.choices == ('one', 'two') and snapshot.value == 0
    document.set_legacy_form_field_value('pick', 0)
    assert document.to_bytes() == before
    updated = document.set_legacy_form_field_value('pick', 1)
    assert updated.value == 1 and snapshot.value == 0
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.legacy_form_fields()[0].value == 1
    root = ET.fromstring(_document_xml(reopened));w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    assert root.find('.//'+w+'ddList/'+w+'result').get(w+'val') == '1'
    assert 'two' in [node.text for node in root.iter(w+'t')]


def test_fx257_word_color_kind_matches_the_pinned_enum():
    import rdocx
    from rdocx.enum.dml import MSO_COLOR_TYPE
    document = rdocx.Document()
    paragraph = document.add_paragraph('held color')
    color = paragraph.runs[0].font.color
    before = document.to_bytes()
    assert color.type is None and document.to_bytes() == before
    assert [(item.name, item.value) for item in MSO_COLOR_TYPE] == [('RGB', 1), ('THEME', 2), ('AUTO', 101)]
    for scalar, expected in [('123456', MSO_COLOR_TYPE.RGB), ('auto', MSO_COLOR_TYPE.AUTO)]:
        if scalar == 'auto':
            document = _replace_document_body(document, '<w:p><w:r><w:rPr><w:color w:val="auto"/></w:rPr><w:t>held color</w:t></w:r></w:p><w:sectPr/>')
            paragraph = document.paragraphs[0]
            color = paragraph.runs[0].font.color
        else:
            paragraph.runs[0].font.color = scalar
        before = document.to_bytes()
        with pytest.raises(ValueError, match='hexadecimal'):
            paragraph.runs[0].font.color = 'auto'
        assert document.to_bytes() == before and paragraph.text == 'held color'
        assert color.type is expected and not isinstance(color.type, str)
        assert document.to_bytes() == before
        assert rdocx.Document.from_bytes(before).paragraphs[0].runs[0].font.color.type is expected
    color.theme_color = 'accent1'
    before = document.to_bytes()
    assert color.type is MSO_COLOR_TYPE.THEME and color.theme_color == 'accent1'
    assert document.to_bytes() == before
    color.theme_color = None
    assert color.type is MSO_COLOR_TYPE.AUTO
    color.rgb = None
    assert color.type is None


@pytest.mark.parametrize('operation', ['space_after', 'character_spacing', 'line_spacing', 'margin_top'])
def test_fx256_issue386_numeric_rows_are_observed_independently(operation):
    import rdocx
    import xml.etree.ElementTree as ET
    document = rdocx.Document()
    document.add_paragraph('held numeric owner')
    document.add_paragraph('unrelated')
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    paragraph, other = document.paragraphs
    run, section = paragraph.runs[0], document.section(0)
    before = document.to_bytes()
    if operation == 'line_spacing':
        paragraph.paragraph_format.line_spacing = 2
        root = ET.fromstring(_document_xml(document))
        spacing = root.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
        word = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        assert spacing is not None and spacing.get(word+'line') == '480'
        assert spacing.get(word+'lineRule') == 'auto'
        assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].paragraph_format.line_spacing == 2.0
    else:
        with pytest.raises(ValueError, match='Pt|Twips|Emu'):
            if operation == 'space_after':
                paragraph.paragraph_format.space_after = 6
            elif operation == 'character_spacing':
                run.font.character_spacing = 2
            else:
                document.update_section(0, margin_top=1440)
        assert document.to_bytes() == before
        assert document.verify_signatures()[0]['cryptographically_valid']
    assert paragraph.text == run.text == 'held numeric owner'
    assert other.text == 'unrelated' and section.xml


@pytest.mark.parametrize('operation', ['space_after', 'character_spacing'])
def test_fx256_spacing_lengths_refuse_quantized_bare_values_atomically(operation):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held')
    document.add_paragraph('other')
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    paragraph, other = document.paragraphs
    run, section = paragraph.runs[0], document.section(0)
    owner = paragraph.paragraph_format if operation == 'space_after' else run.font
    before = document.to_bytes()
    invalid = [1, 6, 634, -1, -634, True, False, 1.5, float('nan'),
               float('inf'), '6', object(), 2**100, rdocx.Emu((2**31)*635)]
    if operation == 'character_spacing':
        invalid += [rdocx.Emu((-2**31-1)*635)]
    else:
        invalid += [-635, rdocx.Emu(-1), rdocx.Pt(-1)]
    for value in invalid:
        with pytest.raises((TypeError, ValueError, OverflowError)):
            setattr(owner, operation, value)
        assert document.to_bytes() == before and document.verify_signatures()[0]['cryptographically_valid']
        assert paragraph.text == run.text == 'held' and other.text == 'other' and section.xml
    setattr(owner, operation, rdocx.Pt(2))
    assert getattr(owner, operation) == rdocx.Pt(2)


@pytest.mark.parametrize('operation', ['space_after', 'character_spacing'])
def test_fx256_spacing_explicit_lengths_keep_quantization_and_clear(operation):
    import rdocx
    import xml.etree.ElementTree as ET
    values = [0, 635, rdocx.Emu(1), rdocx.Emu(634), rdocx.Emu(636),
              rdocx.Pt(6), rdocx.Twips(2**31-1)]
    if operation == 'character_spacing':
        values += [-635, rdocx.Emu(-634), rdocx.Emu(-636), rdocx.Twips(-2**31)]
    for value in values:
        document = rdocx.Document()
        paragraph = document.add_paragraph('value')
        owner = paragraph.paragraph_format if operation == 'space_after' else paragraph.runs[0].font
        setattr(owner, operation, value)
        quotient = abs(int(value)) // 635 * (-1 if value < 0 else 1)
        root = ET.fromstring(_document_xml(document))
        word = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        spacing = root.find('.//'+word+('pPr' if operation == 'space_after' else 'rPr')+'/'+word+'spacing')
        assert spacing is not None and spacing.get(word+('after' if operation == 'space_after' else 'val')) == str(quotient)
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        restored = reopened.paragraphs[0].paragraph_format if operation == 'space_after' else reopened.paragraphs[0].runs[0].font
        assert getattr(restored, operation) == quotient * 635
        setattr(restored, operation, None)
        assert getattr(restored, operation) is None


@pytest.mark.parametrize('absolute', [False, True])
def test_fx256_line_spacing_distinguishes_multiples_and_lengths(absolute):
    import rdocx
    import xml.etree.ElementTree as ET
    limits = (-2**31, 2**31-1)
    values = ([rdocx.Pt(18), rdocx.Emu(634), rdocx.Emu(-634),
               *[rdocx.Emu(n*635) for n in limits],
               rdocx.Emu(limits[1]*635+634), rdocx.Emu(limits[0]*635-634)] if absolute else
              [0, 1, 2, 2.0, 1.75, -1, *[n/240 for n in limits],
               (limits[1]+0.75)/240, (limits[0]-0.75)/240])
    for value in values:
        document = rdocx.Document()
        paragraph = document.add_paragraph('spacing')
        paragraph.paragraph_format.line_spacing = value
        expected = (abs(int(value))//635 * (-1 if value < 0 else 1)) if absolute else int(value*240)
        word = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        spacing = ET.fromstring(_document_xml(document)).find('.//'+word+'spacing')
        assert spacing is not None and spacing.get(word+'line') == str(expected)
        assert spacing.get(word+'lineRule') == ('exact' if absolute else 'auto')
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        formatting = reopened.paragraphs[0].paragraph_format
        assert formatting.line_spacing == (expected*635 if absolute else expected/240)
        formatting.line_spacing = None
        assert formatting.line_spacing is None


def test_fx256_line_spacing_refuses_nonfinite_types_and_overflow_atomically():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held')
    document.add_paragraph('other')
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    held, other = document.paragraphs
    run, section = held.runs[0], document.section(0)
    before = document.to_bytes()
    for value in [True, False, '2', object(), float('nan'), float('inf'), -float('inf'),
                  1e308, 2**10000, (2**31)/240, (-2**31-1)/240,
                  rdocx.Emu((2**31)*635), rdocx.Emu((-2**31-1)*635)]:
        with pytest.raises((TypeError, ValueError, OverflowError)):
            held.paragraph_format.line_spacing = value
        assert document.to_bytes() == before and document.verify_signatures()[0]['cryptographically_valid']
        assert held.text == run.text == 'held' and other.text == 'other' and section.xml
    held.paragraph_format.line_spacing = 2
    assert held.paragraph_format.line_spacing == 2.0


@pytest.mark.parametrize('name', ['page_width', 'page_height', 'margin_top', 'margin_right',
                                 'margin_bottom', 'margin_left', 'gutter', 'column_spacing',
                                 'header_distance', 'footer_distance'])
def test_fx256_update_section_requires_explicit_nonzero_lengths(name):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held')
    document.add_paragraph('other')
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    held, other = document.paragraphs
    run, section = held.runs[0], document.section(0)
    before = document.to_bytes()
    for value in [1440, 635, 1, -1, -635, True, 1.5, float('nan'), float('inf'), 2**100]:
        kwargs = {name: value, 'orientation': 'landscape'}
        with pytest.raises((TypeError, ValueError, OverflowError)):
            document.update_section(0, **kwargs)
        assert document.to_bytes() == before and document.verify_signatures()[0]['cryptographically_valid']
        assert held.text == run.text == 'held' and other.text == 'other' and section.xml
    assert document.update_section(0, **{name: None})
    assert document.to_bytes() == before
    if name in ('page_width', 'page_height'):
        with pytest.raises(rdocx.RdocxError):
            document.update_section(0, **{name: 0})
        assert document.to_bytes() == before
    else:
        assert getattr(document.update_section(0, **{name: 0}), name) == 0
    valid = rdocx.Inches(8.5 if name == 'page_width' else 11) if name in ('page_width', 'page_height') else rdocx.Twips(1440)
    updated = document.update_section(0, **{name: valid})
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert getattr(updated, name) == int(valid) == getattr(reopened.sections[0], name)


def test_fx256_live_section_keeps_its_existing_emu_contract():
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held')
    run, section = paragraph.runs[0], document.section(0)
    for value in [-635, -1, 634, 1440, rdocx.Inches(1)]:
        section.top_margin = value
        expected = abs(int(value))//635 * (-1 if value < 0 else 1) * 635
        assert section.margin_top == expected
        assert rdocx.Document.from_bytes(document.to_bytes()).sections[0].margin_top == expected
        assert document.update_section(0, margin_top=rdocx.Emu(value)).margin_top == expected
    before = document.to_bytes()
    with pytest.raises(ValueError, match='Twips|Emu'):
        document.update_section(0, margin_top=1440, margin_left=rdocx.Inches(2))
    assert document.to_bytes() == before and paragraph.text == run.text == 'held' and section.xml
    assert document.update_section(0, margin_top=rdocx.Twips(1440)).margin_top == rdocx.Inches(1)


@pytest.mark.parametrize('raising', [False, True])
def test_fx256_update_section_refuses_index_protocol_before_conversion(raising):
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph('held')
    run, section = paragraph.runs[0], document.section(0)
    before = document.to_bytes()
    calls = []
    class Unsupported:
        def __index__(self):
            calls.append('called')
            if raising:
                raise RuntimeError('unsupported index hook escaped')
            return 1440
    with pytest.raises(TypeError, match='Length|Twips|Emu'):
        document.update_section(0, margin_top=Unsupported())
    assert calls == []
    assert document.to_bytes() == before and paragraph.text == run.text == 'held' and section.xml


@pytest.mark.parametrize('child', [
    '<w:bogus/>', '<x:foo xmlns:x="urn:example"/>',
    '<m:f xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>',
    '<w:pPr/>',
])
def test_fx252_raw_xml_refusal_preserves_current_owners(child):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held body')
    document.add_paragraph('unrelated body')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'held cell'
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    held, unrelated = document.paragraphs[:2]
    run, section, cell = held.runs[0], document.section(0), document.tables[0].cell(0, 0)
    original = held.xml
    before = document.to_bytes()
    invalid = original.replace('</w:p>', child + '</w:p>')
    with pytest.raises((ValueError, rdocx.RdocxError)):
        held.replace_xml(invalid)
    assert document.to_bytes() == before
    assert held.xml == original and held.text == run.text == 'held body'
    assert unrelated.text == 'unrelated body' and cell.text == 'held cell' and section.xml
    assert document.verify_signatures()[0]['cryptographically_valid']
    # A late native relationship failure retains the same held authorities.
    late = '<w:hyperlink xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="missing"><w:r><w:t>late</w:t></w:r></w:hyperlink>'
    with pytest.raises(rdocx.RdocxError, match='relationship'):
        held.replace_xml(original.replace('</w:p>', late + '</w:p>'))
    assert document.to_bytes() == before and run.text == 'held body'


def test_fx252_foreign_source_and_stale_owner_cannot_lend_opaque_authority():
    import rdocx
    donor = _replace_document_body(rdocx.Document(), '<w:p><w:r><w:t>donor</w:t></w:r><x:opaque xmlns:x="urn:producer" keep="exact"/></w:p>')
    document = rdocx.Document()
    document.add_paragraph('held')
    held = document.paragraphs[0]
    donor_before, before = donor.to_bytes(), document.to_bytes()
    with pytest.raises((ValueError, rdocx.RdocxError)):
        held.replace_xml(donor.paragraphs[0].xml)
    assert document.to_bytes() == before and donor.to_bytes() == donor_before
    assert held.text == 'held' and donor.paragraphs[0].text == 'donor'
    stale = document.add_paragraph('retire me')
    stale_xml = stale.xml
    assert document.remove_content(document.find_content_index(stale))
    held = document.paragraphs[0]
    before = document.to_bytes()
    with pytest.raises(rdocx.StaleElementError):
        stale.replace_xml(stale_xml)
    assert document.to_bytes() == before and held.text == 'held'


@pytest.mark.parametrize('compatibility', ['', ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:MustUnderstand="w"'])
def test_fx252_typed_sdt_metadata_edits_preserve_unmodelled_property_slots(compatibility):
    import rdocx
    opaque = "<f:keep xmlns:f=\"urn:metadata\" data='exact'/>"
    body = '<w:p><w:sdt><w:sdtPr><w:id w:val="10"/><w:alias w:val="Alias"/><w:tag w:val="key"/><w:date w:fullDate="2026-10-16T00:00:00Z"><w:dateFormat w:val="yyyy-MM-dd"/></w:date>' + opaque + '</w:sdtPr><w:sdtContent><w:r><w:t>held</w:t></w:r></w:sdtContent></w:sdt></w:p>'
    body = body.replace('<w:sdtPr>', '<w:sdtPr' + compatibility + '>')
    document = _replace_document_body(rdocx.Document(), body)
    held = document.paragraphs[0]
    original = held.xml
    changed = original.replace('w:val="Alias"', 'w:val="Changed"').replace('yyyy-MM-dd', 'dd/MM/yyyy')
    assert changed != original
    held.replace_xml(changed)
    assert held.text == 'held'
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    xml = reopened.paragraphs[0].xml
    assert 'w:val="Changed"' in xml and 'dd/MM/yyyy' in xml and opaque in xml


def test_fx252_markers_math_and_imported_opaque_scalar_edits_reopen():
    import xml.etree.ElementTree as ET
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held')
    held = document.paragraphs[0]
    original = held.xml
    before = document.to_bytes()
    held.replace_xml(original)
    assert document.to_bytes() == before
    additions = '<w:bookmarkStart w:id="42" w:name="checked"/><w:bookmarkEnd w:id="42"/><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'
    held.replace_xml(original.replace('</w:p>', additions + '</w:p>'))
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    root = ET.fromstring(reopened.paragraphs[0].xml)
    w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    m = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    assert root.find('{' + w + '}bookmarkStart').get('{' + w + '}name') == 'checked'
    assert root.find('{' + w + '}bookmarkEnd').get('{' + w + '}id') == '42'
    assert root.find('{' + m + '}oMath/{' + m + '}r/{' + m + '}t').text == 'x'
    opaque = "<x:opaque xmlns:x=\"urn:producer\" keep='exact'><x:nested/></x:opaque>"
    document = _replace_document_body(rdocx.Document(), '<w:p><w:r><w:t>old</w:t></w:r>' + opaque + '</w:p><w:p><w:r><w:t>other</w:t></w:r></w:p>')
    held, other = document.paragraphs[:2]
    original = held.xml
    before = document.to_bytes()
    held.replace_xml(original)
    assert document.to_bytes() == before
    held.replace_xml(original.replace('>old<', '>new<'))
    assert held.text == 'new' and other.text == 'other'
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert opaque in reopened.paragraphs[0].xml and reopened.paragraphs[1].text == 'other'


@pytest.mark.parametrize('operation', ['pdf', 'duplicate', 'header', 'language', 'font'])
def test_fx251_five_reporter_invalid_values_refuse_atomically(tmp_path, operation):
    from pathlib import Path
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('held body')
    if operation == 'duplicate':
        document.paragraphs[0].paragraph_format.tab_stops.add_tab_stop(rdocx.Inches(1))
    if operation == 'header':
        document.section(0).header.add_paragraph('held header')
    key, certificate = _fx240_material()
    document.sign(key, certificate)
    paragraph, section = document.paragraphs[0], document.section(0)
    run = paragraph.runs[0]
    before = document.to_bytes()
    if operation == 'pdf':
        for name, existing in [('absent.PdF', None), ('existing.pdf', b'keep existing output')]:
            path = tmp_path / name
            if existing is not None:
                path.write_bytes(existing)
            with pytest.raises(ValueError, match='to_pdf'):
                document.save(path)
            assert path.read_bytes() == existing if existing is not None else not path.exists()
    elif operation == 'duplicate':
        tabs = paragraph.paragraph_format.tab_stops
        for value in (rdocx.Inches(1), rdocx.Inches(1) + 634):
            with pytest.raises(ValueError, match='position|duplicate'):
                tabs.add_tab_stop(value)
        assert len(tabs) == 1
    elif operation == 'header':
        header = section.header.paragraphs[-1]
        header_run = header.runs[0]
        with pytest.raises(ValueError, match='header'):
            header_run.add_break('page')
        assert header.text == header_run.text == 'held header'
    elif operation == 'language':
        for invalid in ('not a tag!', '', '-en', 'en-', 'en--US', '1-en', 'en_US', 'é-US', 'en-!'):
            with pytest.raises(ValueError, match='language|tag'):
                run.font.language = invalid
    else:
        font = bytearray((Path(__file__).resolve().parents[2] / 'oxml-layout/fonts/Caladea-Regular.ttf').read_bytes())
        count = int.from_bytes(font[4:6], 'big')
        offset = next(index for index in range(12, 12 + count * 16, 16) if font[index:index + 4] == b'name')
        font[offset:offset + 4] = b'nam_'
        font.extend(b'name')
        head = next(index for index in range(12, 12 + count * 16, 16) if font[index:index + 4] == b'head')
        adjustment = int.from_bytes(font[head + 8:head + 12], 'big') + 8
        font[adjustment:adjustment + 4] = b'\0' * 4
        padded = bytes(font) + b'\0' * (-len(font) % 4)
        checksum = sum(value[0] for value in struct.iter_unpack('>I', padded)) & 0xffffffff
        font[adjustment:adjustment + 4] = ((0xB1B0AFBA - checksum) & 0xffffffff).to_bytes(4, 'big')
        license = rdocx.FontEmbeddingLicense(authorized=True, identity='Apache-2.0: Caladea')
        for explicit in (None, '{00112233-4455-6677-8899-AABBCCDDEEFF}'):
            with pytest.raises(rdocx.RdocxError):
                document.embed_font('Refused font', bytes(font), license=license, font_key=explicit)
    assert document.to_bytes() == before
    assert paragraph.text == run.text == 'held body' and section.xml
    assert document.verify_signatures()[0]['cryptographically_valid']


def test_fx251_valid_paths_tabs_breaks_and_language_remain_supported(tmp_path):
    import xml.etree.ElementTree as ET
    import rdocx
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    document = rdocx.Document()
    document.add_paragraph('valid source')
    for extension in ('.docx', '.dotx', '.docm', '.dotm', '.DOCX', '.custom', ''):
        path = tmp_path / ('output' + extension)
        document.save(path)
        assert rdocx.Document.open(path).paragraphs[0].text == 'valid source'
    class NamedStream(io.BytesIO):
        name = 'a-stream.pdf'
    stream = NamedStream()
    document.save(stream)
    assert not stream.closed and rdocx.Document.open(io.BytesIO(stream.getvalue())).paragraphs[0].text == 'valid source'
    paragraph = document.paragraphs[0]
    tabs = paragraph.paragraph_format.tab_stops
    for position in (-635, 0, 635, 1270):
        tabs.add_tab_stop(position)
    assert [int(tab.position) for tab in tabs] == [-635, 0, 635, 1270]
    before = document.to_bytes()
    with pytest.raises(ValueError):
        tabs.add_tab_stop(-634)
    assert document.to_bytes() == before
    for language in ('en-US', 'de-DE', 'zh-Hans', 'x-example', 'en-001', 'a', 'a-123456789', None):
        paragraph.runs[0].font.language = language
        assert document.paragraphs[0].runs[0].font.language == language
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        assert reopened.paragraphs[0].runs[0].font.language == language
    for host, kind in [('body', 'page'), ('header', 'line'), ('footer', 'page'), ('header', 'column')]:
        owner = rdocx.Document()
        if host == 'body':
            target = owner.add_paragraph('control')
        else:
            target = getattr(owner.section(0), host).add_paragraph('control')
        target.runs[0].add_break(kind)
        with zipfile.ZipFile(io.BytesIO(owner.to_bytes())) as package:
            part = 'word/document.xml' if host == 'body' else next(name for name in package.namelist() if name.startswith('word/' + host) and name.endswith('.xml'))
            root = ET.fromstring(package.read(part))
        br = root.find('.//' + w + 'br')
        assert br is not None and br.get(w + 'type', 'textWrapping') == ('textWrapping' if kind == 'line' else kind)
    malformed_language = _replace_document_body(rdocx.Document(), '<w:p><w:r><w:rPr><w:lang w:val="not a tag!" w:eastAsia="ja-JP" w:bidi="ar-SA"/></w:rPr><w:t>imported</w:t></w:r></w:p><w:sectPr/>')
    assert malformed_language.paragraphs[0].runs[0].font.language == 'not a tag!'
    malformed_language.add_paragraph('unrelated')
    reopened = rdocx.Document.from_bytes(malformed_language.to_bytes())
    assert reopened.paragraphs[0].runs[0].font.language == 'not a tag!'
    assert reopened.paragraphs[0].runs[0].font.language_east_asia == 'ja-JP'
    assert reopened.paragraphs[0].runs[0].font.language_bidi == 'ar-SA'
    # Imported duplicates remain source-owned. Only a newly repeated position refuses.
    duplicate = '<w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="10"/><w:tab w:val="right" w:pos="10"/></w:tabs></w:pPr><w:r><w:t>imported</w:t></w:r></w:p><w:sectPr/>'
    imported = _replace_document_body(rdocx.Document(), duplicate)
    tabs = imported.paragraphs[0].paragraph_format.tab_stops
    assert len(tabs) == 2
    tabs.add_tab_stop(20 * 635)
    assert len(tabs) == 3
    assert [int(tab.position) for tab in tabs] == [10 * 635, 10 * 635, 20 * 635]
    opaque = _replace_document_body(rdocx.Document(), duplicate.replace('w:val="right"', 'w:val="futureAlignment"'))
    opaque.paragraphs[0].paragraph_format.tab_stops.add_tab_stop(20 * 635)
    assert b'futureAlignment' in _document_xml(opaque)


@pytest.mark.parametrize('host', ['header', 'header-cell', 'renamed-header', 'renamed-header-cell', 'renamed-body'])
def test_fx251_page_break_guard_uses_physical_header_owner(host):
    import xml.etree.ElementTree as ET
    import rdocx
    document = rdocx.Document()
    document.add_paragraph('body')
    if host != 'renamed-body':
        header = document.section(0).header
        if host.endswith('cell'):
            header.add_table(1, 1).cell(0, 0).text = 'header cell'
        else:
            header.add_paragraph('header')
    if host.startswith('renamed'):
        source, output = io.BytesIO(document.to_bytes()), io.BytesIO()
        with zipfile.ZipFile(source) as zin, zipfile.ZipFile(output, 'w') as zout:
            old = 'word/document.xml' if host == 'renamed-body' else next(name for name in zin.namelist() if name.startswith('word/header') and name.endswith('.xml'))
            new = 'word/headerPretender.xml' if host == 'renamed-body' else 'word/furnitureOdd.xml'
            for entry in zin.infolist():
                raw = zin.read(entry.filename)
                name = entry.filename
                if name == '[Content_Types].xml':
                    raw = raw.replace(('/' + old).encode(), ('/' + new).encode())
                if name.endswith('.rels'):
                    if host == 'renamed-body' and name == '_rels/.rels':
                        raw = raw.replace(old.encode(), new.encode())
                    elif host != 'renamed-body' and name == 'word/_rels/document.xml.rels':
                        raw = raw.replace(old.removeprefix('word/').encode(), new.removeprefix('word/').encode())
                if name == old:
                    name = new
                elif host == 'renamed-body' and name == 'word/_rels/document.xml.rels':
                    name = 'word/_rels/headerPretender.xml.rels'
                zout.writestr(name, raw)
        document = rdocx.Document.from_bytes(output.getvalue())
    if host == 'renamed-body':
        document.paragraphs[0].runs[0].add_break('page')
        assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == 'body\n'
        return
    header = document.section(0).header
    paragraph = header.tables[0].cell(0, 0).paragraphs[0] if host.endswith('cell') else header.paragraphs[-1]
    run = paragraph.runs[0]
    before, text = document.to_bytes(), paragraph.text
    with pytest.raises(ValueError, match='header'):
        run.add_break('page')
    assert document.to_bytes() == before and paragraph.text == run.text == text
    # An imported header page break survives no-op save and unrelated body work.
    source, output = io.BytesIO(before), io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(output, 'w') as zout:
        part = next(item.part_name.lstrip('/') for item in document.stories if item.kind == 'header')
        for entry in zin.infolist():
            raw = zin.read(entry.filename)
            if entry.filename == part:
                raw = raw.replace(b'</w:r>', b'<w:br w:type="page"/></w:r>', 1)
            zout.writestr(entry, raw)
    imported = rdocx.Document.from_bytes(output.getvalue())
    imported.add_paragraph('unrelated')
    with zipfile.ZipFile(io.BytesIO(imported.to_bytes())) as package:
        assert ET.fromstring(package.read(part)).find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br') is not None


@pytest.mark.parametrize("identifier", ["paraId", "textId", "both"])
def test_fx248_identified_block_control_all_validation_routes(tmp_path, identifier):
    import xml.etree.ElementTree as ET

    import rdocx

    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    producer = "http://schemas.microsoft.com/office/word/2010/wordml"
    attributes = {
        "paraId": 'producer:paraId="388B6825"',
        "textId": 'producer:textId="00AABBCC"',
        "both": 'producer:paraId="388B6825" producer:textId="00AABBCC"',
    }[identifier]
    opaque = '<x:witness x:value="&amp;">keep &amp; exact</x:witness>'
    control = (
        f'<w:sdt><w:sdtPr><w:id w:val="1"/>{opaque}</w:sdtPr><w:sdtContent>'
        f'<w:p {attributes}><w:r><w:t>Inside</w:t></w:r></w:p>'
        f'<w:sdt><w:sdtPr/><w:sdtContent><w:p {attributes}>'
        '<w:r><w:t>Nested</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        '</w:sdtContent></w:sdt>'
    )
    seed = rdocx.Document()
    seed.add_paragraph("Before")
    source = io.BytesIO(seed.to_bytes())
    path = tmp_path / "identified.docx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as output:
        for info in original.infolist():
            payload = original.read(info.filename)
            if info.filename == "word/document.xml":
                xml = payload.decode().replace(
                    "<w:document ",
                    f'<w:document xmlns:producer="{producer}" xmlns:x="urn:opaque" ',
                    1,
                )
                payload = xml.replace("<w:sectPr", control + "<w:sectPr", 1).encode()
            output.writestr(info, payload)
    before = path.read_bytes()
    document = rdocx.Document.open(path)
    document.validate()
    rdocx.validate_file(path)
    assert path.read_bytes() == before
    saved = tmp_path / "saved.docx"
    document.save(saved)
    rdocx.Document.open(saved).validate()
    rdocx.validate_file(saved)
    with zipfile.ZipFile(saved) as package:
        xml = package.read("word/document.xml")
    assert opaque.encode() in xml
    root = ET.fromstring(xml)
    paragraphs = root.findall(f".//{{{word}}}sdtContent/{{{word}}}p")
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        if identifier in ("paraId", "both"):
            assert paragraph.get(f"{{{producer}}}paraId") == "388B6825"
        if identifier in ("textId", "both"):
            assert paragraph.get(f"{{{producer}}}textId") == "00AABBCC"


def _replace_document_body(document, body, root_default=None):
    source = io.BytesIO(document.to_bytes())
    result = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip:
        with zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/document.xml":
                    start = data.index(b"<w:body>") + len(b"<w:body>")
                    end = data.index(b"</w:body>")
                    data = data[:start] + body.encode() + data[end:]
                    if root_default:
                        root = f'<w:document xmlns="{root_default}" '
                        data = data.replace(b"<w:document ", root.encode(), 1)
                result_zip.writestr(info, data)
    return type(document).from_bytes(result.getvalue())


def _document_xml(document):
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        return archive.read("word/document.xml")


def _replace_settings_xml(document, settings):
    source = io.BytesIO(document.to_bytes())
    result = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip:
        with zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/settings.xml":
                    data = settings.encode()
                result_zip.writestr(info, data)
    return type(document).from_bytes(result.getvalue())


def _tracked_document():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("alpha")
    edited = rdocx.Document()
    edited.add_paragraph("alpha beta")
    document = rdocx.Document.from_bytes(original.to_bytes())
    document.compare(edited, "Ada", "2026-01-02T03:04:05Z")
    return document


def test_issue_253_python_render_views():
    from pathlib import Path

    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Keep OLDWORD here.")
    edited = rdocx.Document()
    edited.add_paragraph("Keep NEWWORD here.")
    original.compare(edited, "Reviewer", "2026-09-30T12:00:00Z")
    accepted = original.to_pdf()
    assert accepted == original.to_pdf(revision_view="accepted")
    assert accepted != original.to_pdf(revision_view="tracked")
    fonts = Path(__file__).parents[2] / "oxml-layout" / "fonts"
    assert original.to_pdf(font_dir=fonts) == original.to_pdf(
        font_dir=fonts, revision_view="accepted"
    )
    assert original.to_pdf(font_dir=fonts) != original.to_pdf(
        font_dir=fonts, revision_view="tracked"
    )
    for render in (
        lambda **view: original.render_page_to_png(0, 24.0, **view),
        lambda **view: original.render_all_pages(24.0, **view),
        lambda **view: original.render_pages(dpi=24.0, **view),
    ):
        assert render() == render(revision_view="accepted")
        assert render() != render(revision_view="tracked")
        with pytest.raises(ValueError, match="unknown revision view"):
            render(revision_view="final")
    with pytest.raises(ValueError, match="unknown revision view"):
        original.to_pdf(revision_view="final")


def test_revision_style_colours_only_the_tracked_view():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Keep OLDWORD here.")
    edited = rdocx.Document()
    edited.add_paragraph("Keep NEWWORD here.")
    original.compare(edited, "Reviewer", "2026-09-30T12:00:00Z")
    for render in (
        lambda **view: original.to_pdf(**view),
        lambda **view: original.render_page_to_png(0, 24.0, **view),
        lambda **view: original.render_all_pages(24.0, **view),
        lambda **view: original.render_pages(dpi=24.0, **view),
    ):
        plain = render(revision_view="tracked")
        assert plain == render(revision_view="tracked", revision_style="plain")
        assert plain != render(revision_view="tracked", revision_style="legal")
        assert render() == render(revision_style="plain")
        for view in ({}, {"revision_view": "accepted"}):
            with pytest.raises(ValueError, match="revision_style applies only"):
                render(revision_style="legal", **view)
        with pytest.raises(ValueError, match="unknown revision style"):
            render(revision_view="tracked", revision_style="color")


def test_issue_253_selected_view_controls_page_count():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("OLDWORD " * 1500)
    edited = rdocx.Document()
    edited.add_paragraph("NEWWORD " * 1500)
    original.compare(edited, "Reviewer", "2026-09-30T12:00:00Z")
    accepted = original.render_pages(dpi=24.0)
    tracked = original.render_pages(dpi=24.0, revision_view="tracked")
    assert len(tracked) > len(accepted)
    assert tracked == original.render_all_pages(24.0, revision_view="tracked")


def test_issue_253_pdf_text_view(tmp_path):
    import rdocx

    version = subprocess.run(
        ["pdftotext", "-v"], capture_output=True, check=True, text=True
    )
    assert "pdftotext version 26.01.0" in version.stderr.splitlines()

    original = rdocx.Document()
    original.add_paragraph("Keep OLDWORD here.")
    edited = rdocx.Document()
    edited.add_paragraph("Keep NEWWORD here.")
    original.compare(edited, "Reviewer", "2026-09-30T12:00:00Z")

    def text(view):
        path = tmp_path / "redline.pdf"
        path.write_bytes(original.to_pdf(revision_view=view))
        return subprocess.run(
            ["pdftotext", str(path), "-"], capture_output=True, check=True, text=True
        ).stdout

    accepted = text("accepted")
    tracked = text("tracked")
    assert "NEWWORD" in accepted and "OLDWORD" not in accepted
    assert "NEWWORD" in tracked and "OLDWORD" in tracked


def test_revisions_are_snapshots_and_resolution_reports_counts():
    import rdocx

    document = _tracked_document()
    revisions = document.revisions
    assert revisions
    assert {revision.author for revision in revisions} == {"Ada"}
    kinds = {revision.kind for revision in revisions}
    assert "insertion" in kinds
    assert kinds <= {
        "insertion",
        "deletion",
        "move_from",
        "move_to",
        "run_property_change",
        "paragraph_property_change",
        "table_property_change",
        "section_property_change",
    }
    assert {revision.timestamp for revision in revisions} == {"2026-01-02T03:04:05Z"}

    held = document.paragraphs[0]
    assert document.reject_revisions_by_author("Grace") == 0
    held.text
    with pytest.raises(rdocx.RdocxError):
        document.accept_revisions_in_date_range(start="yesterday", end="2026-01-03T00:00:00Z")
    held.text

    assert (
        document.accept_revisions_in_date_range(
            start="2026-01-01T00:00:00Z", end="2026-01-03T00:00:00Z"
        )
        > 0
    )
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert document.revisions == ()
    assert [paragraph.text for paragraph in document.paragraphs] == ["alpha beta"]


def test_reject_revision_id_then_reject_all_restores_the_original():
    document = _tracked_document()
    first = document.revisions[0]
    assert document.reject_revision_id(first.id) > 0
    assert all(revision.id != first.id for revision in document.revisions)
    document.reject_all()
    assert document.revisions == ()
    assert [paragraph.text for paragraph in document.paragraphs] == ["alpha"]


def test_revisions_list_every_story_and_name_the_story_that_holds_them():
    import rdocx

    def footer_document(text):
        document = rdocx.Document()
        document.add_paragraph("Body.")
        document.set_footer(text)
        return rdocx.Document.from_bytes(document.to_bytes())

    document = footer_document("Footer lorem ipsum")
    document.compare(
        footer_document("Footer lorem IPSUM"), "R", "2026-09-27T12:00:00Z"
    )
    revisions = document.revisions
    footer = next(story for story in document.stories if story.kind == "footer")
    assert footer.part_name == "/word/footer1.xml"
    assert len(revisions) == 2
    assert all(revision.story == footer for revision in revisions)
    assert sorted(revision.kind for revision in revisions) == ["deletion", "insertion"]
    assert {revision.author for revision in revisions} == {"R"}

    body = rdocx.Story(kind="body", part_name="/word/document.xml", owner_index=0)
    assert {revision.story == body for revision in _tracked_document().revisions} == {
        True
    }
    unscoped = rdocx.Revision(id=1, author="R", timestamp=None, kind="insertion")
    assert unscoped.story is None
    scoped = rdocx.Revision(
        id=1, author="R", timestamp=None, kind="insertion", story=footer
    )
    assert scoped.story == footer
    assert scoped != unscoped

    assert document.accept_all() == len(revisions)
    assert document.revisions == ()


def test_counted_replacement_spans_runs_and_a_bad_regex_changes_nothing():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Dear {{na").add_run("me}}, hello")
    held = document.paragraphs[0]
    assert document.try_replace_text("{{missing}}", "x") == 0
    held.text
    assert document.try_replace_text("{{name}}", "Ada") == 1
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert document.paragraphs[0].text == "Dear Ada, hello"

    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="invalid regex"):
        document.replace_all_regex([("hello", "bye"), ("(", "x")])
    assert document.to_bytes() == before
    assert document.replace_all_regex([("h(el)lo", "bye")]) == 1
    assert document.paragraphs[0].text == "Dear Ada, bye"


def test_counted_replacement_checks_expected_counts_before_publishing():
    import pickle

    import rdocx

    document = rdocx.Document()
    document.add_paragraph("{{a}} {{b}} {{b}} {{c}}")
    document.set_header("{{c}} header")
    held = document.paragraphs[0]
    before = document.to_bytes()

    with pytest.raises(rdocx.ReplacementCountError) as raised:
        document.replace_all([("{{a}}", "A", 1), ("{{b}}", "B", 1), ("{{c}}", "C", 2)])
    error = raised.value
    assert isinstance(error, rdocx.RdocxError)
    assert (error.index, error.expected, error.found) == (1, 1, 2)
    assert str(error) == 'pair 1: expected 1 replacement(s) of "{{b}}", found 2'
    # Worker pools pickle exceptions to return them, so the error must survive.
    copy = pickle.loads(pickle.dumps(error))
    assert (type(copy), str(copy), copy.index, copy.expected, copy.found) == (
        rdocx.ReplacementCountError,
        str(error),
        1,
        1,
        2,
    )
    assert document.to_bytes() == before
    assert held.text == "{{a}} {{b}} {{b}} {{c}}"

    with pytest.raises(
        rdocx.ReplacementCountError,
        match=r'^expected 3 replacement\(s\) of "\{\{b\}\}", found 2$',
    ) as raised:
        document.try_replace_text("{{b}}", "B", expect=3)
    assert (raised.value.index, raised.value.expected, raised.value.found) == (None, 3, 2)
    assert document.to_bytes() == before
    assert document.try_replace_text("{{missing}}", "x", expect=0) == 0
    assert document.replace_all([("{{missing}}", "x")]) == (0,)
    assert document.replace_all([]) == ()
    assert held.text == "{{a}} {{b}} {{b}} {{c}}"

    # Pairs run in order, so the second one also replaces what the first wrote.
    assert document.replace_all(
        [("{{a}}", "{{b}}", 1), ("{{b}}", "B", 3), ("{{c}}", "C", None)]
    ) == (1, 3, 2)
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert document.paragraphs[0].text == "B B B C"
    assert _story_paragraph_texts(document, "header") == ["C header"]
    assert document.try_replace_text("B", "b", expect=3) == 3

    for pairs, message in (
        ([("x",)], "pair 0 must be"),
        ([("x", "y"), ("x", "y", -1)], "pair 1 must be"),
        ("xy", "pair 0 must be"),
    ):
        with pytest.raises(TypeError, match=message):
            document.replace_all(pairs)


def test_update_fields_takes_a_keyword_context_and_counts_updates():
    import datetime

    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:fldSimple w:instr=" FILENAME "><w:r><w:t>old.docx</w:t></w:r></w:fldSimple></w:p>'
        '<w:p><w:fldSimple w:instr=" MERGEFIELD Name "><w:r><w:t>Name</w:t></w:r></w:fldSimple></w:p>'
        '<w:p><w:fldSimple w:instr=" DATE \\@ &quot;yyyy-MM-dd&quot; "><w:r><w:t>2000-01-01</w:t></w:r></w:fldSimple></w:p>',
    )
    held = document.paragraphs[0]
    count = document.update_fields(
        now=datetime.datetime(2026, 9, 15, 10, 30),
        file_name="report.docx",
        merge_fields={"Name": "Ada"},
    )
    assert count == 3
    with pytest.raises(rdocx.StaleElementError):
        held.text
    xml = _document_xml(document)
    for cached in (b"report.docx", b"Ada", b"2026-09-15"):
        assert cached in xml
    assert b"old.docx" not in xml and b"2000-01-01" not in xml


def test_python_story_revision_field_and_xml_operations_are_typed_and_atomic():
    import rdocx

    document = _tracked_document()
    revisions = document.revisions
    assert revisions
    assert {revision.author for revision in revisions} == {"Ada"}
    assert document.accept_all() == len(revisions)

    document.set_header("Draft")
    document.set_footer("Page footer")
    header = next(story for story in document.stories if story.kind == "header")
    document.add_hyperlink_to_story(header, "home", "https://example.com/")
    header_item = next(
        item
        for item in document.story_items
        if item.story == header and item.kind == "paragraph"
    )
    assert isinstance(header_item.xml, bytes)
    document.set_story_text(header_item, "Final")

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert "Final" in [item.text for item in reopened.story_items]
    assert [(link.text, link.url) for link in reopened.hyperlinks] == [
        ("home", "https://example.com/")
    ]

    with pytest.raises(rdocx.StaleElementError, match="story item handle"):
        document.set_story_text(header_item, "stale")

    linked = rdocx.Document()
    paragraph = linked.add_paragraph("See ")
    paragraph.add_hyperlink("old link", "https://example.com/old")
    linked_item = next(
        item
        for item in linked.story_items
        if item.story.kind == "body" and item.kind == "paragraph"
    )
    linked.set_story_text(linked_item, "New text without a link")
    assert linked.hyperlinks == ()
    assert b"<w:hyperlink" not in _document_xml(linked)

    commented = rdocx.Document()
    commented.add_paragraph("commented")
    commented.add_paragraph("destination")
    commented.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Ada",
        text="review",
    )
    commented.clone_content(commented.paragraphs[0], 0)
    cloned_xml = _document_xml(commented)
    assert cloned_xml.count(b"commentRangeStart") == 1
    assert cloned_xml.count(b"commentRangeEnd") == 1
    assert cloned_xml.count(b"commentReference") == 1

    lookup = _replace_document_body(
        rdocx.Document(),
        '<w:sdt><w:sdtContent><w:p><w:r><w:t>Background</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        '<w:p><w:r><w:t>Background</w:t></w:r></w:p>',
    )
    assert lookup.find_content_index("Background") == 1
    assert lookup.find_content_indices("Background") == (1, 0)


def test_paragraph_views_include_accepted_nested_runs_in_source_order():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:t xml:space="preserve">Start </w:t></w:r>'
        '<w:ins w:id="1" w:author="Ada"><w:r><w:t xml:space="preserve">inserted </w:t></w:r></w:ins>'
        '<w:sdt><w:sdtPr><w:id w:val="7"/></w:sdtPr><w:sdtContent>'
        '<w:r><w:t xml:space="preserve">controlled </w:t></w:r>'
        '<w:ins w:id="3" w:author="Ada"><w:r><w:t xml:space="preserve">nested </w:t></w:r></w:ins>'
        '</w:sdtContent></w:sdt>'
        '<w:del w:id="2" w:author="Ada"><w:r><w:delText>deleted </w:delText></w:r></w:del>'
        '<w:moveFrom w:id="4" w:author="Ada"><w:r><w:t>moved away </w:t></w:r></w:moveFrom>'
        '<w:r><w:t>end</w:t></w:r></w:p>',
    )
    item = next(
        item
        for item in document.story_items
        if item.story.kind == "body" and item.kind == "paragraph"
    )
    paragraph = document.paragraphs[0]

    assert item.text == "Start inserted controlled nested end"
    assert paragraph.text == item.text
    assert [run.text for run in paragraph.runs] == [
        "Start ",
        "inserted ",
        "controlled ",
        "nested ",
        "end",
    ]

    inserted = paragraph.runs[1]
    controlled = paragraph.runs[2]
    nested = paragraph.runs[3]
    inserted.text = "changed "
    controlled.font.bold = True
    nested.text = "composed "
    assert inserted.text == "changed "
    assert controlled.font.bold is True
    assert nested.text == "composed "

    boundary = document.split_run(body_index=0, run_index=1, character_offset=4)
    assert boundary == 2
    with pytest.raises(rdocx.StaleElementError):
        _ = inserted.text
    assert [run.text for run in document.paragraphs[0].runs] == [
        "Start ",
        "chan",
        "ged ",
        "controlled ",
        "composed ",
        "end",
    ]

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.paragraphs[0].text == "Start changed controlled composed end"
    assert reopened.paragraphs[0].runs[3].font.bold is True
    assert reopened.paragraphs[0].runs[4].text == "composed "


def test_python_story_inventory_scales_linearly():
    import rdocx

    def timed_snapshot(paragraph_count):
        body = []
        for index in range(paragraph_count):
            body.append(
                f'<w:p><w:r><w:t>paragraph {index}</w:t></w:r>'
                f'<w:hyperlink w:anchor="target{index}"><w:r><w:t>link</w:t></w:r></w:hyperlink></w:p>'
            )
            if index % 10 == 0:
                cells = "".join(
                    f'<w:tc><w:p><w:r><w:t>cell {index} {cell}</w:t></w:r></w:p></w:tc>'
                    for cell in range(4)
                )
                body.append(f"<w:tbl><w:tr>{cells}</w:tr></w:tbl>")
        document = _replace_document_body(rdocx.Document(), "".join(body))
        started = time.perf_counter()
        items = document.story_items
        links = document.hyperlinks
        elapsed = time.perf_counter() - started
        return elapsed, len(items), len(links)

    small_elapsed, small_items, small_links = timed_snapshot(50)
    large_elapsed, large_items, large_links = timed_snapshot(100)

    assert large_items == small_items * 2
    assert large_links == small_links * 2
    assert large_elapsed <= small_elapsed * 3.0 + 0.02, (
        f"doubling the story inventory took {large_elapsed / small_elapsed:.2f} times longer"
    )


def test_update_fields_on_open_sets_clears_and_removes_the_setting():
    import rdocx

    document = rdocx.Document()
    assert document.update_fields_on_open is None
    for value in (True, False, None):
        document.update_fields_on_open = value
        assert document.update_fields_on_open is value
        document = rdocx.Document.from_bytes(document.to_bytes())
        assert document.update_fields_on_open is value

    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for settings in (
        f'<w:settings xmlns:w="{word}"><w:updateFields/><w:updateFields w:val="false"/></w:settings>',
        f'<w:settings xmlns:w="{word}"><w:updateFields w:val="invalid"/></w:settings>',
    ):
        document = _replace_settings_xml(rdocx.Document(), settings)
        before = document.to_bytes()
        assert document.update_fields_on_open is None
        with pytest.raises(rdocx.XmlError, match="ambiguous or malformed"):
            document.update_fields_on_open = True
        assert document.to_bytes() == before


def _one_pixel_png(pixel=b"\xff\xff\xff"):
    def chunk(kind, data):
        crc = struct.pack(">I", zlib.crc32(kind + data))
        return struct.pack(">I", len(data)) + kind + data + crc

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(b"\x00" + pixel))
        + chunk(b"IEND", b"")
    )


def _body_images(data):
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        xml = package.read("word/document.xml").decode()
        relationships = package.read("word/_rels/document.xml.rels").decode()
        targets = {
            re.search(r'Id="([^"]+)"', element).group(1): re.search(
                r'Target="([^"]+)"', element
            ).group(1)
            for element in re.findall(r"<Relationship [^>]*>", relationships)
        }
        return [
            package.read(posixpath.join("word", targets[embed]))
            for embed in re.findall(r'r:embed="([^"]+)"', xml)
        ]


@pytest.mark.parametrize(
    ("caption", "size"),
    [
        ("Figure 1. Caption.", (1, 0.66)),
        ("Figure 1. New caption.", (1, 0.66)),
        ("Figure 1. Caption.", (1.2, 0.8)),
    ],
)
def test_compare_records_a_picture_whose_image_changed(caption, size):
    import rdocx

    red, blue = _one_pixel_png(b"\xff\x00\x00"), _one_pixel_png(b"\x00\x00\xff")

    def document(image, caption="Figure 1. Caption.", size=(1, 0.66)):
        document = rdocx.Document()
        document.add_paragraph("Before the figure.")
        document.add_picture(
            image, "figure1.png", rdocx.Inches(size[0]), rdocx.Inches(size[1])
        )
        document.add_paragraph(caption)
        return document

    redline = document(red)
    redline.compare(document(blue, caption, size), "Reviewer", "2026-09-30T12:00:00Z")
    assert redline.revisions
    tracked = redline.to_bytes()
    assert _body_images(tracked) == [red, blue]
    for resolve, image, text in (
        ("accept_all", blue, caption),
        ("reject_all", red, "Figure 1. Caption."),
    ):
        resolved = rdocx.Document.from_bytes(tracked)
        getattr(resolved, resolve)()
        assert _body_images(resolved.to_bytes()) == [image]
        assert [paragraph.text for paragraph in resolved.paragraphs][-1] == text


def test_add_picture_beside_a_content_control_ignores_an_unused_root_default():
    import rdocx

    body = (
        "<w:p><w:r><w:t>Before the control.</w:t></w:r></w:p>"
        '<w:sdt><w:sdtPr><w:tag w:val="goog_rdk_0"/></w:sdtPr><w:sdtContent>'
        "<w:p><w:r><w:t>Inside the control.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
        "<w:p><w:r><w:t>After the control.</w:t></w:r></w:p><w:sectPr/>"
    )
    tasks = "http://schemas.microsoft.com/office/tasks/2019/documenttasks"
    for root_default in (None, tasks):
        document = _replace_document_body(rdocx.Document(), body, root_default)
        control = next(
            item
            for item in document.story_items
            if item.story.kind == "body" and item.kind == "content_control"
        )
        size = {"width": rdocx.Inches(1), "height": rdocx.Inches(1)}
        after = document.add_picture(_one_pixel_png(), "x.png", after=control, **size)
        appended = document.add_picture(_one_pixel_png(), "x.png", **size)
        assert isinstance(after, rdocx.Picture) and isinstance(appended, rdocx.Picture)
        assert after.location.item_kind == appended.location.item_kind == "paragraph"

        xml = _document_xml(rdocx.Document.from_bytes(document.to_bytes()))
        assert xml.count(b"<w:drawing>") == 2
        positions = [
            xml.index(b"Before the control."),
            xml.index(b'<w:tag w:val="goog_rdk_0"/>'),
            xml.index(b"Inside the control."),
            xml.index(b"</w:sdt>"),
            xml.index(b"<w:drawing>"),
            xml.index(b"After the control."),
            xml.rindex(b"<w:drawing>"),
        ]
        assert positions == sorted(positions)


def test_drawings_of_one_part_that_repeat_an_id_open_and_accept_a_new_picture():
    docx = pytest.importorskip("docx")
    import rdocx

    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    for repeated in ("0", "7"):
        source = docx.Document()
        for name in "AB":
            source.add_paragraph(f"Picture {name}: ").add_run().add_picture(
                io.BytesIO(_one_pixel_png()), width=914400
            )
        for doc_pr in source.element.body.iter(wp + "docPr"):
            doc_pr.set("id", repeated)
        buffer = io.BytesIO()
        source.save(buffer)
        with zipfile.ZipFile(buffer) as archive:
            source_xml = archive.read("word/document.xml")

        document = rdocx.Document.from_bytes(buffer.getvalue())
        assert [p.text for p in document.paragraphs] == ["Picture A: ", "Picture B: "]
        assert _document_xml(document) == source_xml

        document.add_picture(_one_pixel_png(), "added.png", 12700, 12700)
        saved = document.to_bytes()
        rdocx.Document.from_bytes(saved)
        with zipfile.ZipFile(io.BytesIO(saved)) as archive:
            xml = archive.read("word/document.xml")
        ids = re.findall(rb'<wp:docPr\b[^>]*?\bid="(\d+)"', xml)
        assert ids[:2] == [repeated.encode()] * 2
        assert len(ids) == 3 and ids[2] != repeated.encode()


def _relationship_target(document, part_name, relationship_id):
    owner = part_name.lstrip("/")
    directory, filename = posixpath.split(owner)
    relationships = posixpath.join(directory, "_rels", f"{filename}.rels")
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        xml = archive.read(relationships)
        match = re.search(
            rb'<Relationship Id="'
            + re.escape(relationship_id.encode())
            + rb'" Type="[^"]+" Target="([^"]+)"',
            xml,
        )
        assert match is not None
        target = match.group(1).decode()
        return posixpath.normpath(posixpath.join(directory, target)).lstrip("/")


def test_replace_image_preserves_drawings_and_story_relationship_ownership():
    docx = pytest.importorskip("docx")
    import rdocx

    source = docx.Document()
    source.add_picture(io.BytesIO(_one_pixel_png()))
    source.sections[0].header.paragraphs[0].add_run().add_picture(
        io.BytesIO(_one_pixel_png())
    )
    source.sections[0].footer.paragraphs[0].add_run().add_picture(
        io.BytesIO(_one_pixel_png())
    )
    buffer = io.BytesIO()
    source.save(buffer)
    document = rdocx.Document.from_bytes(buffer.getvalue())
    held = document.paragraphs[0]
    stories = {story.kind: story for story in document.stories}
    relationship_id = re.search(
        rb'r:embed="([^"]+)"', _document_xml(document)
    ).group(1).decode()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        drawing_xml = {
            story.part_name: archive.read(story.part_name.lstrip("/"))
            for story in (stories["body"], stories["header"], stories["footer"])
        }
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        header_id = re.search(
            rb'r:embed="([^"]+)"', archive.read(stories["header"].part_name.lstrip("/"))
        ).group(1).decode()
        footer_id = re.search(
            rb'r:embed="([^"]+)"', archive.read(stories["footer"].part_name.lstrip("/"))
        ).group(1).decode()

    jpeg = b"\xff\xd8\xff\xd9"
    document.replace_image(relationship_id, jpeg)
    document.replace_image_for_story(stories["header"], header_id, jpeg)
    document.replace_image_for_story(stories["footer"], footer_id, jpeg)
    assert document.image_data(relationship_id) == jpeg
    assert held.text == ""
    assert rdocx.Document.from_bytes(document.to_bytes()).image_data(relationship_id) == jpeg
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert b'ContentType="image/jpeg"' in archive.read("[Content_Types].xml")
        for story, rel_id in [
            (stories["body"], relationship_id),
            (stories["header"], header_id),
            (stories["footer"], footer_id),
        ]:
            target = _relationship_target(document, story.part_name, rel_id)
            assert archive.read(target) == jpeg
            assert archive.read(story.part_name.lstrip("/")) == drawing_xml[story.part_name]

    with pytest.raises(rdocx.RdocxError):
        document.replace_image("rIdMissing", jpeg)
    assert document.image_data("rIdMissing") is None


def _linked_report_docx():
    """A python-docx report with two pictures of one image and hyperlinks in
    the body, a table cell, and the header, two of them on one relationship."""
    docx = pytest.importorskip("docx")
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def add_link(paragraph, text, url, relationship_id=None):
        if relationship_id is None:
            relationship_id = paragraph.part.relate_to(
                url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "Hyperlink")
        properties.append(style)
        properties.append(OxmlElement("w:b"))
        run.append(properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        run.append(text_element)
        link.append(run)
        paragraph._p.append(link)
        return relationship_id

    source = docx.Document()
    source.add_picture(
        io.BytesIO(_one_pixel_png()), width=docx.shared.Inches(2), height=docx.shared.Inches(1)
    )
    source.add_picture(
        io.BytesIO(_one_pixel_png()), width=docx.shared.Inches(4), height=docx.shared.Inches(2)
    )
    paragraph = source.add_paragraph("See ")
    shared = add_link(paragraph, "first", "https://example.com/shared")
    add_link(paragraph, "second", "https://example.com/shared", shared)
    add_link(
        source.add_table(1, 1).cell(0, 0).paragraphs[0], "cell", "https://example.com/cell"
    )
    add_link(
        source.sections[0].header.paragraphs[0], "header", "https://example.com/header"
    )
    buffer = io.BytesIO()
    source.save(buffer)
    return buffer.getvalue()


def test_hyperlinks_are_retargeted_or_removed_in_every_story():
    docx = pytest.importorskip("docx")
    import rdocx

    document = rdocx.Document.from_bytes(_linked_report_docx())
    held = document.paragraphs[2]
    # The first link shares its relationship, so only it moves to a new one.
    document.set_hyperlink_url(document.hyperlinks[0], "https://example.org/new")
    links = document.hyperlinks
    assert [(link.story.kind, link.text, link.url) for link in links] == [
        ("body", "first", "https://example.org/new"),
        ("body", "second", "https://example.com/shared"),
        ("table_cell", "cell", "https://example.com/cell"),
        ("header", "header", "https://example.com/header"),
    ]
    document.set_hyperlink_url(links[3], "https://example.org/header")
    document.remove_hyperlink(links[2])
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="re-fetch"):
        document.remove_hyperlink(links[2])
    assert document.to_bytes() == before
    assert held.text == "See firstsecond"

    saved_bytes = document.to_bytes()
    saved = docx.Document(io.BytesIO(saved_bytes))
    assert [(link.text, link.url) for link in saved.paragraphs[2].hyperlinks] == [
        ("first", "https://example.org/new"),
        ("second", "https://example.com/shared"),
    ]
    cell = saved.tables[0].cell(0, 0).paragraphs[0]
    assert cell.hyperlinks == []
    assert [(run.text, run.bold, run.style.name) for run in cell.runs] == [
        ("cell", True, "Default Paragraph Font")
    ]
    assert [
        (link.text, link.url)
        for link in saved.sections[0].header.paragraphs[0].hyperlinks
    ] == [("header", "https://example.org/header")]
    with zipfile.ZipFile(io.BytesIO(saved_bytes)) as archive:
        assert b"example.com/cell" not in archive.read("word/_rels/document.xml.rels")


def test_a_stale_hyperlink_snapshot_never_edits_another_link():
    import rdocx

    word = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    rel = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    body = (
        '<w:p><w:hyperlink r:id="rIdU"><w:r><w:t>x</w:t></w:r></w:hyperlink>'
        '<w:r><w:t xml:space="preserve"> </w:t></w:r>'
        '<w:hyperlink r:id="rIdU"><w:r><w:rPr><w:b/></w:rPr><w:t>here</w:t></w:r></w:hyperlink>'
        '<w:r><w:t xml:space="preserve"> and </w:t></w:r>'
        '<w:hyperlink r:id="rIdU"><w:r><w:t>here</w:t></w:r></w:hyperlink></w:p><w:sectPr/>'
    )
    source = io.BytesIO(rdocx.Document().to_bytes())
    result = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(result, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "word/document.xml":
                data = f"<w:document {word} {rel}><w:body>{body}</w:body></w:document>".encode()
            elif info.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdU" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://u/" TargetMode="External"/></Relationships>',
                )
            result_zip.writestr(info, data)
    document = rdocx.Document.from_bytes(result.getvalue())

    links = document.hyperlinks
    document.remove_hyperlink(links[0])
    before = document.to_bytes()
    # Both remaining links have equal public fields, but only the original
    # bold physical owner is addressed by the held captured link.
    fresh = document.hyperlinks
    rebuilt = [
        rdocx.Hyperlink(
            story=link.story,
            index_path=link.index_path,
            text=link.text,
            url=link.url,
            anchor=link.anchor,
            relationship_id=link.relationship_id,
        )
        for link in fresh
    ]
    assert rebuilt == list(fresh)
    with pytest.raises(rdocx.RdocxError, match="re-fetch"):
        document.set_hyperlink_url(rebuilt[0], "https://z/")
    assert document.to_bytes() == before
    with pytest.raises(rdocx.StaleElementError, match="Document.remove_hyperlink"):
        document.set_hyperlink_url(links[0], "https://removed/")
    assert document.to_bytes() == before
    document.set_hyperlink_url(links[1], "https://bold/")
    assert [(link.text, link.url) for link in document.hyperlinks] == [
        ("here", "https://bold/"),
        ("here", "https://u/"),
    ]
    saved = document.to_bytes()
    import xml.etree.ElementTree as ET
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        children = ET.fromstring(archive.read("word/document.xml")).findall(".//" + w + "hyperlink")
        assert len(children) == 2
        assert children[0].find(".//" + w + "b") is not None
        assert children[1].find(".//" + w + "b") is None
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in ET.fromstring(archive.read("word/_rels/document.xml.rels"))}
        assert targets[children[0].attrib[r + "id"]] == "https://bold/"
        assert targets[children[1].attrib[r + "id"]] == "https://u/"
    reopened = rdocx.Document.from_bytes(saved)
    assert [(link.text, link.url) for link in reopened.hyperlinks] == [
        ("here", "https://bold/"), ("here", "https://u/"),
    ]


def test_set_picture_size_resizes_every_drawing_of_a_relationship():
    docx = pytest.importorskip("docx")
    import rdocx

    document = rdocx.Document.from_bytes(_linked_report_docx())
    relationship_id = re.search(
        rb'r:embed="([^"]+)"', _document_xml(document)
    ).group(1).decode()
    # Both pictures share one image relationship, so both are resized.
    assert (
        document.set_picture_sizes_for_relationship(
            relationship_id, width=rdocx.Inches(1), height=rdocx.Inches(1)
        )
        == 2
    )
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        document.set_picture_sizes_for_relationship("rIdMissing", width=1, height=1)
    assert document.to_bytes() == before
    saved = docx.Document(io.BytesIO(document.to_bytes()))
    assert [(shape.width, shape.height) for shape in saved.inline_shapes] == [
        (rdocx.Inches(1), rdocx.Inches(1))
    ] * 2


def _document_with_structure_snapshots(document):
    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    body = f"""
        <w:p><w:hyperlink r:id="rIdScoped"><w:r><w:t>body link</w:t></w:r></w:hyperlink></w:p>
        <w:p><w:pPr><w:sectPr>
          <w:headerReference w:type="default" r:id="rIdHeader0"/>
          <w:footerReference w:type="default" r:id="rIdFooter0"/>
          <w:type w:val="nextPage"/>
          <w:pgSz w:w="12240" w:h="15840" w:orient="portrait"/>
          <w:pgMar w:top="1440" w:right="1080" w:bottom="1440" w:left="1080" w:header="720" w:footer="720" w:gutter="120"/>
          <w:pgNumType w:start="3"/><w:cols w:num="2" w:space="360"/>
          <w:titlePg/>
        </w:sectPr></w:pPr><w:r><w:t>section zero</w:t></w:r></w:p>
        <w:p><w:pPr><w:sectPr>
          <w:type w:val="continuous"/>
          <w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>
          <w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720"/>
        </w:sectPr></w:pPr><w:r><w:t>section one</w:t></w:r></w:p>
        <w:p><w:r><w:t>section two</w:t></w:r></w:p>
        <w:sectPr>
          <w:headerReference w:type="default" r:id="rIdHeader2"/>
          <w:footerReference w:type="default" r:id="rIdFooter2"/>
          <w:type w:val="oddPage"/>
          <w:pgSz w:w="11906" w:h="16838" w:orient="portrait"/>
          <w:pgMar w:top="1000" w:right="1000" w:bottom="1000" w:left="1000"/>
        </w:sectPr>
    """
    header0 = f"""<w:hdr xmlns:w="{word}" xmlns:r="{rel}">
      <w:tbl><w:tblPr/><w:tblGrid/><w:tr><w:tc><w:tcPr/><w:p><w:r><w:t>header table</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      <w:p><w:hyperlink r:id="rIdScoped"><w:r><w:t>header zero link</w:t></w:r></w:hyperlink></w:p>
    </w:hdr>"""
    header2 = f"""<w:hdr xmlns:w="{word}" xmlns:r="{rel}">
      <w:p><w:hyperlink r:id="rIdScoped"><w:r><w:t>header two link</w:t></w:r></w:hyperlink></w:p>
    </w:hdr>"""
    footer0 = f"<w:ftr xmlns:w=\"{word}\"><w:p><w:r><w:t>footer zero</w:t></w:r></w:p></w:ftr>"
    footer2 = f"<w:ftr xmlns:w=\"{word}\"><w:p><w:r><w:t>footer two</w:t></w:r></w:p></w:ftr>"
    external_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
      <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
        <Relationship Id="rIdScoped" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="{}" TargetMode="External"/>
      </Relationships>"""
    source = io.BytesIO(document.to_bytes())
    result = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip:
        with zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/document.xml":
                    start = data.index(b"<w:body>") + len(b"<w:body>")
                    end = data.index(b"</w:body>")
                    data = data[:start] + body.encode() + data[end:]
                elif info.filename == "word/_rels/document.xml.rels":
                    additions = """
                      <Relationship Id="rIdScoped" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://body.example/" TargetMode="External"/>
                      <Relationship Id="rIdHeader0" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header0.xml"/>
                      <Relationship Id="rIdFooter0" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer0.xml"/>
                      <Relationship Id="rIdHeader2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header2.xml"/>
                      <Relationship Id="rIdFooter2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer2.xml"/>
                    """.encode()
                    data = data.replace(b"</Relationships>", additions + b"</Relationships>")
                elif info.filename == "[Content_Types].xml":
                    additions = """
                      <Override PartName="/word/header0.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
                      <Override PartName="/word/footer0.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
                      <Override PartName="/word/header2.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>
                      <Override PartName="/word/footer2.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>
                    """.encode()
                    data = data.replace(b"</Types>", additions + b"</Types>")
                elif info.filename == "word/styles.xml":
                    style = b"""<w:style w:type="paragraph" w:styleId="SnapshotStyle"><w:name w:val="Snapshot Style"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="7"/><w:qFormat/></w:style>"""
                    data = data.replace(b"</w:styles>", style + b"</w:styles>")
                result_zip.writestr(info, data)
            result_zip.writestr("word/header0.xml", header0)
            result_zip.writestr("word/header2.xml", header2)
            result_zip.writestr("word/footer0.xml", footer0)
            result_zip.writestr("word/footer2.xml", footer2)
            result_zip.writestr(
                "word/_rels/header0.xml.rels",
                external_rels.format("https://header-zero.example/"),
            )
            result_zip.writestr(
                "word/_rels/header2.xml.rels",
                external_rels.format("https://header-two.example/"),
            )
    return type(document).from_bytes(result.getvalue())


def test_stale_paragraph_after_structural_removal_raises_named_error():
    import rdocx

    doc = rdocx.Document()
    for text in ("zero", "one", "two", "three", "four"):
        doc.add_paragraph(text)

    held = doc.paragraphs[3]
    assert doc.remove_content(1) is True

    with pytest.raises(
        rdocx.StaleElementError,
        match=(
            r"paragraph handle was created at document revision 5, but the document "
            r"is now at revision 6"
        ),
    ):
        _ = held.text


def test_lazy_collections_support_index_slice_and_iteration():
    import rdocx

    doc = rdocx.Document()
    for text in ("alpha", "beta", "gamma"):
        doc.add_paragraph(text)

    paragraphs = doc.paragraphs
    assert len(paragraphs) == 3
    assert paragraphs[-1].text == "gamma"
    assert [paragraph.text for paragraph in paragraphs[0:3:2]] == ["alpha", "gamma"]
    assert [paragraph.text for paragraph in paragraphs] == ["alpha", "beta", "gamma"]

    paragraph = doc.paragraphs[0]
    paragraph.add_run(" one")
    paragraph = doc.paragraphs[0]
    paragraph.add_run(" two")
    runs = doc.paragraphs[0].runs
    assert len(runs) == 3
    assert runs[-1].text == " two"
    assert [run.text for run in runs[0:3:2]] == ["alpha", " two"]
    assert [run.text for run in runs] == ["alpha", " one", " two"]


def test_failed_removal_does_not_stale_live_handles():
    import rdocx

    doc = rdocx.Document()
    doc.add_paragraph("live")
    held = doc.paragraphs[0]

    assert doc.remove_content(99) is False
    assert held.text == "live"


def test_python_indexed_content_mutation_is_counted_and_atomic():
    import rdocx

    document = rdocx.Document()
    first = document.add_paragraph("alpha {{TOKEN")
    first.add_run("}}")
    document.add_table(1, 1)
    document.add_paragraph("omega")

    first = document.paragraphs[0]
    table = document.tables[0]
    assert document.find_content_index(first) == 0
    assert document.find_content_index(table) == 1

    held = document.paragraphs[0]
    inserted = document.insert_paragraph(1, "middle")
    assert inserted.text == "middle"
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text

    fragment = document.pop_content(document.find_content_index(inserted))
    assert fragment.kind == "paragraph"
    with pytest.raises(rdocx.StaleElementError):
        _ = inserted.text

    document.insert_content(2, fragment)
    middle = document.paragraphs[1]
    document.clone_content(middle, 0)
    with pytest.raises(rdocx.StaleElementError):
        _ = middle.text

    table = document.tables[0]
    document.move_content(table, 5)
    with pytest.raises(rdocx.StaleElementError):
        _ = table.style

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [paragraph.text for paragraph in reopened.paragraphs] == [
        "middle",
        "alpha {{TOKEN}}",
        "middle",
        "omega",
    ]
    assert reopened.find_content_index(reopened.tables[0]) == 4

    held = reopened.paragraphs[1]
    assert reopened.try_replace_text("{{TOKEN}}", "done") == 1
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text

    held = reopened.paragraphs[1]
    before = reopened.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        reopened.replace_all_regex([(r"[", "broken")])
    assert reopened.to_bytes() == before
    assert held.text == "alpha done"

    assert reopened.replace_all_regex([(r"middle", "center")]) == 2
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text
    held = reopened.paragraphs[1]
    assert reopened.try_replace_text("not present", "ignored") == 0
    assert held.text == "alpha done"

    before = reopened.to_bytes()
    held = reopened.paragraphs[0]
    with pytest.raises(IndexError):
        reopened.pop_content(99)
    with pytest.raises(IndexError):
        reopened.move_content(reopened.tables[0], 99)
    foreign_document = rdocx.Document()
    foreign = foreign_document.add_paragraph("foreign")
    with pytest.raises(ValueError, match="different document"):
        reopened.clone_content(foreign, 0)
    assert reopened.to_bytes() == before
    assert held.text == "center"

    fragment_document = rdocx.Document()
    fragment_document.add_paragraph("reusable")
    reusable = fragment_document.pop_content(0)
    assert reusable.kind == "paragraph"
    with pytest.raises(TypeError):
        rdocx.ContentFragment()
    fragment_document.insert_content(0, reusable)
    fragment_document.insert_content(1, reusable)
    assert [paragraph.text for paragraph in fragment_document.paragraphs] == [
        "reusable",
        "reusable",
    ]

    nested = _replace_document_body(
        rdocx.Document(),
        "<w:sdt><w:sdtContent><w:p><w:r><w:t>nested</w:t></w:r></w:p></w:sdtContent></w:sdt>",
    )
    nested_paragraph = nested.paragraphs[0]
    with pytest.raises(ValueError, match="not a direct body child"):
        nested.find_content_index(nested_paragraph)
    assert nested_paragraph.text == "nested"


def test_clone_content_scales_linearly_and_names_invalid_arguments():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("source")
    document.add_paragraph("destination")

    with pytest.raises(
        TypeError, match="source must be a Paragraph or Table handle"
    ):
        document.clone_content(0, 1)
    with pytest.raises(
        TypeError, match="destination must be a direct body index integer"
    ):
        document.clone_content(document.paragraphs[0], document.paragraphs[1])

    def clone_elapsed(paragraph_count, destination):
        body = "".join(
            f"<w:p><w:r><w:t>Paragraph {index} with realistic text.</w:t></w:r></w:p>"
            for index in range(paragraph_count)
        )
        candidate = _replace_document_body(rdocx.Document(), body)
        source = candidate.paragraphs[paragraph_count // 2]
        expected = source.text
        started = time.perf_counter()
        destination_index = (
            paragraph_count if destination == "end" else paragraph_count // 2 + 1
        )
        candidate.clone_content(source, destination_index)
        elapsed = time.perf_counter() - started
        assert candidate.paragraphs[destination_index].text == expected
        return elapsed

    for destination in ("middle", "end"):
        small = clone_elapsed(50, destination)
        large = clone_elapsed(100, destination)
        assert large <= small * 4.0 + 0.05, (destination, small, large)


def test_core_text_mutations_survive_bytes_round_trip():
    import rdocx

    doc = rdocx.Document()
    paragraph = doc.add_paragraph("Hello")
    run = paragraph.add_run(" world")
    reopened = rdocx.Document.from_bytes(doc.to_bytes())

    assert reopened.paragraphs[0].text == "Hello world"
    reopened.paragraphs[0].runs[1].text = " Rust"
    reopened_again = rdocx.Document.from_bytes(reopened.to_bytes())
    assert reopened_again.paragraphs[0].text == "Hello Rust"


def test_constructor_accepts_an_optional_input_path(tmp_path):
    import rdocx

    assert len(rdocx.Document().paragraphs) == 0

    path = tmp_path / "input.docx"
    source = rdocx.Document()
    source.add_paragraph("opened by constructor")
    source.save(path)

    reopened = rdocx.Document(path)
    assert reopened.paragraphs[0].text == "opened by constructor"


def test_priority_word_operations_return_typed_snapshots_and_remain_atomic():
    import rdocx

    position = rdocx.RunPosition(body_index=0, run_index=0)
    range_ = rdocx.RunRange(
        start=position,
        end=rdocx.RunPosition(body_index=0, run_index=1),
    )
    with pytest.raises(AttributeError):
        position.run_index = 2

    commented = rdocx.Document()
    commented.add_paragraph("review this")
    held_before_comment = commented.paragraphs[0]
    comment_id = commented.add_comment(
        range_,
        author="Ada",
        text="Please revise",
        initials="AL",
        date="2026-09-16T10:15:30Z",
    )
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 1, but the document is now at revision 2"
    ):
        _ = held_before_comment.text
    held_before_reply = commented.paragraphs[0]
    reply_id = commented.reply_to(
        comment_id,
        author="Grace",
        text="Done",
        date="2026-09-16T11:00:00+01:00",
    )
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 2, but the document is now at revision 3"
    ):
        _ = held_before_reply.text
    held_before_resolve = commented.paragraphs[0]
    assert commented.resolve_comment(comment_id, resolved=True) is True
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 3, but the document is now at revision 4"
    ):
        _ = held_before_resolve.text
    assert commented.comments == (
        rdocx.Comment(
            id=comment_id,
            author="Ada",
            initials="AL",
            date="2026-09-16T10:15:30Z",
            text="Please revise",
            parent_id=None,
            resolved=True,
        ),
        rdocx.Comment(
            id=reply_id,
            author="Grace",
            initials=None,
            date="2026-09-16T11:00:00+01:00",
            text="Done",
            parent_id=comment_id,
            resolved=False,
        ),
    )
    reopened_comments = rdocx.Document.from_bytes(commented.to_bytes())
    assert reopened_comments.comments == commented.comments

    before_invalid_date = reopened_comments.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="invalid RFC 3339 comment timestamp"):
        reopened_comments.add_comment(
            range_, author="Ada", text="invalid", date="2026-02-30"
        )
    assert reopened_comments.to_bytes() == before_invalid_date

    live = reopened_comments.paragraphs[0]
    before_failure = reopened_comments.to_bytes()
    invalid = rdocx.RunRange(
        start=rdocx.RunPosition(body_index=99, run_index=0),
        end=rdocx.RunPosition(body_index=99, run_index=1),
    )
    with pytest.raises(rdocx.RdocxError):
        reopened_comments.add_comment(invalid, author="Ada", text="invalid")
    assert live.text == "review this"
    assert reopened_comments.to_bytes() == before_failure
    held_before_remove = reopened_comments.paragraphs[0]
    assert reopened_comments.remove_comment(reply_id) is True
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 0, but the document is now at revision 1"
    ):
        _ = held_before_remove.text
    live_after_noop = reopened_comments.paragraphs[0]
    assert reopened_comments.remove_comment(reply_id) is False
    assert live_after_noop.text == "review this"

    original = rdocx.Document()
    original.add_paragraph("before")
    edited = rdocx.Document()
    edited.add_paragraph("after")
    held_before_compare = original.paragraphs[0]
    diagnostics = original.compare(
        edited, author="Ada", timestamp="2026-09-14T09:00:00Z"
    )
    assert isinstance(diagnostics, tuple)
    assert all(isinstance(item, rdocx.ComparisonDiagnostic) for item in diagnostics)
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 1, but the document is now at revision 2"
    ):
        _ = held_before_compare.text
    reopened_redline = rdocx.Document.from_bytes(original.to_bytes())
    assert b"before" in _document_xml(reopened_redline)
    assert b"after" in _document_xml(reopened_redline)
    with pytest.raises(rdocx.RdocxError, match="existing modeled revisions"):
        reopened_redline.compare(
            edited, author="Ada", timestamp="2026-09-14T09:01:00Z"
        )

    failed_compare = rdocx.Document()
    failed_compare.add_paragraph("stable")
    live_after_failure = failed_compare.paragraphs[0]
    before_failure = failed_compare.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        failed_compare.compare(edited, author="Ada", timestamp="not-a-timestamp")
    assert live_after_failure.text == "stable"
    assert failed_compare.to_bytes() == before_failure

    unchanged = rdocx.Document()
    unchanged.add_paragraph("same")
    identical = rdocx.Document.from_bytes(unchanged.to_bytes())
    held_before_identical_compare = unchanged.paragraphs[0]
    assert unchanged.compare(
        identical, author="Ada", timestamp="2026-09-14T09:02:00Z"
    ) == ()
    assert held_before_identical_compare.text == "same"

    laid_out = rdocx.Document()
    laid_out.add_paragraph("A deterministic layout fragment")
    fragments = laid_out.layout()
    assert fragments
    assert fragments[0].body_index == 0
    assert fragments[0].physical_page == 1
    assert fragments[0].displayed_page == 1
    assert isinstance(fragments[0].bounds, rdocx.BoundingBox)
    assert fragments[0].bounds.width > 0.0
    page = laid_out.layout_page(0)
    assert page == rdocx.LayoutPage(
        page_number=1,
        displayed_page_number=1,
        width=page.width,
        height=page.height,
    )
    assert laid_out.layout_page(1) is None

    toc_source = rdocx.Document()
    toc_source.add_paragraph("placeholder")
    toc = _replace_document_body(
        toc_source,
        """
        <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>TOC \\o "1-1"</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>
        <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
        <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Heading</w:t></w:r></w:p>
        """,
    )
    held_before_toc = toc.paragraphs[-1]
    report = toc.rebuild_toc()
    assert report == rdocx.TocRebuildReport(
        entry_count=1, bookmark_count=1, diagnostics=()
    )
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 0, but the document is now at revision 1"
    ):
        _ = held_before_toc.text
    reopened_toc = rdocx.Document.from_bytes(toc.to_bytes())
    assert "Heading" in [paragraph.text for paragraph in reopened_toc.paragraphs]

    no_toc = rdocx.Document()
    no_toc.add_paragraph("no table of contents")
    live_after_noop = no_toc.paragraphs[0]
    assert no_toc.rebuild_toc() == rdocx.TocRebuildReport(
        entry_count=0, bookmark_count=0, diagnostics=()
    )
    assert live_after_noop.text == "no table of contents"


_COMPARE_TIMESTAMP = "2026-09-27T12:00:00Z"
_LOREM = (
    "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod "
    "tempor incididunt ut labore et dolore magna aliqua. Ut enim ad minim "
    "veniam, quis nostrud exercitation ullamco."
)


def _tracked_texts(xml):
    xml = xml.decode()
    deleted = [
        "".join(re.findall(r"<w:delText(?: [^>]*)?>([^<]*)</w:delText>", wrapper))
        for wrapper in re.findall(r"<w:del\b[^>]*(?<!/)>.*?</w:del>", xml)
    ]
    inserted = [
        "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", wrapper))
        for wrapper in re.findall(r"<w:ins\b[^>]*(?<!/)>.*?</w:ins>", xml)
    ]
    return deleted, inserted


def _package_part_bytes(document, name):
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        return archive.read(name)


def test_compare_granularity_marks_only_the_changed_word():
    import rdocx

    edited_text = _LOREM.replace("magna", "MAGNA")

    def redline(**options):
        original = rdocx.Document()
        original.add_paragraph(_LOREM)
        edited = rdocx.Document()
        edited.add_paragraph(edited_text)
        assert original.compare(edited, "Ada", _COMPARE_TIMESTAMP, **options) == ()
        return original

    whole_run = ([_LOREM], [edited_text])
    for options, expected in [
        ({}, whole_run),
        ({"granularity": "run"}, whole_run),
        ({"granularity": "word"}, (["magna"], ["MAGNA"])),
        ({"granularity": "character"}, (["magna"], ["MAGNA"])),
    ]:
        compared = redline(**options)
        assert _tracked_texts(_document_xml(compared)) == expected, options
        accepted = rdocx.Document.from_bytes(compared.to_bytes())
        accepted.accept_all()
        assert accepted.paragraphs[0].text == edited_text
        rejected = rdocx.Document.from_bytes(compared.to_bytes())
        rejected.reject_all()
        assert rejected.paragraphs[0].text == _LOREM

    explicit_defaults = redline(
        granularity="run",
        ignore_formatting=False,
        ignore_whitespace=False,
        ignore_fields=False,
        ignore_comments=False,
        ignored_stories=(),
    )
    assert explicit_defaults.to_bytes() == redline().to_bytes()


def test_compare_ignore_options_keep_the_original_side():
    import rdocx

    def compared(original, edited, **options):
        work = rdocx.Document.from_bytes(original.to_bytes())
        assert work.compare(edited, "Ada", _COMPARE_TIMESTAMP, **options) == ()
        return work

    plain = rdocx.Document()
    plain.add_paragraph("plain")
    bold = rdocx.Document.from_bytes(plain.to_bytes())
    bold.paragraphs[0].runs[0].font.bold = True
    assert [item.kind for item in compared(plain, bold).revisions] == [
        "run_property_change"
    ]
    unformatted = compared(plain, bold, ignore_formatting=True)
    assert unformatted.revisions == ()
    assert unformatted.paragraphs[0].runs[0].font.bold is None

    spaced = rdocx.Document()
    spaced.add_paragraph("old  tail")
    single = rdocx.Document()
    single.add_paragraph("old tail")
    assert [item.kind for item in compared(spaced, single).revisions] == [
        "deletion",
        "insertion",
    ]
    unspaced = compared(spaced, single, ignore_whitespace=True)
    assert unspaced.revisions == ()
    assert unspaced.paragraphs[0].text == "old  tail"

    def page_field(result):
        document = rdocx.Document()
        document.add_paragraph("placeholder")
        return _replace_document_body(
            document,
            '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            f"<w:r><w:t>{result}</w:t></w:r>"
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',
        )

    first_page, second_page = page_field("1"), page_field("2")
    assert [item.kind for item in compared(first_page, second_page).revisions] == [
        "deletion",
        "insertion",
    ]
    unfielded = compared(first_page, second_page, ignore_fields=True)
    assert unfielded.revisions == ()
    assert b"<w:t>1</w:t>" in _document_xml(unfielded)

    reviewed = rdocx.Document()
    reviewed.add_paragraph("review this")
    reviewed.add_paragraph("old ending")
    commented = rdocx.Document.from_bytes(reviewed.to_bytes())
    commented.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Bo",
        text="edited side note",
    )
    commented.paragraphs[1].runs[0].text = "new ending"
    redline = compared(reviewed, commented, ignore_comments=True)
    assert redline.comments == ()
    assert b"<w:comment" not in _document_xml(redline)
    assert _tracked_texts(_document_xml(redline)) == (["old ending"], ["new ending"])

    headed = rdocx.Document()
    headed.add_paragraph("old body")
    headed.set_header("old header")
    rewritten = rdocx.Document.from_bytes(headed.to_bytes())
    for story in ("body", "header"):
        item = next(
            item
            for item in rewritten.story_items
            if item.story.kind == story and item.text
        )
        rewritten.set_story_text(item, item.text.replace("old", "new"))
    header_part = next(
        story.part_name for story in headed.stories if story.kind == "header"
    ).lstrip("/")
    tracked = compared(headed, rewritten)
    assert _tracked_texts(_package_part_bytes(tracked, header_part)) == (
        ["old header"],
        ["new header"],
    )
    ignored = compared(headed, rewritten, ignored_stories=("header",))
    assert _package_part_bytes(ignored, header_part) == _package_part_bytes(headed, header_part)
    assert _tracked_texts(_document_xml(ignored)) == (["old body"], ["new body"])
    unbodied = compared(headed, rewritten, ignored_stories=("body",))
    assert _document_xml(unbodied) == _document_xml(headed)
    assert _tracked_texts(_package_part_bytes(unbodied, header_part)) == (
        ["old header"],
        ["new header"],
    )
    for story in ("footer", "comment", "text_box", "footnote", "endnote"):
        untouched = compared(headed, rewritten, ignored_stories=(story,))
        assert _tracked_texts(_document_xml(untouched)) == (
            ["old body"],
            ["new body"],
        ), story
        assert _tracked_texts(_package_part_bytes(untouched, header_part)) == (
            ["old header"],
            ["new header"],
        ), story


def test_compare_rejects_unknown_options_before_mutation():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("before")
    edited = rdocx.Document()
    edited.add_paragraph("after")
    live = original.paragraphs[0]
    before = original.to_bytes()
    for options, error, message in [
        ({"granularity": "words"}, rdocx.RdocxError, 'granularity "words"'),
        ({"ignored_stories": ["main"]}, rdocx.RdocxError, 'story "main"'),
        ({"ignored_stories": ["table_cell"]}, rdocx.RdocxError, 'story "table_cell"'),
        (
            {"ignored_stories": ["header", "header"]},
            rdocx.RdocxError,
            "duplicate ignored story",
        ),
        ({"ignored_stories": "header"}, TypeError, None),
    ]:
        with pytest.raises(error, match=message):
            original.compare(edited, "Ada", _COMPARE_TIMESTAMP, **options)
        assert original.to_bytes() == before
        assert live.text == "before"
    with pytest.raises(TypeError):
        original.compare(edited, "Ada", _COMPARE_TIMESTAMP, "word")
    assert live.text == "before"


def test_issue_161_compare_carries_edited_side_comment_thread():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Review this heading")
    edited = rdocx.Document.from_bytes(original.to_bytes())
    edited.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Bo",
        text="Edited side note",
    )
    redline = rdocx.Document.from_bytes(original.to_bytes())
    assert redline.compare(edited, "Ada", _COMPARE_TIMESTAMP) == ()
    assert [comment.text for comment in redline.comments] == ["Edited side note"]
    accepted = rdocx.Document.from_bytes(redline.to_bytes())
    accepted.accept_all()
    assert [comment.text for comment in accepted.comments] == ["Edited side note"]
    rejected = rdocx.Document.from_bytes(redline.to_bytes())
    rejected.reject_all()
    assert rejected.comments == ()


@pytest.mark.parametrize("edit", ["add", "remove", "reply", "resolve", "redate"])
def test_issue_161_comment_edits_resolve_to_each_input(edit):
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Review this heading")
    if edit != "add":
        original.add_comment(
            rdocx.RunRange(
                start=rdocx.RunPosition(body_index=0, run_index=0),
                end=rdocx.RunPosition(body_index=0, run_index=1),
            ),
            author="Bo",
            text="Original note",
            date="2026-09-16T10:15:30Z",
        )
    edited = rdocx.Document.from_bytes(original.to_bytes())
    if edit == "add":
        edited.add_comment(
            rdocx.RunRange(
                start=rdocx.RunPosition(body_index=0, run_index=0),
                end=rdocx.RunPosition(body_index=0, run_index=1),
            ),
            author="Bo",
            text="Added note",
        )
    elif edit == "remove":
        assert edited.remove_comment(edited.comments[0].id)
    elif edit == "reply":
        edited.reply_to(edited.comments[0].id, author="Ada", text="Agreed")
    elif edit == "resolve":
        assert edited.resolve_comment(edited.comments[0].id, resolved=True)
    else:
        source = io.BytesIO(edited.to_bytes())
        result = io.BytesIO()
        with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/comments.xml":
                    data = data.replace(b"2026-09-16T10:15:30Z", b"2026-09-17T10:15:30Z")
                result_zip.writestr(info, data)
        edited = rdocx.Document.from_bytes(result.getvalue())

    redline = rdocx.Document.from_bytes(original.to_bytes())
    assert redline.compare(edited, "Ada", _COMPARE_TIMESTAMP) == ()
    assert redline.comments == edited.comments
    if edit == "redate":
        assert len(redline.revisions) == 1
        selected = rdocx.Document.from_bytes(redline.to_bytes())
        assert selected.accept_revision_id(redline.revisions[0].id) == 1
        assert selected.comments == edited.comments
    accepted = rdocx.Document.from_bytes(redline.to_bytes())
    assert accepted.accept_all() > 0
    assert accepted.comments == edited.comments
    rejected = rdocx.Document.from_bytes(redline.to_bytes())
    assert rejected.reject_all() > 0
    assert rejected.comments == original.comments


def test_issue_161_edited_comment_relationship_survives_redline():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Review this")
    edited = rdocx.Document.from_bytes(original.to_bytes())
    edited.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Bo",
        text="See link",
    )
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(edited.to_bytes())) as source_zip:
        with zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/comments.xml":
                    data = data.replace(
                        b"</w:p>",
                        b'<w:hyperlink xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId42"><w:r><w:t>linked</w:t></w:r></w:hyperlink></w:p>',
                        1,
                    )
                if info.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"</Types>",
                        b'<Override PartName="/word/media/review.bin" ContentType="application/octet-stream"/></Types>',
                    )
                result_zip.writestr(info, data)
            result_zip.writestr(
                "word/_rels/comments.xml.rels",
                b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId42" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/review" TargetMode="External"/><Relationship Id="rId43" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/review.bin"/></Relationships>',
            )
            result_zip.writestr("word/media/review.bin", b"edited comment asset")
    edited = rdocx.Document.from_bytes(result.getvalue())
    redline = rdocx.Document.from_bytes(original.to_bytes())
    assert redline.compare(edited, "Ada", _COMPARE_TIMESTAMP) == ()
    with zipfile.ZipFile(io.BytesIO(redline.to_bytes())) as archive:
        assert b"https://example.com/review" in archive.read("word/_rels/comments.xml.rels")
        assert archive.read("word/media/review.bin") == b"edited comment asset"
    accepted = rdocx.Document.from_bytes(redline.to_bytes())
    accepted.accept_all()
    with zipfile.ZipFile(io.BytesIO(accepted.to_bytes())) as archive:
        assert b"https://example.com/review" in archive.read("word/_rels/comments.xml.rels")
        assert archive.read("word/media/review.bin") == b"edited comment asset"
    rejected = rdocx.Document.from_bytes(redline.to_bytes())
    rejected.reject_all()
    with zipfile.ZipFile(io.BytesIO(rejected.to_bytes())) as archive:
        assert "word/_rels/comments.xml.rels" not in archive.namelist()
        assert "word/media/review.bin" not in archive.namelist()


@pytest.mark.parametrize("options", [{"ignore_comments": True}, {"ignored_stories": ["comment"]}])
def test_issue_161_ignored_comments_skip_missing_related_target(options):
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Unchanged body")
    document.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Bo",
        text="Ignored note",
    )
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as source_zip:
        with zipfile.ZipFile(output, "w") as result_zip:
            for info in source_zip.infolist():
                result_zip.writestr(info, source_zip.read(info.filename))
            result_zip.writestr(
                "word/_rels/comments.xml.rels",
                b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId42" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/missing.png"/></Relationships>',
            )
    damaged = rdocx.Document.from_bytes(output.getvalue())
    original = rdocx.Document.from_bytes(damaged.to_bytes())
    assert original.compare(damaged, "Ada", _COMPARE_TIMESTAMP, **options) == ()


def test_issue_161_comment_asset_collision_preserves_body_image():
    import rdocx

    original = rdocx.Document()
    original.add_paragraph("Review this")
    original.add_picture(_one_pixel_png(), "body.png")
    edited = rdocx.Document.from_bytes(original.to_bytes())
    edited.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Bo",
        text="New comment",
    )
    output = io.BytesIO()
    body_asset = None
    with zipfile.ZipFile(io.BytesIO(edited.to_bytes())) as source_zip:
        body_asset = source_zip.read("word/media/image1.png")
        with zipfile.ZipFile(output, "w") as result_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/_rels/document.xml.rels":
                    data = data.replace(b'media/image1.png', b'media/image2.png')
                if info.filename == "word/media/image1.png":
                    data = b"comment asset with a colliding producer name"
                result_zip.writestr(info, data)
            result_zip.writestr("word/media/image2.png", body_asset)
            result_zip.writestr(
                "word/_rels/comments.xml.rels",
                b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId88" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/></Relationships>',
            )
    edited = rdocx.Document.from_bytes(output.getvalue())
    redline = rdocx.Document.from_bytes(original.to_bytes())
    assert redline.compare(edited, "Ada", _COMPARE_TIMESTAMP) == ()
    with zipfile.ZipFile(io.BytesIO(redline.to_bytes())) as archive:
        assert archive.read("word/media/image1.png") == body_asset
        assert archive.read("word/media/image1-rdocx-comment-1.png") == b"comment asset with a colliding producer name"
    accepted = rdocx.Document.from_bytes(redline.to_bytes())
    accepted.accept_all()
    assert accepted.compare(edited, "Ada", _COMPARE_TIMESTAMP) == ()
    rejected = rdocx.Document.from_bytes(redline.to_bytes())
    rejected.reject_all()
    assert rejected.compare(original, "Ada", _COMPARE_TIMESTAMP) == ()


def test_word_default_toc_switch_rebuilds_and_reports_ordered_diagnostics():
    import rdocx

    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(
        source,
        """
        <w:p><w:fldSimple w:instr="TOC"><w:r><w:t>simple cache</w:t></w:r></w:fldSimple></w:p>
        <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>TOC \\x</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>
        <w:p><w:r><w:t>complex cache</w:t></w:r></w:p>
        <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
        """,
    )

    report = document.rebuild_toc()
    assert report.diagnostics == (
        "simple table of contents fields are not rebuilt, stored display retained",
        "field TOC uses unsupported switch \\x, stored display retained",
    )
    assert report.diagnostic_count == len(report.diagnostics) == 2
    with pytest.raises(AttributeError):
        report.diagnostics = ()


def test_rebuild_toc_accepts_identity_attributes_on_the_field_runs():
    import rdocx

    # Word writes w:rsidR and w:rsidRPr on the runs it saves, and Google Docs
    # exports write them on every run, the runs of the TOC field code included.
    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(
        source,
        """
        <w:p><w:r w:rsidR="00A1B2C3"><w:fldChar w:fldCharType="begin"/></w:r><w:r w:rsidRPr="00A1B2C3"><w:instrText>TOC \\o "1-1"</w:instrText></w:r><w:r w:rsidDel="00A1B2C3"><w:fldChar w:fldCharType="separate"/></w:r></w:p>
        <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
        <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Heading</w:t></w:r></w:p>
        """,
    )

    assert document.rebuild_toc() == rdocx.TocRebuildReport(
        entry_count=1, bookmark_count=1, diagnostics=()
    )
    saved = _document_xml(document)
    for identity in (b'w:rsidR="00A1B2C3"', b'w:rsidRPr="00A1B2C3"', b'w:rsidDel="00A1B2C3"'):
        assert identity in saved


def test_rebuild_toc_keeps_the_field_markers_in_the_first_and_last_entries():
    import rdocx

    # Word and Google Docs write the field begin in the first entry paragraph
    # and the end in the last one: a rebuild keeps them there.
    def entry(text, first=False, last=False):
        begin = (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            if first
            else ""
        )
        end = '<w:r><w:fldChar w:fldCharType="end"/></w:r>' if last else ""
        return (
            f'<w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr>{begin}'
            f"<w:r><w:t>{text}</w:t></w:r>{end}</w:p>"
        )

    headings = ("One", "Two", "Three")
    body = entry("One", first=True) + entry("Two") + entry("Three", last=True)
    body += "".join(
        f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{heading}</w:t></w:r></w:p>'
        for heading in headings
    )
    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(source, body)
    before = len(document.paragraphs)

    assert document.rebuild_toc().entry_count == 3
    assert len(document.paragraphs) == before
    toc = [
        paragraph.text
        for paragraph in document.paragraphs
        if (paragraph.style or "").startswith("TOC")
    ]
    assert toc == ["One\t1", "Two\t1", "Three\t1"]


def test_update_page_fields_writes_layout_page_numbers():
    import rdocx

    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(
        source,
        """
        <w:p><w:r><w:t>Page one.</w:t></w:r></w:p>
        <w:p><w:pPr><w:pageBreakBefore/></w:pPr><w:r><w:t xml:space="preserve">Page </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> of </w:t></w:r><w:fldSimple w:instr=" NUMPAGES "><w:r><w:t>9</w:t></w:r></w:fldSimple></w:p>
        """,
    )
    held_before_update = document.paragraphs[1]
    assert document.update_page_fields() == 2
    with pytest.raises(
        rdocx.StaleElementError, match=r"revision 0, but the document is now at revision 1"
    ):
        _ = held_before_update.text
    xml = _document_xml(document)
    assert b"<w:t>7</w:t>" not in xml
    assert b"<w:t>9</w:t>" not in xml
    assert xml.count(b"<w:t>2</w:t>") == 2

    no_page_fields = rdocx.Document()
    no_page_fields.add_paragraph("no page fields")
    live_after_noop = no_page_fields.paragraphs[0]
    assert no_page_fields.update_page_fields() == 0
    assert live_after_noop.text == "no page fields"


def test_update_layout_backed_fields_returns_owned_report():
    import rdocx

    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(
        source,
        """
        <w:p><w:fldSimple w:instr="PAGE"><w:r><w:t>stale page</w:t></w:r></w:fldSimple><w:fldSimple w:instr="NUMPAGES"><w:r><w:t>stale count</w:t></w:r></w:fldSimple></w:p>
        <w:p><w:fldSimple w:instr="SECTION"><w:r><w:t>stale section</w:t></w:r></w:fldSimple><w:fldSimple w:instr="SECTIONPAGES"><w:r><w:t>stale section count</w:t></w:r></w:fldSimple></w:p>
        <w:p><w:fldSimple w:instr="PAGEREF destination"><w:r><w:t>stale target</w:t></w:r></w:fldSimple></w:p>
        <w:p><w:pPr><w:pageBreakBefore/></w:pPr><w:bookmarkStart w:id="7" w:name="destination"/><w:r><w:t>Target</w:t></w:r><w:bookmarkEnd w:id="7"/></w:p>
        """,
    )

    report = document.update_layout_backed_fields()
    assert report == rdocx.LayoutBackedFieldUpdateReport(
        page_fields=1,
        num_pages_fields=1,
        page_reference_fields=1,
        diagnostics=report.diagnostics,
        section_fields=1,
        section_pages_fields=1,
    )
    assert report.updated_count == 5
    assert report.diagnostic_count == len(report.diagnostics)
    with pytest.raises(AttributeError):
        report.page_fields = 0
    xml = _document_xml(document)
    assert b"stale page" not in xml
    assert b"stale count" not in xml
    assert b"stale target" not in xml
    assert b"stale section" not in xml
    legacy = rdocx.LayoutBackedFieldUpdateReport(page_fields=0, num_pages_fields=0, page_reference_fields=0, diagnostics=())
    assert legacy.section_fields == legacy.section_pages_fields == 0
    assert legacy.updated_count == 0
    with pytest.raises(AttributeError):
        report.section_fields = 0


def test_pageref_to_a_bookmarked_run_range_is_filled_from_the_layout():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("See page ")
    document.add_table(2, 2)
    target = document.add_paragraph("The target phrase ends here.")
    target.paragraph_format.page_break_before = True
    index = document.find_content_index("The target phrase")
    assert index == 2
    assert document.split_run(index, 0, len("The target phrase")) == 1
    range_ = rdocx.RunRange(
        start=rdocx.RunPosition(body_index=index, run_index=0),
        end=rdocx.RunPosition(body_index=index, run_index=1),
    )
    held = document.paragraphs[0].runs[0]

    bookmark_id = document.add_bookmark("target", range_)
    assert document.bookmarks == (
        rdocx.Bookmark(
            id=bookmark_id,
            name="target",
            # The recursive range counts the four cell paragraphs of the table.
            range=rdocx.RunRange(
                start=rdocx.RunPosition(body_index=5, run_index=0),
                end=rdocx.RunPosition(body_index=5, run_index=1),
            ),
            direct_range=range_,
            text="The target phrase",
            issue=None,
        ),
    )
    later = next(
        item
        for item in document.story_items
        if item.story.kind == "body" and item.text == "The target phrase ends here."
    )
    held.add_field("PAGEREF target \\h", "?")
    # The field is a story item of its own, so later items move and go stale.
    for stale in (lambda: held.text, lambda: document.set_story_text(later, "x")):
        with pytest.raises(rdocx.StaleElementError):
            stale()
    assert document.paragraphs[0].runs[0].text == "See page "

    report = document.update_layout_backed_fields()
    assert report.page_reference_fields == 1
    field = re.search(
        rb'w:instr="PAGEREF target \\h"[^>]*>\s*<w:r>\s*<w:t>([^<]*)</w:t>',
        _document_xml(document),
    )
    assert field is not None and field.group(1) == b"2"

    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="already exists"):
        document.add_bookmark("target", range_)
    table_range = rdocx.RunRange(
        start=rdocx.RunPosition(body_index=1, run_index=0),
        end=rdocx.RunPosition(body_index=1, run_index=0),
    )
    with pytest.raises(rdocx.RdocxError, match="is not a paragraph"):
        document.add_bookmark("in_table", table_range)
    run = document.paragraphs[0].runs[0]
    with pytest.raises(rdocx.RdocxError, match="field name"):
        run.add_field("  ")
    assert run.text == "See page "
    assert document.to_bytes() == before


def test_bookmarks_report_nested_ranges_without_a_direct_range():
    import rdocx

    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(
        source,
        """
        <w:sdt><w:sdtContent><w:p><w:bookmarkStart w:id="1" w:name="inside"/><w:r><w:t>Inside</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p></w:sdtContent></w:sdt>
        <w:p><w:bookmarkStart w:id="2" w:name="open"/><w:r><w:t>Open</w:t></w:r></w:p>
        """,
    )
    inside, unmatched = document.bookmarks
    assert inside.name == "inside"
    assert inside.range == rdocx.RunRange(
        start=rdocx.RunPosition(body_index=0, run_index=0),
        end=rdocx.RunPosition(body_index=0, run_index=1),
    )
    assert inside.direct_range is None
    assert inside.text == "Inside"
    assert unmatched.range is None and unmatched.direct_range is None
    assert unmatched.issue == "bookmark id 2 has 1 start markers and 0 end markers"


def test_insert_toc_links_headings_and_tab_and_fields_extend_a_run():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Intro")
    document.add_paragraph("Chapter one").style = "Heading1"
    document.add_paragraph("Section one point one").style = "Heading2"
    document.add_paragraph("Deep detail").style = "Heading3"
    held = document.paragraphs[0]

    before = document.to_bytes()
    with pytest.raises(IndexError):
        document.insert_toc(5)
    with pytest.raises(ValueError, match="max_level"):
        document.insert_toc(0, max_level=0)
    assert document.to_bytes() == before
    assert held.text == "Intro"

    assert document.insert_toc(0, max_level=2) is None
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert [paragraph.text for paragraph in document.paragraphs[:3]] == [
        "Table of Contents",
        "Chapter one\t",
        "Section one point one\t",
    ]
    assert [(link.text, link.anchor) for link in document.hyperlinks] == [
        ("Chapter one", "_Toc1"),
        ("Section one point one", "_Toc2"),
    ]
    assert [(bookmark.name, bookmark.text) for bookmark in document.bookmarks] == [
        ("_Toc1", "Chapter one"),
        ("_Toc2", "Section one point one"),
    ]

    run = document.paragraphs[3].runs[0]
    run.add_tab()
    assert run.text == "Intro\t"
    run.add_field("PAGE", "1")
    with pytest.raises(rdocx.StaleElementError):
        run.text
    xml = _document_xml(document)
    assert b'<w:fldSimple w:instr="PAGE"' in xml


def test_word_structure_snapshots_preserve_order_ownership_and_types():
    import rdocx

    document = _document_with_structure_snapshots(rdocx.Document())
    reopened = rdocx.Document.from_bytes(document.to_bytes())

    assert len(reopened.sections) == 3
    assert reopened.sections[0] == rdocx.Section(
        ordinal=0,
        is_final=False,
        orientation="portrait",
        page_width=7_772_400,
        page_height=10_058_400,
        margin_top=914_400,
        margin_right=685_800,
        margin_bottom=914_400,
        margin_left=685_800,
        gutter=76_200,
        column_count=2,
        column_spacing=228_600,
        page_number_start=3,
        header_distance=457_200,
        footer_distance=457_200,
        different_first_page=True,
        break_type="nextPage",
    )
    assert [section.ordinal for section in reopened.sections] == [0, 1, 2]
    assert [section.is_final for section in reopened.sections] == [False, False, True]
    with pytest.raises(AttributeError):
        reopened.sections[0].orientation = "landscape"

    assert reopened.styles[-1] == rdocx.Style(
        style_id="SnapshotStyle",
        name="Snapshot Style",
        based_on="Normal",
        style_type="paragraph",
        linked_style=None,
        next_style="Normal",
        priority=7,
        auto_redefine=None,
        hidden=None,
        semi_hidden=None,
        unhide_when_used=None,
        quick_format=True,
        locked=None,
        is_default=False,
    )
    assert [style.style_id for style in reopened.styles][-1] == "SnapshotStyle"

    variants = reopened.header_footer_variants
    assert len(variants) == 18
    assert variants[0] == rdocx.HeaderFooterVariant(
        section_index=0,
        kind="header",
        variant="default",
        story=rdocx.Story(
            kind="header", part_name="/word/header0.xml", owner_index=0
        ),
        source_section=0,
        inherited=False,
    )
    assert variants[6].story == variants[0].story
    assert variants[6].source_section == 0
    assert variants[6].inherited is True
    assert variants[12].story == rdocx.Story(
        kind="header", part_name="/word/header2.xml", owner_index=0
    )
    assert variants[12].source_section == 2
    assert variants[12].inherited is False
    assert variants[1].story is None

    header_items = [
        item for item in reopened.story_items if item.story == variants[0].story
    ]
    assert [(item.kind, item.index_path, item.text) for item in header_items] == [
        ("table", (0,), None),
        ("paragraph", (1,), "header zero link"),
    ]
    assert all(isinstance(item, rdocx.StoryItem) for item in reopened.story_items)
    assert all(isinstance(story, rdocx.Story) for story in reopened.stories)
    assert all(item.direct_body_index is None for item in header_items)
    body_items = [item for item in reopened.story_items if item.story.kind == "body"]
    assert [item.direct_body_index for item in body_items] == [0, 1, 2, 3, None]

    owned = _replace_document_body(
        rdocx.Document(),
        """
        <w:p><w:r><w:t>first</w:t></w:r>
          <w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>
          <w:r><w:drawing xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><wp:inline><wp:docPr id="41" name="item"/></wp:inline></w:drawing></w:r>
          <w:sdt><w:sdtContent><w:r><w:t>nested control</w:t></w:r></w:sdtContent></w:sdt>
        </w:p>
        <w:tbl><w:tblPr/><w:tblGrid/><w:tr><w:tc><w:tcPr/><w:p><w:r><w:t>cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
        <w:sdt><w:sdtContent><w:p><w:r><w:t>body control</w:t></w:r></w:p></w:sdtContent></w:sdt>
        <w:p><w:r><w:t>second</w:t></w:r></w:p>
        <w:sectPr/>
        """,
    )
    owned_body_items = [
        item for item in owned.story_items if item.story.kind == "body"
    ]
    assert [item.direct_body_index for item in owned_body_items] == [
        0,
        0,
        0,
        0,
        1,
        2,
        3,
        None,
    ]

    assert [link.relationship_id for link in reopened.hyperlinks] == [
        "rIdScoped",
        "rIdScoped",
        "rIdScoped",
    ]
    assert [link.url for link in reopened.hyperlinks] == [
        "https://body.example/",
        "https://header-zero.example/",
        "https://header-two.example/",
    ]
    assert [link.text for link in reopened.hyperlinks] == [
        "body link",
        "header zero link",
        "header two link",
    ]
    assert [link.index_path for link in reopened.hyperlinks] == [(0,), (1,), (0,)]

    nested = _replace_document_body(
        rdocx.Document(),
        """
        <w:sdt><w:sdtContent><w:sdt><w:sdtContent><w:p>
          <w:hyperlink w:anchor="nested-target"><w:r><w:t>nested link</w:t></w:r></w:hyperlink>
        </w:p></w:sdtContent></w:sdt><w:p>
          <w:hyperlink w:anchor="after-target"><w:r><w:t>after link</w:t></w:r></w:hyperlink>
        </w:p></w:sdtContent></w:sdt>
        <w:sectPr/>
        """,
    )
    assert [
        (link.text, link.anchor, link.index_path) for link in nested.hyperlinks
    ] == [
        ("nested link", "nested-target", (1,)),
        ("after link", "after-target", (0,)),
    ]

    captured_items = reopened.story_items
    captured_links = reopened.hyperlinks
    reopened.add_paragraph("later body content")
    assert captured_items[-1].text != "later body content"
    assert len(reopened.story_items) == len(captured_items) + 1
    assert reopened.hyperlinks == captured_links


def test_update_section_margins_and_orientation_reach_the_layout():
    import rdocx
    from rdocx import Inches

    document = rdocx.Document()
    document.add_paragraph("body text")
    held = document.paragraphs[0]
    assert (document.layout_page(0).width, document.layout_page(0).height) == (612.0, 792.0)

    updated = document.update_section(
        0, orientation="landscape", margin_top=Inches(0.5), margin_left=Inches(2)
    )

    assert updated == document.sections[0]
    assert (updated.orientation, updated.page_width, updated.page_height) == (
        "landscape",
        Inches(11),
        Inches(8.5),
    )
    assert (
        updated.margin_top,
        updated.margin_right,
        updated.margin_bottom,
        updated.margin_left,
    ) == (Inches(0.5), Inches(1), Inches(1), Inches(2))
    assert held.text == "body text"
    assert (document.layout_page(0).width, document.layout_page(0).height) == (792.0, 612.0)
    bounds = document.layout()[0].bounds
    assert (bounds.x, bounds.y, bounds.width) == (144.0, 36.0, 576.0)
    xml = _document_xml(document)
    assert b'<w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>' in xml
    assert b'<w:pgMar w:top="720" w:right="1440" w:bottom="1440" w:left="2880"' in xml
    assert rdocx.Document.from_bytes(document.to_bytes()).sections[0] == updated


def test_update_section_keeps_unnamed_partners_and_rejects_atomically():
    import rdocx
    from rdocx import Inches

    document = rdocx.Document()
    document.add_paragraph("body")
    updated = document.update_section(
        0,
        page_height=Inches(14),
        header_distance=Inches(0.25),
        gutter=Inches(0.1),
        column_count=2,
        column_spacing=Inches(0.5),
        page_number_start=5,
        different_first_page=True,
        break_type="oddPage",
    )
    assert (updated.page_width, updated.page_height) == (Inches(8.5), Inches(14))
    assert (updated.header_distance, updated.footer_distance) == (Inches(0.25), Inches(0.5))
    assert (updated.gutter, updated.column_count, updated.column_spacing) == (
        Inches(0.1),
        2,
        Inches(0.5),
    )
    assert (
        updated.page_number_start,
        updated.different_first_page,
        updated.break_type,
    ) == (5, True, "oddPage")
    assert document.update_section(0, column_count=3).column_spacing == Inches(0.5)

    before = document.to_bytes()
    with pytest.raises(ValueError, match="orientation must be portrait or landscape"):
        document.update_section(0, margin_top=0, orientation="sideways")
    with pytest.raises(ValueError, match="break type"):
        document.update_section(0, break_type="page")
    with pytest.raises(rdocx.RdocxError, match="gutter must be nonnegative"):
        document.update_section(0, margin_top=Inches(3), orientation="landscape", gutter=rdocx.Emu(-635))
    with pytest.raises(rdocx.RdocxError, match="page width must be positive"):
        document.update_section(0, page_width=0)
    with pytest.raises(rdocx.RdocxError, match="page-number start must be positive"):
        document.update_section(0, margin_left=0, page_number_start=0)
    with pytest.raises(IndexError, match="section index out of range"):
        document.update_section(1, gutter=0)
    assert document.to_bytes() == before


def test_update_section_never_rewrites_unequal_width_columns():
    import rdocx
    from rdocx import Inches

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:t>body</w:t></w:r></w:p><w:sectPr>'
        '<w:cols w:num="2" w:space="720" w:equalWidth="0">'
        '<w:col w:w="3000" w:space="720"/><w:col w:w="5000"/></w:cols>'
        "</w:sectPr>",
    )
    section = document.sections[0]
    assert (section.column_count, section.column_spacing) == (None, None)

    before = document.to_bytes()
    with pytest.raises(ValueError, match="require both"):
        document.update_section(0, column_count=3)
    with pytest.raises(ValueError, match="require both"):
        document.update_section(0, column_spacing=Inches(0.25))
    assert document.to_bytes() == before

    updated = document.update_section(0, column_count=3, column_spacing=Inches(0.25))
    assert (updated.column_count, updated.column_spacing) == (3, Inches(0.25))
    assert b"<w:col " not in _document_xml(document)


def test_update_section_uses_layout_defaults_for_missing_partners():
    import rdocx
    from rdocx import Inches

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:t>body</w:t></w:r></w:p><w:sectPr/>',
    )
    updated = document.update_section(
        0,
        page_width=Inches(9),
        margin_left=Inches(2),
        header_distance=Inches(0.25),
        column_count=2,
    )
    assert (updated.page_width, updated.page_height) == (Inches(9), Inches(11))
    assert (
        updated.margin_top,
        updated.margin_right,
        updated.margin_bottom,
        updated.margin_left,
    ) == (Inches(1), Inches(1), Inches(1), Inches(2))
    assert (updated.header_distance, updated.footer_distance) == (
        Inches(0.25), Inches(0.5)
    )
    assert (updated.column_count, updated.column_spacing) == (2, Inches(0.5))
    assert rdocx.Document.from_bytes(document.to_bytes()).sections[0] == updated


def test_insert_and_remove_section_restructure_the_body():
    import rdocx
    from rdocx import Inches

    document = rdocx.Document()
    document.add_paragraph("first")
    held = document.paragraphs[0]
    document.insert_section(1)
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text

    assert [section.is_final for section in document.sections] == [False, True]
    assert document.sections[1].margin_top == Inches(1)
    assert document.update_section(1, margin_top=Inches(1)).margin_right == Inches(1)
    document.update_section(
        1,
        orientation="landscape",
        page_width=Inches(8.5),
        page_height=Inches(11),
        margin_top=Inches(1),
        margin_right=Inches(1),
        margin_bottom=Inches(1),
        margin_left=Inches(1),
    )
    document.add_paragraph("second")
    assert [fragment.physical_page for fragment in document.layout()] == [1, 1, 2]
    assert document.layout_page(0).width == 612.0
    assert document.layout_page(1).width == 792.0

    with pytest.raises(IndexError, match="section index out of range"):
        document.insert_section(3)
    with pytest.raises(IndexError, match="section index out of range"):
        document.remove_section(2)
    document.remove_section(1)
    assert len(document.sections) == 1
    assert [paragraph.text for paragraph in document.paragraphs] == ["first", "second"]
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="sole document section"):
        document.remove_section(0)
    assert document.to_bytes() == before


def test_section_edits_write_the_native_body():
    # section_edits_write_the_body_the_python_binding_pins in
    # crates/rdocx/tests/integration_test.rs pins the same body for the native
    # calls, so CI checks that these calls write what the native ones write.
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("first")
    document.insert_section(1)
    document.update_section(
        0,
        orientation="landscape",
        margin_top=rdocx.Emu(457200),
        margin_left=rdocx.Emu(1828800),
        page_number_start=3,
        footer_distance=rdocx.Emu(228600),
        different_first_page=True,
    )
    document.update_section(
        1,
        page_width=rdocx.Emu(7772400),
        page_height=rdocx.Emu(10058400),
        margin_top=rdocx.Emu(914400),
        margin_right=rdocx.Emu(1143000),
        margin_bottom=rdocx.Emu(685800),
        margin_left=rdocx.Emu(1371600),
        gutter=rdocx.Emu(127000),
        column_count=2,
        column_spacing=rdocx.Emu(457200),
        break_type="continuous",
    )
    document.add_paragraph("second")

    xml = _document_xml(document).decode()
    body = xml[xml.index("<w:body>") : xml.index("</w:body>") + len("</w:body>")]
    assert "".join(line.strip() for line in body.splitlines()) == (
        '<w:body><w:p><w:r><w:t>first</w:t></w:r></w:p><w:p><w:pPr>'
        '<w:sectPr><w:pgSz w:w="15840" w:h="12240" w:orient="landscape"/>'
        '<w:pgMar w:top="720" w:right="1440" w:bottom="1440" w:left="2880"'
        ' w:gutter="0" w:header="720" w:footer="360"/>'
        '<w:pgNumType w:start="3"/><w:titlePg/></w:sectPr></w:pPr></w:p>'
        '<w:p><w:r><w:t>second</w:t></w:r></w:p><w:sectPr>'
        '<w:type w:val="continuous"/><w:pgSz w:w="12240" w:h="15840"/>'
        '<w:pgMar w:top="1440" w:right="1800" w:bottom="1080" w:left="2160"'
        ' w:gutter="200" w:header="720" w:footer="720"/>'
        '<w:cols w:num="2" w:space="720"/></w:sectPr></w:body>'
    )


def _story_paragraph_texts(document, kind):
    return [
        item.text
        for item in document.story_items
        if item.story.kind == kind and item.kind == "paragraph"
    ]


def test_header_footer_and_story_text_edit_the_section_stories():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body")
    held = document.paragraphs[0]
    document.set_header("Draft")
    document.set_footer("Page footer")
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert _story_paragraph_texts(document, "header") == ["Draft"]
    assert _story_paragraph_texts(document, "footer") == ["Page footer"]

    header_item = next(
        item
        for item in document.story_items
        if item.story.kind == "header" and item.kind == "paragraph"
    )
    document.set_story_text(header_item, "Final")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert _story_paragraph_texts(reopened, "header") == ["Final"]
    assert _story_paragraph_texts(reopened, "body") == ["body"]


@pytest.mark.parametrize(
    ("separator", "expected"),
    [("add_tab", "left\tright"), ("add_break", "left\nright")],
)
def test_story_items_keep_run_tabs_and_breaks_in_the_body_header_and_footer(
    separator, expected
):
    docx = pytest.importorskip("docx")
    import rdocx

    source = docx.Document()
    section = source.sections[0]
    for paragraph in (
        source.add_paragraph(),
        section.header.paragraphs[0],
        section.footer.paragraphs[0],
    ):
        paragraph.add_run("left")
        getattr(paragraph.add_run(), separator)()
        paragraph.add_run("right")
    buffer = io.BytesIO()
    source.save(buffer)

    document = rdocx.Document.from_bytes(buffer.getvalue())
    kinds = ("body", "header", "footer")
    for kind in kinds:
        assert _story_paragraph_texts(document, kind) == [expected], kind

    # The new text replaces the tab or the break too, and a tab or a line
    # feed in it reads back as written.
    for value in ("new", "a\tb\nc"):
        for kind in kinds:
            item = next(
                item
                for item in document.story_items
                if item.story.kind == kind and item.kind == "paragraph"
            )
            document.set_story_text(item, value)
        for kind in kinds:
            assert _story_paragraph_texts(document, kind) == [value], kind
        assert document.paragraphs[0].text == value


def test_set_story_text_rejects_a_story_that_is_not_in_the_document():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body")
    missing = rdocx.StoryItem(
        story=rdocx.Story(kind="header", part_name="/word/missing.xml", owner_index=0),
        kind="paragraph",
        index_path=(0,),
        text="stale",
        xml=b"",
    )
    held = document.paragraphs[0]
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="no header story"):
        document.set_story_text(missing, "edited")
    assert document.to_bytes() == before
    assert held.text == "body"


def _two_section_document():
    import rdocx

    section = (
        '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" '
        'w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" '
        'w:gutter="0"/></w:sectPr>'
    )
    source = rdocx.Document()
    source.add_paragraph("placeholder")
    return _replace_document_body(
        source,
        f"<w:p><w:pPr>{section}</w:pPr><w:r><w:t>First section.</w:t></w:r></w:p>"
        f"<w:p><w:r><w:t>Second section.</w:t></w:r></w:p>{section}",
    )


def _default_footer(document, section_index):
    return next(
        variant.story
        for variant in document.header_footer_variants
        if variant.section_index == section_index
        and variant.kind == "footer"
        and variant.variant == "default"
    )


def _story_part(document, story):
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        return archive.read(story.part_name.lstrip("/"))


def test_footer_with_text_tab_and_page_fields_is_built_for_one_section():
    import rdocx

    document = _two_section_document()
    held = document.paragraphs[0]
    footer = document.create_section_story(1, "footer", "default")
    assert footer.kind == "footer"
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert _default_footer(document, 0) is None
    assert _default_footer(document, 1) == footer

    # Author the paragraph in the body with the typed API, then move it.
    paragraph = document.add_paragraph("Confidential")
    paragraph.runs[0].add_tab()
    document.paragraphs[2].add_run("Page ").add_field("PAGE", "1")
    document.paragraphs[2].add_run(" of ").add_field("NUMPAGES", "1")
    fragment = document.pop_content(document.find_content_index("Confidential"))
    with pytest.raises(rdocx.StaleElementError):
        document.insert_content(footer, fragment)
    footer = _default_footer(document, 1)
    document.insert_content(footer, fragment)
    footer = _default_footer(document, 1)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "First section.",
        "Second section.",
    ]
    assert _story_paragraph_texts(document, "footer") == ["Confidential\tPage 1 of 1"]

    report = document.update_layout_backed_fields()
    assert (report.page_fields, report.num_pages_fields) == (0, 0)
    footer = _default_footer(document, 1)
    xml = _story_part(document, footer)
    assert re.search(rb"<w:t>Confidential</w:t>\s*<w:tab/>", xml)
    for instruction in (b"PAGE", b"NUMPAGES"):
        field = re.search(
            rb'w:instr="' + instruction + rb'"[^>]*>\s*<w:r>\s*<w:t>([^<]*)</w:t>', xml
        )
        assert field is not None and field.group(1) == b"1"
    body = _document_xml(document)
    first_section = body[: body.index(b"First section.")]
    assert b"footerReference" not in first_section
    assert body.count(b"footerReference") == 1
    assert document.to_pdf().startswith(b"%PDF")


def test_section_stories_link_unlink_and_reject_bad_names_atomically():
    import rdocx

    document = _two_section_document()
    footer = document.create_section_story(1, "footer", "default")
    document.add_paragraph("Footer text")
    fragment = document.pop_content(document.find_content_index("Footer text"))
    with pytest.raises(rdocx.StaleElementError):
        document.insert_content(footer, fragment)
    footer = _default_footer(document, 1)
    document.insert_content(footer, fragment)
    footer = _default_footer(document, 1)
    before = document.to_bytes()
    for arguments, error in (
        ((1, "side", "default"), ValueError),
        ((1, "footer", "odd"), ValueError),
        ((2, "footer", "default"), IndexError),
    ):
        with pytest.raises(error):
            document.create_section_story(*arguments)
        with pytest.raises(error):
            document.unlink_section_story(*arguments)
        with pytest.raises(error):
            document.link_section_story(*arguments, footer)
    with pytest.raises(rdocx.RdocxError):
        document.link_section_story(0, "header", "default", footer)
    with pytest.raises(TypeError, match="destination must be an int, a StoryItem or a Story"):
        document.insert_content("footer", fragment)
    with pytest.raises(TypeError, match="index must be an int or a StoryItem"):
        document.pop_content(footer)
    assert document.to_bytes() == before

    linked = document.link_section_story(0, "footer", "default", footer)
    footer = _default_footer(document, 1)
    assert linked == footer
    assert _default_footer(document, 0) == footer
    copy = document.unlink_section_story(0, "footer", "default")
    footer = _default_footer(document, 1)
    assert copy.part_name != footer.part_name
    assert _default_footer(document, 0) == copy
    assert _story_paragraph_texts(document, "footer") == ["Footer text", "Footer text"]

    [item] = [item for item in document.story_items if item.story == copy]
    copy_part_name = copy.part_name
    popped = document.pop_content(item)
    copy = next(story for story in document.stories if story.part_name == copy_part_name)
    footer = _default_footer(document, 1)
    assert popped.kind == "paragraph"
    assert [item for item in document.story_items if item.story == copy] == []
    assert _default_footer(document, 1) == footer
    assert _story_paragraph_texts(document, "footer") == ["Footer text"]


def test_hyperlinks_are_added_to_paragraphs_and_stories():
    import rdocx

    document = rdocx.Document()
    run = document.add_paragraph("See ").add_hyperlink("docs", "https://example.com/docs")
    run.font.bold = True
    document.set_header("Header")
    header = next(story for story in document.stories if story.kind == "header")
    document.add_hyperlink_to_story(header, "home", "https://example.com/")

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [(link.story.kind, link.text, link.url) for link in reopened.hyperlinks] == [
        ("body", "docs", "https://example.com/docs"),
        ("header", "home", "https://example.com/"),
    ]
    assert reopened.paragraphs[0].runs[1].font.bold is True


def test_complex_field_story_snapshots_preserve_cached_text_and_revision():
    import rdocx

    field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>7</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    seed = rdocx.Document()
    seed.set_header("header")
    seed.set_footer("footer")
    section_xml = "<w:sectPr" + _document_xml(seed).decode().split("<w:sectPr", 1)[1].split("</w:body>", 1)[0]
    document = _replace_document_body(
        seed, f'<w:p>{field}</w:p><w:sdt><w:sdtPr/><w:sdtContent>'
        f'<w:p>{field}</w:p></w:sdtContent></w:sdt>'
        + section_xml,
    )
    related = {story.part_name.lstrip("/"): story.kind for story in document.stories
               if story.kind in ("header", "footer")}
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as source, zipfile.ZipFile(result, "w") as output:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in related:
                root = "hdr" if related[info.filename] == "header" else "ftr"
                data = (f'<w:{root} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        f'<w:p>{field}</w:p></w:{root}>').encode()
            output.writestr(info, data)
    document = rdocx.Document.from_bytes(result.getvalue())
    before = document.to_bytes()
    held = document.paragraphs[0]
    snapshots = document.story_items
    fields = [item for item in snapshots if item.kind == "field"]
    assert [(item.story.kind, item.text) for item in fields] == [
        ("body", "7"), ("body", "7"), ("header", "7"), ("footer", "7")]
    assert [item.direct_body_index for item in fields] == [0, 1, None, None]
    assert all(field.encode() in item.xml for item in fields)
    assert len({item.revision for item in snapshots}) == 1
    assert document.story_items == snapshots
    assert document.to_bytes() == before
    assert held.text == "7"
    with zipfile.ZipFile(io.BytesIO(before)) as original:
        expected_parts = {name: original.read(name) for name in original.namelist()}
    reopened = rdocx.Document.from_bytes(before)
    assert [(item.story.kind, item.index_path, item.text, item.xml, item.direct_body_index)
            for item in reopened.story_items] == [
        (item.story.kind, item.index_path, item.text, item.xml, item.direct_body_index)
        for item in snapshots]
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as saved:
        assert {name: saved.read(name) for name in saved.namelist()} == expected_parts


def test_complex_field_nested_snapshots_keep_links_and_read_handles():
    import rdocx

    nested = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> IF </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>7</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        '<w:r><w:instrText> = 7 yes no </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>yes</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    seed = rdocx.Document()
    paragraph = seed.add_paragraph("")
    paragraph.add_hyperlink("seed", "https://example.test/field-owner")
    relationship_id = seed.hyperlinks[0].relationship_id
    document = _replace_document_body(seed,
        f'<w:p><w:hyperlink r:id="{relationship_id}">{nested}</w:hyperlink>'
        '<w:hyperlink w:anchor="neighbor"><w:r><w:t>neighbor</w:t></w:r></w:hyperlink></w:p>')
    before = document.to_bytes()
    held = document.paragraphs[0]
    snapshots = document.story_items
    fields = [item for item in snapshots if item.kind == "field"]
    assert [item.text for item in fields] == ["yes", "7"]
    assert all(item.revision == snapshots[0].revision for item in snapshots)
    links = document.hyperlinks
    assert [(link.text, link.url, link.anchor) for link in links] == [
        ("yes", "https://example.test/field-owner", None), ("neighbor", None, "neighbor")]
    assert links[0].relationship_id == relationship_id
    assert document.hyperlinks == links
    assert document.story_items == snapshots
    assert held.text == "yesneighbor"
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert [(item.index_path, item.text, item.xml) for item in reopened.story_items] == [
        (item.index_path, item.text, item.xml) for item in snapshots]
    assert [(link.index_path, link.text, link.url, link.anchor) for link in reopened.hyperlinks] == [
        (link.index_path, link.text, link.url, link.anchor) for link in links]


@pytest.mark.parametrize("default_binding,ancestor_word", [(False, False), (False, True), (True, False), (True, True)])
def test_complex_field_first_run_namespace_lifetimes_do_not_escape(default_binding, ancestor_word):
    import rdocx

    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ancestor = word if ancestor_word else "urn:producer"
    local = "urn:producer" if ancestor_word else word
    declaration = "xmlns" if default_binding else "xmlns:x"
    prefix = "" if default_binding else "x:"
    content = "8" if ancestor_word else "SHADOW"
    expected = "87" if ancestor_word else "7"
    header_xml = (f'<q:hdr xmlns:q="{word}" {declaration}="{ancestor}"><q:p>'
        f'<q:r {declaration}="{local}"><q:fldChar q:fldCharType="begin"/></q:r>'
        '<q:r><q:instrText> PAGE </q:instrText></q:r>'
        '<q:r><q:fldChar q:fldCharType="separate"/></q:r>'
        f'<q:r><{prefix}t>{content}</{prefix}t><q:t>7</q:t></q:r>'
        '<q:r><q:fldChar q:fldCharType="end"/></q:r></q:p></q:hdr>').encode()
    seed = rdocx.Document()
    seed.add_paragraph("held")
    seed.set_header("source")
    header = next(story for story in seed.stories if story.kind == "header")
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(seed.to_bytes())) as source, zipfile.ZipFile(result, "w") as output:
        for info in source.infolist():
            data = header_xml if info.filename == header.part_name.lstrip("/") else source.read(info.filename)
            output.writestr(info, data)
    document = rdocx.Document.from_bytes(result.getvalue())
    before = document.to_bytes()
    held = document.paragraphs[0]
    snapshots = document.story_items
    fields = [item for item in snapshots if item.kind == "field" and item.story.kind == "header"]
    assert len(fields) == 1
    assert all(item.revision == snapshots[0].revision for item in snapshots)
    assert document.story_items == snapshots
    assert held.text == "held"
    assert document.to_bytes() == before
    with zipfile.ZipFile(io.BytesIO(before)) as archive:
        assert archive.read(header.part_name.lstrip("/")) == header_xml
    reopened = rdocx.Document.from_bytes(before)
    assert [(item.story.kind, item.index_path, item.xml) for item in reopened.story_items] == [
        (item.story.kind, item.index_path, item.xml) for item in snapshots]
    assert fields[0].text == expected


def test_story_item_xml_is_a_detached_snapshot():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("hello")
    item = next(
        item
        for item in document.story_items
        if item.story.kind == "body" and item.kind == "paragraph"
    )
    assert isinstance(item.xml, bytes)
    assert b"hello" in item.xml
    document.add_paragraph("later")
    assert b"later" not in item.xml


def test_split_run_enables_exact_comment_ranges_without_losing_content():
    import re

    import rdocx

    document = rdocx.Document()
    run = document.add_paragraph("").add_run("Hello 🐝brave world")
    run.font.bold = True
    before_no_ops = document.to_bytes()
    paragraph = document.paragraphs[0]
    assert document.split_run(0, 0, 0) == 0
    assert document.split_run(0, 0, 18) == 1
    assert document.to_bytes() == before_no_ops
    assert paragraph.text == "Hello 🐝brave world"

    brave_start = document.split_run(0, 0, 6)
    with pytest.raises(rdocx.StaleElementError):
        run.text
    brave_end = document.split_run(0, brave_start, 6)
    assert (brave_start, brave_end) == (1, 2)
    runs = document.paragraphs[0].runs
    assert [item.text for item in runs] == ["Hello ", "🐝brave", " world"]
    assert [item.font.bold for item in runs] == [True, True, True]

    brave = rdocx.RunRange(
        start=rdocx.RunPosition(body_index=0, run_index=brave_start),
        end=rdocx.RunPosition(body_index=0, run_index=brave_end),
    )
    comment_id = document.add_comment(brave, author="Ada", text="Which one?")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    xml = _document_xml(reopened).decode()
    start = xml.index(f'commentRangeStart w:id="{comment_id}"')
    end = xml.index(f'commentRangeEnd w:id="{comment_id}"')
    assert re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml[start:end]) == [
        "🐝brave"
    ]


def test_split_run_rejects_bad_coordinates_without_changing_the_document():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("abc")
    before = document.to_bytes()
    for args in ((1, 0, 1), (0, 1, 1), (0, 0, 4)):
        with pytest.raises(rdocx.RdocxError):
            document.split_run(*args)
        assert document.to_bytes() == before


def test_split_run_takes_the_direct_body_index_after_a_table():
    import rdocx

    # GitHub issue #163: the index find_content_index returns must address
    # the same paragraph in split_run.
    document = rdocx.Document()
    document.add_paragraph("Alpha paragraph before the table.")
    document.add_table(1, 1).cell(0, 0).text = "cell"
    document.add_paragraph("Beta paragraph after the table.")
    document.add_paragraph("Gamma paragraph at the end.")
    target = next(p for p in document.paragraphs if p.text.startswith("Beta"))
    body_index = document.find_content_index(target)
    assert body_index == 2

    assert document.split_run(body_index, 0, 4) == 1
    assert [[run.text for run in p.runs] for p in document.paragraphs] == [
        ["Alpha paragraph before the table."],
        ["Beta", " paragraph after the table."],
        ["Gamma paragraph at the end."],
    ]

    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="is a table, not a paragraph"):
        document.split_run(1, 0, 1)
    with pytest.raises(TypeError, match="must be an int or a Paragraph handle"):
        document.split_run("Beta", 0, 1)
    with pytest.raises(OverflowError):
        document.split_run(-1, 0, 1)
    assert document.to_bytes() == before


def test_split_run_accepts_a_paragraph_handle_in_a_control_or_the_body():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        "<w:p><w:r><w:t>Alpha.</w:t></w:r></w:p>"
        "<w:sdt><w:sdtContent>"
        "<w:p><w:r><w:t>Control one.</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Control two.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
        '<w:tbl><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>'
        "<w:tr><w:tc>"
        "<w:sdt><w:sdtContent>"
        "<w:p><w:r><w:t>In control.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
        "<w:p><w:r><w:t>Cell text.</w:t></w:r></w:p>"
        "</w:tc></w:tr>"
        "</w:tbl>"
        "<w:p><w:r><w:t>Beta.</w:t></w:r></w:p>",
    )
    control_two = document.paragraphs[2]
    assert control_two.text == "Control two."
    run = control_two.runs[0]
    before_no_op = document.to_bytes()
    assert document.split_run(control_two, 0, 0) == 0
    assert document.to_bytes() == before_no_op
    assert run.text == "Control two."

    assert document.split_run(control_two, 0, 7) == 1
    with pytest.raises(rdocx.StaleElementError):
        run.text
    with pytest.raises(rdocx.StaleElementError):
        document.split_run(control_two, 0, 1)
    assert [r.text for r in document.paragraphs[2].runs] == ["Control", " two."]

    # A cell handle is refused, because its index counts the paragraph
    # inside the cell's content control and the cell writer does not.
    cell = document.tables[0].rows[0].cells[0]
    assert [p.text for p in cell.paragraphs] == ["In control.", "Cell text."]
    before_cell = document.to_bytes()
    for cell_paragraph in cell.paragraphs:
        with pytest.raises(ValueError, match="table cell paragraph handle"):
            document.split_run(cell_paragraph, 0, 2)
    assert document.to_bytes() == before_cell

    beta = document.paragraphs[3]
    assert document.find_content_index(beta) == 3
    beta_run = beta.runs[0]
    assert document.split_run(beta, 0, 5) == 1
    assert beta_run.text == "Beta."
    assert document.split_run(beta, 0, 4) == 1
    assert [r.text for r in document.paragraphs[3].runs] == ["Beta", "."]

    other = rdocx.Document()
    other.add_paragraph("elsewhere")
    before = document.to_bytes()
    with pytest.raises(ValueError, match="different document"):
        document.split_run(other.paragraphs[0], 0, 1)
    assert document.to_bytes() == before


def test_run_remove_keeps_markers_and_drops_emptied_wrappers():
    import rdocx

    # GitHub issue #168 asks for the python-docx idiom
    # run._r.getparent().remove(run._r) as a method of the run.
    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:bookmarkStart w:id="1" w:name="kept"/>'
        '<w:r><w:t xml:space="preserve">Keep </w:t></w:r>'
        "<w:r><w:t>drop</w:t></w:r><w:bookmarkEnd w:id=\"1\"/>"
        '<w:hyperlink w:anchor="kept"><w:r><w:t>link</w:t></w:r></w:hyperlink>'
        '<w:ins w:id="2" w:author="Ada"><w:r><w:t>inserted</w:t></w:r></w:ins>'
        "</w:p>",
    )
    paragraph = document.paragraphs[0]
    runs = paragraph.runs
    assert [run.text for run in runs] == ["Keep ", "drop", "link", "inserted"]
    kept = runs[0]
    runs[1].remove()
    with pytest.raises(rdocx.StaleElementError):
        kept.text
    with pytest.raises(rdocx.StaleElementError):
        paragraph.text
    assert [run.text for run in document.paragraphs[0].runs] == [
        "Keep ",
        "link",
        "inserted",
    ]
    document.paragraphs[0].runs[1].remove()
    document.paragraphs[0].runs[1].remove()
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [run.text for run in reopened.paragraphs[0].runs] == ["Keep "]
    xml = _document_xml(reopened).decode()
    assert '<w:bookmarkStart w:id="1" w:name="kept"/>' in xml
    assert '<w:bookmarkEnd w:id="1"/>' in xml
    assert "hyperlink" not in xml
    assert "w:ins" not in xml

    table = document.add_table(1, 1)
    table.cell(0, 0).text = "cell"
    document.tables[0].cell(0, 0).paragraphs[0].runs[0].remove()
    assert document.tables[0].cell(0, 0).paragraphs[0].text == ""


def test_run_remove_refuses_part_of_a_field_or_a_comment_reference():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>Entry</w:t></w:r></w:p>"
        '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        "<w:p><w:r><w:t>Commented</w:t></w:r></w:p>",
    )
    comment_range = rdocx.RunRange(
        start=rdocx.RunPosition(body_index=2, run_index=0),
        end=rdocx.RunPosition(body_index=2, run_index=1),
    )
    document.add_comment(comment_range, author="Ada", text="Here")
    before = document.to_bytes()
    for paragraph, run, reason in (
        (0, 0, "part of a complex field"),
        (0, 2, "part of a complex field"),
        (1, 0, "part of a complex field"),
        (2, 1, "the reference of comment 0"),
    ):
        handle = document.paragraphs[paragraph].runs[run]
        with pytest.raises(rdocx.RdocxError, match=reason):
            handle.remove()
        # A refused removal leaves the handles valid.
        assert handle.text == ""
    assert document.to_bytes() == before

    document.paragraphs[0].runs[3].remove()
    assert document.paragraphs[0].text == ""
    assert [comment.text for comment in document.comments] == ["Here"]


def test_story_comment_after_a_block_content_control_anchors_on_its_paragraph():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        "<w:p><w:r><w:t>Alpha.</w:t></w:r></w:p>"
        "<w:sdt><w:sdtContent>"
        "<w:p><w:r><w:t>Control one.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
        "<w:p><w:r><w:t>Beta paragraph.</w:t></w:r></w:p>",
    )
    item = next(
        item
        for item in document.story_items
        if item.story.kind == "body"
        and item.kind == "paragraph"
        and item.text == "Beta paragraph."
    )
    comment_id = document.add_comment(
        rdocx.StoryRunRange(
            start=rdocx.StoryRunPosition(item=item, run_index=0),
            end=rdocx.StoryRunPosition(item=item, run_index=1),
        ),
        author="Ada",
        text="Here",
    )
    xml = _document_xml(document).decode()
    start = xml.index(f'commentRangeStart w:id="{comment_id}"')
    end = xml.index(f'commentRangeEnd w:id="{comment_id}"')
    assert re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml[start:end]) == [
        "Beta paragraph."
    ]


# GitHub issue #172: a run index read from Paragraph.runs anchors on that run,
# including the runs inside an inline content control.
_INLINE_CONTROL_PARAGRAPH = (
    '<w:p><w:r><w:t xml:space="preserve">before </w:t></w:r>'
    "<w:sdt><w:sdtPr/><w:sdtContent>{runs}</w:sdtContent></w:sdt>"
    '<w:r><w:t xml:space="preserve"> after</w:t></w:r></w:p>'
)


def _anchored_texts(document, comment_id):
    xml = _document_xml(document).decode()
    start = xml.index(f'commentRangeStart w:id="{comment_id}"')
    end = xml.index(f'commentRangeEnd w:id="{comment_id}"')
    return re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml[start:end])


def _run_range(start, end, body_index=0):
    import rdocx

    return rdocx.RunRange(
        start=rdocx.RunPosition(body_index=body_index, run_index=start),
        end=rdocx.RunPosition(body_index=body_index, run_index=end),
    )


def test_comment_run_index_counts_the_runs_of_an_inline_content_control():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        _INLINE_CONTROL_PARAGRAPH.format(runs="<w:r><w:t>TARGET</w:t></w:r>"),
    )
    runs = [run.text for run in document.paragraphs[0].runs]
    assert runs == ["before ", "TARGET", " after"]

    target = document.add_comment(_run_range(1, 2), author="A", text="x")
    assert _anchored_texts(document, target) == ["TARGET"]
    # The reference run of the first comment is run 2 now.
    assert [run.text for run in document.paragraphs[0].runs][3] == " after"
    last = document.add_comment(_run_range(3, 4), author="A", text="y")
    assert _anchored_texts(document, last) == [" after"]

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [comment.text for comment in reopened.comments] == ["x", "y"]
    assert _anchored_texts(reopened, target) == ["TARGET"]


def test_story_comment_run_index_counts_the_runs_of_an_inline_content_control():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        _INLINE_CONTROL_PARAGRAPH.format(runs="<w:r><w:t>TARGET</w:t></w:r>"),
    )
    item = next(
        item
        for item in document.story_items
        if item.story.kind == "body" and item.kind == "paragraph"
    )
    comment_id = document.add_comment(
        rdocx.StoryRunRange(
            start=rdocx.StoryRunPosition(item=item, run_index=1),
            end=rdocx.StoryRunPosition(item=item, run_index=2),
        ),
        author="A",
        text="x",
    )
    assert _anchored_texts(document, comment_id) == ["TARGET"]


def test_comment_ranges_that_cannot_be_anchored_exactly_are_refused():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        _INLINE_CONTROL_PARAGRAPH.format(
            runs="<w:r><w:t>A</w:t></w:r><w:r><w:t>B</w:t></w:r>"
        ),
    )
    before = document.to_bytes()
    for start, end in ((0, 2), (2, 4)):
        with pytest.raises(
            rdocx.RdocxError, match="crosses the edge of an inline content control"
        ):
            document.add_comment(_run_range(start, end), author="A", text="x")
    assert document.to_bytes() == before

    comment_id = document.add_comment(_run_range(2, 3), author="A", text="x")
    assert _anchored_texts(document, comment_id) == ["B"]
    assert document.remove_comment(comment_id)
    xml = _document_xml(document).decode()
    assert "commentRange" not in xml and "commentReference" not in xml
    assert [run.text for run in document.paragraphs[0].runs] == [
        "before ",
        "A",
        "B",
        " after",
    ]


def test_story_run_position_takes_a_paragraph_handle_inside_a_block_control():
    import rdocx

    # GitHub issue #163: a paragraph inside a block content control has no
    # story item of its own, and a handle reaches it.
    document = _replace_document_body(
        rdocx.Document(),
        "<w:p><w:r><w:t>Alpha.</w:t></w:r></w:p>"
        "<w:sdt><w:sdtContent>"
        "<w:p><w:r><w:t>Control one.</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Control two.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
        '<w:tbl><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>'
        "<w:tr><w:tc><w:p><w:r><w:t>Cell.</w:t></w:r></w:p></w:tc></w:tr>"
        "</w:tbl>"
        "<w:p><w:r><w:t>Beta.</w:t></w:r></w:p>",
    )
    control_two = document.paragraphs[2]
    assert control_two.text == "Control two."
    position = rdocx.StoryRunPosition(paragraph=control_two, run_index=0)
    assert len(position.item.index_path) == 2
    assert position.item.text == "Control two."
    assert position.item.direct_body_index == 1
    comment_id = document.add_comment(
        rdocx.StoryRunRange(
            start=position,
            end=rdocx.StoryRunPosition(paragraph=control_two, run_index=1),
        ),
        author="Ada",
        text="Here",
    )
    assert _anchored_texts(document, comment_id) == ["Control two."]
    xml = _document_xml(document).decode()
    assert xml.index("<w:sdtContent>") < xml.index("commentRangeStart")
    assert xml.index("commentRangeEnd") < xml.index("</w:sdtContent>")
    with pytest.raises(rdocx.StaleElementError):
        rdocx.StoryRunPosition(paragraph=control_two, run_index=0)

    beta = document.paragraphs[3]
    beta_item = next(item for item in document.story_items if item.text == "Beta.")
    assert rdocx.StoryRunPosition(paragraph=beta, run_index=0).item == beta_item

    cell_paragraph = document.tables[0].rows[0].cells[0].paragraphs[0]
    with pytest.raises(ValueError, match="table cell paragraph handle"):
        rdocx.StoryRunPosition(paragraph=cell_paragraph, run_index=0)
    with pytest.raises(TypeError, match="exactly one of item and paragraph"):
        rdocx.StoryRunPosition(item=beta_item, paragraph=beta, run_index=0)
    with pytest.raises(TypeError, match="exactly one of item and paragraph"):
        rdocx.StoryRunPosition(run_index=0)


def test_add_comment_on_text_anchors_the_occurrence_after_a_table():
    import rdocx

    # GitHub issue #163 asks for a helper that comments on a piece of text.
    # This is the fixture of the issue.
    def fixture():
        document = rdocx.Document()
        document.add_paragraph("Alpha paragraph before the table.")
        document.add_table(1, 1).cell(0, 0).text = "cell"
        document.add_paragraph("Beta paragraph after the table.")
        document.add_paragraph("Gamma paragraph at the end.")
        return document

    document = fixture()
    held = document.paragraphs[1]
    comment_id = document.add_comment_on_text("Beta", author="Ada", text="Here")
    assert _anchored_texts(document, comment_id) == ["Beta"]
    assert [run.text for run in document.paragraphs[1].runs if run.text] == [
        "Beta",
        " paragraph after the table.",
    ]
    with pytest.raises(rdocx.StaleElementError):
        held.text

    document = fixture()
    comment_id = document.add_comment_on_text(
        "paragraph",
        author="Ada",
        text="Second",
        occurrence=1,
        initials="AL",
        date="2026-09-27T10:00:00Z",
    )
    assert _anchored_texts(document, comment_id) == ["paragraph"]
    assert [run.text for run in document.paragraphs[1].runs if run.text] == [
        "Beta ",
        "paragraph",
        " after the table.",
    ]
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    [comment] = reopened.comments
    assert (comment.text, comment.initials, comment.date) == (
        "Second",
        "AL",
        "2026-09-27T10:00:00Z",
    )

    before = document.to_bytes()
    for anchor, occurrence in (("beta", 0), ("paragraph", 3), ("", 0)):
        with pytest.raises(rdocx.RdocxError):
            document.add_comment_on_text(
                anchor, author="Ada", text="x", occurrence=occurrence
            )
    assert document.to_bytes() == before


def test_python_round_three_authoring_and_inspection_is_typed_and_lossless():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("before")
    anchor = next(
        item
        for item in document.story_items
        if item.story.kind == "body" and item.kind == "paragraph"
    )
    picture = document.add_picture(
        _one_pixel_png(), "pixel.png", rdocx.Inches(1), rdocx.Inches(1), after=anchor
    )
    assert isinstance(picture, rdocx.Picture)
    assert picture.location.item_kind == "paragraph"
    owner = next(item for item in document.story_items if item.story.kind == picture.story.kind and item.index_path == picture.location.index_path)
    assert b"<w:drawing>" in owner.xml
    before_bad_picture = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        document.add_picture(
            _one_pixel_png(), "pixel.png", width=rdocx.Inches(1), after=picture
        )
    assert document.to_bytes() == before_bad_picture

    table = document.add_table(1, 1)
    table.cell(0, 0).text = "cell text"
    cell_item = next(
        item
        for item in document.story_items
        if item.story.kind == "table_cell"
        and item.kind == "paragraph"
        and item.text == "cell text"
    )
    before_bad_comment = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        document.add_comment(
            rdocx.StoryRunRange(
                start=rdocx.StoryRunPosition(item=cell_item, run_index=0),
                end=rdocx.StoryRunPosition(item=cell_item, run_index=9),
            ),
            author="Ada",
            text="invalid",
        )
    assert document.to_bytes() == before_bad_comment
    comment_id = document.add_comment(
        rdocx.StoryRunRange(
            start=rdocx.StoryRunPosition(item=cell_item, run_index=0),
            end=rdocx.StoryRunPosition(item=cell_item, run_index=1),
        ),
        author="Ada",
        text="Check this cell",
    )

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.comments[0].id == comment_id
    xml = _document_xml(reopened)
    assert b"pixel.png" not in xml
    assert f'w:id="{comment_id}"'.encode() in xml
    assert b"cell text" in xml


def test_paragraph_text_setter_replaces_runs_and_keeps_format_and_comments():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Alpha ")
    document.paragraphs[0].add_run("beta").font.bold = True
    document.paragraphs[0].style = "Heading1"
    document.paragraphs[0].alignment = rdocx.WD_ALIGN_PARAGRAPH.CENTER
    comment_id = document.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=1),
            end=rdocx.RunPosition(body_index=0, run_index=2),
        ),
        author="Ada",
        text="Bold?",
    )
    paragraph = document.paragraphs[0]
    run = paragraph.runs[0]

    paragraph.text = "Omega\tone\ntwo"

    assert paragraph.text == "Omega\tone\ntwo"
    with pytest.raises(rdocx.StaleElementError):
        _ = run.text
    paragraph = document.paragraphs[0]
    assert paragraph.text == "Omega\tone\ntwo"
    assert paragraph.style == "Heading1"
    assert paragraph.alignment == rdocx.WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.runs[0].text == "Omega\tone\ntwo"
    assert paragraph.runs[0].font.bold is None
    xml = re.sub(r">\s+<", "><", _document_xml(document).decode())
    assert "<w:t>Omega</w:t><w:tab/><w:t>one</w:t><w:br/><w:t>two</w:t>" in xml
    start = xml.index(f'<w:commentRangeStart w:id="{comment_id}"/>')
    end = xml.index(f'<w:commentRangeEnd w:id="{comment_id}"/>')
    assert start < xml.index("Omega") < end
    assert f'<w:commentReference w:id="{comment_id}"/>' in xml[end:]
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [comment.text for comment in reopened.comments] == ["Bold?"]

    reopened.paragraphs[0].text = None
    assert reopened.paragraphs[0].text == ""

    document.add_table(1, 1)
    document.tables[0].rows[0].cells[0].paragraphs[0].text = "cell"
    assert document.tables[0].rows[0].cells[0].text == "cell"


def test_paragraph_text_setter_rejects_every_paragraph_of_a_multi_paragraph_field():
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3"</w:instrText></w:r></w:p>'
        '<w:p><w:r><w:instrText xml:space="preserve"> \\h </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>Entry</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Last entry</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',
    )
    held = list(document.paragraphs)
    texts = [paragraph.text for paragraph in held]
    before = document.to_bytes()

    for paragraph in held:
        with pytest.raises(rdocx.RdocxError, match="field"):
            paragraph.text = "Rewritten"

    assert document.to_bytes() == before
    assert [paragraph.text for paragraph in held] == texts


_CORE_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"
)


def _rewrite_package(document, rewrite):
    source = io.BytesIO(document.to_bytes())
    result = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip:
        with zipfile.ZipFile(result, "w") as result_zip:
            for info in source_zip.infolist():
                data = rewrite(info.filename, source_zip.read(info.filename))
                if data is not None:
                    result_zip.writestr(info, data)
    return type(document).from_bytes(result.getvalue())


def _package_part(document, name):
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        return archive.read(name).decode()


def test_core_properties_use_python_docx_names_and_round_trip(tmp_path):
    import datetime

    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body")
    held = document.paragraphs[0]
    core = document.core_properties
    assert (core.title, core.author, core.revision, core.created) == ("", "", 0, None)

    utc = datetime.timezone.utc
    plus_two = datetime.timezone(datetime.timedelta(hours=2))
    texts = {
        "author": "Ada",
        "category": "Reports",
        "comments": "Reviewed twice",
        "content_status": "Draft",
        "identifier": "DOC-7",
        "keywords": "alpha, beta",
        "language": "en-GB",
        "last_modified_by": "Grace",
        "subject": "Quarterly figures",
        "title": "A title",
        "version": "1.4",
    }
    for name, value in texts.items():
        setattr(core, name, value)
    core.revision = 3
    core.created = datetime.datetime(2026, 9, 1, 10, 0, 0, tzinfo=plus_two)
    core.modified = datetime.datetime(2026, 9, 2, 8, 30, 0)
    core.last_printed = datetime.datetime(2026, 9, 3, 7, 0, 0, tzinfo=utc)
    assert held.text == "body"

    path = tmp_path / "core.docx"
    document.save(path)
    reopened = rdocx.Document(path).core_properties
    for name, value in texts.items():
        assert getattr(reopened, name) == value, name
    assert reopened.revision == 3
    assert reopened.created == datetime.datetime(2026, 9, 1, 8, 0, 0, tzinfo=utc)
    assert reopened.modified == datetime.datetime(2026, 9, 2, 8, 30, 0, tzinfo=utc)
    assert reopened.last_printed == datetime.datetime(2026, 9, 3, 7, 0, 0, tzinfo=utc)
    xml = _package_part(document, "docProps/core.xml")
    assert '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-01T08:00:00Z<' in xml
    assert "<cp:lastPrinted>2026-09-03T07:00:00Z</cp:lastPrinted>" in xml

    docx = pytest.importorskip("docx")
    oracle = docx.Document(str(path)).core_properties
    for name, value in texts.items():
        assert getattr(oracle, name) == value, name
    assert oracle.revision == 3
    assert oracle.created == reopened.created

    core.title = None
    core.comments = ""
    core.revision = None
    core.created = None
    assert (core.title, core.comments, core.revision, core.created) == ("", "", 0, None)
    xml = _package_part(document, "docProps/core.xml")
    for element in ("dc:title", "dc:description", "cp:revision", "dcterms:created"):
        assert element not in xml, element


def test_core_properties_reject_bad_values_without_changing_the_document():
    import datetime

    import rdocx

    document = rdocx.Document()
    document.core_properties.title = "Kept"
    before = document.to_bytes()
    for name, value, error in (
        ("revision", 0, ValueError),
        ("title", "x" * 256, ValueError),
        ("created", "2026-09-01", TypeError),
        ("last_printed", datetime.date(2026, 9, 1), TypeError),
    ):
        with pytest.raises(error):
            setattr(document.core_properties, name, value)
    assert document.to_bytes() == before
    document.core_properties.title = "x" * 255
    assert len(document.core_properties.title) == 255


def test_core_properties_read_w3cdtf_dates_like_python_docx():
    import datetime

    import rdocx

    def with_core(created, modified, last_printed):
        def core_xml(name, data):
            if name != "docProps/core.xml":
                return data
            return (
                '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/'
                '2006/metadata/core-properties" xmlns:dcterms="http://purl.org/dc/terms/" '
                'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
                f'<dcterms:created xsi:type="dcterms:W3CDTF">{created}</dcterms:created>'
                f'<dcterms:modified xsi:type="dcterms:W3CDTF">{modified}</dcterms:modified>'
                f"<cp:lastPrinted>{last_printed}</cp:lastPrinted>"
                "<cp:revision>seven</cp:revision>"
                "</cp:coreProperties>"
            ).encode()

        return _rewrite_package(rdocx.Document(), core_xml).core_properties

    utc = datetime.timezone.utc
    core = with_core("2024-01-15T10:30:00-05:30", "2024-02", "not a date")
    assert core.created == datetime.datetime(2024, 1, 15, 16, 0, 0, tzinfo=utc)
    assert core.modified == datetime.datetime(2024, 2, 1, tzinfo=utc)
    assert core.last_printed is None
    assert core.revision == 0

    core = with_core(
        "2024-01-15T10:30:00+99:99", "0001-01-01T00:30:00+01:00", "2024-13-01"
    )
    assert (core.created, core.modified, core.last_printed) == (None, None, None)


def test_core_properties_create_a_missing_part_with_its_relationship():
    import rdocx

    def drop_core(name, data):
        if name == "docProps/core.xml":
            return None
        if name == "_rels/.rels":
            return re.sub(
                rb'<Relationship [^>]*metadata/core-properties"[^>]*/>', b"", data
            )
        if name == "[Content_Types].xml":
            return re.sub(rb'<Override PartName="/docProps/core.xml"[^>]*/>', b"", data)
        return data

    document = _rewrite_package(rdocx.Document(), drop_core)
    assert _CORE_RELATIONSHIP not in _package_part(document, "_rels/.rels")
    assert document.core_properties.title == ""

    document.core_properties.title = "Created"

    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.core_properties.title == "Created"
    relationships = _package_part(document, "_rels/.rels")
    assert f'Type="{_CORE_RELATIONSHIP}" Target="docProps/core.xml"' in relationships
    assert (
        '<Override PartName="/docProps/core.xml" '
        'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
    ) in _package_part(document, "[Content_Types].xml")


def test_fx153_run_remove_survives_save_and_reopen():
    import rdocx

    document = rdocx.Document()
    paragraph = document.add_paragraph("Keep")
    paragraph.add_run(" Drop").remove()
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.paragraphs[0].text == "Keep"


def test_issue_168_complete_edit_save_reopen_layout_and_render(tmp_path):
    import importlib.metadata
    from pathlib import Path

    import docx
    from lxml import etree
    import rdocx

    assert importlib.metadata.version("python-docx") == "1.2.0"
    document = rdocx.Document()
    document.add_paragraph("Report {{name}}. See ")
    document.add_paragraph("Target phrase {{date}}")
    table = document.insert_table(1, 3, 3)
    table.set_cell_grid_span(0, 0, 2)
    table = document.tables[0]
    table.set_cell_vertical_merge(1, 2, "restart")
    table.set_cell_vertical_merge(2, 2, "continue")
    table.set_borders("single", size=4, color="000000")
    table.set_cell_margins(top=0, right=63500, bottom=12700, left=127000)
    table.grid_widths = [1371600, 1828800, 1828800]
    table.cell(1, 0).shading = "D9D9D9"
    table.cell(1, 0).set_margins(top=12700, right=0, bottom=25400, left=6350)
    table.cell(1, 0).set_border("bottom", "double", size=6, color="auto")
    table.rows[0].height = 254000
    table.rows[0].cant_split = True
    table.rows[0].is_header = True

    document.add_style("Report Note", based_on="Normal", bold=True)
    definition = document.add_numbering_definition([rdocx.ListLevel(format="decimal")])
    instance = document.add_numbering_instance(definition)
    document.link_style_to_numbering("Report Note", instance, 0)
    document.paragraphs[1].style = "Report Note"
    document.set_default_style("Report Note")
    before_style_render = document.render_page_to_png(0, 72)
    document.set_style("Report Note", bold=False, space_after=rdocx.Pt(6))
    assert document.render_page_to_png(0, 72) != before_style_render
    before_invalid_style = document.to_bytes()
    with pytest.raises(KeyError):
        document.set_style("Missing", bold=True)
    assert document.to_bytes() == before_invalid_style
    assert document.remove_style("Subtitle") is True
    document.update_section(0, orientation="landscape", margin_left=rdocx.Inches(1.5))
    document.insert_section(1)
    document.add_paragraph("Second section {{state}}")
    footer = document.create_section_story(1, "footer", "default")
    footer_part = footer.part_name
    document.add_paragraph("Confidential").runs[0].add_tab()
    document.paragraphs[-1].add_run("Page ").add_field("PAGE", "1")
    document.paragraphs[-1].add_run(" of ").add_field("NUMPAGES", "1")
    fragment = document.pop_content(document.find_content_index("Confidential"))
    footer = next(story for story in document.stories if story.part_name == footer_part)
    document.insert_content(footer, fragment)

    target_index = document.find_content_index("Target phrase")
    document.split_run(target_index, 0, len("Target phrase"))
    document.add_bookmark(
        "target",
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=target_index, run_index=0),
            end=rdocx.RunPosition(body_index=target_index, run_index=1),
        ),
    )
    document.paragraphs[0].runs[0].add_field("PAGEREF target \\h", "?")
    document.paragraphs[0].add_hyperlink("source", "https://example.com/old")
    document.set_hyperlink_url(document.hyperlinks[0], "https://example.com/new")
    document.paragraphs[0].add_hyperlink("remove", "https://example.com/remove")
    document.remove_hyperlink(document.hyperlinks[-1])
    document.add_picture(_one_pixel_png(), "pixel.png")
    relationship_id = re.search(rb'r:embed="([^"]+)"', _document_xml(document)).group(1).decode()
    assert document.set_picture_sizes_for_relationship(
        relationship_id, width=rdocx.Inches(1), height=rdocx.Inches(1)
    ) == 1

    held = document.paragraphs[0]
    before = document.to_bytes()
    with pytest.raises(rdocx.ReplacementCountError) as mismatch:
        document.replace_all(
            [("{{name}}", "Ada", 1), ("{{date}}", "October", 2), ("{{state}}", "Done", 1)]
        )
    assert mismatch.value.index == 1
    assert document.to_bytes() == before
    assert held.text.startswith("Report {{name}}")
    assert document.try_replace_text("{{missing}}", "x", expect=0) == 0
    assert document.replace_all(
        [("{{name}}", "Ada", 1), ("{{date}}", "October", 1), ("{{state}}", "Done", 1)]
    ) == (1, 1, 1)
    document.paragraphs[-1].text = "Second section Done"
    document.paragraphs[1].add_run(" remove").remove()
    comment_id = document.add_comment_on_text("Target phrase", author="Ada", text="Review")
    assert _anchored_texts(document, comment_id) == ["Target phrase"]
    report = document.update_layout_backed_fields()
    assert (report.page_fields, report.num_pages_fields, report.page_reference_fields) == (
        0, 0, 1
    )

    path = tmp_path / "issue-168.docx"
    document.save(path)
    # The accepted package-write escape hatch is an external lxml ZIP step.
    # Change an existing unmodelled part without adding a typed binding writer.
    changed = tmp_path / "issue-168-package-edited.docx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(changed, "w") as output:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                application = root.find(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Application"
                )
                assert application is not None
                application.text = "Issue 168 workflow"
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            output.writestr(member, data)
    path = changed
    reopened = rdocx.Document(path)
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        assert b"<w:tab/>" in package.read(footer.part_name.lstrip("/"))
        assert b"Issue 168 workflow" in package.read("docProps/app.xml")
    assert [style.style_id for style in reopened.styles if style.is_default] == ["ReportNote"]
    assert docx.Document(str(path)).styles["Report Note"].font.bold is False
    assert [(link.text, link.url) for link in reopened.hyperlinks] == [
        ("source", "https://example.com/new")
    ]
    assert reopened.comments[0].text == "Review"
    assert reopened.tables[0].rows[0].is_header is True
    assert reopened.layout_page(0).width == 792.0
    fonts = Path(__file__).parents[2] / "oxml-layout" / "fonts"
    assert reopened.to_pdf(font_dir=fonts).startswith(b"%PDF")
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as package:
        assert b"Issue 168 workflow" in package.read("docProps/app.xml")
    oracle = docx.Document(str(path))
    assert oracle.paragraphs[0].text.startswith("Report Ada. See ")
    assert [(link.text, link.url) for link in oracle.paragraphs[0].hyperlinks] == [
        ("source", "https://example.com/new")
    ]
    assert oracle.tables[0].cell(1, 0).text == ""
    assert oracle.sections[1].footer.paragraphs[0].text.startswith("Confidential")


# Issue 282, hadim: ownership checks happen before publication or invalidation.
def test_comment_removal_paths_preserve_atomicity_and_revisions():
    import rdocx

    def commented(kind, wrapped=False):
        document = rdocx.Document()
        document.add_paragraph("body anchor")
        document.add_paragraph("retained")
        if kind == "cell":
            document.add_table(2, 1).cell(0, 0).text = "cell anchor"
        item = next(item for item in document.story_items
                    if item.kind == "paragraph"
                    and item.story.kind == ("table_cell" if kind == "cell" else "body")
                    and item.text == ("cell anchor" if kind == "cell" else "body anchor"))
        root = document.add_comment(rdocx.StoryRunRange(
            start=rdocx.StoryRunPosition(item=item, run_index=0),
            end=rdocx.StoryRunPosition(item=item, run_index=1),
        ), author="Ada", text="root")
        reply = document.reply_to(root, author="Ben", text="reply")
        document.reply_to(reply, author="Cyd", text="grandchild")
        if wrapped:
            xml = _document_xml(document).decode()
            body = xml[xml.index("<w:body>") + len("<w:body>"):xml.index("</w:body>")]
            start = body.index("<w:p")
            end = body.index("</w:p>", start) + len("</w:p>")
            body = (body[:start] + '<w:sdt><w:sdtPr><w:tag w:val="goog_rdk_ownership"/></w:sdtPr>'
                    '<w:sdtContent>' + body[start:end] + '</w:sdtContent></w:sdt>' + body[end:])
            document = _replace_document_body(document, body)
        return document

    for wrapped in (False, True):
        document = commented("body", wrapped)
        held = document.paragraphs[-1]
        before = document.to_bytes()
        with pytest.raises(rdocx.RdocxError, match="comment"):
            document.pop_content(0)
        assert document.to_bytes() == before
        assert held.text == "retained"
        assert document.remove_content(0) is True
        assert document.comments == ()
        with pytest.raises(rdocx.StaleElementError):
            held.text
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        assert reopened.comments == ()
        assert [p.text for p in reopened.paragraphs] == ["retained"]

    document = commented("cell")
    held = document.paragraphs[0]
    removed_cell = document.tables[0].cell(0, 0)
    removed_run = removed_cell.paragraphs[0].runs[0]
    document.tables[0].remove_row(0)
    assert document.comments == ()
    assert held.text == "body anchor"
    for retired in (lambda: removed_cell.text, lambda: removed_run.text):
        with pytest.raises(rdocx.StaleElementError):
            retired()
    assert rdocx.Document.from_bytes(document.to_bytes()).comments == ()

    document = commented("cell")
    held = document.tables[0].cell(0, 0)
    replaced_run = held.paragraphs[0].runs[0]
    held.text = "replacement"
    assert len(document.comments) == 3
    assert held.text == "replacement"
    with pytest.raises(rdocx.StaleElementError):
        _ = replaced_run.text
    assert document.tables[0].cell(0, 0).text == "replacement"
    assert len(rdocx.Document.from_bytes(document.to_bytes()).comments) == 3

    document = rdocx.Document()
    document.add_paragraph("first")
    document.add_paragraph("last")
    document.add_comment(rdocx.RunRange(
        start=rdocx.RunPosition(body_index=0, run_index=0),
        end=rdocx.RunPosition(body_index=1, run_index=1),
    ), author="Ada", text="partial")
    held = document.paragraphs[0]
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="comment"):
        document.remove_content(0)
    assert document.to_bytes() == before
    assert held.text == "first"


def test_comment_anchor_fields_preserve_constructor_compatibility():
    import rdocx

    detached = rdocx.Comment(id=7, author="Ada", initials=None, date=None,
                             text="Review", parent_id=None, resolved=False)
    assert detached.anchor_text is None
    assert detached.anchor is None
    document = rdocx.Document()
    document.add_paragraph("Alpha stays.")
    document.add_paragraph("Delta goes away.")
    document.add_comment(rdocx.RunRange(
        start=rdocx.RunPosition(body_index=1, run_index=0),
        end=rdocx.RunPosition(body_index=1, run_index=1),
    ), author="Ada", text="Review Delta")
    comment = document.comments[0]
    assert comment.anchor_text == "Delta goes away."
    assert isinstance(comment.anchor, rdocx.StoryRunRange)
    assert comment.anchor.start.item.direct_body_index == 1

    with pytest.raises(AttributeError):
        comment.anchor_text = "changed"
    explicit = rdocx.Comment(id=8, author=None, initials=None, date=None,
                             text="detached", parent_id=None, resolved=False,
                             anchor_text=comment.anchor_text, anchor=comment.anchor)
    assert explicit.anchor == comment.anchor
    assert explicit.anchor_text == comment.anchor_text
    preserved = comment.anchor_text
    document.add_comment(comment.anchor, author="Ben", text="Same span")
    assert document.comments[-1].anchor_text == preserved
    assert comment.anchor_text == preserved
    with pytest.raises(rdocx.StaleElementError, match="document revision"):
        document.add_comment(comment.anchor, author="Ben", text="Stale span")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [item.anchor_text for item in reopened.comments] == [preserved, preserved]


@pytest.mark.parametrize("kind", ["body", "table_cell", "header", "footer"])
def test_comment_anchor_typed_related_ranges_are_owned_and_reusable(kind):
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body anchor")
    document.add_table(1, 1).cell(0, 0).text = "cell anchor"
    document.set_header("header anchor")
    document.set_footer("footer anchor")
    item = next(item for item in document.story_items
                if item.story.kind == kind and item.kind == "paragraph")
    span = rdocx.StoryRunRange(
        start=rdocx.StoryRunPosition(item=item, run_index=0),
        end=rdocx.StoryRunPosition(item=item, run_index=1))
    document.add_comment(span, author="Ada", text="Review")
    before = document.to_bytes()
    comment = document.comments[0]
    assert comment.anchor_text == item.text
    assert comment.anchor.start.item.story.kind == kind
    assert comment.anchor.start.item.direct_body_index == (0 if kind == "body" else None)
    assert document.to_bytes() == before
    document.add_comment(comment.anchor, author="Ben", text="Same span")
    assert document.comments[-1].anchor_text == comment.anchor_text


@pytest.mark.parametrize("multi", [False, True])
def test_comment_anchor_block_control_paragraph_endpoints_are_real_typed_ranges(multi):
    import rdocx

    seed = rdocx.Document()
    seed.add_paragraph("seed")
    identity = seed.add_comment(rdocx.RunRange(
        start=rdocx.RunPosition(body_index=0, run_index=0),
        end=rdocx.RunPosition(body_index=0, run_index=1)), author="Ada", text="Review")
    middle = "</w:p><w:p/><w:p>" if multi else ""
    body = (f'<w:sdt><w:sdtPr><w:tag w:val="goog_rdk_test"/></w:sdtPr><w:sdtContent>'
            f'<w:p><w:commentRangeStart w:id="{identity}"/><w:r><w:t>A</w:t></w:r>'
            f'{middle}<w:r><w:t>B</w:t></w:r><w:commentRangeEnd w:id="{identity}"/>'
            f'<w:r><w:commentReference w:id="{identity}"/></w:r></w:p></w:sdtContent></w:sdt>')
    document = _replace_document_body(seed, body)
    before = document.to_bytes()
    comment = document.comments[0]
    assert comment.anchor_text == ("A\n\nB" if multi else "AB")
    assert comment.anchor.start.item.index_path == (0, 0)
    assert comment.anchor.end.item.index_path == (0, 2 if multi else 0)
    assert comment.anchor.start.item.direct_body_index == 0
    assert comment.anchor.start.item.kind == "paragraph"
    assert isinstance(comment.anchor.start.item.xml, bytes)
    assert document.to_bytes() == before
    document.add_comment(comment.anchor, author="Ben", text="Same source")
    assert document.comments[-1].anchor_text == comment.anchor_text

    # Issue289: every path2 endpoint carries the actual paragraph snapshot.
    import xml.etree.ElementTree as ET
    prefixed = _replace_document_body(seed,
        '<w:p><w:r><w:t>Plain paragraph.</w:t></w:r></w:p>' + body)
    original = prefixed.to_bytes()
    for current in [prefixed, rdocx.Document.from_bytes(original)]:
        bounds = current.comments[0].anchor
        for endpoint, path, text in [
            (bounds.start, (1, 0), "A" if multi else "AB"),
            (bounds.end, (1, 2 if multi else 0), "B" if multi else "AB"),
        ]:
            item = endpoint.item
            assert item.story.kind == "body" and item.kind == "paragraph"
            assert item.index_path == path and item.direct_body_index == 1
            assert item.text == text and item.xml
            assert ET.fromstring(item.xml).tag == (
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
        held = current.paragraphs[1]
        position = rdocx.StoryRunPosition(paragraph=held, run_index=0)
        assert position.item == bounds.start.item
        assert position.item.text == ("A" if multi else "AB")
        assert position.item.xml
        assert current.to_bytes() == original
        retired_run = held.runs[0]
        current.add_paragraph("Revision changes")
        survivor = rdocx.StoryRunPosition(paragraph=held, run_index=0)
        assert survivor.item.text == ("A" if multi else "AB")
        assert survivor.item.xml == position.item.xml
        with pytest.raises(rdocx.StaleElementError, match="document revision"):
            _ = retired_run.text


def test_comment_anchor_header_block_control_preserves_nonbody_identity():
    import rdocx

    seed = rdocx.Document()
    seed.add_paragraph("body")
    seed.set_header("header")
    identity = seed.add_comment(rdocx.RunRange(
        start=rdocx.RunPosition(body_index=0, run_index=0),
        end=rdocx.RunPosition(body_index=0, run_index=1)), author="Ada", text="Review")
    header = next(story for story in seed.stories if story.kind == "header")
    body = _document_xml(seed)
    for local in ["commentRangeStart", "commentRangeEnd", "commentReference"]:
        body = body.replace(f'<w:{local} w:id="{identity}"/>'.encode(), b"")
    xml = (f'<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:sdt><w:sdtPr><w:tag w:val="goog_rdk_header"/></w:sdtPr><w:sdtContent>'
           f'<w:p><w:commentRangeStart w:id="{identity}"/><w:r><w:t>HEADER</w:t></w:r>'
           f'<w:commentRangeEnd w:id="{identity}"/><w:r><w:commentReference w:id="{identity}"/>'
           f'</w:r></w:p></w:sdtContent></w:sdt></w:hdr>').encode()
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(seed.to_bytes())) as archive, zipfile.ZipFile(result, "w") as output:
        for info in archive.infolist():
            data = archive.read(info.filename)
            if info.filename == "word/document.xml":
                data = body
            elif info.filename == header.part_name.lstrip("/"):
                data = xml
            output.writestr(info, data)
    document = rdocx.Document.from_bytes(result.getvalue())
    before = document.to_bytes()
    comment = document.comments[0]
    assert comment.anchor_text == "HEADER"
    assert comment.anchor.start.item.index_path == (0, 0)
    assert comment.anchor.start.item.story.kind == "header"
    assert comment.anchor.start.item.story.part_name == header.part_name
    assert comment.anchor.start.item.direct_body_index is None
    assert document.to_bytes() == before
    document.add_comment(comment.anchor, author="Ben", text="Same header")
    assert document.comments[-1].anchor_text == "HEADER"


def test_comment_moves_refuse_unknown_reply_and_invalid_ranges_atomically():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("SOURCE")
    document.add_paragraph("TARGET TARGET")
    identity = document.add_comment_on_text("SOURCE", author="Ada", initials="A",
        text="Review", date="2026-10-08T12:00:00Z")
    reply = document.reply_to(identity, author="Ben", text="Reply")
    document.resolve_comment(identity)
    metadata = document.comments
    target = document.paragraphs[1]
    bounds = rdocx.StoryRunRange(
        start=rdocx.StoryRunPosition(paragraph=target, run_index=0),
        end=rdocx.StoryRunPosition(paragraph=target, run_index=1))
    held = document.paragraphs[0]
    document.move_comment(identity, bounds)
    assert document.comments == metadata
    assert document.comments[0].anchor_text == "TARGET TARGET"
    with pytest.raises(rdocx.StaleElementError):
        held.text
    document.move_comment_to_text(identity, "TARGET", occurrence=1)
    assert document.comments == metadata
    assert document.comments[0].anchor_text == "TARGET"
    before = document.to_bytes()
    current = document.comments[0].anchor
    for selected in [999, reply]:
        with pytest.raises(rdocx.RdocxError):
            document.move_comment(selected, current)
        assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError):
        document.move_comment_to_text(identity, "absent")
    assert document.to_bytes() == before
    with pytest.raises(rdocx.StaleElementError):
        document.move_comment(identity, bounds)
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert reopened.comments == metadata
    assert reopened.comments[0].anchor_text == "TARGET"


def test_scoped_replacement_count_guards_are_atomic():
    import pickle
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Ver").add_run("sion head")
    document.add_paragraph("Version head")
    held = document.paragraphs[0]
    run = held.runs[0]
    before = document.to_bytes()
    for old, new, expect in [("absent", "x", 0), ("", "x", 0)]:
        assert held.replace_text(old, new, expect=expect) == 0
        assert held.text == "Version head"
        assert run.text == "Ver"
        assert document.to_bytes() == before
    with pytest.raises(rdocx.ReplacementCountError) as raised:
        held.replace_text("Version", "Release", expect=2)
    error = raised.value
    assert (error.index, error.expected, error.found) == (None, 2, 1)
    copied = pickle.loads(pickle.dumps(error))
    assert (copied.index, copied.expected, copied.found) == (None, 2, 1)
    assert document.to_bytes() == before
    assert held.text == "Version head"
    with pytest.raises(rdocx.RdocxError):
        held.replace_text("absent", "\x01")
    assert document.to_bytes() == before
    assert held.replace_text(old="Version", new="Release", expect=1) == 1
    for stale in (held, run):
        with pytest.raises(rdocx.StaleElementError):
            stale.text
    assert [p.text for p in document.paragraphs] == ["Release head", "Version head"]
    with pytest.raises(rdocx.StaleElementError):
        held.replace_text("Release", "Again")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [p.text for p in reopened.paragraphs] == ["Release head", "Version head"]


def test_scoped_cell_and_story_item_replacements_have_local_counts():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Version outside")
    table = document.add_table(1, 2)
    table.cell(0, 0).text = "Version cell"
    document.tables[0].cell(0, 1).text = "Version neighbor"
    cell = document.tables[0].cell(0, 0)
    cell.add_paragraph("Version second")
    cell = document.tables[0].cell(0, 0)
    paragraph = cell.paragraphs[1]
    before = document.to_bytes()
    assert cell.replace_text("absent", "x", expect=0) == 0
    assert cell.text == "Version cell\nVersion second"
    with pytest.raises(rdocx.ReplacementCountError) as raised:
        cell.replace_text("Version", "Release", expect=1)
    assert (raised.value.index, raised.value.found) == (None, 2)
    assert document.to_bytes() == before
    assert paragraph.replace_text("Version", "Release", expect=1) == 1
    assert document.tables[0].cell(0, 0).text == "Version cell\nRelease second"
    cell = document.tables[0].cell(0, 0)
    assert cell.replace_text("Version", "Release", expect=1) == 1
    assert document.tables[0].cell(0, 1).text == "Version neighbor"
    assert document.paragraphs[0].text == "Version outside"
    document.set_header("Version header")
    item = next(item for item in document.story_items if item.story.kind == "header" and item.kind == "paragraph")
    held = document.paragraphs[0]
    before = document.to_bytes()
    assert document.replace_text_at(item, "absent", "x", expect=0) == 0
    assert document.to_bytes() == before
    assert held.text == "Version outside"
    assert document.replace_text_at(item, "Version", "Release", expect=1) == 1
    with pytest.raises(rdocx.StaleElementError):
        document.replace_text_at(item, "Release", "Again")
    assert _story_paragraph_texts(document, "header") == ["Release header"]
    assert document.paragraphs[0].text == "Version outside"


def test_scoped_control_paragraph_handles_use_direct_owner_ordinals():
    import rdocx

    body = '<w:sdt><w:sdtPr/><w:sdtContent><w:sdt><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>Version nested</w:t></w:r></w:p></w:sdtContent></w:sdt><w:p><w:r><w:t>Version A</w:t></w:r></w:p><w:p><w:r><w:t>Version B</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    document = _replace_document_body(rdocx.Document(), body)
    nested = rdocx.Document.from_bytes(document.to_bytes())
    nested_before = nested.to_bytes()
    held_nested = nested.paragraphs[0]
    nested_unsupported = False
    try:
        nested.paragraphs[0].replace_text("Version", "Release", expect=1)
    except IndexError:
        nested_unsupported = True
    count = document.paragraphs[1].replace_text("Version", "Release", expect=1)
    texts = [paragraph.text for paragraph in document.paragraphs]
    print(f"count={count}, nested_unsupported={nested_unsupported}, texts={texts!r}")
    assert count == 1
    assert texts == ["Version nested", "Release A", "Version B"]
    assert nested_unsupported
    assert nested.to_bytes() == nested_before
    assert held_nested.text == "Version nested"


@pytest.mark.parametrize("kind", ["header", "footer"])
def test_whole_story_comment_refusals_preserve_bytes_and_revision(kind):
    import rdocx

    def fixture(partial=False, malformed=False):
        document = rdocx.Document()
        document.add_paragraph("main retained")
        getattr(document, f"set_{kind}")("old story")
        item = next(item for item in document.story_items
                    if item.kind == "paragraph" and item.story.kind == kind)
        root = document.add_comment(rdocx.StoryRunRange(
            start=rdocx.StoryRunPosition(item=item, run_index=0),
            end=rdocx.StoryRunPosition(item=item, run_index=1)),
            author="Ada", text="root")
        reply = document.reply_to(root, author="Ben", text="reply")
        document.reply_to(reply, author="Cyd", text="grandchild")
        if not partial and not malformed:
            return rdocx.Document.from_bytes(document.to_bytes())
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as source:
            with zipfile.ZipFile(output, "w") as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if partial:
                        marker = f'<w:commentRangeStart w:id="{root}"/>'.encode()
                        if info.filename.startswith(f"word/{kind}") and info.filename.endswith(".xml"):
                            data = data.replace(marker, b"")
                        if info.filename == "word/document.xml":
                            data = data.replace(b"<w:p>", b"<w:p>" + marker, 1)
                    if malformed and info.filename == "word/commentsExtended.xml":
                        data = re.sub(rb'(w15:paraId=")[^"]+', rb'\g<1>DEADBEEF', data, count=1)
                    target.writestr(info, data)
        return rdocx.Document.from_bytes(output.getvalue())

    document = fixture()
    held = document.paragraphs[0]
    getattr(document, f"set_{kind}")("new story")
    assert document.comments == ()
    with pytest.raises(rdocx.StaleElementError, match="revision 0.*revision 1"):
        held.text
    assert rdocx.Document.from_bytes(document.to_bytes()).comments == ()
    for partial, malformed, text in ((True, False, "new"),
                                     (False, True, "new"),
                                     (False, False, "invalid \x01")):
        document = fixture(partial, malformed)
        held = document.paragraphs[0]
        before = document.to_bytes()
        with pytest.raises(rdocx.RdocxError):
            getattr(document, f"set_{kind}")(text)
        assert document.to_bytes() == before
        assert held.text == "main retained"


def test_scoped_alternate_text_box_refusal_is_rdocxerror_without_mutation():
    import rdocx

    content = '<w:txbxContent><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>box NEEDLE</w:t></w:r></w:p></w:txbxContent>'
    body = (
        '<w:p><w:r><w:t>held outside</w:t></w:r></w:p><w:p><w:r>'
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        '<mc:Choice Requires="wps"><w:drawing><wp:inline><wp:extent cx="914400" cy="457200"/>'
        '<wp:docPr id="1" name="scoped box"/><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:txbx>{content}</wps:txbx></wps:wsp>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape id="fallback"><v:textbox>{content}</v:textbox></v:shape>'
        '</w:pict><!--preserve alternate producer--></mc:Fallback></mc:AlternateContent></w:r></w:p>'
    )
    document = _replace_document_body(rdocx.Document(), body)
    item = next(item for item in document.story_items
                if item.story.kind == "text_box" and item.kind == "paragraph")
    held = document.paragraphs[0]
    before = document.to_bytes()
    assert document.replace_text_at(item, "absent", "PIN", expect=0) == 0
    with pytest.raises(rdocx.ReplacementCountError) as raised:
        document.replace_text_at(item, "NEEDLE", "PIN", expect=2)
    assert (raised.value.expected, raised.value.found) == (2, 1)
    assert document.to_bytes() == before
    assert held.text == "held outside"
    assert document.replace_text_at(item, "NEEDLE", "PIN", expect=1) == 1
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text
    assert document.paragraphs[0].text == "held outside"
    saved = document.to_bytes()
    xml = _document_xml(document)
    assert xml.count(b"box PIN") == 2 and b"NEEDLE" not in xml
    assert b"<!--preserve alternate producer-->" in xml
    reopened = rdocx.Document.from_bytes(saved)
    assert _story_paragraph_texts(reopened, "text_box") == ["box PIN"]
    assert _document_xml(reopened) == xml
    fresh = next(entry for entry in reopened.story_items
                 if entry.story.kind == "text_box" and entry.kind == "paragraph")
    assert reopened.replace_text_at(fresh, "absent", "PIN", expect=0) == 0
    assert reopened.to_bytes() == saved

    # The corresponding source slot is not proved when the fallback text differs.
    asymmetric = body.replace(
        f"<v:textbox>{content}</v:textbox>",
        f"<v:textbox>{content.replace('box NEEDLE', 'different text')}</v:textbox>",
    )
    assert asymmetric != body
    refused = _replace_document_body(rdocx.Document(), asymmetric)
    selected = next(entry for entry in refused.story_items
                    if entry.story.kind == "text_box" and entry.kind == "paragraph")
    original = refused.to_bytes()
    outside = refused.paragraphs[0]
    with pytest.raises(rdocx.RdocxError, match="asymmetric"):
        refused.replace_text_at(selected, "NEEDLE", "PIN", expect=1)
    assert refused.to_bytes() == original and outside.text == "held outside"


def test_rebuild_toc_single_entry_preserves_opaque_properties_and_reopens():
    import rdocx

    body = (
        '<w:p xmlns:f="urn:producer" w:rsidR="00ABCDEF">'
        "<w:pPr f:keep='&#65;'><w:pStyle w:val=\"TOC1\"></w:pStyle>"
        "<f:payload data='&#66;'/></w:pPr>"
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>Old entry</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Only</w:t></w:r></w:p>'
    )
    source = rdocx.Document()
    source.add_paragraph("placeholder")
    document = _replace_document_body(source, body)
    for _ in range(2):
        assert document.rebuild_toc().entry_count == 1
        assert len(document.paragraphs) == 2
        assert document.paragraphs[0].text == "Only\t1"
        saved = _document_xml(document)
        assert b"f:keep='&#65;'" in saved
        assert b"<f:payload data='&#66;'/>" in saved
        assert saved.count(b'w:fldCharType="begin"') == 1
        assert saved.count(b'w:fldCharType="end"') == 1
        document = rdocx.Document.from_bytes(document.to_bytes())
        assert _document_xml(document) == saved


def test_fx201_existing_section_geometry_positive_and_atomic_index():
    import rdocx
    from rdocx import Inches

    document = rdocx.Document()
    document.add_paragraph("section geometry")
    updated = document.update_section(
        0, page_width=Inches(8.5), page_height=Inches(11),
        orientation="landscape", margin_top=Inches(0.75),
        margin_right=Inches(0.8), margin_bottom=Inches(0.9),
        margin_left=Inches(1.1), gutter=Inches(0.2),
        header_distance=Inches(0.3), footer_distance=Inches(0.4),
    )
    assert (updated.page_width, updated.page_height) == (Inches(11), Inches(8.5))
    assert updated.orientation == "landscape"
    assert rdocx.Document.from_bytes(document.to_bytes()).sections[0] == updated
    before = document.to_bytes()
    with pytest.raises(IndexError, match="section index out of range"):
        document.insert_section(2)
    assert document.to_bytes() == before


def test_fx201_inserted_section_inherits_resolved_geometry():
    import rdocx
    from rdocx import Inches

    document = rdocx.Document()
    document.add_paragraph("first section")
    previous = document.update_section(
        0, page_width=Inches(8.5), page_height=Inches(11),
        orientation="landscape", margin_top=Inches(0.75),
        margin_right=Inches(0.8), margin_bottom=Inches(0.9),
        margin_left=Inches(1.1), gutter=Inches(0.2),
        header_distance=Inches(0.3), footer_distance=Inches(0.4),
        column_count=2, page_number_start=7, different_first_page=True,
    )
    assert previous.page_width == Inches(11)
    assert previous.orientation == "landscape"
    document.insert_section(1)
    assert len(document.sections) == 2
    inserted = document.sections[1]
    fields = ("page_width", "page_height", "orientation", "margin_top",
              "margin_right", "margin_bottom", "margin_left", "gutter",
              "header_distance", "footer_distance")
    assert tuple(getattr(inserted, f) for f in fields) == tuple(
        getattr(previous, f) for f in fields
    ), "F201 inserted section must inherit resolved geometry"
    assert inserted.column_count is None
    assert inserted.page_number_start is None
    assert not inserted.different_first_page
    assert rdocx.Document.from_bytes(document.to_bytes()).sections[1] == inserted


def test_fx201_even_footer_creation_enables_setting_atomically():
    import rdocx
    from xml.etree import ElementTree as ET

    document = rdocx.Document()
    document.add_paragraph("section footer")
    default = document.create_section_story(0, "footer", "default")
    assert default.kind == "footer"
    assert _story_part(document, default)
    before = document.to_bytes()
    with pytest.raises(IndexError, match="section index out of range"):
        document.create_section_story(9, "footer", "even")
    assert document.to_bytes() == before
    even = document.create_section_story(0, "footer", "even")
    assert even.kind == "footer"
    assert _story_part(document, even)
    variant = next(v for v in document.header_footer_variants
                   if v.section_index == 0 and v.kind == "footer" and v.variant == "even")
    assert variant.story == even and not variant.inherited
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        settings = ET.fromstring(archive.read("word/settings.xml"))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    enabled = settings.find(word + "evenAndOddHeaders")
    assert enabled is not None and enabled.get(word + "val", "1") not in (
        "0", "false", "off"
    ), "F201 even footer creation must enable evenAndOddHeaders"
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert any(v.section_index == 0 and v.variant == "even" and v.story is not None
               for v in reopened.header_footer_variants)


def test_fx201_twips_constructor_and_section_length_equivalence():
    import rdocx

    assert rdocx.Twips(1440) == rdocx.Inches(1)
    assert rdocx.Twips(1.75) == 1111
    assert rdocx.Twips(-1.75) == -1111
    document = rdocx.Document()
    document.add_paragraph("length control")
    section = document.update_section(0, margin_top=rdocx.Twips(1440))
    assert section.margin_top == rdocx.Pt(72)
    before = document.to_bytes()
    for bad in (True, 1.5, 2**64):
        with pytest.raises((TypeError, OverflowError, ValueError, rdocx.RdocxError)):
            document.update_section(0, margin_top=bad)
        assert document.to_bytes() == before
    assert document.update_section(0, margin_top=rdocx.Emu(-635)).margin_top == -635
    assert document.update_section(0, margin_top=rdocx.Emu(-1)).margin_top == 0
    assert document.update_section(0, margin_top=rdocx.Emu(634)).margin_top == 0
    before = document.to_bytes()
    for name in ("page_width", "page_height", "margin_top", "margin_right", "margin_bottom", "margin_left", "gutter", "header_distance", "footer_distance"):
        with pytest.raises(TypeError):
            setattr(document.section(0), name, None)
        assert document.to_bytes() == before
    with pytest.raises(TypeError):
        document.section(0).orientation = None
    assert document.to_bytes() == before


def test_fx201_live_section_footer_uses_actual_related_owner():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body retained")
    section = document.section(0)
    assert section.footer.paragraphs == ()
    section.footer.is_linked_to_previous = False
    footer = document.section(0).footer
    paragraph = footer.add_paragraph("Page ")
    paragraph.runs[0].font.bold = True
    paragraph.add_run("1").add_field("PAGE", "1")
    footer = document.section(0).footer
    assert footer.paragraphs[0].text == "Page 11"
    assert document.paragraphs[0].text == "body retained"
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        footers = [n for n in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", n)]
        assert len(footers) == 1
        assert b"PAGE" in archive.read(footers[0])
        assert b"Page " not in archive.read("word/document.xml")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.section(0).footer.paragraphs[0].text == "Page 11"


def test_fx201_note_reference_creation_uses_checked_paragraph():
    import rdocx

    document = rdocx.Document()
    reference = document.add_paragraph("reference")
    retired_run = reference.runs[0]
    note = document.add_footnote(reference, "note content")
    assert note.paragraphs[0].text == "note content"
    assert reference.text == "reference"
    with pytest.raises(rdocx.StaleElementError):
        _ = retired_run.text
    saved = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        assert b"note content" in archive.read("word/footnotes.xml")
        assert b"footnoteReference" in archive.read("word/document.xml")
    note.delete()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert b"footnoteReference" not in archive.read("word/document.xml")
    before = document.to_bytes()
    foreign = rdocx.Document().add_paragraph("foreign")
    with pytest.raises(ValueError, match="different document"):
        document.add_footnote(foreign, "refused")
    assert document.to_bytes() == before


def test_fx201_settings_and_run_break_are_owned():
    import rdocx
    from xml.etree import ElementTree as ET

    document = rdocx.Document()
    document.add_paragraph("break before").runs[0].add_break("page")
    document.settings.odd_and_even_pages_header_footer = True
    assert document.settings.odd_and_even_pages_header_footer
    saved = document.to_bytes()
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        main = ET.fromstring(archive.read("word/document.xml"))
        assert any(e.get(word + "type") == "page" for e in main.iter(word + "br"))
        assert ET.fromstring(archive.read("word/settings.xml")).find(word + "evenAndOddHeaders") is not None
    reopened = rdocx.Document.from_bytes(saved)
    assert reopened.settings.odd_and_even_pages_header_footer


def test_fx201_note_policy_scopes_preserve_notes_and_refuse_atomically():
    import rdocx
    from xml.etree import ElementTree as ET

    document = rdocx.Document()
    reference = document.add_paragraph("note reference")
    document.add_footnote(reference, "retained footnote")
    document.set_note_policy("footnote", format="lowerRoman", start=4, restart="eachPage", placement="beneathText")
    section = document.section(0)
    section.set_note_policy("endnote", format="upperLetter", start=2, restart="eachSect", placement="sectEnd")
    saved = document.to_bytes()
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        notes = archive.read("word/footnotes.xml")
        assert b"retained footnote" in notes
        settings = ET.fromstring(archive.read("word/settings.xml"))
        policy = settings.find(word + "footnotePr")
        assert policy is not None
        assert policy.find(word + "numFmt").get(word + "val") == "lowerRoman"
        assert policy.find(word + "numStart").get(word + "val") == "4"
        assert policy.find(word + "numRestart").get(word + "val") == "eachPage"
        assert policy.find(word + "pos").get(word + "val") == "beneathText"
        main = ET.fromstring(archive.read("word/document.xml"))
        override = main.find(".//" + word + "sectPr/" + word + "endnotePr")
        assert override is not None
        assert override.find(word + "numFmt").get(word + "val") == "upperLetter"
        assert override.find(word + "numStart").get(word + "val") == "2"
    for bad in (True, 1.5, -1, 2**64, 0):
        with pytest.raises((TypeError, ValueError, OverflowError, rdocx.RdocxError)):
            document.set_note_policy("footnote", start=bad)
        assert document.to_bytes() == saved
    with pytest.raises(rdocx.RdocxError, match="placement is invalid"):
        section.set_note_policy("footnote", placement="docEnd")
    assert document.to_bytes() == saved
    reopened = rdocx.Document.from_bytes(saved)
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as archive:
        assert archive.read("word/footnotes.xml") == notes


def test_fx201_nested_tables_use_actual_body_and_footer_cell_owners():
    import rdocx
    from xml.etree import ElementTree as ET

    document = rdocx.Document()
    document.add_paragraph("body retained")
    table = document.add_table(1, 1)
    nested = table.cell(0, 0).add_table(1, 1)
    nested.cell(0, 0).text = "nested body"
    assert document.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "nested body"
    footer = document.section(0).footer
    table = footer.add_table(1, 1)
    nested = table.cell(0, 0).add_table(1, 1)
    nested.cell(0, 0).text = "nested footer"
    before = document.to_bytes()
    cell = document.section(0).footer.tables[0].cell(0, 0)
    with pytest.raises(rdocx.RdocxError, match="dimensions must be positive"):
        cell.add_table(0, 1)
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert reopened.paragraphs[0].text == "body retained"
    assert reopened.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "nested body"
    assert reopened.section(0).footer.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "nested footer"
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(io.BytesIO(before)) as archive:
        footers = [n for n in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", n)]
        assert len(footers) == 1
        for name in ("word/document.xml", footers[0]):
            for cell in ET.fromstring(archive.read(name)).iter(word + "tc"):
                assert list(cell)[-1].tag == word + "p"


def test_fx201_live_page_border_edges_and_page_color_preserve_source():
    import rdocx
    from rdocx import Document
    document = Document()
    document.add_paragraph("page body remains")
    section = document.section(0)
    for edge in ("top", "left", "bottom", "right"):
        section.set_page_border(edge, style="double", width=16, space=4, color="123456")
    document.set_page_color("AABBCC")
    saved = document.to_bytes()
    from lxml import etree
    import zipfile
    import io
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(saved)).read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert root.find("w:background", ns).get("{%s}color" % ns["w"]) == "AABBCC"
    borders = root.find("w:body/w:sectPr/w:pgBorders", ns)
    assert borders is not None
    assert [etree.QName(child).localname for child in borders] == ["top", "left", "bottom", "right"]
    for border in borders:
        assert border.get("{%s}val" % ns["w"]) == "double"
        assert border.get("{%s}sz" % ns["w"]) == "16"
    section = document.section(0)
    with pytest.raises(rdocx.RdocxError, match="page border width"):
        section.set_page_border("top", width=1)
    assert document.to_bytes() == saved
    assert Document.from_bytes(saved).to_bytes() == saved


def test_fx201_section_drawingml_text_watermark_preserves_real_text_and_refusals():
    import rdocx
    from rdocx import Document
    import io
    import zipfile
    from lxml import etree
    document = Document()
    document.add_paragraph("body stays")
    section = document.section(0)
    section.header.add_paragraph("other header")
    section = document.section(0)
    section.set_text_watermark("CONFIDENTIAL", width=4572000, height=1143000,
                               font_size=609600, font_family="Carlito", opacity=0.5)
    saved = document.to_bytes()
    package = zipfile.ZipFile(io.BytesIO(saved))
    header = next(name for name in package.namelist() if name.startswith("word/header") and name.endswith(".xml"))
    root = etree.fromstring(package.read(header))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
          "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
          "v": "urn:schemas-microsoft-com:vml"}
    assert root.xpath("count(.//w:t[text()='CONFIDENTIAL'])", namespaces=ns) == 1
    assert root.xpath("count(.//w14:alpha[@w14:val='50000'])", namespaces=ns) == 1
    assert root.xpath("count(.//wp:anchor[@behindDoc='1'])", namespaces=ns) == 1
    choices = root.xpath(".//mc:AlternateContent/mc:Choice", namespaces=ns)
    assert [choice.get("Requires") for choice in choices] == ["wps w16"]
    assert choices[0].nsmap["w16"] == "http://schemas.microsoft.com/office/word/2018/wordml"
    assert choices[0].xpath("count(.//w14:solidFill)", namespaces=ns) == 1
    assert root.xpath(".//mc:Fallback//v:textpath/@string", namespaces=ns) == ["CONFIDENTIAL"]
    assert root.xpath(".//mc:Fallback//v:fill/@opacity", namespaces=ns) == ["0.5"]
    assert root.xpath("count(.//mc:Fallback//v:shapetype[@o:spt='136'])", namespaces={**ns, "o": "urn:schemas-microsoft-com:office:office"}) == 1
    assert "other header" in "".join(root.xpath(".//w:t/text()", namespaces=ns))
    assert Document.from_bytes(saved).to_bytes() == saved
    for unsupported in (0.0, 0.25, 1.0):
        with pytest.raises(rdocx.RdocxError, match="^DrawingML text watermark opacity other than 0\\.5 is unsupported by the qualified native text-fill representation$"):
            section.set_text_watermark("bad", width=4572000, height=1143000, font_size=609600, opacity=unsupported)
        assert document.to_bytes() == saved
        absent = Document()
        absent.add_paragraph("no header")
        absent_saved = absent.to_bytes()
        with pytest.raises(rdocx.RdocxError, match="^DrawingML text watermark opacity other than 0\\.5 is unsupported by the qualified native text-fill representation$"):
            absent.section(0).set_text_watermark("bad", width=4572000, height=1143000, font_size=609600, opacity=unsupported)
        assert absent.to_bytes() == absent_saved
    for invalid in (float("nan"), -0.1, 1.1):
        with pytest.raises(rdocx.RdocxError, match="watermark opacity"):
            section.set_text_watermark("bad", width=4572000, height=1143000, font_size=609600, opacity=invalid)
        assert document.to_bytes() == saved
    with pytest.raises(TypeError):
        section.set_text_watermark("bad", width=True, height=1143000, font_size=609600)
    assert document.to_bytes() == saved
    section.set_text_watermark("REPLACED", width=4572000, height=1143000, font_size=609600)
    replaced = document.to_bytes()
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(replaced)).read(header))
    assert root.xpath("count(.//w:t[text()='REPLACED'])", namespaces=ns) == 1
    assert not root.xpath(".//w:t[text()='CONFIDENTIAL']", namespaces=ns)


@pytest.mark.parametrize("style_mode", ("standard", "localized", "absent"))
def test_fx201_all_six_story_slots_resolve_real_paragraph_style_ids(style_mode):
    import rdocx
    from rdocx import Document
    from lxml import etree
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document = Document()
    document.add_paragraph("body unchanged")
    source = io.BytesIO(document.to_bytes())
    result = io.BytesIO()
    header_id = "Kopfzeile" if style_mode == "localized" else "Header"
    footer_id = "PiedDePage" if style_mode == "localized" else "Footer"
    title_id = "TitreLocal" if style_mode == "localized" else "Title"
    title_name = "Titre local" if style_mode == "localized" else "Title"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(result, "w") as changed:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "word/styles.xml":
                root = etree.fromstring(data)
                for style in list(root):
                    if style.tag == word + "style" and style.get(word + "styleId") in ("Header", "Footer", "Title"):
                        root.remove(style)
                values = [(title_id, title_name)]
                if style_mode != "absent":
                    values.extend(((header_id, "Header"), (footer_id, "Footer")))
                for style_id, name in values:
                    style = etree.SubElement(root, word + "style", {word + "type": "paragraph", word + "styleId": style_id})
                    etree.SubElement(style, word + "name", {word + "val": name})
                data = etree.tostring(root)
            changed.writestr(info, data)
    document = Document.from_bytes(result.getvalue())
    slots = ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer")
    for slot in slots:
        owner = getattr(document.section(0), slot)
        assert owner.paragraphs == ()
        owner.add_paragraph(slot + " default")
        owner = getattr(document.section(0), slot)
        owner.add_paragraph(slot + " explicit", style=title_name)
        saved = document.to_bytes()
        owner = getattr(document.section(0), slot)
        with pytest.raises(KeyError, match="missing paragraph style"):
            owner.add_paragraph("not published", style="missing paragraph style")
        assert document.to_bytes() == saved
    saved = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        all_paragraphs = []
        for name in archive.namelist():
            if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                all_paragraphs.extend(etree.fromstring(archive.read(name)).iter(word + "p"))
        for slot in slots:
            for suffix in ("default", "explicit"):
                expected = slot + " " + suffix
                matches = [p for p in all_paragraphs if "".join(p.itertext()) == expected]
                assert len(matches) == 1
                style = matches[0].find(word + "pPr/" + word + "pStyle")
                expected_id = title_id if suffix == "explicit" else None if style_mode == "absent" else header_id if "header" in slot else footer_id
                assert (None if style is None else style.get(word + "val")) == expected_id
        assert b"body unchanged" in archive.read("word/document.xml")
    reopened = Document.from_bytes(saved)
    assert reopened.to_bytes() == saved
    for slot in slots:
        texts = [p.text for p in getattr(reopened.section(0), slot).paragraphs]
        assert slot + " default" in texts and slot + " explicit" in texts


@pytest.mark.parametrize("host", ("footer", "footer-cell", "footnote", "endnote", "body-cell"))
def test_fx201_owned_rich_picture_is_inserted_with_its_owner_relationship(host):
    import rdocx
    from rdocx import Document
    from lxml import etree
    document = Document()
    document.add_paragraph("body retained")
    if host == "footer":
        paragraph = document.section(0).footer.add_paragraph("owner text")
    elif host == "footer-cell":
        table = document.section(0).footer.add_table(1, 1)
        table.cell(0, 0).text = "owner text"
        paragraph = document.section(0).footer.tables[0].cell(0, 0).paragraphs[0]
    elif host == "body-cell":
        table = document.add_table(1, 1)
        table.cell(0, 0).text = "owner text"
        paragraph = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        reference = document.paragraphs[0].runs[0]
        note = getattr(document, "add_" + host)(reference, "owner text")
        paragraph = note.paragraphs[0]
    run = paragraph.runs[0]
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="image"):
        run.add_picture(b"not an image", "invalid.png", width=127000, height=127000)
    assert document.to_bytes() == before
    image = _one_pixel_png(b"\xff\x00\x00")
    run.add_picture(image, "owner.png", width=127000, height=127000)
    with pytest.raises(rdocx.StaleElementError):
        _ = run.text
    saved = document.to_bytes()
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    relationship_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        owner = "word/document.xml" if host == "body-cell" else "word/footnotes.xml" if host == "footnote" else "word/endnotes.xml" if host == "endnote" else next(n for n in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", n))
        root = etree.fromstring(archive.read(owner))
        embeds = root.xpath(".//a:blip/@r:embed", namespaces=ns)
        assert len(embeds) == 1
        rel_name = posixpath.dirname(owner) + "/_rels/" + posixpath.basename(owner) + ".rels"
        rels = etree.fromstring(archive.read(rel_name))
        matching = [rel for rel in rels if rel.tag == relationship_ns + "Relationship" and rel.get("Id") == embeds[0]]
        assert len(matching) == 1 and matching[0].get("Type").endswith("/image")
        target = posixpath.normpath(posixpath.join(posixpath.dirname(owner), matching[0].get("Target")))
        assert archive.read(target) == image
        assert b"owner text" in archive.read(owner)
        if host != "body-cell":
            assert b"owner text" not in archive.read("word/document.xml")
        assert b"body retained" in archive.read("word/document.xml")
    assert Document.from_bytes(saved).to_bytes() == saved


def test_fx201_rich_footer_keeps_page_and_total_fields_separate():
    import rdocx
    from rdocx import Document
    from lxml import etree
    document = Document()
    document.add_paragraph("body retained")
    document.section(0).footer.add_paragraph("Page ")
    paragraph = document.section(0).footer.paragraphs[0]
    paragraph.runs[0].font.bold = True
    paragraph.add_run("").add_field("PAGE", "1")
    paragraph = document.section(0).footer.paragraphs[0]
    paragraph.add_run(" of ")
    paragraph = document.section(0).footer.paragraphs[0]
    paragraph.add_run("").add_field("NUMPAGES", "1")
    paragraph = document.section(0).footer.paragraphs[0]
    paragraph.add_run(" logo").add_picture(_one_pixel_png(b"\xff\x00\x00"), "footer-logo.png", width=127000, height=127000)
    document.section(0).footer.add_table(1, 1).cell(0, 0).text = "footer table"
    document.add_footnote(document.paragraphs[0].runs[0], "footnote retained")
    document.add_endnote(document.paragraphs[0].runs[0], "endnote retained")
    saved = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        footer = next(n for n in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", n))
        word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        root = etree.fromstring(archive.read(footer))
        instructions = [node.get(word + "instr").strip() for node in root.iter(word + "fldSimple")]
        assert instructions == ["PAGE", "NUMPAGES"]
        assert ["".join(node.itertext()) for node in root.iter(word + "fldSimple")] == ["1", "1"]
        assert len(list(root.iter(word + "r"))) >= 7
        assert root.find(".//" + word + "rPr/" + word + "b") is not None
        assert b"PAGE" not in archive.read("word/document.xml")
    assert Document.from_bytes(saved).to_bytes() == saved
    svg = document.render_page_to_svg(0)
    assert svg is not None and "Page" in svg.svg and "footer" in svg.svg and "table" in svg.svg
    import os
    directory = os.environ.get("RDOCX_FX201_NATIVE_OUTPUT_DIR")
    if directory:
        from pathlib import Path
        target = Path(directory)
        target.mkdir(parents=True, exist_ok=True)
        (target / "installed-rich-footer-notes.docx").write_bytes(saved)
        (target / "installed-rich-footer-notes.svg").write_text(svg.svg)


def test_fx201_image_watermark_repeat_removal_and_invalid_opacity_are_owned():
    import rdocx
    from rdocx import Document
    document = Document()
    document.add_paragraph("body retained")
    document.section(0).header.add_paragraph("unrelated header")
    section = document.section(0)
    image = _one_pixel_png()
    for opacity in (0.5, 0.0, 1.0):
        section.set_image_watermark(image, "mark.png", width=1524000, height=762000, rotation=30.0, opacity=opacity)
        saved = document.to_bytes()
        with zipfile.ZipFile(io.BytesIO(saved)) as archive:
            header = next(n for n in archive.namelist() if re.fullmatch(r"word/header\d+\.xml", n))
            xml = archive.read(header)
            assert b"unrelated header" in xml
            assert f'amt="{round(opacity * 100000)}"'.encode() in xml
            assert b'rot="1800000"' in xml
            assert len([n for n in archive.namelist() if n.startswith("word/media/")]) == 1
        assert Document.from_bytes(saved).to_bytes() == saved
    for opacity in (float("nan"), -0.1, 1.1):
        with pytest.raises(rdocx.RdocxError, match="geometry or opacity"):
            section.set_image_watermark(image, "bad.png", width=1524000, height=762000, opacity=opacity)
        assert document.to_bytes() == saved
    section.remove_watermark()
    removed = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(removed)) as archive:
        assert b"unrelated header" in archive.read(header)
        assert b"RDOCX DrawingML Watermark" not in archive.read(header)
        assert not [n for n in archive.namelist() if n.startswith("word/media/")]
    section.remove_watermark()
    assert document.to_bytes() == removed


def test_fx201_all_six_story_slots_keep_inherited_and_unlinked_physical_owners():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("body retained")
    document.insert_section(1)
    document.insert_section(2)
    slots = ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer")
    for slot in slots:
        getattr(document.section(0), slot).add_paragraph(slot + " original")
        inherited = getattr(document.section(1), slot)
        assert inherited.is_linked_to_previous
        assert slot + " original" in [p.text for p in inherited.paragraphs]
        inherited.add_paragraph(slot + " shared")
        assert slot + " shared" in [p.text for p in getattr(document.section(0), slot).paragraphs]
        held = getattr(document.section(0), slot).paragraphs[0]
        local = getattr(document.section(1), slot)
        local.is_linked_to_previous = False
        assert not getattr(document.section(1), slot).is_linked_to_previous
        getattr(document.section(1), slot).add_paragraph(slot + " local")
        assert slot + " local" not in [p.text for p in getattr(document.section(0), slot).paragraphs]
        assert slot + " local" in [p.text for p in getattr(document.section(2), slot).paragraphs]
        with pytest.raises(rdocx.StaleElementError):
            _ = held.text
        local = getattr(document.section(1), slot)
        local.is_linked_to_previous = True
        assert getattr(document.section(1), slot).is_linked_to_previous
        assert slot + " local" not in [p.text for p in getattr(document.section(2), slot).paragraphs]
    saved = document.to_bytes()
    reopened = rdocx.Document.from_bytes(saved)
    assert reopened.to_bytes() == saved
    for slot in slots:
        assert slot + " shared" in [p.text for p in getattr(reopened.section(2), slot).paragraphs]
    assert reopened.settings.odd_and_even_pages_header_footer


@pytest.mark.parametrize("family", ("footnote", "endnote"))
def test_fx201_checked_note_cell_boundary_preserves_other_owners_and_refuses_stale_and_header(family):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("body retained")
    table = document.add_table(1, 1)
    table.cell(0, 0).text = "cell reference"
    reference = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
    note = getattr(document, "add_" + family)(reference, family + " content")
    assert note.paragraphs[0].text == family + " content"
    before = document.to_bytes()
    with pytest.raises(rdocx.StaleElementError):
        getattr(document, "add_" + family)(reference, "stale refused")
    assert document.to_bytes() == before
    note.delete()
    assert family.encode() + b"Reference" not in _document_xml(document)
    header = document.section(0).header.add_paragraph("not a note host")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="body or body-table"):
        getattr(document, "add_" + family)(header, "header refused")
    assert document.to_bytes() == before
    foreign = rdocx.Document().add_paragraph("foreign")
    with pytest.raises(ValueError, match="different document"):
        getattr(document, "add_" + family)(foreign, "foreign refused")
    assert document.to_bytes() == before
    assert rdocx.Document.from_bytes(before).to_bytes() == before


@pytest.mark.parametrize("operation", [
    "to_html", "to_markdown", "text", "text_report", "to_odt", "to_rtf", "to_epub",
    "word_count", "character_count", "page_count", "validate", "content_controls",
])
def test_fx203_installed_document_operation_has_deciding_result(operation):
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Alpha beta")
    result = getattr(document, operation)()
    if operation in ("to_html", "to_markdown", "text"):
        assert "Alpha beta" in result
    elif operation == "text_report":
        assert "Alpha beta" in result.text
        assert result.revision_view == "accepted"
        assert isinstance(result.diagnostics, tuple)
    elif operation in ("to_odt", "to_rtf", "to_epub"):
        assert isinstance(result.bytes, bytes) and result.bytes
        assert isinstance(result.diagnostics, tuple)
    elif operation == "word_count":
        assert result == 2
    elif operation == "character_count":
        assert result == 10
    elif operation == "page_count":
        assert result == 1
    elif operation == "validate":
        assert result.is_valid
        assert isinstance(result.findings, tuple)
    else:
        assert result == ()


def test_fx203_installed_template_and_checked_insertion_have_real_results():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Hello {{ name }}")
    assert document.render_template({"name": "Ada"}) == 1
    assert document.paragraphs[0].text == "Hello Ada"
    source = rdocx.Document()
    source.add_paragraph("Imported")
    before = source.to_bytes()
    document.insert_document(source, at=1)
    assert [p.text for p in document.paragraphs] == ["Hello Ada", "Imported"]
    assert source.to_bytes() == before


def test_fx203_installed_properties_and_tracking_are_live_and_typed():
    import rdocx

    document = rdocx.Document()
    document.custom_properties.set("Approved", True)
    assert document.custom_properties.get("Approved").value is True
    assert document.custom_properties.get("Approved").type == "bool"
    document.app_properties.company = "Example"
    assert document.app_properties.company == "Example"
    document.settings.track_revisions = True
    assert document.settings.track_revisions is True
    document.settings.track_revisions = None
    assert document.settings.track_revisions is None


def test_fx203_installed_equation_snapshot_keeps_conversion_boundary():
    import rdocx

    document = rdocx.Document()
    paragraph = document.add_paragraph("")
    inserted = paragraph.add_equation("x+1", format="latex")
    assert isinstance(inserted.diagnostics, tuple)
    equations = document.paragraphs[0].equations()
    assert len(equations) == 1
    result = equations[0].to_latex()
    assert result.is_display is False
    assert len(result.equations) == 1
    assert "x" in result.equations[0].value


def test_fx203_installed_settings_tracking_has_presence_and_live_owner():
    from rdocx import Document
    document = Document()
    paragraph = document.add_paragraph("kept")
    settings = document.settings
    assert settings.track_revisions is None
    settings.track_revisions = False
    assert settings.track_revisions is False
    settings.track_revisions = True
    assert document.settings.track_revisions is True
    settings.track_revisions = None
    assert settings.track_revisions is None
    assert paragraph.text == "kept"


def test_fx203_installed_app_properties_are_live_and_counts_are_metadata():
    from rdocx import Document
    document = Document()
    document.add_paragraph("one two")
    properties = document.app_properties
    properties.company = "Example"
    properties.manager = "Owner"
    assert document.app_properties.company == "Example"
    assert document.app_properties.manager == "Owner"
    assert document.word_count() == 2


def test_fx203_installed_module_validation_does_not_require_open_document():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("visible")
    before = document.to_bytes()
    report = rdocx.validate_bytes(before)
    assert report.is_valid
    assert isinstance(report.findings, tuple)
    assert document.to_bytes() == before


def test_fx203_installed_fragment_policy_exposes_three_independent_toggles():
    import rdocx
    policy = rdocx.FragmentConflictPolicy.reuse_equivalent()
    assert policy.with_style_reuse(False) is not policy
    assert policy.with_numbering_reuse(False) is not policy
    assert policy.with_related_part_reuse(False) is not policy
    assert rdocx.FragmentConflictPolicy.rename_all() is not policy


def test_fx203_installed_content_control_setter_owns_qualified_text():
    from rdocx import Document
    document = _replace_document_body(Document(), '<w:sdt><w:sdtPr><w:tag w:val="customer"/><w:alias w:val="Name"/><w:text/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>old</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    assert document.set_content_control_value_by_tag("customer", "new") == 1
    assert document.set_content_control_value_by_alias("Name", "final") == 1
    assert document.content_controls()[0].text == "final"


def test_fx203_custom_property_snapshots_dates_bounds_and_atomic_noops():
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph("live")
    properties = document.custom_properties
    assert properties.get("missing") is None
    before = document.to_bytes()
    assert properties.delete("missing") is False
    assert document.to_bytes() == before
    properties.set("Approved", True)
    original = properties.get("Approved")
    assert original.type == "bool" and original.value is True
    properties.set("Approved", False)
    assert original.value is True
    assert properties.get("Approved").pid == original.pid
    assert properties.get("Approved").fmtid == original.fmtid
    date = "2024-02-29T23:59:59.123+05:30"
    properties.set("When", rdocx.CustomFileTime(date))
    assert properties.get("When").type == "filetime"
    assert properties.get("When").value == date
    for date in ["2023-02-29T00:00:00Z", "2024-01-01T00:00:60Z", "2024-01-01T00:00:00", "0000-01-01T00:00:00Z"]:
        before = document.to_bytes()
        with pytest.raises(ValueError):
            properties.set("When", rdocx.CustomFileTime(date))
        assert document.to_bytes() == before
    for value, error in [(2**31, OverflowError), (-2**31 - 1, OverflowError), (float("inf"), ValueError), (float("nan"), ValueError), (object(), TypeError), (None, TypeError)]:
        before = document.to_bytes()
        with pytest.raises(error):
            properties.set("Bad", value)
        assert document.to_bytes() == before
    before = document.to_bytes()
    properties.set("Approved", False)
    assert document.to_bytes() == before
    assert paragraph.text == "live"
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.custom_properties.get("When").value == "2024-02-29T23:59:59.123+05:30"
    assert properties.delete("Approved") is True
    assert properties.get("Approved") is None


def test_fx203_app_property_optional_latest_owner_and_failure_atomicity():
    from rdocx import Document, RdocxError
    document = Document()
    paragraph = document.add_paragraph("live")
    first = document.app_properties
    second = document.app_properties
    first.company = "Example"
    second.manager = "Owner"
    assert first.manager == "Owner" and second.company == "Example"
    first.template = ""
    assert second.template == ""
    before = document.to_bytes()
    second.company = "Example"
    assert document.to_bytes() == before
    with pytest.raises(RdocxError, match="application version"):
        first.application_version = "invalid"
    assert document.to_bytes() == before
    first.company = None
    assert second.company is None
    assert paragraph.text == "live"
    reopened = Document.from_bytes(document.to_bytes())
    assert reopened.app_properties.manager == "Owner"
    assert reopened.app_properties.template == ""


def test_fx203_fragment_snapshot_physical_boundaries_and_empty_source_validation():
    import rdocx
    source = rdocx.Document()
    source.add_paragraph("first")
    source.add_paragraph("second")
    body = next(story for story in source.stories if story.kind == "body")
    start = rdocx.ContentLocation(body, "paragraph", [0])
    end = rdocx.ContentLocation(body, "paragraph", [1])
    fragment = rdocx.DocumentFragment.from_range(source, start, end)
    source.paragraphs[0].text = "changed"
    destination = rdocx.Document()
    destination.add_paragraph("kept")
    body = next(story for story in destination.stories if story.kind == "body")
    destination.import_fragment(rdocx.ContentLocation.end(body), fragment)
    assert [p.text for p in destination.paragraphs] == ["kept", "first"]
    before = destination.to_bytes()
    empty = rdocx.Document()
    with pytest.raises(IndexError):
        destination.insert_document(empty, at=99)
    assert destination.to_bytes() == before
    with pytest.raises(OverflowError):
        destination.insert_document(source, at=-1)
    assert destination.to_bytes() == before
    with pytest.raises(TypeError):
        destination.insert_document(source, at=True)
    assert destination.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert [p.text for p in reopened.paragraphs] == ["kept", "first"]


def test_fx203_content_control_snapshots_atomic_noop_and_duplicate_inventory():
    from rdocx import Document
    xml = '<w:sdt><w:sdtPr><w:tag w:val="same"/><w:alias w:val="Name"/><w:text/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>old</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    document = _replace_document_body(Document(), xml + xml)
    controls = document.content_controls()
    assert isinstance(controls, tuple) and len(controls) == 2
    before = document.to_bytes()
    with pytest.raises(KeyError, match="absent"):
        document.set_content_control_value_by_tag("absent", "new")
    assert document.to_bytes() == before
    assert document.set_content_control_value_by_tag("same", "old") == 2
    assert document.to_bytes() == before
    assert document.set_content_control_value_by_alias("Name", "new") == 2
    assert [control.text for control in controls] == ["old", "old"]
    assert [control.text for control in document.content_controls()] == ["new", "new"]
    assert [control.text for control in Document.from_bytes(document.to_bytes()).content_controls()] == ["new", "new"]


def test_fx203_equation_display_children_losses_and_source_snapshot_are_ordered():
    import rdocx
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    one = '<m:oMath f:keep="yes"><m:r><m:t>x</m:t></m:r><f:opaque/></m:oMath>'
    body = f'<w:p><m:oMathPara xmlns:m="{ns}" xmlns:f="urn:foreign" f:keep="yes"><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>{one}{one}<m:oMath><m:r><m:t>y</m:t></m:r></m:oMath></m:oMathPara></w:p>'
    document = _replace_document_body(rdocx.Document(), body)
    before = document.to_bytes()
    equations = document.paragraphs[0].equations()
    assert isinstance(equations, tuple) and len(equations) == 1
    equation = equations[0]
    assert equation.is_display is True
    for conversion in [equation.to_latex(), equation.to_mathml()]:
        assert conversion.is_display is True
        assert isinstance(conversion.equations, tuple) and len(conversion.equations) == 3
        assert conversion.equations[0].value == conversion.equations[1].value
        assert "x" in conversion.equations[0].value and "y" in conversion.equations[2].value
        assert conversion.equations[0].diagnostics and conversion.equations[1].diagnostics
        assert any("wrapper" in loss.message for loss in conversion.diagnostics)
        assert any("justification" in loss.message for loss in conversion.diagnostics)
    assert document.to_bytes() == before
    document.paragraphs[0].add_equation("z+1", format="latex")
    assert len(equation.to_latex().equations) == 3
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [eq.is_display for eq in reopened.paragraphs[0].equations()] == [True, False]
    before = document.to_bytes()
    with pytest.raises(ValueError, match="format"):
        document.paragraphs[0].add_equation("bad", format="unknown")
    assert document.to_bytes() == before


def test_fx203_checkbox_font_glyph_and_opaque_owner_are_atomic_and_repeatable():
    import rdocx
    body = '<w:sdt><w:sdtPr><w:tag w:val="check"/><x:checkbox xmlns:x="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:f="urn:foreign"><f:checked f:val="opaque"/><x:checked x:val="0" f:keep="yes"/><x:checkedState x:val="0041" x:font="Carlito"/><x:uncheckedState x:val="0042" x:font="Selawik"/></x:checkbox></w:sdtPr><w:sdtContent><w:p><w:r><w:t>B</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    document = _replace_document_body(rdocx.Document(), body)
    snapshot = document.content_controls()[0]
    settings = document.settings
    assert document.set_content_control_value_by_tag("check", "true") == 1
    assert snapshot.text == "B"
    assert document.content_controls()[0].text == "A"
    xml = _document_xml(document)
    assert b'<f:checked f:val="opaque"/>' in xml and b'f:keep="yes"' in xml
    assert b'w:ascii="Carlito"' in xml and b'w:hAnsi="Carlito"' in xml
    before = document.to_bytes()
    assert document.set_content_control_value_by_tag("check", "true") == 1
    assert document.to_bytes() == before
    assert settings.track_revisions is None
    assert document.set_content_control_value_by_tag("check", "false") == 1
    assert document.content_controls()[0].text == "B"
    assert b'w:ascii="Selawik"' in _document_xml(document)
    assert rdocx.Document.from_bytes(document.to_bytes()).content_controls()[0].text == "B"


@pytest.mark.parametrize("invalid", ["D800", "0000"])
def test_fx203_checkbox_late_invalid_state_preserves_package_and_snapshots(invalid):
    import rdocx
    def control(code):
        return f'<w:sdt><w:sdtPr><w:tag w:val="check"/><x:checkbox xmlns:x="http://schemas.microsoft.com/office/word/2010/wordml"><x:checkedState x:val="{code}" x:font="Carlito"/></x:checkbox></w:sdtPr><w:sdtContent><w:p><w:r><w:t>old</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    document = _replace_document_body(rdocx.Document(), control("0041") + control(invalid))
    controls = document.content_controls()
    paragraph = document.paragraphs[0]
    before = document.to_bytes()
    with pytest.raises(ValueError, match="glyph"):
        document.set_content_control_value_by_tag("check", "true")
    assert document.to_bytes() == before
    assert paragraph.text == "old"
    assert [c.text for c in controls] == ["old", "old"]
    assert [c.text for c in document.content_controls()] == ["old", "old"]


def test_fx203_checkbox_default_namespace_has_qualified_glyph_font_and_repeat():
    import rdocx
    body = '<w:sdt><w:sdtPr><w:tag w:val="check"/><checkbox xmlns="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:q="http://schemas.microsoft.com/office/word/2010/wordml"><checked q:val="0"/><checkedState q:val="0044" q:font="Carlito"/></checkbox></w:sdtPr><w:sdtContent><w:p><w:r><w:t>old</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    document = _replace_document_body(rdocx.Document(), body)
    assert document.content_controls()[0].type == "checkbox"
    assert document.set_content_control_value_by_tag("check", "true") == 1
    assert document.content_controls()[0].text == "D"
    assert b'w:ascii="Carlito"' in _document_xml(document)
    before = document.to_bytes()
    assert document.set_content_control_value_by_tag("check", "true") == 1
    assert document.to_bytes() == before
    assert rdocx.Document.from_bytes(before).content_controls()[0].text == "D"


def test_fx203_template_json_conversion_is_typed_and_atomically_preflighted():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("{{ flag }} {{ large }} {{ name }}")
    assert document.render_template({"flag": True, "large": 2**64 - 1, "name": "Ada"}) == 3
    assert document.paragraphs[0].text == "true 18446744073709551615 Ada"
    document = rdocx.Document()
    document.add_paragraph("{{ name }}")
    held = document.paragraphs[0]
    cycle = []
    cycle.append(cycle)
    class Arbitrary:
        def __str__(self):
            raise AssertionError("arbitrary string conversion must never run")
    for data, error in [
        ({"name": cycle}, ValueError), ({"name": float("nan")}, ValueError),
        ({"name": float("inf")}, ValueError), ({"name": 2**64}, ValueError),
        ({"name": -(2**63)-1}, ValueError), ({"name": Arbitrary()}, TypeError),
        ({1: "not a string key"}, TypeError), (["not a dict"], TypeError),
    ]:
        before = document.to_bytes()
        with pytest.raises(error):
            document.render_template(data)
        assert document.to_bytes() == before
        assert held.text == "{{ name }}"
    assert document.render_template({"name": "Ada"}) == 1
    assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == "Ada"


def test_fx203_exports_keep_owned_payload_and_full_loss_records():
    import rdocx
    from xml.etree import ElementTree
    document = rdocx.Document()
    document.add_paragraph("Export α beta")
    document.add_table(1, 1).cell(0, 0).text = "Cell Ω"
    before = document.to_bytes()
    odt, rtf, epub = document.to_odt(), document.to_rtf(), document.to_epub()
    with zipfile.ZipFile(io.BytesIO(odt.bytes)) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
        text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
        for element in root.iter(text_ns + "s"):
            element.text = " " * int(element.get(text_ns + "c", "1"))
        content = "".join(root.itertext())
        assert "Export α beta" in content and "Cell Ω" in content
    assert rtf.bytes.startswith(b"{\\rtf")
    with zipfile.ZipFile(io.BytesIO(epub.bytes)) as archive:
        assert archive.read("mimetype") == b"application/epub+zip"
        content = "".join("".join(ElementTree.fromstring(archive.read(name)).itertext()) for name in archive.namelist() if name.endswith(".xhtml"))
        assert "Export α beta" in content and "Cell Ω" in content
    for result in (odt, rtf, epub):
        assert isinstance(result.diagnostics, tuple)
        for diagnostic in result.diagnostics:
            assert isinstance(diagnostic.message, str) and diagnostic.message
    assert document.to_bytes() == before


@pytest.mark.parametrize("operation", ["inventory", "options", "cell", "heading_link"])
def test_fx202_picture_and_heading_surface_on_installed_base(operation):
    from rdocx import Document
    document = Document()
    if operation == "inventory":
        document.add_picture(_one_pixel_png(), "pixel.png", width=127000, height=127000)
        pictures = document.pictures()
        assert len(pictures) == 1
        document.set_picture_size(pictures[0], 254000, 254000)
        assert len(Document.from_bytes(document.to_bytes()).pictures()) == 1
    elif operation == "options":
        from rdocx import PictureOptions
        picture = document.add_picture(_one_pixel_png(), "pixel.png", width=127000, height=127000,
                                       options=PictureOptions(description="owned description", title="owned title"))
        assert picture.description == "owned description"
        assert picture.title == "owned title"
    elif operation == "cell":
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        picture = cell.add_picture(_one_pixel_png(), width=127000, height=127000)
        assert picture.content_type == "image/png"
    else:
        document.add_paragraph("Target").style = "Heading1"
        document.add_paragraph("See ")
        document.add_heading_hyperlink(document.paragraphs[1], document.paragraphs[0], "target", tooltip="owned tip")
        xml = _document_xml(document)
        assert b'anchor="Heading_' in xml and b'owned tip' in xml

@pytest.mark.parametrize("host", ("body", "header", "body-cell", "header-cell"))
def test_fx202_configured_physical_picture_owner_and_frozen_mutation(host):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("body retained")
    if host == "body":
        paragraph = document.paragraphs[0]
    elif host == "header":
        paragraph = document.section(0).header.add_paragraph("owner text")
    elif host == "body-cell":
        document.add_table(1, 1).cell(0, 0).text = "owner text"
        paragraph = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        document.section(0).header.add_table(1, 1).cell(0, 0).text = "owner text"
        paragraph = document.section(0).header.tables[0].cell(0, 0).paragraphs[0]
    run = paragraph.runs[0]
    options = rdocx.PictureOptions(name="owned picture", title="literal <title>", description="owned & description", decorative=True, crop=(.1, 0., .2, 0.))
    image = _one_pixel_png(b"\xff\x00\x00")
    picture = run.add_picture(image, width=254000, height=127000, options=options)
    assert isinstance(picture, rdocx.Picture)
    assert (picture.width, picture.height) == (254000, 127000)
    assert picture.bytes == image and picture.content_type == "image/png"
    assert picture.title == "literal <title>" and picture.decorative
    assert not picture.anchored
    assert len(document.pictures()) == 1
    saved = document.to_bytes()
    with pytest.raises(rdocx.StaleElementError, match="different document"):
        rdocx.Document().set_picture_size(picture, 1, 1)
    assert document.to_bytes() == saved
    document.set_picture_size(picture, 508000, 254000)
    current = document.pictures()[0]
    assert (current.width, current.height) == (508000, 254000)
    assert current.bytes == image and current.title == picture.title
    before = document.to_bytes()
    with pytest.raises(rdocx.StaleElementError):
        document.set_picture_size(picture, 1, 1)
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert reopened.pictures()[0].bytes == image
    assert b"body retained" in _document_xml(reopened)

@pytest.mark.parametrize("kwargs", [dict(crop=(float("nan"),0.,0.,0.)), dict(crop=(True,0.,0.,0.)), dict(placement="through"), dict(horizontal_offset=True), dict(placement="inline", horizontal_offset=1), dict(placement="square", horizontal_offset=1, horizontal_alignment="left")])
def test_fx202_picture_options_reject_unsupported_operands(kwargs):
    import rdocx
    with pytest.raises((TypeError, ValueError, rdocx.RdocxError)):
        rdocx.PictureOptions(**kwargs)

def test_fx202_heading_links_reuse_bookmark_and_reject_foreign_stale_targets():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("Target").style = "Heading1"
    document.add_paragraph("See ")
    heading, source = document.paragraphs[0], document.paragraphs[1]
    inserted = document.add_heading_hyperlink(source, heading, "target", tooltip="owned tip")
    assert inserted.text == "target"
    with pytest.raises(rdocx.StaleElementError):
        document.add_heading_hyperlink(source, heading, "stale")
    document.add_heading_hyperlink(document.paragraphs[1], document.paragraphs[0], "again")
    assert len(document.bookmarks) == 1
    source = document.paragraphs[1]
    name = document.bookmarks[0].name
    source.add_hyperlink("anchor", "#" + name, tooltip="internal tip")
    xml = _document_xml(document)
    assert xml.count(b'w:anchor="Heading_1"') == 3
    assert b"owned tip" in xml and b"internal tip" in xml
    before = document.to_bytes()
    foreign = rdocx.Document().add_paragraph("foreign")
    with pytest.raises(ValueError, match="different document"):
        document.add_heading_hyperlink(document.paragraphs[1], foreign, "bad")
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(before)
    assert reopened.paragraphs[1].text == "See targetagainanchor"

@pytest.mark.parametrize("host", ("header", "footer", "body-cell", "header-cell"))
def test_fx202_related_source_heading_and_internal_links_publish_once(host):
    import rdocx
    from xml.etree import ElementTree
    document = rdocx.Document()
    document.add_paragraph("Target").style = "Heading1"
    if host == "header":
        source = document.section(0).header.add_paragraph("See ")
    elif host == "footer":
        source = document.section(0).footer.add_paragraph("See ")
    elif host == "body-cell":
        document.add_table(1, 1).cell(0, 0).text = "See "
        source = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        document.section(0).header.add_table(1, 1).cell(0, 0).text = "See "
        source = document.section(0).header.tables[0].cell(0, 0).paragraphs[0]
    inserted = document.add_heading_hyperlink(source, document.paragraphs[0], "target", tooltip="related tip")
    assert inserted.text == "target"
    assert len(document.bookmarks) == 1
    saved = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        owner = "word/document.xml" if host == "body-cell" else next(n for n in archive.namelist() if re.fullmatch("word/" + ("footer" if host == "footer" else "header") + r"\d+\.xml", n))
        root = ElementTree.fromstring(archive.read(owner))
        links = root.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")
        assert len(links) == 1
        assert links[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor") == document.bookmarks[0].name
        assert links[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tooltip") == "related tip"
    assert rdocx.Document.from_bytes(saved).bookmarks[0].name == document.bookmarks[0].name

def test_fx202_picture_snapshot_has_no_public_constructor():
    import rdocx
    with pytest.raises(TypeError):
        rdocx.Picture()

@pytest.mark.parametrize("host", ("header", "body-cell", "footer-cell"))
def test_fx202_heading_destination_uses_supported_physical_owner(host):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("Source ")
    if host == "header":
        target = document.section(0).header.add_paragraph("Related target")
    elif host == "body-cell":
        document.add_table(1, 1).cell(0, 0).text = "Related target"
        target = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        document.section(0).footer.add_table(1, 1).cell(0, 0).text = "Related target"
        target = document.section(0).footer.tables[0].cell(0, 0).paragraphs[0]
    target.style = "Heading1"
    if host == "header":
        target = document.section(0).header.paragraphs[-1]
    elif host == "body-cell":
        target = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        target = document.section(0).footer.tables[0].cell(0, 0).paragraphs[0]
    inserted = document.add_heading_hyperlink(document.paragraphs[0], target, "target", tooltip="supported owner")
    assert inserted.text == "target"
    assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == "Source target"

@pytest.mark.parametrize("host", ("body", "header"))
def test_fx202_external_tooltip_xml_refusal_preserves_package_and_handle(host):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("Source ")
    paragraph = document.paragraphs[0] if host == "body" else document.section(0).header.add_paragraph("Related ")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        paragraph.add_hyperlink("invalid", "https://example.com/#fragment", tooltip="invalid\x00tooltip")
    assert document.to_bytes() == before
    assert paragraph.text in ("Source ", "Related ")
    paragraph.add_hyperlink("valid", "https://example.com/#fragment", tooltip="valid & tooltip")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert (reopened.paragraphs[0].text if host == "body" else reopened.section(0).header.paragraphs[-1].text).endswith("valid")

@pytest.mark.parametrize("host", ("document", "run", "cell"))
def test_fx202_malformed_explicit_size_picture_preserves_held_owner(host):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("retained")
    if host == "run":
        owner = document.paragraphs[0].runs[0]
    elif host == "cell":
        document.add_table(1, 1)
        owner = document.tables[0].cell(0, 0)
    else:
        owner = document
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        owner.add_picture(b"not an image", width=127000, height=127000)
    assert document.to_bytes() == before
    assert len(document.pictures()) == 0
    if host == "run":
        assert owner.text == "retained"
    elif host == "cell":
        assert owner.text == ""


def test_fx204_counted_regex_preserves_failed_noop_and_empty_handles():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("a12 a34")
    held = document.paragraphs[0]
    before = document.to_bytes()
    assert document.replace_all_regex_expected([]) == ()
    assert document.to_bytes() == before
    assert held.text == "a12 a34"
    assert document.replace_all_regex_expected([(r"missing", "x", 0)]) == (0,)
    assert held.text == "a12 a34"
    assert document.replace_all_regex_expected([(r"a([0-9]+)", "a$1", 2)]) == (2,)
    assert held.text == "a12 a34"
    with pytest.raises(rdocx.ReplacementCountError):
        document.replace_all_regex_expected([(r"a([0-9]+)", "b$1", 2), ("b", "c", 1)])
    assert document.to_bytes() == before
    assert held.text == "a12 a34"
    with pytest.raises(rdocx.RdocxError):
        document.replace_all_regex_expected([("a", "b"), ("[", "c")])
    assert document.to_bytes() == before
    assert held.text == "a12 a34"
    assert document.replace_all_regex_expected([(r"a([0-9]+)", "b$1", 2), ("b", "c", 2)]) == (2, 2)
    assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == "c12 c34"
    with pytest.raises(rdocx.StaleElementError):
        _ = held.text


def _content_location(document, predicate):
    import rdocx
    item = next(item for item in document.story_items if predicate(item))
    return rdocx.ContentLocation(item.story, item.kind, item.index_path)


@pytest.mark.parametrize(
    ("payload", "refusal"),
    [
        ("&#65;<w:p/>", "not a block sequence"),
        ("&lt;<w:p/>", "not a block sequence"),
        ("&#32;<w:p/>", None),
        (
            '<w:sdt><w:sdtPr><w:id w:val="2"/><w:richText/></w:sdtPr>'
            '<w:sdtContent>&#65;<w:p/></w:sdtContent></w:sdt>',
            "not a block sequence",
        ),
        (
            '<w:sdt><w:sdtPr><w:id w:val="2"/><w:richText/></w:sdtPr>'
            '<w:sdtContent><w:p/></w:sdtContent></w:sdt>',
            None,
        ),
    ],
)
def test_f286_repeating_section_checks_decoded_boundary_text(payload, refusal):
    import rdocx

    document = _replace_document_body(
        rdocx.Document(),
        '<w:sdt><w:sdtPr><w:tag w:val="probe"/><w:id w:val="1"/>'
        '<w:richText/></w:sdtPr><w:sdtContent>' + payload
        + '</w:sdtContent></w:sdt><w:p><w:r><w:t>tail</w:t></w:r></w:p>',
    )
    location = _content_location(document, lambda item: item.kind == "content_control")
    before = document.to_bytes()
    if refusal:
        with pytest.raises(rdocx.RdocxError, match=refusal):
            document.set_content_control_type(
                location, rdocx.ContentControlTypePayload.repeating_section()
            )
        assert document.to_bytes() == before
    else:
        document.set_content_control_type(
            location, rdocx.ContentControlTypePayload.repeating_section()
        )
        assert document.to_bytes() != before
        xml = _document_xml(document)
        assert b"<w:sdtContent>" + payload.encode() + b"</w:sdtContent>" in xml
        assert xml.count(b"repeatingSectionItem") == 1
        assert b">tail<" in xml


@pytest.fixture(scope="module")
def _f286_native_word_input():
    import hashlib
    import json
    import os
    from pathlib import Path

    receipt = os.environ.get("RDOCX_F286_WORD_PROVENANCE")
    if receipt is None:
        pytest.skip("F-286 needs the separately retained native Word provenance bundle")
    provenance = json.loads(Path(receipt).read_text())
    assert (provenance["producer"], provenance["version"], provenance["build"]) == (
        "Microsoft Word for Mac", "16.114", "26101123"
    )
    inputs = {Path(row["source"]).name: row for row in provenance["files"]}

    def read(name):
        row = inputs[name]
        raw = Path(row["source"]).read_bytes()
        assert hashlib.sha256(raw).hexdigest() == row["sha256"], name
        return raw

    return read


def test_f286_native_word_replacement_preserves_typed_payloads(_f286_native_word_input):
    import xml.etree.ElementTree as ET
    import rdocx

    read = _f286_native_word_input
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    types = {
        "richText", "text", "date", "checkbox", "comboBox", "dropDownList",
        "picture", "group", "repeatingSection", "repeatingSectionItem",
    }

    def facts(raw):
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            tree = ET.fromstring(package.read("word/document.xml"))

        def node(element):
            return (
                element.tag, sorted(element.attrib.items()), element.text or "",
                [node(child) for child in element],
            )

        result = []
        for control in tree.iter(w + "sdt"):
            properties = control.find(w + "sdtPr")
            payload = next(
                (child for child in properties if child.tag.rsplit("}", 1)[-1] in types),
                None,
            )
            result.append((
                properties.find(w + "id").get(w + "val"),
                node(payload) if payload is not None else None,
                "".join(child.text or "" for child in control.iter(w + "t")),
                len(list(control.iter(w + "drawing"))),
            ))
        return result

    for stem, value in [
        ("scalar-0", "Replacement"), ("scalar-1", "Replacement"),
        ("scalar-3", "two"), ("scalar-4", "Second"), ("scalar-8", "true"),
        ("date-refined-1", "2026-10-19"), ("date-refined-2", "2026-10-19"),
    ]:
        document = rdocx.Document.from_bytes(read(stem + "-before.docx"))
        assert document.set_content_control_value_by_tag(stem, value) == 1
        assert facts(document.to_bytes()) == facts(read(stem + "-after.docx")), stem

    for name, tag in [("picture-before.docx", "picture")] + [
        (f"group-native-{index}-before.docx", "group-native") for index in (1, 2, 3)
    ]:
        document = rdocx.Document.from_bytes(read(name))
        before = document.to_bytes()
        with pytest.raises(rdocx.RdocxError):
            document.set_content_control_value_by_tag(tag, "Replacement")
        assert document.to_bytes() == before


def test_f286_approved_repeat_matches_native_equivalent_operations(_f286_native_word_input):
    import xml.etree.ElementTree as ET
    import rdocx

    read = _f286_native_word_input
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    w15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"

    def items(raw):
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            tree = ET.fromstring(package.read("word/document.xml"))
        result = []
        for control in tree.iter(w + "sdt"):
            properties = control.find(w + "sdtPr")
            if properties.find(w15 + "repeatingSectionItem") is None:
                continue
            runs = []
            for run in control.iter(w + "r"):
                bold = run.find(w + "rPr/" + w + "b")
                runs.append((
                    "".join(child.text or "" for child in run.iter(w + "t")),
                    bold is not None and bold.get(w + "val", "1") not in ("0", "false", "off"),
                ))
            result.append((properties.find(w + "id").get(w + "val"), runs))
        return result

    # Native equivalent is append after the last item, then copy the first
    # item's FormattedText. Word's default addressed-item clone differs.
    before = read("repeat-native-first-v3-before.docx")
    document = rdocx.Document.from_bytes(before)
    section = _content_location(
        document,
        lambda item: item.kind == "content_control"
        and b'w:val="repeat-clone-first"' in item.xml,
    )
    document.repeat_content_control(section)
    original, actual = items(before), items(document.to_bytes())
    native = items(read("repeat-native-first-v3-after.docx"))
    assert len(original) == 3 and len(actual) == 4
    assert [row[1] for row in actual] == [row[1] for row in native]
    assert actual[:3] == original
    assert len({row[0] for row in actual}) == 4
    assert actual[-1][0] not in {row[0] for row in original}
    assert actual[-1][1] == original[0][1] != original[-1][1]

    def locations():
        return [
            rdocx.ContentLocation(item.story, item.kind, item.index_path)
            for item in document.story_items
            if item.kind == "content_control"
            and b"repeatingSectionItem" in item.xml
            and b'w:val="repeat-clone-first"' not in item.xml
        ]

    document.remove_content_control(locations()[1])
    removed = items(document.to_bytes())
    native_removed = items(read("repeat-native-first-v3-remove-middle.docx"))
    assert [row[1] for row in removed] == [row[1] for row in native_removed]
    assert removed == [actual[index] for index in (0, 2, 3)]
    while len(locations()) > 1:
        document.remove_content_control(locations()[-1])
    before_refusal = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="last repeating"):
        document.remove_content_control(locations()[0])
    assert document.to_bytes() == before_refusal
    # Observed native last-item deletion removes the entire section. This
    # asserts the difference, not approval of a blanket Word-parity claim.
    assert items(read("repeat-last-refusal.docx")) == []


def test_f286_typed_content_control_payload_repeat_and_refusal():
    import rdocx
    document = rdocx.Document()
    body = next(story for story in document.stories if story.kind == "body")
    document.insert_content_control(
        rdocx.ContentLocation.end(body), "block", "rich_text", tag="typed"
    )
    location = _content_location(document, lambda item: item.kind == "content_control")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="empty"):
        document.set_content_control_type(
            location, rdocx.ContentControlTypePayload.drop_down_list([])
        )
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError, match="repeated"):
        document.set_content_control_type(
            location,
            rdocx.ContentControlTypePayload.combo_box([("One", "same"), ("Two", "same")]),
        )
    assert document.to_bytes() == before
    document.set_content_control_type(location, rdocx.ContentControlTypePayload.checkbox())
    document = rdocx.Document.from_bytes(document.to_bytes())
    assert document.content_controls()[0].type == "checkbox"
    document.set_content_control_value_by_tag("typed", "yes")
    assert document.content_controls()[0].text == "\u2612"
    location = _content_location(document, lambda item: item.kind == "content_control")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="repeating section"):
        document.repeat_content_control(location)
    assert document.to_bytes() == before

    empty = _replace_document_body(
        rdocx.Document(),
        '<w:sdt><w:sdtPr><w:tag w:val="empty"/><w:id w:val="1"/><w:richText/></w:sdtPr><w:sdtContent/></w:sdt>',
    )
    location = _content_location(empty, lambda item: item.kind == "content_control")
    before = empty.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="no item"):
        empty.set_content_control_type(
            location, rdocx.ContentControlTypePayload.repeating_section()
        )
    assert empty.to_bytes() == before

    section = rdocx.Document()
    body = next(story for story in section.stories if story.kind == "body")
    section.insert_content_control(
        rdocx.ContentLocation.end(body), "block", "rich_text", tag="section"
    )
    location = _content_location(section, lambda item: item.kind == "content_control")
    section.set_content_control_type(
        location, rdocx.ContentControlTypePayload.repeating_section()
    )
    location = _content_location(
        section, lambda item: item.kind == "content_control" and b'w:val="section"' in item.xml
    )
    section.repeat_content_control(location)
    section = rdocx.Document.from_bytes(section.to_bytes())
    items = [
        item
        for item in section.story_items
        if item.kind == "content_control"
        and b"repeatingSectionItem" in item.xml
        and b'w:val="section"' not in item.xml
    ]
    assert len(items) == 2
    section.remove_content_control(
        rdocx.ContentLocation(items[0].story, items[0].kind, items[0].index_path)
    )
    survivor = _content_location(
        section,
        lambda item: item.kind == "content_control"
        and b"repeatingSectionItem" in item.xml
        and b'w:val="section"' not in item.xml,
    )
    before = section.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="last repeating"):
        section.remove_content_control(survivor)
    assert section.to_bytes() == before


def test_f285_content_control_shell_metadata_and_refusal():
    import rdocx
    document = rdocx.Document()
    body = next(story for story in document.stories if story.kind == "body")
    end = rdocx.ContentLocation.end(body)
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="does not match"):
        document.insert_content_control(end, "inline", "rich_text", tag="refused")
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError, match="F-286"):
        document.insert_content_control(end, "block", "checkbox", tag="later")
    assert document.to_bytes() == before
    with pytest.raises(TypeError):
        document.insert_content_control(end, "block", "rich_text", placeholder=1)
    assert document.to_bytes() == before
    document.insert_content_control(
        end,
        "block",
        "rich_text",
        tag="fill-me",
        alias="shown",
        lock="contentLocked",
        appearance="boundingBox",
        placeholder=True,
    )
    item = next(item for item in document.story_items if item.kind == "content_control")
    location = rdocx.ContentLocation(item.story, item.kind, item.index_path)
    document.set_content_control_tag(location, "renamed")
    item = next(item for item in document.story_items if b"renamed" in item.xml)
    location = rdocx.ContentLocation(item.story, item.kind, item.index_path)
    document.fill_content_control_tags([("renamed", "hello")])
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    control = reopened.content_controls()[0]
    assert control.tag == "renamed"
    assert control.alias == "shown"
    assert control.type == "rich_text"
    assert control.text == "hello"
    assert control.id == 1
    item = next(item for item in reopened.story_items if item.kind == "content_control")
    reopened.remove_content_control(
        rdocx.ContentLocation(item.story, item.kind, item.index_path)
    )
    assert reopened.content_controls() == ()
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as package:
        assert b"hello" not in package.read("word/document.xml")


def test_f286_type_namespace_identity_survives_staging_and_refusal():
    from xml.etree import ElementTree as ET
    import rdocx

    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    extension = "http://schemas.microsoft.com/office/word/2012/wordml"
    opaque = '<w:customXmlPr><w:attr w:name="x" w:val="kept"/></w:customXmlPr>'
    metadata = '<w:customXml w:uri="u" w:element="kept">' + opaque + '</w:customXml>'
    base = rdocx.Document().to_bytes()

    def package(prefix, scope, shadow, direct):
        declaration = f' xmlns:{prefix}="{extension}"'

        def shell(tag, identifier, kind, inner):
            properties = declaration if scope == "properties" and kind else ""
            marker = (
                f'<{prefix}:{kind}' + (declaration if scope == "marker" else "") + '/>'
                if kind else '<w:richText/>'
            )
            return (
                f'<w:sdt><w:sdtPr{properties}><w:id w:val="{identifier}"/>'
                f'<w:tag w:val="{tag}"/>{marker}</w:sdtPr>'
                f'<w:sdtContent>{inner}</w:sdtContent></w:sdt>'
            )

        item = shell("item", 3, "repeatingSectionItem", metadata)
        section = shell("section", 2, "repeatingSection", item)
        outer = shell("outer", 1, None, section)
        if scope == "outer":
            outer = outer.replace('<w:sdt>', '<w:sdt' + declaration + '>', 1)
        if shadow:
            outer = outer.replace('<w:sdt', '<w:sdt xmlns:w15="urn:opaque"', 1)
        body = '<w:p><w:r><w:t>host</w:t></w:r></w:p>'
        body += outer if direct else shell("container", 4, None, outer)
        result = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(base)) as source, zipfile.ZipFile(result, "w") as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == "word/document.xml":
                    start = data.index(b"<w:body>") + len(b"<w:body>")
                    end = data.index(b"</w:body>")
                    data = data[:start] + body.encode() + data[end:]
                    if scope == "root":
                        data = data.replace(b"<w:document", ('<w:document' + declaration).encode(), 1)
                    assert len(list(ET.fromstring(data).iter("{" + extension + "}repeatingSection"))) == 1
                target.writestr(entry, data)
        return rdocx.Document.from_bytes(result.getvalue())

    def by_tag(document, tag):
        needle = f'w:val="{tag}"'.encode()
        return _content_location(
            document,
            lambda item: item.kind == "content_control" and needle in item.xml.split(b"<w:sdtContent", 1)[0],
        )

    def assert_types(document, count):
        xml = _document_xml(document)
        root = ET.fromstring(xml)
        for kind in ("repeatingSection", "repeatingSectionItem"):
            assert len(list(root.iter("{" + extension + "}" + kind))) == count
        assert xml.count(opaque.encode()) == count
        identifiers = [int(node.attrib["{" + word + "}val"]) for node in root.iter("{" + word + "}id")]
        assert len(identifiers) == len(set(identifiers))
        return set(identifiers)

    variants = [(prefix, scope, False) for prefix in ("rep", "w15") for scope in ("outer", "root", "properties", "marker")]
    variants.append(("rep", "outer", True))
    for prefix, scope, shadow in variants:
        for operation in ("move_content_control", "clone_content_control"):
            document = package(prefix, scope, shadow, direct=False)
            before = document.to_bytes()
            assert_types(document, 1)  # Raw passthrough is only a control.
            source = by_tag(document, "outer")
            host = _content_location(document, lambda item: item.kind == "paragraph" and item.text == "host")
            with pytest.raises(rdocx.RdocxError, match="placement"):
                getattr(document, operation)(source, host)
            assert document.to_bytes() == before
            assert_types(document, 1)

            # A direct source inserts at a body boundary, not inside a paragraph.
            document = package(prefix, scope, shadow, direct=True)
            source = by_tag(document, "outer")
            body = next(story for story in document.stories if story.kind == "body")
            getattr(document, operation)(source, rdocx.ContentLocation.end(body))
            count = 2 if operation == "clone_content_control" else 1
            identifiers = assert_types(document, count)
            expected = set(range(1, 7)) if count == 2 else {1, 2, 3}
            assert identifiers == expected
            document = rdocx.Document.from_bytes(document.to_bytes())
            document.set_content_control_alias(by_tag(document, "outer"), "retained")
            assert_types(document, count)
            assert any(control.alias == "retained" for control in document.content_controls())


def test_f286_repeating_types_keep_physical_and_owner_context():
    import rdocx
    from xml.etree import ElementTree as ET

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    w15_uri = "http://schemas.microsoft.com/office/word/2012/wordml"
    metadata = '<w:customXml w:uri="urn:valid" w:element="kept"><w:customXmlPr><w:attr w:name="m" w:val="kept"/></w:customXmlPr></w:customXml>'
    paragraph = '<w:p><w:r><w:t>kept</w:t></w:r></w:p>'
    cell = '<w:tc><w:tcPr><w:tcW w:w="1000" w:type="dxa"/></w:tcPr>' + paragraph + '</w:tc>'
    row = '<w:tr>' + cell + '</w:tr>'
    grid = '<w:tblPr/><w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid>'

    def shell(tag, ident, kind, inner):
        marker = '<w:richText/>' if kind == 'rich' else f'<w15:{kind} xmlns:w15="{w15_uri}"/>'
        return f'<w:sdt><w:sdtPr><w:id w:val="{ident}"/><w:tag w:val="{tag}"/>{marker}</w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>'

    def document(inner):
        return _replace_document_body(rdocx.Document(), '<w:p><w:r><w:t>host</w:t></w:r></w:p>' + inner + '<w:tbl>' + grid + row + '</w:tbl>')

    def location(doc, tag):
        def match(item):
            if item.kind != 'content_control':
                return False
            raw = item.xml
            at = raw.find(f'w:val="{tag}"'.encode())
            return 0 <= at < raw.find(b'<w:sdtContent')
        return _content_location(doc, match)

    def xml(doc):
        with zipfile.ZipFile(io.BytesIO(doc.to_bytes())) as package:
            return package.read('word/document.xml')

    def refuse(doc, operation, label):
        before = doc.to_bytes()
        try:
            operation()
        except rdocx.RdocxError:
            assert doc.to_bytes() == before, label
        else:
            pytest.fail(label + ' published an invalid repeating context')

    # Begin valid: the inline rich shell is inside a real block item/section.
    inline = shell('target', 3, 'rich', '<w:r><w:t>inline</w:t></w:r>')
    source = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', '<w:p>' + inline + '</w:p>'))
    doc = document(source)
    refuse(doc, lambda: doc.set_content_control_type(location(doc, 'target'), rdocx.ContentControlTypePayload.repeating_section_item()), 'inline item type')

    # A nonempty wrapper with empty payload fits both grammars, not both slots.
    for inner in (shell('empty', 4, 'rich', ''), metadata):
        doc = document('<w:p>' + shell('target', 3, 'rich', inner) + '</w:p>')
        refuse(doc, lambda: doc.set_content_control_type(location(doc, 'target'), rdocx.ContentControlTypePayload.repeating_section()), 'inline empty section')
        doc = document(shell('target', 3, 'rich', inner))
        doc.set_content_control_type(location(doc, 'target'), rdocx.ContentControlTypePayload.repeating_section())
        assert b'repeatingSectionItem' in xml(doc)
        if inner == metadata:
            assert metadata.encode() in xml(rdocx.Document.from_bytes(doc.to_bytes()))

    # Same type retirement would either orphan items or replace a required item.
    simple = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', paragraph))
    for tag in ('section', 'item'):
        doc = document(simple)
        refuse(doc, lambda: doc.set_content_control_type(location(doc, tag), rdocx.ContentControlTypePayload.rich_text()), 'retire ' + tag)
        doc = document(simple)
        refuse(doc, lambda: doc.unwrap_content_control(location(doc, tag)), 'unwrap ' + tag)

    # Removing a section deletes its owned items, so it does not orphan them.
    doc = document(simple)
    doc.remove_content_control(location(doc, 'section'))
    assert b'repeatingSection' not in xml(doc)

    # An enclosing section can retain ownership across an inner marker/shell loss.
    nested = shell('section', 1, 'repeatingSection', shell('outer-item', 2, 'repeatingSectionItem', shell('inner-section', 3, 'repeatingSection', shell('inner-item', 4, 'repeatingSectionItem', paragraph))))
    for unwrap in (False, True):
        doc = document(nested)
        if unwrap:
            doc.unwrap_content_control(location(doc, 'inner-section'))
        else:
            doc.set_content_control_type(location(doc, 'inner-section'), rdocx.ContentControlTypePayload.rich_text())
        tree = ET.fromstring(xml(doc))
        assert len(list(tree.iter('{' + w15_uri + '}repeatingSection'))) == 1
        assert len(list(tree.iter('{' + w15_uri + '}repeatingSectionItem'))) == 2
        assert any(control.id == 4 for control in doc.content_controls())

    # Flattening an item into an item-only sequence keeps a legal section.
    flattened = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', shell('a', 3, 'repeatingSectionItem', paragraph) + shell('b', 4, 'repeatingSectionItem', paragraph)))
    doc = document(flattened)
    doc.unwrap_content_control(location(doc, 'item'))
    assert {control.id for control in doc.content_controls()} == {1, 3, 4}

    # Existing legal block, row and cell items/sections retain their operations.
    for slot, payload in [('block', paragraph), ('row', row), ('cell', cell)]:
        section = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', payload))
        if slot == 'row':
            section = '<w:tbl>' + grid + section + row + '</w:tbl>'
        elif slot == 'cell':
            section = '<w:tbl>' + grid + '<w:tr>' + section + cell + '</w:tr></w:tbl>'
        doc = document(section)
        doc.set_content_control_type(location(doc, 'section'), rdocx.ContentControlTypePayload.repeating_section())
        doc.set_content_control_type(location(doc, 'item'), rdocx.ContentControlTypePayload.repeating_section_item())
        doc.repeat_content_control(location(doc, 'section'))
        reopened = rdocx.Document.from_bytes(doc.to_bytes())
        assert len(list(ET.fromstring(xml(reopened)).iter('{' + w15_uri + '}repeatingSectionItem'))) == 2
        assert len({control.id for control in reopened.content_controls()}) == 3

    # Block item authoring requires containment, not an immediate section parent.
    contained = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', shell('target', 3, 'rich', paragraph)))
    doc = document(contained)
    doc.set_content_control_type(location(doc, 'target'), rdocx.ContentControlTypePayload.repeating_section_item())
    assert len(list(ET.fromstring(xml(doc)).iter('{' + w15_uri + '}repeatingSectionItem'))) == 2

    # Republish the complete type, not only an empty payload's generic grammar.
    section = shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', metadata))
    for operation in ('clone_content_control', 'move_content_control'):
        doc = document(shell('container', 3, 'rich', section))
        host = _content_location(doc, lambda item: item.kind == 'paragraph' and item.text == 'host')
        refuse(doc, lambda: getattr(doc, operation)(location(doc, 'section'), host), operation + ' inline section')
        # Row sections carry their own item owner into a compatible table.
        row_section = '<w:tbl>' + grid + shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', row)) + row + '</w:tbl>'
        doc = document(row_section)
        target = [item for item in doc.story_items if item.kind == 'table'][-1]
        getattr(doc, operation)(location(doc, 'section'), rdocx.ContentLocation(target.story, target.kind, target.index_path))
        sections = [control for control in doc.content_controls() if control.tag == 'section']
        assert len(sections) == (2 if operation.startswith('clone') else 1)
        assert len({control.id for control in doc.content_controls()}) == (4 if operation.startswith('clone') else 2)
        if operation.startswith('clone'):
            assert 1 in {control.id for control in sections}
        else:
            assert sections[0].id == 1

    # An item copied outside all sections loses ownership even at a legal level.
    row_section = '<w:tbl>' + grid + shell('section', 1, 'repeatingSection', shell('item', 2, 'repeatingSectionItem', row) + shell('other-item', 3, 'repeatingSectionItem', row)) + row + '</w:tbl>'
    for operation in ('clone_content_control', 'move_content_control'):
        doc = document(row_section)
        target = [item for item in doc.story_items if item.kind == 'table'][-1]
        refuse(doc, lambda: getattr(doc, operation)(location(doc, 'item'), rdocx.ContentLocation(target.story, target.kind, target.index_path)), operation + ' orphan item')


def test_f285_custom_xml_wrappers_use_destination_grammar():
    import rdocx
    from xml.etree import ElementTree as ET

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = '<w:p><w:r><w:t>kept</w:t></w:r></w:p>'
    cell = '<w:tc><w:tcPr><w:tcW w:w="1000" w:type="dxa"/></w:tcPr>' + paragraph + '</w:tc>'
    row = '<w:tr>' + cell + '</w:tr>'
    grid = '<w:tblPr/><w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid>'

    def shell(tag, ident, inner):
        return f'<w:sdt><w:sdtPr><w:id w:val="{ident}"/><w:tag w:val="{tag}"/><w:richText/></w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>'

    def custom(inner):
        return '<w:customXml w:uri="urn:valid" w:element="retained"><w:customXmlPr><w:attr w:name="marker" w:val="kept"/></w:customXmlPr>' + inner + '</w:customXml>'

    def document(slot, nested, direct_block=False):
        payload = {"block": paragraph, "row": row, "cell": cell, "inline": '<w:r><w:t>kept</w:t></w:r>'}[slot]
        payload = custom(payload)
        if nested:
            payload = custom(shell("inner", 2, payload))
        control = shell("context", 1, payload)
        owner = {
            "block": control if direct_block else shell("container", 3, control),
            "row": '<w:tbl>' + grid + control + row + '</w:tbl>',
            "cell": '<w:tbl>' + grid + '<w:tr>' + control + cell + '</w:tr></w:tbl>',
            "inline": '<w:p>' + control + '</w:p>',
        }[slot]
        return _replace_document_body(rdocx.Document(), '<w:p><w:r><w:t>host</w:t></w:r></w:p>' + owner + '<w:tbl>' + grid + row + '</w:tbl>')

    def source(doc):
        def is_source(item):
            if item.kind != "content_control":
                return False
            raw = item.xml
            at = raw.find(b'w:val="context"')
            return 0 <= at < raw.find(b"<w:sdtContent")
        return _content_location(doc, is_source)

    def destination(doc, slot, compatible):
        host = _content_location(doc, lambda item: item.kind == "paragraph" and item.text == "host")
        table = [item for item in doc.story_items if item.kind == "table"][-1]
        table_location = rdocx.ContentLocation(table.story, table.kind, table.index_path)
        if not compatible:
            return host if slot in ("block", "row") else table_location
        if slot == "block":
            return rdocx.ContentLocation.end(next(story for story in doc.stories if story.kind == "body"))
        if slot == "inline":
            return host
        if slot == "cell":
            return rdocx.ContentLocation(table.story, "table", list(table.index_path) + [0])
        return table_location

    def raw_xml(doc):
        with zipfile.ZipFile(io.BytesIO(doc.to_bytes())) as package:
            return package.read("word/document.xml")

    def xml(doc):
        return ET.fromstring(raw_xml(doc))

    metadata = b'<w:customXmlPr><w:attr w:name="marker" w:val="kept"/></w:customXmlPr>'

    for slot in ("row", "cell", "block", "inline"):
        for nested in (False, True):
            for operation in ("clone_content_control", "move_content_control"):
                doc = document(slot, nested)
                before = doc.to_bytes()
                try:
                    getattr(doc, operation)(source(doc), destination(doc, slot, False))
                except rdocx.RdocxError as error:
                    assert "does not match" in str(error), (slot, nested, operation, error)
                else:
                    pytest.fail(f"{slot}/{nested}/{operation} accepted incompatible content")
                assert doc.to_bytes() == before, (slot, nested, operation)

                doc = document(slot, nested, direct_block=True)
                getattr(doc, operation)(source(doc), destination(doc, slot, True))
                controls = [control for control in doc.content_controls() if control.tag == "context"]
                ids = [control.id for control in controls]
                assert len(ids) == (2 if operation.startswith("clone") else 1)
                assert 1 in ids and len(set(ids)) == len(ids)
                tree = xml(doc)
                wrappers = list(tree.iter(w + "customXml"))
                assert wrappers and all(wrapper.find(w + "customXmlPr/" + w + "attr").get(w + "val") == "kept" for wrapper in wrappers)
                assert all(wrapper.get(w + "uri") == "urn:valid" for wrapper in wrappers)
                assert raw_xml(doc).count(metadata) == len(wrappers)
                reopened = rdocx.Document.from_bytes(doc.to_bytes())
                assert raw_xml(reopened).count(metadata) == len(wrappers)

            doc = document(slot, nested)
            doc.unwrap_content_control(source(doc))
            assert not any(control.tag == "context" for control in doc.content_controls())
            assert list(xml(doc).iter(w + "customXml"))

            doc = document(slot, nested)
            before = doc.to_bytes()
            if slot == "block":
                doc.set_content_control_type(source(doc), rdocx.ContentControlTypePayload.repeating_section())
                assert any(element.tag == "{http://schemas.microsoft.com/office/word/2012/wordml}repeatingSectionItem" for element in xml(doc).iter())
                assert list(xml(doc).iter(w + "customXml"))
            else:
                with pytest.raises(rdocx.RdocxError, match="block sequence|placement"):
                    doc.set_content_control_type(source(doc), rdocx.ContentControlTypePayload.repeating_section())
                assert doc.to_bytes() == before, (slot, nested, "repeating")


def test_f285_move_and_clone_content_control_identity_and_refusal():
    import rdocx

    def located(document, predicate):
        return _content_location(document, predicate)

    def by_tag(document, tag):
        needle = f'w:val="{tag}"'.encode()

        def predicate(item):
            if item.kind != "content_control":
                return False
            xml = item.xml
            tag_at = xml.find(needle)
            content_at = xml.find(b"<w:sdtContent")
            return tag_at != -1 and (content_at == -1 or tag_at < content_at)

        return located(document, predicate)

    def tables(document):
        return [
            rdocx.ContentLocation(item.story, item.kind, item.index_path)
            for item in document.story_items
            if item.kind == "table"
        ]

    def paragraph(document, text):
        return located(
            document,
            lambda item: item.kind == "paragraph" and item.text == text,
        )

    def tag_ids(document, tag):
        return [control.id for control in document.content_controls() if control.tag == tag]

    def table_xml(document, index):
        return [item for item in document.story_items if item.kind == "table"][index].xml

    document = rdocx.Document()
    body = next(story for story in document.stories if story.kind == "body")
    document.add_paragraph("host")
    document.add_paragraph("other")
    document.add_table(1, 1)
    document.add_table(1, 1)
    document.insert_content_control(tables(document)[0], "row", "rich_text", tag="row-shell")
    table = tables(document)[0]
    document.insert_content_control(
        rdocx.ContentLocation(table.story, "table", list(table.index_path) + [0]),
        "cell",
        "plain_text",
        tag="cell-shell",
    )
    document.insert_content_control(paragraph(document, "host"), "inline", "rich_text", tag="inline-shell")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="does not match"):
        document.clone_content_control(by_tag(document, "row-shell"), paragraph(document, "host"))
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError, match="does not match"):
        document.clone_content_control(by_tag(document, "cell-shell"), tables(document)[0])
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError, match="does not match"):
        document.move_content_control(by_tag(document, "cell-shell"), tables(document)[1])
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError, match="does not match"):
        document.clone_content_control(by_tag(document, "inline-shell"), tables(document)[1])
    assert document.to_bytes() == before
    with pytest.raises(rdocx.RdocxError):
        document.move_content_control(
            rdocx.ContentLocation(body, "content_control", [9999]),
            tables(document)[1],
        )
    assert document.to_bytes() == before

    row_id = tag_ids(document, "row-shell")[0]
    document.move_content_control(by_tag(document, "row-shell"), tables(document)[1])
    assert tag_ids(document, "row-shell") == [row_id]
    assert f'w:id w:val="{row_id}"'.encode() in table_xml(document, 1)
    assert f'w:id w:val="{row_id}"'.encode() not in table_xml(document, 0)
    inline_id = tag_ids(document, "inline-shell")[0]
    document.clone_content_control(by_tag(document, "inline-shell"), paragraph(document, "other"))
    inline_ids = tag_ids(document, "inline-shell")
    assert inline_id in inline_ids and len(inline_ids) == 2 and len(set(inline_ids)) == 2
    cell_id = tag_ids(document, "cell-shell")[0]
    table = tables(document)[0]
    document.clone_content_control(
        by_tag(document, "cell-shell"),
        rdocx.ContentLocation(table.story, "table", list(table.index_path) + [0]),
    )
    cell_ids = tag_ids(document, "cell-shell")
    assert cell_id in cell_ids and len(cell_ids) == 2 and len(set(cell_ids)) == 2


def test_f285_nested_transparent_payload_matches_destination():
    import rdocx

    row = '<w:tr><w:tc><w:tcPr><w:tcW w:w="0" w:type="auto"/></w:tcPr><w:p/></w:tc></w:tr>'
    run = "<w:r><w:t>kept</w:t></w:r>"

    def shell(tag, identifier, inner):
        return (
            f'<w:sdt><w:sdtPr><w:tag w:val="{tag}"/><w:id w:val="{identifier}"/>'
            f"</w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>"
        )

    grid = '<w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid>'
    row_nest = shell("row-nest", 1, shell("row-inner", 2, row))
    inline_ok = shell("inline-ok", 5, shell("inline-inner", 6, run))
    inline_bad = shell("inline-bad", 7, shell("inline-bad-inner", 8, row))
    body = (
        "<w:p><w:r><w:t>host</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>other</w:t></w:r></w:p>"
        f"<w:tbl>{grid}{row_nest}</w:tbl>"
        f"<w:tbl>{grid}<w:tr><w:tc><w:tcPr><w:tcW w:w=\"0\" w:type=\"auto\"/></w:tcPr><w:p/></w:tc></w:tr></w:tbl>"
        f"<w:p>{inline_ok}</w:p><w:p>{inline_bad}</w:p>"
    )

    def located(document, predicate):
        return _content_location(document, predicate)

    def by_tag(document, tag):
        needle = f'w:val="{tag}"'.encode()

        def predicate(item):
            if item.kind != "content_control":
                return False
            xml = item.xml
            tag_at = xml.find(needle)
            content_at = xml.find(b"<w:sdtContent")
            return tag_at != -1 and (content_at == -1 or tag_at < content_at)

        return located(document, predicate)

    def paragraph(document, text):
        return located(document, lambda item: item.kind == "paragraph" and item.text == text)

    def tables(document):
        return [
            rdocx.ContentLocation(item.story, item.kind, item.index_path)
            for item in document.story_items
            if item.kind == "table"
        ]

    failures = []
    for label, operation, tag, target in [
        ("clone row", "clone", "row-nest", "host"),
        ("move row", "move", "row-nest", "host"),
        ("clone inline", "clone", "inline-ok", "table"),
    ]:
        document = _replace_document_body(rdocx.Document(), body)
        before = document.to_bytes()
        source = by_tag(document, tag)
        destination = paragraph(document, "host") if target == "host" else tables(document)[1]
        try:
            if operation == "clone":
                document.clone_content_control(source, destination)
            else:
                document.move_content_control(source, destination)
        except rdocx.RdocxError as error:
            if "does not match" not in str(error) or document.to_bytes() != before:
                failures.append(f"{label}: {error} unchanged={document.to_bytes() == before}")
        else:
            failures.append(f"{label}: published unchanged={document.to_bytes() == before}")

    document = _replace_document_body(rdocx.Document(), body)
    before = document.to_bytes()
    try:
        document.unwrap_content_control(by_tag(document, "inline-bad"))
    except rdocx.RdocxError as error:
        if "parent cannot hold" not in str(error) or document.to_bytes() != before:
            failures.append(f"unwrap: {error} unchanged={document.to_bytes() == before}")
    else:
        failures.append(f"unwrap: published unchanged={document.to_bytes() == before}")

    document = _replace_document_body(rdocx.Document(), body)
    document.clone_content_control(by_tag(document, "row-nest"), tables(document)[1])
    ids = [control.id for control in document.content_controls() if control.tag == "row-nest"]
    if sorted(ids) == [1] or len(set(ids)) != 2:
        failures.append(f"legal row clone ids: {ids}")
    document.unwrap_content_control(by_tag(document, "inline-ok"))
    xml = _document_xml(document)
    if b"inline-ok" in xml or b"inline-inner" not in xml or b"kept" not in xml:
        failures.append("legal inline unwrap lost the nested run")
    assert failures == []


def test_f285_reference_text_matches_payload_boundary():
    import rdocx

    def shell(tag, identifier, inner):
        return (
            f'<w:sdt><w:sdtPr><w:tag w:val="{tag}"/><w:id w:val="{identifier}"/>'
            f"</w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>"
        )

    def document_for(inner):
        body = (
            "<w:p><w:r><w:t>host</w:t></w:r></w:p><w:p>"
            + shell("outer", 1, shell("inner", 2, inner))
            + "</w:p>"
        )
        return _replace_document_body(rdocx.Document(), body)

    def by_tag(document, tag):
        needle = f'w:val="{tag}"'.encode()

        def predicate(item):
            if item.kind != "content_control":
                return False
            xml = item.xml
            tag_at = xml.find(needle)
            content_at = xml.find(b"<w:sdtContent")
            return tag_at != -1 and (content_at == -1 or tag_at < content_at)

        return _content_location(document, predicate)

    def host(document):
        return _content_location(
            document, lambda item: item.kind == "paragraph" and item.text == "host"
        )

    failures = []
    for label, operation, inner, needle in [
        ("clone decimal", "clone", "&#65;", "does not match"),
        ("move hexadecimal", "move", "&#x41;", "does not match"),
        ("clone predefined", "clone", "&lt;", "does not match"),
        ("unwrap decimal", "unwrap", "&#65;", "parent cannot hold"),
    ]:
        document = document_for(inner)
        before = document.to_bytes()
        try:
            if operation == "clone":
                document.clone_content_control(by_tag(document, "outer"), host(document))
            elif operation == "move":
                document.move_content_control(by_tag(document, "outer"), host(document))
            else:
                document.unwrap_content_control(by_tag(document, "outer"))
        except rdocx.RdocxError as error:
            if needle not in str(error) or document.to_bytes() != before:
                failures.append(f"{label}: {error} unchanged={document.to_bytes() == before}")
        else:
            failures.append(f"{label}: published unchanged={document.to_bytes() == before}")

    document = document_for("&#32;")
    document.clone_content_control(by_tag(document, "outer"), host(document))
    ids = [control.id for control in document.content_controls() if control.tag == "outer"]
    if len(set(ids)) != 2 or 1 not in ids:
        failures.append(f"legal space clone ids: {ids}")
    document = document_for("<w:r><w:t>&#65;</w:t></w:r>")
    document.clone_content_control(by_tag(document, "outer"), host(document))
    if len({control.id for control in document.content_controls() if control.tag == "outer"}) != 2:
        failures.append("legal run reference clone did not keep two ids")
    assert failures == []


def test_f287_custom_xml_reference_forms_and_position():
    import rdocx

    document = rdocx.Document()
    failures = []
    for label, item, needle in [
        ("decimal before root", b"&#65;<root/>", "fragment"),
        ("predefined after root", b"<root/>&lt;", "fragment"),
        ("cdata space before root", b"<![CDATA[ ]]><root/>", "fragment"),
        ("uppercase attribute", b'<root a="&#X41;"/>', "not XML"),
        ("double hyphen in comment", b"<root><!--a--b--></root>", "not XML"),
    ]:
        fresh = rdocx.Document()
        before = fresh.to_bytes()
        try:
            fresh.add_custom_xml_store(item)
        except rdocx.RdocxError as error:
            if needle not in str(error) or fresh.to_bytes() != before:
                failures.append(f"{label}: {error} unchanged={fresh.to_bytes() == before}")
        else:
            failures.append(f"{label}: published unchanged={fresh.to_bytes() == before}")
    assert failures == []
    kept = [
        b"<!--note--><root/>",
        b"<root><!--note--></root>",
        b"<root>&#x41;</root>",
        b'<root a:b="1"/>',
        b"<root>&#32;</root>",
        b"<root><![CDATA[A]]></root>",
    ]
    for item in kept:
        document.add_custom_xml_store(item)
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as archive:
        for index, item in enumerate(kept, start=1):
            assert archive.read(f"customXml/item{index}.xml") == item


def test_f287_custom_xml_document_events_are_owned_and_atomic():
    import io
    import zipfile

    import rdocx

    for item in [
        b"<root><!DOCTYPE root></root>",
        b"<!DOCTYPE root><root/>",
        b"<root/><!DOCTYPE root>",
        b'<root><?xml version="1.0"?></root>',
        b' <?xml version="1.0"?><root/>',
        b'<?xml version="1.0" standalone="maybe"?><root/>',
        b"<?XmL note?><root/>",
        b"<root><?1bad note?></root>",
        b"<root>text]]></root>",
        b'<?xml version="1.0" encoding="UTF-16"?><root/>',
        "<?xml version='1.0' encoding='US-ASCII'?><root>é</root>".encode(),
        b'<root a="1"b="2"/>',
        b'<?xml version="1.0"encoding="UTF-8"?><root/>',
    ]:
        document = rdocx.Document()
        before = document.to_bytes()
        with pytest.raises(rdocx.RdocxError, match="not XML"):
            document.add_custom_xml_store(item)
        assert document.to_bytes() == before

    for item in [
        b'<?xml version="1.0" encoding="UTF-8"?><root/>',
        b"<?before ok?><root><?inside ok?></root><?after ok?>",
        b"<!--before--><root/><!--after-->",
        b'<root a:b:c="&lt;"/>',
    ]:
        document = rdocx.Document()
        document.add_custom_xml_store(item)
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as archive:
            assert archive.read("customXml/item1.xml") == item


def test_f287_custom_xml_names_values_and_references():
    import rdocx

    document = rdocx.Document()
    failures = []
    for label, item in [
        ("literal less-than", b'<root a="<"/>'),
        ("attribute name", b'<root 1bad="x"/>'),
        ("attribute entity", b'<root a="&bad;"/>'),
        ("element name", b"<1bad/>"),
        ("text entity", b"<root>&bad;</root>"),
        ("character reference", b'<root a="&#x0;"/>'),
    ]:
        before = document.to_bytes()
        try:
            document.add_custom_xml_store(item)
        except rdocx.RdocxError as error:
            if "not XML" not in str(error) or document.to_bytes() != before:
                failures.append(f"{label}: {error} unchanged={document.to_bytes() == before}")
        else:
            failures.append(f"{label}: published unchanged={document.to_bytes() == before}")
    assert failures == []
    opaque = b'<root a:b="1"/>'
    entity = b'<root a="&lt;">&#65;</root>'
    document.add_custom_xml_store(opaque)
    document.add_custom_xml_store(entity)
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert archive.read("customXml/item1.xml") == opaque
        assert archive.read("customXml/item2.xml") == entity


def test_fx204_tag_fill_keeps_alias_and_late_refusal_atomic():
    import rdocx
    document = _replace_document_body(rdocx.Document(), '<w:sdt><w:sdtPr><w:tag w:val="client"/><w:alias w:val="alias"/><w:text/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>before</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    before = document.to_bytes()
    for pairs in [[("client", "x"), ("missing", "y")], [("client", "x"), ("client", "y")], [("alias", "x")]]:
        with pytest.raises(rdocx.RdocxError):
            document.fill_content_control_tags(pairs)
        assert document.to_bytes() == before
    assert document.fill_content_control_tags([]) == ()
    assert document.to_bytes() == before
    assert document.fill_content_control_tags([("client", "Ada")]) == (1,)
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.content_controls()[0].text == "Ada"


def test_fx204_combined_fields_keeps_explicit_and_omitted_time_policy():
    import datetime
    import rdocx
    document = rdocx.Document()
    paragraph = document.add_paragraph("Page ")
    paragraph.add_run("").add_field("PAGE", "9")
    document.paragraphs[0].add_run("").add_field("NUMPAGES", "9")
    document.paragraphs[0].add_run("").add_field('DATE \\@ "yyyy-MM-dd"', "1999-01-01")
    raw = document.to_bytes()
    omitted = rdocx.Document.from_bytes(raw)
    report = omitted.update_all_fields()
    ordinary, pages = report.ordinary_updated, report.layout
    assert isinstance(report, rdocx.FieldsUpdateReport)
    assert isinstance(ordinary, int)
    assert isinstance(pages, rdocx.LayoutBackedFieldUpdateReport)
    assert report.total_updated == ordinary + pages.updated_count
    assert [evaluation.field_index for evaluation in report.evaluations] == [0, 1, 2]
    assert [evaluation.cached_result for evaluation in report.evaluations] == ["9", "9", "1999-01-01"]
    assert [evaluation.outcome["kind"] for evaluation in report.evaluations] == ["deferred_pagination", "deferred_pagination", "keep_stored"]
    assert report.diagnostics and "header-footer" in report.saved_cache_scope
    assert pages.page_fields == 1 and pages.num_pages_fields == 1
    assert omitted.paragraphs[0].text.endswith("1999-01-01")
    explicit = rdocx.Document.from_bytes(raw)
    report = explicit.update_all_fields(now=datetime.datetime(2026, 10, 10, 8, 30))
    ordinary, pages = report.ordinary_updated, report.layout
    assert ordinary >= 1 and pages.page_fields == 1
    assert report.total_updated == ordinary + pages.updated_count
    assert report.evaluations[2].outcome == {"kind": "resolved", "text": "2026-10-10"}
    assert isinstance(report.evaluations[2], rdocx.FieldEvaluation)
    assert explicit.paragraphs[0].text == "Page 112026-10-10"
    assert rdocx.Document.from_bytes(explicit.to_bytes()).paragraphs[0].text == "Page 112026-10-10"


def test_fx204_picture_media_facts_match_saved_owner():
    from rdocx import Document
    document = Document()
    image = _one_pixel_png()
    document.add_picture(image, "pixel.png", width=127000, height=127000)
    picture = document.pictures()[0]
    assert picture.media_part_name and picture.external_url is None
    assert picture.bytes == image
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert archive.read(picture.media_part_name.lstrip("/")) == image
    reopened = Document.from_bytes(document.to_bytes()).pictures()[0]
    assert reopened.media_part_name == picture.media_part_name
    assert reopened.external_url is None and reopened.bytes == image


def test_fx204_field_report_allocation_failure_discards_candidate(monkeypatch):
    import datetime
    import json
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("retained ").add_run("").add_field('DATE \\@ "yyyy-MM-dd"', "1999-01-01")
    held = document.paragraphs[0]
    before, text = document.to_bytes(), held.text

    def refuse_report_allocation(*args, **kwargs):
        raise MemoryError("injected prepublication evaluation result allocation")

    with monkeypatch.context() as patch:
        patch.setattr(json, "loads", refuse_report_allocation)
        with pytest.raises(MemoryError, match="prepublication"):
            document.update_all_fields(now=datetime.datetime(2026, 10, 10))
        assert document.to_bytes() == before and held.text == text

    report = document.update_all_fields(now=datetime.datetime(2026, 10, 10))
    assert isinstance(report, rdocx.FieldsUpdateReport)
    assert report.evaluations[0].cached_result == "1999-01-01"
    assert report.evaluations[0].outcome == {"kind": "resolved", "text": "2026-10-10"}
    assert rdocx.Document.from_bytes(document.to_bytes()).paragraphs[0].text == "retained 2026-10-10"

    empty = rdocx.Document()
    empty.add_paragraph("still live")
    held = empty.paragraphs[0]
    before = empty.to_bytes()
    report = empty.update_all_fields()
    assert report.total_updated == 0 and report.evaluations == ()
    assert held.text == "still live" and empty.to_bytes() == before


def test_fx222_public_empty_comment_and_noop_fields_preserve_source():
    import xml.etree.ElementTree as ET
    import rdocx

    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    seed = rdocx.Document()
    seed.add_paragraph("Commented text")
    comment_id = seed.add_comment_on_text("Commented", author="Ada", text="note")
    assert len(seed.comments) == 1
    source = io.BytesIO(seed.to_bytes())
    output = io.BytesIO()
    original_comments = None
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "word/comments.xml":
                root_start = data.index(b"<w:comments")
                content_start = data.index(b">", root_start) + 1
                empty = (
                    f'<q:comment xmlns:q="{namespace}" xmlns:x="urn:opaque" '
                    f'q:id="{comment_id + 1}" q:author="Ada" x:keep="original"/>'
                ).encode()
                data = data[:content_start] + empty + data[content_start:]
                parsed = ET.fromstring(data)
                assert len(parsed.findall(f"{{{namespace}}}comment")) == 2
                original_comments = data
            result_zip.writestr(info, data)
    assert original_comments is not None
    document = rdocx.Document.from_bytes(output.getvalue())
    held = document.paragraphs[0]
    before = document.to_bytes()
    comments = {comment.id: comment for comment in document.comments}
    assert set(comments) == {comment_id, comment_id + 1}
    assert comments[comment_id].text == "note"
    assert comments[comment_id + 1].text == ""
    assert comments[comment_id + 1].anchor is None
    assert document.story_items
    assert held.text == "Commented text"
    assert document.update_page_fields() == 0
    assert document.update_fields() == 0
    assert document.rebuild_toc().entry_count == 0
    assert held.text == "Commented text"
    assert document.to_bytes() == before
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert {comment.id for comment in reopened.comments} == set(comments)
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as saved:
        assert saved.read("word/comments.xml") == original_comments


def test_fx217_public_html_markdown_construct_open_and_diagnostics(tmp_path):
    import rdocx
    import xml.etree.ElementTree as ET

    image = _one_pixel_png()
    html = '<h1>Title</h1><p>Plain <strong>bold</strong> <em>italic</em> <mark>marked</mark></p><table><tr><td>cell</td></tr></table><p><img src="face.png" alt="specimen"></p>'
    document, diagnostics = rdocx.Document.from_html(html, images=(("face.png", image, "face.png"),))
    assert diagnostics == () and isinstance(diagnostics, tuple)
    assert len(document.tables) == 1 and document.tables[0].cell(0, 0).text == "cell"
    assert document.paragraphs[0].text == "Title"
    assert len(document.pictures()) == 1 and document.pictures()[0].bytes == image
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert root.find('.//w:b', ns) is not None and root.find('.//w:i', ns) is not None
    shade = root.find('.//w:shd', ns)
    assert shade is not None and shade.get('{'+ns['w']+'}fill') == 'FFFF00'
    source = tmp_path / "source.html"
    source.write_text(html)
    opened, _ = rdocx.Document.open_html(source, images=(("face.png", image, "face.png"),))
    assert opened.text() == document.text() and opened.pictures()[0].bytes == image
    markdown = '# Title\n\n> Quote **bold**\n\n3. third\n4. fourth\n\n| A | B |\n| --- | --- |\n| x | y |\n\n```rust\nlet x = 1;\n```\n\n![specimen](face.png)\n\n<br>\n'
    imported, notes = rdocx.Document.from_markdown(markdown, images=(("face.png", image, "face.png"),))
    assert imported.paragraphs[0].text == "Title" and len(imported.tables) == 1
    assert imported.pictures()[0].bytes == image and notes
    assert all(isinstance(note, rdocx.HtmlDiagnostic) for note in notes)
    assert all(note.location.startswith("bytes ") for note in notes)
    with pytest.raises(AttributeError):
        notes[0].message = "changed"
    source = tmp_path / "source.md"
    source.write_text(markdown)
    opened, opened_notes = rdocx.Document.open_markdown(source, images=(("face.png", image, "face.png"),))
    assert opened.text() == imported.text() and opened_notes == notes
    canonical = rdocx.Document.from_bytes(imported.to_bytes()).to_markdown()
    again, _ = rdocx.Document.from_markdown(canonical, images=(("face.png", image, "face.png"),))
    assert rdocx.Document.from_bytes(again.to_bytes()).to_markdown() == canonical


@pytest.mark.parametrize("owner", ("body", "cell", "header", "footer"))
def test_fx217_public_insertion_has_checked_owner_and_atomic_failure(owner):
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("Target")
    document.add_table(1, 1).cell(0, 0).text = "Cell target"
    document.add_paragraph("End")
    if owner == "body":
        paragraph = document.paragraphs[0]
    elif owner == "cell":
        paragraph = document.tables[0].cell(0, 0).paragraphs[0]
    else:
        story = getattr(document.section(0), owner)
        story.add_paragraph("Related target")
        paragraph = story.paragraphs[-1]
    before, text = document.to_bytes(), paragraph.text
    foreign = rdocx.Document().add_paragraph("Foreign")
    with pytest.raises(ValueError, match="live paragraph"):
        document.insert_html("<p>must not insert</p>", after=foreign)
    with pytest.raises(rdocx.RdocxError):
        document.insert_html("<p>must not insert</p>", after=paragraph, images=(("bad", b"bad image", "bad.png"),))
    assert document.to_bytes() == before and paragraph.text == text
    assert document.insert_html("", after=paragraph) == ()
    assert document.to_bytes() == before and paragraph.text == text
    assert document.insert_html('<p>Inserted <strong>new</strong></p>', after=paragraph) == ()
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    if owner == "body":
        assert [p.text for p in reopened.paragraphs] == ["Target", "Inserted new", "End"]
        assert reopened.tables[0].cell(0, 0).text == "Cell target"
    elif owner == "cell":
        assert [p.text for p in reopened.tables[0].cell(0, 0).paragraphs] == ["Cell target", "Inserted new"]
    else:
        assert [p.text for p in getattr(reopened.section(0), owner).paragraphs][-2:] == ["Related target", "Inserted new"]
    with pytest.raises(rdocx.StaleElementError):
        _ = paragraph.text


def test_fx217_public_markdown_template_preserves_graph_and_inputs(tmp_path):
    import rdocx
    template = rdocx.Document()
    template.add_paragraph("Discard template body")
    template.section(0).header.add_paragraph("Keep header")
    template.section(0).footer.add_paragraph("Keep footer")
    before = template.to_bytes()
    path = tmp_path / "template.docx"
    path.write_bytes(before)
    for supplied in (template, path):
        imported, _ = rdocx.Document.from_markdown("# Replacement\n\nBody", template=supplied)
        assert "Discard template body" not in imported.text() and imported.paragraphs[0].text == "Replacement"
        assert imported.section(0).header.paragraphs[-1].text == "Keep header"
        assert imported.section(0).footer.paragraphs[-1].text == "Keep footer"
        with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(io.BytesIO(imported.to_bytes())) as new:
            for name in old.namelist():
                if name.startswith(("word/header", "word/footer")):
                    assert old.read(name) == new.read(name)
    assert template.to_bytes() == before and path.read_bytes() == before
    invalid = tmp_path / "invalid.md"
    invalid.write_bytes(b"\xff")
    with pytest.raises(ValueError, match="UTF-8"):
        rdocx.Document.open_markdown(invalid, template=template)
    assert template.to_bytes() == before


def test_fx220_word_python_font_graph_snapshots_inputs_and_removal(tmp_path):
    from pathlib import Path
    from rdocx import Document, FontDefinition, FontEmbeddingLicense

    source = Path(__file__).resolve().parents[2] / "oxml-layout/fonts/Selawik-Regular.ttf"
    body = source.read_bytes()
    license = FontEmbeddingLicense(authorized=True, identity="SIL Open Font License 1.1")
    document = Document()
    document.add_paragraph("font control")
    document.set_font(FontDefinition("FX220 Python Word", alternate_name="control alias", family="swiss", pitch="variable"))
    for kind, value in zip(("regular", "bold", "italic", "bold_italic"), (body, source, io.BytesIO(body), body)):
        document.embed_font("FX220 Python Word", value, kind, license=license, subsetted=kind == "italic")
    font = next(font for font in document.fonts if font.name == "FX220 Python Word")
    assert isinstance(document.fonts, tuple) and isinstance(font.embedded_fonts, tuple)
    assert font.alternate_name == "control alias" and font.family == "swiss" and font.pitch == "variable"
    assert tuple(face.kind for face in font.embedded_fonts) == ("regular", "bold", "italic", "bold_italic")
    for face in font.embedded_fonts:
        assert face.data == body and face.license.authorized and face.license.identity == license.identity
        assert face.font_key.startswith("{") and face.font_key.endswith("}")
        with pytest.raises(AttributeError):
            face.subsetted = False
    assert font.embedded_fonts[2].subsetted
    saved = document.to_bytes()
    package = zipfile.ZipFile(io.BytesIO(saved))
    font_parts = [name for name in package.namelist() if name.startswith("word/fonts/")]
    assert len(font_parts) == 4 and all(package.read(name) != body for name in font_parts)
    reopened = Document.from_bytes(saved)
    assert next(font for font in reopened.fonts if font.name == "FX220 Python Word").embedded_fonts[0].data == body
    held = document.paragraphs[0]
    document.embed_font("FX220 Python Word", body, license=license)
    assert document.to_bytes() == saved and held.text == "font control"
    assert document.remove_embedded_font("FX220 Python Word", "italic").data == body
    assert document.remove_embedded_font("FX220 Python Word", "italic") is None
    removed = document.remove_font("FX220 Python Word")
    assert len(removed.embedded_fonts) == 3
    assert document.remove_font("FX220 Python Word") is None
    output = tmp_path / "removed.docx"
    document.save(output)
    assert all(font.name != "FX220 Python Word" for font in Document(output).fonts)


def test_fx220_word_python_license_flags_and_refusal_are_atomic():
    from pathlib import Path
    from rdocx import Document, FontDefinition, FontEmbeddingLicense

    body = (Path(__file__).resolve().parents[2] / "oxml-layout/fonts/Selawik-Regular.ttf").read_bytes()
    document = Document()
    document.add_paragraph("held")
    for name in ("embed_true_type_fonts", "embed_system_fonts", "save_subset_fonts"):
        assert getattr(document, name) is None
        for value in (False, True, None):
            setattr(document, name, value)
            assert getattr(document, name) is value
    baseline = document.to_bytes()
    held = document.paragraphs[0]
    good = FontEmbeddingLicense(authorized=True, identity="OFL 1.1")
    for value, kind, license, key in ((body, "regular", FontEmbeddingLicense(authorized=False, identity="OFL"), None), (body, "regular", FontEmbeddingLicense(authorized=True, identity=" "), None), (b"not a font", "regular", good, None), (body, "unknown", good, None), (body, "regular", good, "bad-key"), (io.StringIO("not bytes"), "regular", good, None)):
        with pytest.raises(Exception):
            document.embed_font("FX220 refusal", value, kind, license=license, font_key=key)
        assert document.to_bytes() == baseline and held.text == "held"
    class OversizedReader:
        def read(self, size):
            assert 0 < size <= 8192
            return b"x" * (64 * 1024 * 1024 + 1)
    with pytest.raises(ValueError):
        document.embed_font("FX220 refusal", OversizedReader(), license=good)
    assert document.to_bytes() == baseline and held.text == "held"
    document.set_font(FontDefinition("FX220 explicit"))
    document.embed_font("FX220 explicit", body, license=good, font_key="{00112233-4455-6677-8899-AABBCCDDEEFF}")
    assert next(font for font in document.fonts if font.name == "FX220 explicit").embedded_fonts[0].font_key == "{00112233-4455-6677-8899-AABBCCDDEEFF}"


@pytest.mark.parametrize("mode", ["short_chunks", "later_type", "later_error", "cumulative_bound"])
def test_fx220_rdocx_binary_readers_reach_real_eof_before_publication(mode):
    from pathlib import Path
    import rdocx

    body = (Path(__file__).resolve().parents[2] / "oxml-layout/fonts/Selawik-Regular.ttf").read_bytes()
    value = rdocx.Document()
    value.add_paragraph("binary owner")
    held = value.paragraphs[0]
    baseline = value.to_bytes()
    license = rdocx.FontEmbeddingLicense(authorized=True, identity="OFL 1.1")
    class Reader:
        def __init__(self):
            self.calls = 0
            self.offset = 0
        def read(self, size):
            self.calls += 1
            if mode == "short_chunks":
                part = body[self.offset:self.offset + min(size, 8192)]
                self.offset += len(part)
                return part
            if self.calls == 1:
                return body
            if mode == "later_type":
                return "invalid later text"
            if mode == "later_error":
                raise OSError("later font read failed")
            return b"x" * (1024 * 1024)
    reader = Reader()
    if mode == "short_chunks":
        value.embed_font("FX220 binary stream", reader, license=license)
        assert reader.calls > 2 and reader.offset == len(body)
        assert next(font for font in value.fonts if font.name == "FX220 binary stream").embedded_fonts[0].data == body
    else:
        if mode == "later_error":
            expected, match = OSError, "later font read failed"
        elif mode == "later_type":
            expected, match = TypeError, "reader must return bytes"
        else:
            expected, match = ValueError, "64 MiB"
        with pytest.raises(expected, match=match):
            value.embed_font("FX220 binary stream", reader, license=license)
        assert reader.calls > 1
        assert value.to_bytes() == baseline and held.text == "binary owner"

@pytest.mark.parametrize("method", ["pdf", "png", "svg", "all_pages", "selected_pages"])
def test_fx220_word_embedded_faces_and_caller_override_reach_exports(method):
    import base64
    import re
    from pathlib import Path
    from rdocx import Document, FontEmbeddingLicense

    directory = Path(__file__).resolve().parents[2] / "oxml-layout/fonts"
    kinds = ("regular", "bold", "italic", "bold_italic")
    styles = ("Regular", "Bold", "Italic", "BoldItalic")
    embedded = [(directory / f"Caladea-{style}.ttf").read_bytes() for style in styles]
    caller = [(directory / f"LiberationMono-{style}.ttf").read_bytes() for style in styles]
    name = "FX220 deterministic export face"
    document = Document()
    for index in range(4):
        run = document.add_paragraph("").add_run("iiiiWWWW font route")
        run.font.name = name
        run.font.bold = bool(index & 1)
        run.font.italic = bool(index & 2)
    license = FontEmbeddingLicense(authorized=True, identity="Apache License 2.0")
    for kind, body in zip(kinds, embedded):
        document.embed_font(name, body, kind, license=license)
    source = document.to_bytes()
    before = document.to_bytes()
    held = document.paragraphs[0]
    disabled = Document.from_bytes(source)
    disabled.embed_true_type_fonts = False
    missing = Document.from_bytes(source)
    missing.remove_font(name)
    reference = Document.from_bytes(missing.to_bytes())
    for value in (document, disabled, missing, reference):
        assert value.paragraphs[0].text == "iiiiWWWW font route"

    def invoke(value, fonts):
        if method == "pdf": return value.to_pdf(fonts=fonts)
        if method == "png": return value.render_page_to_png(0, dpi=72.0, fonts=fonts)
        if method == "svg": return value.render_page_to_svg(0, fonts=fonts).svg
        if method == "all_pages": return value.render_all_pages(dpi=72.0, fonts=fonts)
        return value.render_pages(pages=[0], dpi=72.0, fonts=fonts)

    positive = invoke(document, [])
    off = invoke(disabled, [])
    absent = invoke(missing, [])
    override = invoke(document, list(zip([name] * 4, caller)))
    expected = invoke(reference, list(zip([name] * 4, embedded)))
    expected_override = invoke(reference, list(zip([name] * 4, caller)))
    assert positive != off and positive != override
    assert off == absent
    if method in ("png", "all_pages", "selected_pages"):
        assert positive == expected
        assert override == expected_override
    elif method == "pdf":
        assert positive.startswith(b"%PDF") and b"Caladea" in positive
        assert b"LiberationMono" in override and b"Caladea" not in override
        assert b"Caladea" not in off
    else:
        def bodies(svg):
            return {base64.b64decode(value) for value in re.findall(r"base64,([A-Za-z0-9+/=]+)", svg)}
        assert set(embedded) <= bodies(positive)
        assert set(caller) <= bodies(override)
        assert not set(embedded) & bodies(off)
    assert document.to_bytes() == before and held.text == "iiiiWWWW font route"


@pytest.mark.parametrize("kind", [51, 65])
def test_fx215_word_public_creation_after_owner_and_atomic_refusal(kind):
    from rdocx import Document, CategoryChartData, Inches, Paragraph, StaleElementError
    data = CategoryChartData()
    data.categories = ["East", "West"]
    data.add_series("Sales", [11.0, 23.0])
    doc = Document()
    doc.add_paragraph("before")
    doc.add_table(1, 1)
    doc.add_paragraph("after")
    held_tail = doc.paragraphs[1]
    result = doc.add_chart(kind, data, Inches(4), Inches(2), after=doc.paragraphs[0])
    assert isinstance(result, Paragraph)
    assert result.element is result
    returned_xml = result.element.xml
    package = doc.to_bytes()
    assert result.element.replace_xml(returned_xml) is None
    assert result.element.xml == returned_xml and doc.to_bytes() == package
    with pytest.raises(StaleElementError, match="Document.add_chart"):
        _ = held_tail.text
    assert doc.paragraphs[2].text == "after"
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        names = archive.namelist()
        charts = [name for name in names if name.startswith("word/charts/chart") and name.endswith(".xml")]
        assert len(charts) == 1
        assert b"Sales" in archive.read(charts[0]) and b"23" in archive.read(charts[0])
        assert any(name.startswith("word/embeddings/") for name in names)
        from xml.etree import ElementTree as ET
        body = ET.fromstring(archive.read("word/document.xml")).find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
        assert [element.tag.rsplit("}", 1)[-1] for element in list(body)[:4]] == ["p", "p", "tbl", "p"]
    foreign = Document(); foreign.add_paragraph("foreign")
    with pytest.raises(ValueError):
        doc.add_chart(kind, data, Inches(4), Inches(2), after=foreign.paragraphs[0])
    assert doc.to_bytes() == package
    bad = CategoryChartData(); bad.categories = ["East"]; bad.add_series("Bad", [float("nan")])
    with pytest.raises(Exception):
        doc.add_chart(kind, bad, Inches(4), Inches(2))
    assert doc.to_bytes() == package

    assert result.element.xml == returned_xml
    reopened = Document.from_bytes(package)
    assert reopened.paragraphs[1].element.xml == returned_xml
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["before", "", "after"]

    appended = Document()
    held = appended.add_paragraph("append survivor")
    held_xml = held.element.xml
    appended_result = appended.add_chart(kind, data, Inches(4), Inches(2))
    assert held.text == "append survivor" and held.element.xml == held_xml
    appended_xml = appended_result.element.xml
    appended_bytes = appended.to_bytes()
    assert appended_result.element.replace_xml(appended_xml) is None
    assert appended.to_bytes() == appended_bytes
    with pytest.raises(Exception):
        appended.add_chart(kind, bad, Inches(4), Inches(2))
    assert appended.to_bytes() == appended_bytes
    assert held.element.xml == held_xml and appended_result.element.xml == appended_xml
    reopened_append = Document.from_bytes(appended_bytes)
    assert reopened_append.paragraphs[0].text == "append survivor"
    assert reopened_append.paragraphs[1].element.xml == appended_xml


def test_fx215_word_chart_refuses_retired_and_non_body_after_handles():
    from rdocx import Document, CategoryChartData, Inches, StaleElementError
    data = CategoryChartData(); data.categories = ["A"]; data.add_series("Series", [7.0])
    doc = Document(); doc.add_paragraph("body witness"); doc.add_paragraph("retire")
    retired = doc.paragraphs[1]
    assert doc.remove_content(doc.find_content_index(retired))
    before = doc.to_bytes()
    with pytest.raises(StaleElementError):
        doc.add_chart(51, data, Inches(4), Inches(2), after=retired)
    assert doc.to_bytes() == before
    doc.add_table(1, 1)
    cell = doc.tables[0].cell(0, 0).paragraphs[0]
    before = doc.to_bytes()
    with pytest.raises(ValueError, match="direct body"):
        doc.add_chart(51, data, Inches(4), Inches(2), after=cell)
    assert doc.to_bytes() == before
    header = doc.section(0).header.add_paragraph("header witness")
    before = doc.to_bytes()
    with pytest.raises(ValueError, match="direct body"):
        doc.add_chart(51, data, Inches(4), Inches(2), after=header)
    assert doc.to_bytes() == before


@pytest.mark.parametrize("kind", ["paragraph", "run", "table", "cell", "section"])
def test_fx211a_five_checked_xml_owners_preserve_noop_and_atomic_refusal(kind):
    import xml.etree.ElementTree as ET
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("before")
    document.add_table(1, 1)
    document.tables[0].cell(0, 0).text = "cell-before"
    owner = {
        "paragraph": lambda: document.paragraphs[0],
        "run": lambda: document.paragraphs[0].runs[0],
        "table": lambda: document.tables[0],
        "cell": lambda: document.tables[0].cell(0, 0),
        "section": lambda: document.section(0),
    }[kind]()
    expected = {"paragraph": "p", "run": "r", "table": "tbl", "cell": "tc", "section": "sectPr"}[kind]
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    assert owner.element is owner
    original_xml = owner.element.xml
    assert isinstance(original_xml, str)
    assert ET.fromstring(original_xml).tag == "{" + namespace + "}" + expected
    before = document.to_bytes()
    assert owner.element.replace_xml(original_xml) is None
    assert owner.element.xml == original_xml
    assert document.to_bytes() == before
    snapshot = document.story_items[0]
    malformed = [
        '<x:' + expected + ' xmlns:x="urn:foreign"/>',
        original_xml + '<w:p xmlns:w="' + namespace + '"/>',
        '<!DOCTYPE x [<!ENTITY injected "external">]>' + original_xml,
        original_xml[:-1],
    ]
    for xml in malformed:
        with pytest.raises(rdocx.RdocxError):
            owner.element.replace_xml(xml)
        assert document.to_bytes() == before
        assert owner.element.xml == original_xml
    if kind == "section":
        # Actual valid geometry mutation through the selected section owner.
        xml = re.sub(r'(w:pgSz\b[^>]*w:w=")[0-9]+', r'\g<1>13000', original_xml, count=1)
        assert xml != original_xml
    else:
        old = "cell-before" if kind in {"table", "cell"} else "before"
        xml = original_xml.replace(old, "after")
        assert xml != original_xml
    owner.element.replace_xml(xml)
    assert owner.element.xml != original_xml
    with pytest.raises(rdocx.StaleElementError):
        document.set_story_text(snapshot, "rejected stale snapshot")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    if kind in {"paragraph", "run"}:
        assert reopened.paragraphs[0].text == "after"
    elif kind in {"table", "cell"}:
        assert reopened.tables[0].cell(0, 0).text == "after"
    else:
        assert reopened.section(0).page_width == 13000 * 635
    assert rdocx.Document.from_bytes(reopened.to_bytes()).to_bytes() == reopened.to_bytes()


def test_fx211a_xml_child_views_retire_without_retiring_true_ancestors():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("selected")
    document.add_paragraph("sibling")
    paragraph = document.paragraphs[0]
    sibling = document.paragraphs[1]
    run = paragraph.runs[0]
    font = run.font
    color = font.color
    properties = paragraph.paragraph_format
    tabs = properties.tab_stops
    xml = paragraph.element.xml
    paragraph.element.replace_xml(xml)
    assert run.text == "selected"
    assert font.name is None
    assert properties.alignment is None
    with pytest.raises(rdocx.RdocxError):
        paragraph.element.replace_xml('<w:p xmlns:w="urn:wrong"/>')
    assert run.text == "selected"
    paragraph.element.replace_xml(xml.replace("selected", "replacement"))
    assert paragraph.text == "replacement"
    assert sibling.text == "sibling"
    for getter in [lambda: run.text, lambda: font.name, lambda: color.rgb, lambda: properties.alignment, lambda: len(tabs)]:
        with pytest.raises(rdocx.StaleElementError):
            getter()
    fresh = paragraph.runs[0]
    assert fresh.text == "replacement"
    fresh_font = fresh.font
    fresh.text = "ordinary"
    assert paragraph.text == "ordinary"
    assert fresh_font.name is None


def test_fx211a_section_furniture_retirement_preserves_shared_physical_part():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body")
    document.insert_section(1)
    document.section(0).header.add_paragraph("shared header")
    story = next(story for story in document.stories if story.kind == "header")
    document.link_section_story(1, "header", "default", story)
    first = document.section(0)
    second = document.section(1)
    first_header = first.header
    second_header = second.header
    selected = first_header.paragraphs[0]
    unselected = second_header.paragraphs[0]
    selected_run = selected.runs[0]
    unselected_run = unselected.runs[0]
    properties = selected.paragraph_format
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        physical_before = archive.read(story.part_name.lstrip("/"))
    xml = first.element.xml
    first.element.replace_xml(xml)
    assert selected.text == unselected.text == "shared header"
    with pytest.raises(rdocx.RdocxError):
        first.element.replace_xml('<w:sectPr xmlns:w="urn:foreign"/>')
    assert selected_run.text == unselected_run.text == "shared header"
    replacement, count = re.subn(r'<w:headerReference\b[^>]*/>', "", xml, count=1)
    assert count == 1
    first.element.replace_xml(replacement)
    for getter in [lambda: first_header.paragraphs, lambda: selected.text, lambda: selected_run.text, lambda: properties.alignment]:
        with pytest.raises(rdocx.StaleElementError):
            getter()
    assert second.page_width is not None
    assert second_header.paragraphs[0].text == "shared header"
    assert unselected.text == unselected_run.text == "shared header"
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert archive.read(story.part_name.lstrip("/")) == physical_before
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.section(0).header.paragraphs == ()
    assert reopened.section(1).header.paragraphs[0].text == "shared header"


def test_fx211a_run_text_retires_picture_with_first_cause_and_keeps_rpr():
    import rdocx

    def chunk(kind, body):
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xffffffff)

    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b"")
    document = rdocx.Document()
    document.add_paragraph("before")
    document.add_paragraph("sibling")
    picture = document.paragraphs[0].runs[0].add_picture(png, "red.png", width=914400, height=914400)
    paragraph = document.paragraphs[0]
    sibling = document.paragraphs[1]
    run = paragraph.runs[0]
    font = run.font
    font.bold = True
    before = document.to_bytes()
    with pytest.raises(TypeError):
        run.text = object()
    assert document.to_bytes() == before
    assert run.text == "before"
    assert font.bold is True
    run.text = "replacement"
    assert paragraph.text == run.text == "replacement"
    assert sibling.text == "sibling"
    assert font.bold is True
    with pytest.raises(rdocx.StaleElementError, match="Run.text"):
        document.set_picture_size(picture, width=457200, height=457200)
    unchanged = document.to_bytes()
    run.text = "replacement"
    assert document.to_bytes() == unchanged
    assert run.text == "replacement"
    assert font.bold is True
    current_run = document.add_paragraph("appended").runs[0]
    with pytest.raises(rdocx.StaleElementError, match="Document.add_paragraph"):
        _ = run.text
    assert current_run.text == "appended"
    assert paragraph.text == "replacement"
    assert sibling.text == "sibling"


def test_fx211a_imported_local_picture_run_xml_preserves_owned_media_and_opaque():
    import xml.etree.ElementTree as ET
    import rdocx

    image = _one_pixel_png()
    seed = rdocx.Document()
    seed.add_paragraph("seed").runs[0].add_picture(image, "owned.png", width=152400, height=152400)
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
    relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    seed_root = ET.fromstring(_document_xml(seed))
    rel_id = seed_root.find(".//{" + drawing + "}blip").get("{" + relationships + "}embed")
    assert rel_id
    opaque = '<x:payload xmlns:x="urn:producer" x:keep=\'yes\'/>'
    source = f'<w:document xmlns:w="{namespace}"><w:body><w:p><w:r><w:t>before</w:t></w:r><w:r><w:drawing xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="{drawing}" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" xmlns:r="{relationships}"><wp:inline><wp:extent cx="152400" cy="152400"/><wp:docPr id="1" name="owned.png"/>{opaque}<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="1" name="owned.png"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="152400" cy="152400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p><w:sectPr/></w:body></w:document>'
    package = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(seed.to_bytes())) as original:
        with zipfile.ZipFile(package, "w") as output:
            for member in original.infolist():
                output.writestr(member, source.encode() if member.filename == "word/document.xml" else original.read(member.filename))
    document = rdocx.Document.from_bytes(package.getvalue())
    paragraph = document.paragraphs[0]
    run = paragraph.runs[1]
    xml = run.element.xml
    root = ET.fromstring(xml)
    assert root.tag == "{" + namespace + "}r"
    assert root.find(".//{" + drawing + "}blip").get("{" + relationships + "}embed") == rel_id
    assert root.find("{" + namespace + "}rPr") is None
    assert opaque in xml
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        run.element.replace_xml(xml.replace(rel_id, "rId_missing"))
    assert document.to_bytes() == before
    insertion = xml.index(">") + 1
    changed = xml[:insertion] + "<w:rPr><w:b/></w:rPr>" + xml[insertion:]
    run.element.replace_xml(changed)
    assert run.font.bold is True
    assert paragraph.text == "before"
    for _ in range(2):
        payload = document.to_bytes()
        assert opaque.encode() in _document_xml(document)
        document = rdocx.Document.from_bytes(payload)
        assert document.paragraphs[0].runs[1].font.bold is True
        assert len(document.pictures()) == 1 and document.pictures()[0].bytes == image


@pytest.mark.parametrize("structured", [False, True])
def test_fx211a_constructor_inventory_preserves_direct_control_and_cell_joins(structured):
    import rdocx

    def paragraph(text):
        return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"

    middle = paragraph("middle")
    if structured:
        middle = "<w:sdt><w:sdtPr/><w:sdtContent>" + middle + "</w:sdtContent></w:sdt>"
    inner = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w='1000'/></w:tblGrid><w:tr><w:tc><w:tcPr/>" + paragraph("inner") + "</w:tc></w:tr></w:tbl>"
    field = "<w:p><w:fldSimple w:instr='PAGE'><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>"
    table = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w='1000'/></w:tblGrid><w:tr><w:tc><w:tcPr/>" + field + paragraph("cell-second") + inner + paragraph("cell-last") + "</w:tc></w:tr></w:tbl>"
    document = _replace_document_body(rdocx.Document(), paragraph("first") + middle + table + paragraph("tail") + "<w:sectPr/>")
    first, middle_owner, tail = document.paragraphs
    outer = document.tables[0]
    cell = outer.cell(0, 0)
    second, last = cell.paragraphs[1:]
    inner_owner = cell.tables[0].cell(0, 0).paragraphs[0]
    assert [first.text, middle_owner.text, tail.text] == ["first", "middle", "tail"]
    assert [second.text, last.text, inner_owner.text] == ["cell-second", "cell-last", "inner"]
    document.add_paragraph("appended")
    assert [first.text, middle_owner.text, tail.text] == ["first", "middle", "tail"]
    assert [second.text, last.text, inner_owner.text] == ["cell-second", "cell-last", "inner"]
    second.element.replace_xml(second.element.xml.replace("cell-second", "changed-second"))
    assert first.text == "first" and middle_owner.text == "middle" and tail.text == "tail"
    assert last.text == "cell-last" and inner_owner.text == "inner"
    assert outer.cell(0, 0).paragraphs[1].text == "changed-second"
    for _ in range(2):
        document = rdocx.Document.from_bytes(document.to_bytes())
        assert [p.text for p in document.paragraphs] == ["first", "middle", "tail", "appended"]
        cell = document.tables[0].cell(0, 0)
        assert [p.text for p in cell.paragraphs] == ["1", "changed-second", "cell-last"]
        assert cell.tables[0].cell(0, 0).paragraphs[0].text == "inner"



def test_fx211c_word_hyperlink_capture_survives_unrelated_append_and_reopens():
    from rdocx import Document
    document = Document()
    paragraph = document.add_paragraph("")
    paragraph.add_hyperlink("one", "https://example.com/one")
    held = document.hyperlinks[0]
    paragraph.add_hyperlink("two", "https://example.com/two")
    document.set_hyperlink_url(held, "https://example.com/edited")
    reopened = Document.from_bytes(document.to_bytes())
    assert [link.url for link in reopened.hyperlinks] == ["https://example.com/edited", "https://example.com/two"]


def test_fx211c_word_identical_hyperlink_removal_cannot_resurrect_capture():
    from rdocx import Document
    document = Document()
    paragraph = document.add_paragraph("")
    paragraph.add_hyperlink("same", "https://example.com/same")
    paragraph.add_hyperlink("same", "https://example.com/same")
    removed, survivor = document.hyperlinks
    document.remove_hyperlink(removed)
    paragraph.add_hyperlink("same", "https://example.com/same")
    before = document.to_bytes()
    with pytest.raises(Exception, match="removed|stale|hyperlink"):
        document.set_hyperlink_url(removed, "https://example.com/unsafe")
    assert document.to_bytes() == before
    document.set_hyperlink_url(survivor, "https://example.com/survivor")
    reopened = Document.from_bytes(document.to_bytes())
    assert [link.url for link in reopened.hyperlinks] == ["https://example.com/survivor", "https://example.com/same"]


def test_fx211c_word_table_append_preserves_owner_and_retires_old_grid():
    from rdocx import Document, StaleElementError
    document = Document()
    table = document.add_table(1, 2)
    cell = table.cell(0, 0)
    cell.text = "before"
    new_row = table.add_row()
    assert len(table.rows) == 2
    assert len(new_row.cells) == 2
    before = document.to_bytes()
    with pytest.raises(StaleElementError, match="Table.add_row"):
        cell.text = "unsafe"
    assert document.to_bytes() == before
    reopened = Document.from_bytes(before)
    assert len(reopened.tables[0].rows) == 2
    assert reopened.tables[0].cell(0, 0).text == "before"



def test_fx211c_word_captured_link_survives_append_with_fresh_paragraph_setup():
    from rdocx import Document
    document = Document()
    document.add_paragraph("").add_hyperlink("one", "https://example.com/one")
    held = document.hyperlinks[0]
    document.paragraphs[0].add_hyperlink("two", "https://example.com/two")
    document.set_hyperlink_url(held, "https://example.com/edited")
    assert [link.url for link in document.hyperlinks] == ["https://example.com/edited", "https://example.com/two"]



def test_fx211c_word_hyperlink_foreign_manual_and_noop_authority():
    from rdocx import Document, Hyperlink
    document = Document()
    document.add_paragraph("").add_hyperlink("same", "https://example.com/same")
    held = document.hyperlinks[0]
    manual = Hyperlink(story=held.story, index_path=held.index_path, text=held.text, url=held.url,
                      anchor=held.anchor, relationship_id=held.relationship_id)
    assert manual == held
    before = document.to_bytes()
    document.set_hyperlink_url(manual, "https://example.com/same")
    assert document.to_bytes() == before
    foreign = Document()
    foreign.add_paragraph("").add_hyperlink("same", "https://example.com/same")
    foreign_before = foreign.to_bytes()
    with pytest.raises(ValueError, match="different document"):
        foreign.set_hyperlink_url(held, "https://example.com/unsafe")
    assert foreign.to_bytes() == foreign_before
    document.set_hyperlink_url(held, "https://example.com/changed")
    document.set_hyperlink_url(held, "https://example.com/changed-again")
    assert Document.from_bytes(document.to_bytes()).hyperlinks[0].url == "https://example.com/changed-again"


def test_fx211c_word_removed_link_first_cause_survives_identical_replacements():
    from rdocx import Document, StaleElementError
    document = Document()
    paragraph = document.add_paragraph("")
    paragraph.add_hyperlink("same", "https://example.com/same")
    removed = document.hyperlinks[0]
    document.remove_hyperlink(removed)
    paragraph.add_hyperlink("same", "https://example.com/same")
    for i in range(32):
        document.add_paragraph(str(i))
    before = document.to_bytes()
    with pytest.raises(StaleElementError) as error:
        document.set_hyperlink_url(removed, "https://example.com/unsafe")
    assert "Document.remove_hyperlink" in str(error.value)
    assert "Document.add_paragraph" not in str(error.value)
    assert document.to_bytes() == before
    assert Document.from_bytes(before).hyperlinks[0].url == "https://example.com/same"


def test_fx211c_word_run_keeps_first_global_cause_and_manual_item_equality():
    from rdocx import Document, StaleElementError, StoryItem
    document = Document()
    paragraph = document.add_paragraph("old")
    run = paragraph.add_run(" run")
    item = document.story_items[0]
    manual = StoryItem(story=item.story, kind=item.kind, index_path=item.index_path,
                       direct_body_index=item.direct_body_index, text=item.text, xml=item.xml, revision=item.revision)
    assert manual == item
    before = document.to_bytes()
    assert document.replace_text_at(manual, "absent", "unused", expect=0) == 0
    with pytest.raises(Exception, match="count|expected"):
        document.replace_text_at(item, "old", "unsafe", expect=2)
    assert document.to_bytes() == before
    foreign = Document()
    foreign.add_paragraph("old run")
    foreign_before = foreign.to_bytes()
    with pytest.raises(ValueError, match="different document"):
        foreign.replace_text_at(item, "old", "unsafe", expect=1)
    assert foreign.to_bytes() == foreign_before
    document.add_paragraph("first invalidating commit")
    document.add_table(1, 1)
    before = document.to_bytes()
    with pytest.raises(StaleElementError) as error:
        run.text = "unsafe"
    assert "Document.add_paragraph" in str(error.value)
    assert "Document.add_table" not in str(error.value)
    assert document.to_bytes() == before
    assert paragraph.text == "old run"
    reopened = Document.from_bytes(before)
    assert [p.text for p in reopened.paragraphs] == ["old run", "first invalidating commit"]


def test_fx211c_word_grid_retirement_is_local_and_invalid_access_atomic():
    from rdocx import Document, StaleElementError
    document = Document()
    first = document.add_table(1, 2)
    second = document.add_table(1, 1)
    second_cell = second.cell(0, 0)
    second_cell.text = "unrelated"
    old_cell = first.cell(0, 0)
    first.add_row()
    assert len(first.rows) == 2
    assert second_cell.text == "unrelated" and len(second.rows) == 1
    before = document.to_bytes()
    with pytest.raises(StaleElementError, match="Table.add_row"):
        old_cell.text = "unsafe"
    assert document.to_bytes() == before
    first.cell(1, 1).text = "new row"
    before = document.to_bytes()
    with pytest.raises(Exception):
        first.cell(-100, 0).text = "unsafe"
    assert document.to_bytes() == before
    assert second_cell.text == "unrelated"
    reopened = Document.from_bytes(before)
    assert reopened.tables[0].cell(1, 1).text == "new row"
    assert reopened.tables[1].cell(0, 0).text == "unrelated"


@pytest.mark.parametrize("kind", ["table", "cell", "section"])
def test_fx211d_opaque_word_owners_preserve_corpus_fragments(kind):
    import xml.etree.ElementTree as ET
    from rdocx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "XML seed"
    document.add_paragraph("XML sibling preserved")
    owner = table if kind == "table" else table.cell(0, 0) if kind == "cell" else document.section(0)
    owner.element.replace_xml(ET.tostring(ET.fromstring(owner.element.xml), encoding="unicode"))
    original = owner.element.xml
    opening = re.search(r"<(?![?!])[^>]+>", original)
    stamp = ' xmlns:corpus="urn:s91:corpus" corpus:stamp="opaque-attribute"'
    seeded = original[:opening.end() - 1] + stamp + original[opening.end() - 1:]
    opaque = '<corpus:keep key="one"/><corpus:keep key="two"/>'
    closing = seeded.rfind("</")
    owner.element.replace_xml(seeded[:closing] + opaque + seeded[closing:])
    exact = owner.element.xml
    assert 'corpus:stamp="opaque-attribute"' in exact and opaque in exact
    if kind == "section":
        changed = re.sub(r'((?:w:)?left=")[^"]+(" )', r'\g<1>1600\2', exact, count=1)
        assert changed != exact
    else:
        changed = exact.replace("XML seed", "XML changed", 1)
    owner.element.replace_xml(changed)
    assert opaque in owner.element.xml
    saved = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(saved)) as archive:
        body = archive.read("word/document.xml")
    assert opaque.encode() in body and b'corpus:stamp="opaque-attribute"' in body
    assert b"XML sibling preserved" in body
    reopened = Document.from_bytes(saved)
    selected = reopened.tables[0] if kind == "table" else reopened.tables[0].cell(0, 0) if kind == "cell" else reopened.section(0)
    assert opaque in selected.element.xml


@pytest.mark.parametrize("case", ["margin", "header", "notes", "equation", "merge", "style"])
def test_fx211d_word_current_compatibility_repairs(case):
    from rdocx import Document, Inches

    document = Document()
    if case == "margin":
        section = document.section(0)
        section.left_margin = Inches(1.25)
        assert int(section.left_margin) == int(Inches(1.25))
        reopened = Document.from_bytes(document.to_bytes())
        assert int(reopened.section(0).left_margin) == int(Inches(1.25))
    elif case == "header":
        section = document.section(0)
        section.different_first_page_header_footer = True
        section.header.add_paragraph("Default header")
        section.first_page_header.add_paragraph("First header")
        section.footer.add_paragraph("Footer")
        with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
            assert any(b"Default header" in archive.read(name) for name in archive.namelist() if name.startswith("word/header"))
            assert any(b"Footer" in archive.read(name) for name in archive.namelist() if name.startswith("word/footer"))
    elif case == "notes":
        paragraph = document.add_paragraph("Host ")
        run = paragraph.add_run("reference")
        document.add_footnote(run, "Footnote corpus")
        document.add_endnote(paragraph, "Endnote corpus")
        assert paragraph.text.startswith("Host")
        with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
            assert b"Footnote corpus" in archive.read("word/footnotes.xml")
            assert b"Endnote corpus" in archive.read("word/endnotes.xml")
    elif case == "equation":
        paragraph = document.add_paragraph("Equation ")
        assert paragraph.add_equation("x^2", format="latex").diagnostics == ()
        assert len(paragraph.equations()) == 1
        with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
            assert b"oMath" in archive.read("word/document.xml")
    elif case == "merge":
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "origin"
        merged = table.cell(0, 0).merge(table.cell(1, 1))
        merged.split()
        reopened = Document.from_bytes(document.to_bytes())
        assert len(reopened.tables[0].rows) == 2
        assert all(len(row.cells) == 2 for row in reopened.tables[0].rows)
        assert "origin" in reopened.tables[0].cell(0, 0).text
    else:
        table = document.add_table(rows=2, cols=2)
        table.style = "Light Shading Accent 1"
        assert table.style == "LightShading-Accent1"
        reopened = Document.from_bytes(document.to_bytes())
        assert reopened.tables[0].style == "LightShading-Accent1"
        with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
            assert b'"LightShading-Accent1"' in archive.read("word/styles.xml")


@pytest.mark.parametrize("kind", ["properties", "borders"])
def test_fx211d_merge_opaque_refusal_preserves_package_and_owners(kind):
    from rdocx import Document, RdocxError

    document = Document()
    table = document.add_table(1, 2)
    first = table.cell(0, 0)
    first.text = "first"
    second = table.cell(0, 1)
    second.text = "second"
    opaque = '<x:keep xmlns:x="urn:opaque" marker="preserve"/>'
    properties = '<w:tcBorders>' + opaque + '</w:tcBorders>' if kind == "borders" else opaque
    xml = '<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:tcPr>' + properties + '</w:tcPr><w:p><w:r><w:t>second</w:t></w:r></w:p></w:tc>'
    second.element.replace_xml(xml)
    first = table.cell(0, 0)
    before = document.to_bytes()
    assert second.merge(second) is second
    assert document.to_bytes() == before
    with pytest.raises(RdocxError, match="opaque cell metadata"):
        first.merge(second)
    assert document.to_bytes() == before
    assert first.text == "first" and second.text == "second"
    assert len(table.rows[0].cells) == 2
    assert opaque in second.element.xml
    reopened = Document.from_bytes(before)
    assert opaque in reopened.tables[0].cell(0, 1).element.xml


@pytest.mark.parametrize("properties", [
    '<w:tag w:val="goog_rdk_1"/><w:id w:val="1209466531"/>',
    '<w:id w:val="-1482736121"/><w:docPartObj>'
    '<w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj>',
])
def test_preserved_control_paragraphs_keep_checked_source_and_ordinary_reads(properties):
    import xml.etree.ElementTree as ET

    import rdocx

    body = (
        '<w:p><w:r><w:t>outside</w:t></w:r></w:p>'
        f'<w:sdt><w:sdtPr>{properties}</w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:t>preserved</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'
    )
    document = _replace_document_body(rdocx.Document(), body)
    before = document.to_bytes()
    before_xml = _document_xml(document)
    assert [paragraph.text for paragraph in document.paragraphs] == ["outside", "preserved"]
    held = document.paragraphs[1]
    outside = document.paragraphs[0]
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    original_xml = held.element.xml
    root = ET.fromstring(original_xml)
    assert root.tag == "{" + namespace + "}p"
    assert "".join(root.itertext()) == "preserved"
    preserved_properties = re.search(rb"<w:sdtPr>.*?</w:sdtPr>", before_xml, re.S).group(0)
    assert held.element.replace_xml(original_xml) is None
    assert document.to_bytes() == before
    for invalid in (
        '<x:p xmlns:x="urn:foreign"/>',
        original_xml + '<w:p xmlns:w="' + namespace + '"/>',
    ):
        with pytest.raises(rdocx.RdocxError):
            held.element.replace_xml(invalid)
        assert document.to_bytes() == before
        assert held.element.xml == original_xml
        assert held.text == "preserved" and outside.text == "outside"
    # The physical two-segment route owns this paragraph, not its sdtPr wrapper.
    alias = (
        '<x:p xmlns:x="' + namespace + '">'
        '<x:r><x:t>updated</x:t></x:r></x:p>'
    )
    assert held.element.replace_xml(alias) is None
    assert held.text == "updated" and outside.text == "outside"
    after_xml = _document_xml(document)
    assert re.search(rb"<w:sdtPr>.*?</w:sdtPr>", after_xml, re.S).group(0) == preserved_properties
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["outside", "updated"]
    assert _document_xml(reopened) == after_xml
    run = held.runs[0]
    document.add_paragraph("later")
    assert outside.text == "outside" and held.text == "updated"
    with pytest.raises(rdocx.StaleElementError, match="Document.add_paragraph"):
        _ = run.text
    assert [paragraph.text for paragraph in document.paragraphs] == ["outside", "updated", "later"]


def test_fx231_installed_word_html_whitespace_entrypoints(tmp_path):
    import rdocx
    import xml.etree.ElementTree as ET

    html = '<p>a <b>b </b> c&nbsp;&nbsp;d</p><p><a href="https://example.test/">x </a><i> </i> y</p>'
    source = tmp_path / "source.html"
    source.write_text(html, encoding="utf-8")
    inserted = rdocx.Document()
    anchor = inserted.add_paragraph("anchor")
    assert inserted.insert_html(html, after=anchor) == ()
    created, diagnostics = rdocx.Document.from_html(html)
    opened, opened_diagnostics = rdocx.Document.open_html(source)
    assert diagnostics == opened_diagnostics == ()
    for document in (created, opened, inserted):
        output = tmp_path / "saved.docx"
        document.save(output)
        reopened = rdocx.Document(output)
        with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            relations = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = root.findall("./w:body/w:p", ns)[-2:]
        first = [("".join(run.itertext()), run.find("w:rPr/w:b", ns) is not None) for run in paragraphs[0].findall("w:r", ns)]
        assert first == [("a ", False), ("b ", True), ("c\u00a0\u00a0d", False)]
        link = paragraphs[1].find("w:hyperlink", ns)
        assert "".join(link.itertext()) == "x "
        relation_id = link.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        assert next(item.attrib["Target"] for item in relations if item.attrib["Id"] == relation_id) == "https://example.test/"
        assert "".join(paragraphs[1].itertext()) == "x y"


def test_fx237_table_open_edit_and_fill_costs():
    import hashlib
    import json
    import os
    import platform
    import statistics
    import sys
    from pathlib import Path
    import rdocx

    output = os.environ.get("FX237_TIMING_OUTPUT")
    results = {}
    for cells in (60, 120, 240):
        rows = cells // 6
        fills = []
        # One discarded warm-up, then three measured fills for qualification.
        for trial in range(4 if output else 1):
            document = rdocx.Document()
            document.add_table(rows, 6)
            started = time.perf_counter()
            for row in range(rows):
                for column in range(6):
                    document.tables[0].cell(row, column).text = f"r{row} c{column}"
            elapsed = time.perf_counter() - started
            if trial:
                fills.append(elapsed)
            data = document.to_bytes()
            reopened = rdocx.Document.from_bytes(data)
            assert [[reopened.tables[0].cell(row, column).text for column in range(6)]
                    for row in range(rows)] == [[f"r{row} c{column}" for column in range(6)]
                                               for row in range(rows)]
        fill_output_sha256 = hashlib.sha256(data).hexdigest()
        input_directory = os.environ.get("FX237_TIMING_INPUTS") if output else None
        fixture_path = None
        if input_directory:
            directory = Path(input_directory)
            directory.mkdir(parents=True, exist_ok=True)
            fixture_path = directory / f"{cells}-cells.docx"
            if not fixture_path.exists():
                assert not os.environ.get("FX237_TIMING_BASELINE"), "candidate must consume authenticated baseline input"
                fixture_path.write_bytes(data)
            data = fixture_path.read_bytes()
        input_sha256 = hashlib.sha256(data).hexdigest()
        opens, edits = [], []
        for trial in range(6 if output else 1):
            started = time.perf_counter()
            document = rdocx.Document.from_bytes(data)
            opened = time.perf_counter() - started
            started = time.perf_counter()
            document.tables[0].cell(0, 0).text = "edited"
            edited = time.perf_counter() - started
            assert rdocx.Document.from_bytes(document.to_bytes()).tables[0].cell(0, 0).text == "edited"
            if trial:
                opens.append(opened)
                edits.append(edited)
        if output:
            results[str(cells)] = {"fill": fills, "open": opens, "edit": edits,
                                   "input_sha256": input_sha256, "input_path": str(fixture_path) if fixture_path else None,
                                   "fill_output_sha256": fill_output_sha256,
                                   "medians": {"fill": statistics.median(fills),
                                               "open": statistics.median(opens),
                                               "edit": statistics.median(edits)}}
    if output:
        controls = {"no_table_open": [], "unrelated_replace_text": []}
        plain = rdocx.Document()
        plain.add_paragraph("control")
        plain_data = plain.to_bytes()
        related = rdocx.Document.from_bytes(data)
        related.add_paragraph("unrelated control")
        related_data = related.to_bytes()
        for trial in range(6):
            started = time.perf_counter()
            assert rdocx.Document.from_bytes(plain_data).paragraphs[0].text == "control"
            plain_time = time.perf_counter() - started
            related = rdocx.Document.from_bytes(related_data)
            paragraph = related.paragraphs[-1]
            started = time.perf_counter()
            paragraph.replace_text("unrelated", "changed")
            related_time = time.perf_counter() - started
            assert related.paragraphs[-1].text == "changed control"
            if trial:
                controls["no_table_open"].append(plain_time)
                controls["unrelated_replace_text"].append(related_time)
        from importlib.metadata import version
        payload = {"version": version("rdocx"), "module": rdocx.__file__,
                   "python": sys.version, "executable": sys.executable,
                   "platform": platform.platform(), "machine": platform.machine(),
                   "results": results, "controls": controls}
        baseline_path = os.environ.get("FX237_TIMING_BASELINE")
        if baseline_path:
            baseline = json.loads(Path(baseline_path).read_text())
            payload["baseline_comparison"] = {
                cells: {phase: {"limit": max(3 * baseline["results"][cells]["medians"][phase], floor),
                                "within_limit": actual["medians"][phase] <= max(3 * baseline["results"][cells]["medians"][phase], floor),
                                "enforced": phase != "fill"}
                        for phase, floor in (("open", .010), ("edit", .030), ("fill", .500))}
                for cells, actual in results.items()
            }
        Path(output).write_text(json.dumps(payload, indent=2) + "\n")
        if baseline_path:
            baseline = json.loads(Path(baseline_path).read_text())
            assert baseline["version"] == "0.16.0"
            assert baseline["python"] == payload["python"]
            assert baseline["platform"] == payload["platform"]
            for cells, actual in results.items():
                assert actual["input_sha256"] == baseline["results"][cells]["input_sha256"]
                for phase, floor in (("open", .010), ("edit", .030), ("fill", .500)):
                    limit = max(3 * baseline["results"][cells]["medians"][phase], floor)
                    # Fill retains its former 3x comparison as diagnostic data.
                    # The approved plan records the exact user-accepted result.
                    if phase != "fill":
                        assert actual["medians"][phase] <= limit, (cells, phase, actual, limit)


def test_fx237_table_owner_inventory_preserves_physical_routes_and_generations():
    import rdocx

    document = rdocx.Document()
    paragraph = document.add_paragraph("unrelated")
    outer = document.add_table(2, 2)
    outer.cell(0, 0).text = "first"
    outer.cell(0, 1).text = "same"
    outer.cell(1, 1).text = "same"
    outer.cell(0, 0).add_table(1, 1).cell(0, 0).text = "nested"
    with pytest.raises(Exception, match="Cell.add_table|stale"):
        _ = outer.cell(0, 0)
    outer = document.tables[0]
    paragraph = document.paragraphs[0]
    nested = outer.cell(0, 0).tables[0].cell(0, 0).paragraphs[0]
    sibling = outer.cell(0, 1).paragraphs[0]
    selected = outer.cell(1, 1)
    removed = selected.paragraphs[0]
    old_run = sibling.runs[0]
    selected.text = "changed"
    assert selected.text == "changed" and sibling.text == "same"
    assert nested.text == "nested" and paragraph.text == "unrelated"
    with pytest.raises(Exception, match="Cell.text|stale"):
        _ = removed.text
    with pytest.raises(Exception, match="Cell.text|stale"):
        _ = old_run.text
    sibling.replace_text("same", "retained")
    with pytest.raises(Exception, match="Paragraph.replace_text|stale"):
        _ = nested.text
    assert document.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "nested"
    assert document.paragraphs[0].text == "unrelated"
    for _ in range(2):
        document = rdocx.Document.from_bytes(document.to_bytes())
        table = document.tables[0]
        assert table.cell(1, 1).text == "changed"
        assert table.cell(0, 1).text == "retained"
        assert table.cell(0, 0).tables[0].cell(0, 0).text == "nested"
        assert document.paragraphs[0].text == "unrelated"


def test_fx237_owner_inventory_refusal_is_atomic(monkeypatch):
    import datetime
    import json
    import rdocx

    document = rdocx.Document()
    table = document.add_table(1, 2)
    table.cell(0, 0).text = "first"
    table.cell(0, 1).text = "second"
    first, second = table.cell(0, 0), table.cell(0, 1)
    paragraph = first.paragraphs[0]
    opaque = '<x:keep xmlns:x="urn:fx237" marker="preserve"/>'
    xml = second.element.xml
    start = xml.index(">") + 1
    second.element.replace_xml(xml[:start] + "<w:tcPr>" + opaque + "</w:tcPr>" + xml[start:])
    assert opaque in second.element.xml
    document.add_paragraph("field ").add_run("").add_field('DATE \\@ "yyyy-MM-dd"', "1999-01-01")
    # Refresh after the conservative append, then force a real report allocation
    # refusal after native candidate preparation on a document with cell owners.
    first, second = document.tables[0].cell(0, 0), document.tables[0].cell(0, 1)
    paragraph = first.paragraphs[0]
    baseline = document.to_bytes()
    def refuse_report(*args, **kwargs):
        raise MemoryError("fx237 prepublication report allocation")

    with monkeypatch.context() as patch:
        patch.setattr(json, "loads", refuse_report)
        with pytest.raises(MemoryError, match="fx237 prepublication"):
            document.update_all_fields(now=datetime.datetime(2026, 10, 10))
    assert document.to_bytes() == baseline and paragraph.text == "first"
    with pytest.raises(Exception):
        first.merge(second)
    assert document.to_bytes() == baseline
    assert first.text == "first" and paragraph.text == "first"
    assert opaque in second.element.xml
    before = second.element.xml
    second.element.replace_xml(before)
    assert document.to_bytes() == baseline and paragraph.text == "first"
    foreign = rdocx.Document().add_table(1, 1).cell(0, 0)
    with pytest.raises(Exception):
        first.merge(foreign)
    assert document.to_bytes() == baseline


@pytest.mark.parametrize("slot", ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"))
@pytest.mark.parametrize("cell", (False, True))
def test_fx232_checked_header_footer_paragraph_and_run_reopen_atomically(slot, cell):
    import rdocx
    import xml.etree.ElementTree as ET

    document = rdocx.Document()
    document.add_paragraph("body witness")
    owner = getattr(document.section(0), slot)
    if cell:
        owner.add_table(1, 1).cell(0, 0).text = "old owner text"
    else:
        owner.add_paragraph("old owner text")

    def selected(doc):
        story = getattr(doc.section(0), slot)
        return story.tables[0].cell(0, 0).paragraphs[0] if cell else story.paragraphs[-1]

    paragraph = selected(document)
    assert ET.fromstring(paragraph.xml).tag.endswith("}p")
    paragraph.replace_xml(paragraph.xml.replace("old owner text", "new owner text"))
    run = selected(document).runs[0]
    assert ET.fromstring(run.xml).tag.endswith("}r")
    run.replace_xml(run.xml.replace("new owner text", "final owner text"))
    for getter in (lambda: selected(document), lambda: selected(document).runs[0]):
        target = getter()
        before = document.to_bytes()
        original = target.xml
        target.replace_xml(original)
        assert document.to_bytes() == before
        with pytest.raises(rdocx.RdocxError):
            target.replace_xml(original.replace("final owner text", "<unbound:bad/>"))
        assert document.to_bytes() == before
        assert getter().xml == original
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert selected(reopened).text == "final owner text"
    assert reopened.paragraphs[0].text == "body witness"


@pytest.mark.parametrize("where", ("svgBlip", "w:p"))
def test_fx232_issue373_svg_namespace_ancestor_retains_the_picture_reference(where):
    import rdocx
    import xml.etree.ElementTree as ET

    document = rdocx.Document()
    document.add_picture(_one_pixel_png(b"\xff\x00\x00"), width=rdocx.Inches(.5), height=rdocx.Inches(.5))
    document.add_picture(_one_pixel_png(b"\xff\x00\x00"), width=rdocx.Inches(.5), height=rdocx.Inches(.5))
    first, second = [p.xml for p in document.paragraphs[:2]]
    relationship = re.search(r'r:embed="([^"]+)"', second).group(1)
    declaration = 'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main"'
    extension = ('<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
                 '<asvg:svgBlip %s r:embed="%s"/></a:ext></a:extLst>' %
                 (declaration if where == "svgBlip" else "", relationship))
    if "</a:blip>" in first:
        replacement = first.replace("</a:blip>", extension + "</a:blip>", 1)
    else:
        replacement = re.sub(r"(<a:blip\b[^>]*?)/>", lambda m: m.group(1) + ">" + extension + "</a:blip>", first, count=1)
    if where == "w:p":
        replacement = re.sub(r"^<w:p\b", "<w:p " + declaration, replacement, count=1)
    document.paragraphs[0].replace_xml(replacement)
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    root = ET.fromstring(reopened.paragraphs[0].xml)
    svg = root.find(".//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip")
    assert svg is not None
    assert svg.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed") == relationship
    before = reopened.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        reopened.paragraphs[0].replace_xml(replacement.replace(declaration, ""))
    assert reopened.to_bytes() == before


def test_fx246_word_legend_controls_live_owner_defaults_and_atomicity():
    import rdocx
    data = rdocx.CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("North", [12, 14])
    data.add_series("South", [9, 11])
    document = rdocx.Document()
    paragraph = document.add_chart(51, data, rdocx.Inches(6), rdocx.Inches(2.6))
    expected = {"visible": True, "position": "r", "overlay": False}
    assert document.chart_legend(paragraph) == expected
    before = document.to_bytes()
    document.set_chart_legend(paragraph)
    document.set_chart_legend(paragraph, visible=True, position="r", overlay=False)
    assert document.to_bytes() == before
    assert paragraph.xml
    for position in ("r", "l", "t", "b", "tr"):
        paragraph = document.paragraphs[0]
        document.set_chart_legend(paragraph, position=position, overlay=True)
        assert document.chart_legend(document.paragraphs[0]) == dict(expected, position=position, overlay=True)
    paragraph = document.paragraphs[0]
    document.set_chart_legend(paragraph, visible=False)
    paragraph = document.paragraphs[0]
    before = document.to_bytes()
    document.set_chart_legend(paragraph, position="l", overlay=True)
    assert document.to_bytes() == before
    assert document.chart_legend(paragraph) == dict(expected, visible=False)
    document.set_chart_legend(paragraph, visible=True, position="b", overlay=True)
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.chart_legend(reopened.paragraphs[0]) == dict(expected, position="b", overlay=True)
    paragraph = document.paragraphs[0]
    before = document.to_bytes()
    for kwargs, error in (({"visible": 1}, TypeError), ({"overlay": "yes"}, TypeError), ({"position": "bad"}, ValueError), ({"position": 1}, TypeError)):
        with pytest.raises(error):
            document.set_chart_legend(paragraph, **kwargs)
        assert document.to_bytes() == before
    foreign = rdocx.Document()
    foreign_paragraph = foreign.add_chart(51, data, rdocx.Inches(6), rdocx.Inches(2))
    with pytest.raises(ValueError):
        document.set_chart_legend(foreign_paragraph, visible=False)
    assert document.to_bytes() == before
    plain = document.add_paragraph("no chart")
    before = document.to_bytes()
    with pytest.raises(Exception):
        document.set_chart_legend(plain, visible=True)
    assert document.to_bytes() == before


def test_fx246_word_legend_refuses_retired_related_and_ambiguous_owners():
    import re
    import rdocx
    data = rdocx.CategoryChartData()
    data.categories = ["A"]
    data.add_series("North", [7])
    document = rdocx.Document()
    paragraph = document.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    # One-series authoring retains its existing absent legend default.
    assert document.chart_legend(paragraph) == {"visible": False, "position": "r", "overlay": False}
    retired = paragraph
    assert document.remove_content(document.find_content_index(retired))
    before = document.to_bytes()
    with pytest.raises(rdocx.StaleElementError):
        document.set_chart_legend(retired, visible=True)
    assert document.to_bytes() == before
    cell = document.add_table(1, 1).cell(0, 0).paragraphs[0]
    header = document.section(0).header.add_paragraph("header witness")
    for owner in (cell, header):
        before = document.to_bytes()
        with pytest.raises(ValueError, match="direct body"):
            document.set_chart_legend(owner, visible=True)
        assert document.to_bytes() == before
    paragraph = document.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    xml = paragraph.xml
    drawing = re.search(r"<w:r>.*?</w:r>", xml, re.S).group(0)
    drawing = re.sub(r'(wp:docPr\s+id=")[0-9]+', r'\g<1>999', drawing, count=1)
    paragraph.replace_xml(xml.replace("</w:p>", drawing + "</w:p>"))
    paragraph = document.paragraphs[-1]
    before = document.to_bytes()
    with pytest.raises(Exception, match="exactly one"):
        document.set_chart_legend(paragraph, visible=True)
    assert document.to_bytes() == before


@pytest.mark.parametrize("with_owned", (False, True))
def test_fx246_word_opaque_ancestors_never_grant_chart_authority(with_owned):
    import copy
    import io
    import zipfile
    import xml.etree.ElementTree as ET
    import rdocx
    data = rdocx.CategoryChartData()
    data.categories = ["A"]
    data.add_series("North", [7])
    document = rdocx.Document()
    document.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    tree = ET.fromstring(parts["word/document.xml"])
    run = tree.find(".//" + w + "r")
    drawing = run.find(w + "drawing")
    opaque = ET.Element("{urn:fx246:opaque}holder")
    if with_owned:
        extra = copy.deepcopy(drawing)
        for node in extra.iter(wp + "docPr"):
            node.set("id", "999")
        opaque.append(extra)
    else:
        run.remove(drawing)
        opaque.append(drawing)
    run.append(opaque)
    parts["word/document.xml"] = ET.tostring(tree)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        for name, body in parts.items():
            package.writestr(name, body)
    document = rdocx.Document.from_bytes(buffer.getvalue())
    paragraph = document.paragraphs[0]
    before = document.to_bytes()
    opaque_before = ET.tostring(ET.fromstring(paragraph.xml).find(".//{urn:fx246:opaque}holder"))
    if with_owned:
        document.set_chart_legend(paragraph, visible=True, position="b")
        assert document.chart_legend(document.paragraphs[0])["position"] == "b"
        assert ET.tostring(ET.fromstring(document.paragraphs[0].xml).find(".//{urn:fx246:opaque}holder")) == opaque_before
    else:
        for operation in (lambda: document.chart_legend(paragraph), lambda: document.set_chart_legend(paragraph, visible=True)):
            with pytest.raises(Exception, match="exactly one"):
                operation()
            assert document.to_bytes() == before


def test_fx246_word_shared_chart_retargets_only_selected_owned_attribute():
    import copy
    import io
    import zipfile
    import xml.etree.ElementTree as ET
    import rdocx
    data = rdocx.CategoryChartData()
    data.categories = ["A"]
    data.add_series("North", [7])
    document = rdocx.Document()
    document.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    c = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    tree = ET.fromstring(parts["word/document.xml"])
    body = tree.find(w + "body")
    paragraph = body.find(w + "p")
    relationship = paragraph.find(".//" + c + "chart").get(r + "id")
    other = copy.deepcopy(paragraph)
    for node in other.iter(wp + "docPr"):
        node.set("id", "999")
    body.insert(1, other)
    ET.SubElement(paragraph, "{urn:fx246:foreign}chart", {r + "id": relationship, "marker": "untouched"})
    parts["word/document.xml"] = ET.tostring(tree)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        for name, value in parts.items():
            package.writestr(name, value)
    document = rdocx.Document.from_bytes(buffer.getvalue())
    document.set_chart_legend(document.paragraphs[0], visible=True, position="b")
    current = ET.fromstring(document.paragraphs[0].xml)
    assert current.find("{urn:fx246:foreign}chart").get(r + "id") == relationship
    assert current.find(".//" + c + "chart").get(r + "id") != relationship
    assert document.chart_legend(document.paragraphs[1])["visible"] is False
    assert document.chart_legend(document.paragraphs[0])["position"] == "b"


@pytest.mark.parametrize("trailer", (b"NOT_XML_DOCUMENT_CONTENT", b"<![CDATA[tail]]>", b"&amp;"))
def test_fx246_word_malformed_chart_document_is_refused_atomically(trailer):
    import io
    import zipfile
    import rdocx
    data = rdocx.CategoryChartData()
    data.categories = ["A"]
    data.add_series("North", [7])
    document = rdocx.Document()
    document.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            value = source.read(info.filename)
            if info.filename.startswith("word/charts/chart") and info.filename.endswith(".xml"):
                value += trailer
            target.writestr(info, value)
    document = rdocx.Document.from_bytes(buffer.getvalue())
    paragraph = document.paragraphs[0]
    before = document.to_bytes()
    for operation in (lambda: document.chart_legend(paragraph), lambda: document.set_chart_legend(paragraph, visible=True)):
        with pytest.raises(Exception):
            operation()
        assert document.to_bytes() == before and paragraph.xml



def _fx247_oracle_and_cases():
    import importlib.util
    import json
    from pathlib import Path

    root = Path(__file__).resolve().parents[3]
    spec = importlib.util.spec_from_file_location("fx247_oracle", root / "scripts/officemath_oracle_harness.py")
    oracle = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(oracle)
    manifest = json.loads((root / "scripts/officemath_oracle_manifest.json").read_text())
    return oracle, manifest["fx247"]["cases"]


def _fx247_document():
    import rdocx

    oracle, cases = _fx247_oracle_and_cases()
    document = rdocx.Document()
    for source in cases:
        converted = document.add_paragraph("").add_equation(source, format="latex")
        assert converted.diagnostics == (), (source, converted.diagnostics)
    return oracle, cases, document


def test_fx247_all73_saved_semantics_and_public_reopen(tmp_path):
    import rdocx

    oracle, cases, document = _fx247_document()
    path = tmp_path / "all73.docx"
    document.save(path)
    first = oracle.verify_fx247_package(path)
    assert len(first["rows"]) == len(cases) == 73
    reopened = rdocx.Document(path)
    assert len(reopened.paragraphs) == 73
    for paragraph in reopened.paragraphs:
        equations = paragraph.equations()
        assert len(equations) == 1
        assert equations[0].to_latex().equations[0].diagnostics == ()
        assert equations[0].to_mathml().equations[0].diagnostics == ()
    reopened.save(path)
    assert [row["semantics"] for row in oracle.verify_fx247_package(path)["rows"]] == [row["semantics"] for row in first["rows"]]


@pytest.mark.parametrize("mutation,row", [
    ("swap-fraction", 58), ("binomial-bar", 54), ("phi-variant", 9),
    ("detach-script", 72), ("italic-function", 47), ("strip-text-space", 53),
    ("collapse-quad", 57), ("collapse-semicolon", 71),
])
def test_fx247_saved_semantic_observer_rejects_real_mutations(tmp_path, mutation, row):
    from xml.etree import ElementTree as ET

    oracle, _, document = _fx247_document()
    original = io.BytesIO(document.to_bytes())
    result = tmp_path / (mutation + ".docx")
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(result, "w") as output:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename == "word/document.xml":
                root = ET.fromstring(raw)
                equation = root.findall(".//" + ns + "oMath")[row - 1]
                if mutation == "swap-fraction":
                    fraction = equation.find(ns + "f")
                    numerator = fraction.find(ns + "num")
                    denominator = fraction.find(ns + "den")
                    numerator.tag, denominator.tag = denominator.tag, numerator.tag
                elif mutation == "binomial-bar":
                    equation.find(".//" + ns + "type").set(ns + "val", "bar")
                elif mutation == "phi-variant":
                    equation.find(".//" + ns + "t").text = "φ"
                elif mutation == "detach-script":
                    equation.find(".//" + ns + "sSup/" + ns + "e/" + ns + "r/" + ns + "t").text = "4π"
                elif mutation == "italic-function":
                    equation.find(".//" + ns + "nor").set(ns + "val", "0")
                elif mutation == "strip-text-space":
                    equation.find(".//" + ns + "t").text = "if"
                else:
                    text = equation.find(".//" + ns + "t")
                    text.text = text.text.replace("\u2001", "\u2009").replace("\u2004", "\u2009")
                raw = ET.tostring(root, encoding="utf-8")
            output.writestr(info, raw)
    with pytest.raises(ValueError, match=f"row{row} saved semantics differ"):
        oracle.verify_fx247_package(result)


def test_fx247_literal_caps_refuse_atomically_and_opaque_source_survives(tmp_path):
    import rdocx

    document = rdocx.Document()
    paragraph = document.add_paragraph("before")
    before = document.to_bytes()
    for source in [r"\text{bad{nested}}", r"\text{bad\unknown}", r"\text{" + "x" * 100_001 + "}"]:
        with pytest.raises(rdocx.RdocxError):
            paragraph.add_equation(source, format="latex")
        assert document.to_bytes() == before
    opaque = '<x:opaque xmlns:x="urn:fx247" data="untouched">  kept </x:opaque>'
    body = '<w:p><w:r><w:t>before</w:t></w:r>' + opaque + '</w:p>'
    document = _replace_document_body(document, body)
    document.paragraphs[0].add_equation(r"\text{if }x+\binom{n}{k}", format="latex")
    document = rdocx.Document.from_bytes(document.to_bytes())
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as archive:
        assert opaque.encode() in archive.read("word/document.xml")



def test_fx247_synthesized_mathml_text_space_cap_is_atomic(tmp_path):
    import rdocx

    document = rdocx.Document()
    paragraph = document.add_paragraph("kept")
    before = document.to_bytes()
    namespace = "http://www.w3.org/1998/Math/MathML"
    for prefix, suffix in [("", '<mspace width="0.333em"/>'), ('<mspace width="0.333em"/>', "")]:
        source = f'<math xmlns="{namespace}">{prefix}<mtext>{"x" * 524_288}</mtext>{suffix}</math>'
        with pytest.raises(rdocx.RdocxError):
            paragraph.add_equation(source, format="mathml")
        assert document.to_bytes() == before
    source = f'<math xmlns="{namespace}"><mtext>{"x" * 524_287}</mtext><mspace width="0.333em"/></math>'
    assert paragraph.add_equation(source, format="mathml").diagnostics == ()
    path = tmp_path / "at-limit.docx"
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        from xml.etree import ElementTree as ET
        root = ET.fromstring(archive.read("word/document.xml"))
        texts = root.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}t")
        assert sum(len(t.text or "") for t in texts) == 524_288
        assert texts[-1].text.endswith(" ")


def _fx241_legacy_text(name="legacy", text="old", maximum=40):
    return (f'<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="{name}"/>'
            f'<w:enabled/><w:calcOnExit w:val="0"/><w:textInput><w:default w:val="{text}"/>'
            f'<w:maxLength w:val="{maximum}"/></w:textInput></w:ffData></w:fldChar></w:r>'
            '<w:r><w:instrText> FORMTEXT </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r><w:t>{text}</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')


def _fx241_control(tag, properties, text="old", alias=None):
    alias_xml = f'<w:alias w:val="{alias}"/>' if alias else ''
    return (f'<w:sdt><w:sdtPr><w:tag w:val="{tag}"/>{alias_xml}{properties}</w:sdtPr>'
            f'<w:sdtContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:sdtContent></w:sdt>')


def _fx241_five_kind_document():
    import rdocx
    body = _fx241_control("plain", "<w:text/>", alias="Plain alias")
    body += _fx241_control("check", '<w14:checkbox xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w14:checked w14:val="0"/><w14:checkedState w14:val="2612" w14:font="MS Gothic"/><w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox>', "☐")
    body += _fx241_control("choice", '<w:dropDownList w:lastValue="one"><w:listItem w:displayText="One" w:value="one"/><w:listItem w:displayText="Two" w:value="two"/></w:dropDownList>', "one")
    body += _fx241_control("date", '<w:date w:fullDate="2026-09-15T00:00:00Z"><w:dateFormat w:val="yyyy-MM-dd"/></w:date>', "2026-09-15")
    return _replace_document_body(rdocx.Document(), body + _fx241_legacy_text() + '<w:sectPr/>')


def test_fx249_individual_errors_noops_and_bulk_contracts():
    import rdocx
    import xml.etree.ElementTree as ET
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    body = _fx241_control('plain', '<w:text/>', 'same', alias='Plain alias')
    body += _fx241_control('date', '<w:date w:fullDate="2026-10-31T00:00:00Z"><w:dateFormat w:val="dd/MM/yyyy"/></w:date>', '31/10/2026')
    body += _fx241_control('choice', '<w:dropDownList><w:listItem w:displayText="Draft" w:value="draft"/><w:listItem w:displayText="Final" w:value="final"/></w:dropDownList>', 'Draft')
    body += '<w:p><w:r><w:t>held</w:t></w:r></w:p><w:sectPr/>'
    document = _replace_document_body(rdocx.Document(), body)
    paragraph, section = document.paragraphs[-1], document.section(0)
    run = paragraph.runs[0]
    before = document.to_bytes()
    assert document.set_content_control_value_by_tag('plain', 'same') == 1
    assert document.set_content_control_value_by_alias('Plain alias', 'same') == 1
    assert document.set_content_control_value_by_tag('date', '2026-10-31') == 1
    assert document.bind_content_controls({'missing': 'ignored', 'plain': 'same'}) == 1
    assert document.to_bytes() == before
    for method in (document.set_content_control_value_by_tag, document.set_content_control_value_by_alias):
        with pytest.raises(KeyError):
            method('missing', 'value')
        for value in (None, 12, True):
            with pytest.raises(TypeError):
                method('plain', value)
    for tag, value in [('date', '2026-02-29'), ('choice', 'Archived'), ('plain', '\x01')]:
        with pytest.raises(ValueError):
            document.set_content_control_value_by_tag(tag, value)
    with pytest.raises(rdocx.XmlError, match='Draft|Final'):
        document.bind_content_controls({'plain': 'early', 'choice': 'Archived'})
    assert document.to_bytes() == before
    assert paragraph.text == 'held' and run.text == 'held' and section.xml
    assert document.set_content_control_value_by_tag('choice', 'final') == 1
    assert document.content_controls()[2].text == 'Final'
    with pytest.raises(rdocx.StaleElementError):
        _ = paragraph.text

    # A successful metadata-only fill is a real conservative publication.
    placeholder = _fx241_control('plain', '<w:showingPlcHdr/><p:showingPlcHdr xmlns:p="urn:producer"/><w:text/>', 'same')
    document = _replace_document_body(rdocx.Document(), placeholder + '<w:p><w:r><w:t>held</w:t></w:r></w:p><w:sectPr/>')
    paragraph = document.paragraphs[-1]
    key, cert = _fx240_material()
    document.sign(key, cert)
    before = document.to_bytes()
    assert document.set_content_control_value_by_tag('plain', 'same') == 1
    assert document.to_bytes() != before
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as package:
        root = ET.fromstring(package.read('word/document.xml'))
    assert root.find('.//' + w + 'showingPlcHdr') is None
    assert root.find('.//{urn:producer}showingPlcHdr') is not None
    assert not document.verify_signatures()[0]['cryptographically_valid']
    with pytest.raises(rdocx.StaleElementError):
        _ = paragraph.text

    # Missing external bindings and unsupported kinds retain their generic family.
    for owner, tag in [(_fx241_bound_document(missing=True), 'bound'),
                       (_replace_document_body(rdocx.Document(), _fx241_control('picture', '<w:picture/>') + '<w:sectPr/>'), 'picture')]:
        before = owner.to_bytes()
        with pytest.raises(rdocx.RdocxError) as caught:
            owner.set_content_control_value_by_tag(tag, 'value')
        assert not isinstance(caught.value, (KeyError, ValueError))
        assert owner.to_bytes() == before

    # Selected malformed InvalidValue payloads use the individual local mapping.
    malformed = _fx241_control('bad', '<w14:checkbox xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w14:checkedState w14:val="not-hex"/></w14:checkbox>')
    owner = _replace_document_body(rdocx.Document(), malformed + '<w:sectPr/>')
    before = owner.to_bytes()
    with pytest.raises(ValueError, match='Unicode scalar'):
        owner.set_content_control_value_by_tag('bad', 'true')
    with pytest.raises(rdocx.XmlError, match='Unicode scalar'):
        owner.bind_content_controls({'bad': 'true'})
    assert owner.to_bytes() == before


def test_fx241_python_fill_all_five_reported_form_kinds():
    import rdocx
    document = _fx241_five_kind_document()
    frozen = document.content_controls()
    old_legacy = document.legacy_form_fields()[0]
    assert document.bind_content_controls({"plain": "Ada", "check": "true", "choice": "two", "date": "2026-10-16"}) == 4
    result = document.set_legacy_form_field_value("legacy", "filled")
    assert isinstance(result, rdocx.LegacyFormField)
    assert result.value == "filled" and old_legacy.value == "old"
    assert result.type == "text" and result.ordinal == 0
    assert result.source_part == "/word/document.xml" and result.choices == ()
    assert result.enabled and not result.calculate_on_exit
    assert frozen[0].text == "old"
    with pytest.raises(AttributeError):
        result.value = "mutate"
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    controls = {item.tag: item for item in reopened.content_controls()}
    assert [controls[key].text for key in ("plain", "choice", "date")] == ["Ada", "Two", "2026-10-16"]
    assert reopened.legacy_form_fields()[0].value == "filled"
    _fx241_assert_five_kind_source(reopened.to_bytes())


def _fx241_assert_five_kind_source(raw):
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(io.BytesIO(raw)) as package:
        root = ET.fromstring(package.read('word/document.xml'))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "w14": "http://schemas.microsoft.com/office/word/2010/wordml"}
    w = '{'+ns['w']+'}'
    assert root.find('.//w14:checked', ns).get('{'+ns['w14']+'}val') == '1'
    check = next(c for c in root.iter(w+'sdt') if c.find('w:sdtPr/w:tag', ns).get(w+'val') == 'check')
    assert '☒' in ''.join(check.itertext())
    fonts = check.find('.//w:rFonts', ns)
    assert fonts.get(w+'ascii') == fonts.get(w+'hAnsi') == 'MS Gothic'
    assert root.find('.//w:date', ns).get(w+'fullDate') == '2026-10-16T00:00:00Z'
    assert root.find('.//w:dateFormat', ns).get(w+'val') == 'yyyy-MM-dd'
    assert root.find('.//w:dropDownList', ns).get(w+'lastValue') == 'one'
    assert [x.get(w+'value') for x in root.findall('.//w:dropDownList/w:listItem', ns)] == ['one', 'two']
    legacy = root.find('.//w:ffData/w:textInput/w:default', ns)
    assert legacy.get(w+'val') == 'filled'
    field_paragraph = next(p for p in root.iter(w+'p') if p.find('.//w:ffData', ns) is not None)
    assert ''.join(x.text or '' for x in field_paragraph.iter(w+'t')) == 'filled'


def _fx241_legacy_typed(name, kind):
    properties, instruction, cached = {
        "checkbox": ('<w:checkBox><w:sizeAuto/><w:default w:val="0"/><w:checked w:val="0"/></w:checkBox>', "FORMCHECKBOX", "☐"),
        "dropdown": ('<w:ddList><w:result w:val="0"/><w:listEntry w:val="one"/><w:listEntry w:val="two"/></w:ddList>', "FORMDROPDOWN", "one"),
    }[kind]
    return (f'<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="{name}"/>{properties}</w:ffData></w:fldChar></w:r>'
            f'<w:r><w:instrText> {instruction} </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r><w:t>{cached}</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')


def test_fx241_legacy_snapshots_typed_values_choices_and_qualified_unnamed():
    import rdocx
    import xml.etree.ElementTree as ET
    body = _fx241_legacy_text() + _fx241_legacy_typed("flag", "checkbox") + _fx241_legacy_typed("pick", "dropdown")
    body += _fx241_legacy_text().replace('<w:name w:val="legacy"/>', '')
    document = _replace_document_body(rdocx.Document(), body + '<w:sectPr/>')
    snapshots = document.legacy_form_fields()
    assert isinstance(snapshots, tuple)
    assert [x.type for x in snapshots] == ["text", "checkbox", "dropdown", "text"]
    assert [type(x.value) for x in snapshots] == [str, bool, int, str]
    assert snapshots[2].choices == ("one", "two") and snapshots[3].name is None
    assert document.set_legacy_form_field_value("flag", True).value is True
    assert document.set_legacy_form_field_value("pick", 1).value == 1
    unnamed = document.set_legacy_form_field_value(None, "unnamed", source_part=snapshots[3].source_part, ordinal=3)
    assert unnamed.name is None and unnamed.value == "unnamed"
    assert snapshots[1].value is False and snapshots[2].value == 0
    with pytest.raises(TypeError):
        rdocx.LegacyFormField()
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [x.value for x in reopened.legacy_form_fields()] == ["old", True, 1, "unnamed"]
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    root = ET.fromstring(_document_xml(reopened))
    assert root.find('.//'+w+'checkBox/'+w+'checked').get(w+'val') == '1'
    assert root.find('.//'+w+'ddList/'+w+'result').get(w+'val') == '1'
    assert 'two' in [x.text for x in root.iter(w+'t')]


@pytest.mark.parametrize("name,value,kwargs", [
    ("missing", "x", {}), (None, "x", {}), ("legacy", "x", {"source_part": "/word/document.xml"}),
    ("legacy", "x", {"ordinal": 0}), ("legacy", "x", {"source_part": "/word/header9.xml", "ordinal": 0}),
    ("stale", "x", {"source_part": "/word/document.xml", "ordinal": 0}),
    ("legacy", "x", {"source_part": "/word/document.xml", "ordinal": True}),
    ("legacy", "x", {"source_part": "/word/document.xml", "ordinal": -1}),
    ("legacy", 1, {}), ("legacy", True, {}), ("legacy", "longer-than-four", {}),
    ("flag", 1, {}), ("flag", "true", {}), ("pick", True, {}), ("pick", -1, {}), ("pick", 2, {}), ("pick", "one", {}),
])
def test_fx241_legacy_identity_and_value_refusals_are_atomic(name, value, kwargs):
    import rdocx
    document = _replace_document_body(rdocx.Document(), _fx241_legacy_text(maximum=4) + _fx241_legacy_typed("flag", "checkbox") + _fx241_legacy_typed("pick", "dropdown") + '<w:p><w:r><w:t>held</w:t></w:r></w:p><w:sectPr/>')
    paragraph, section = document.paragraphs[-1], document.section(0)
    story = next(x for x in document.stories if x.kind == "body")
    before = document.to_bytes()
    with pytest.raises((ValueError, TypeError, OverflowError, rdocx.RdocxError)):
        document.set_legacy_form_field_value(name, value, **kwargs)
    assert document.to_bytes() == before
    assert paragraph.text == "held" and section.xml and story.kind == "body"
    assert document.set_legacy_form_field_value("legacy", "safe").value == "safe"
    assert paragraph.text == "held"


def test_fx241_duplicate_names_refuse_but_current_qualified_identity_selects():
    import rdocx
    document = _replace_document_body(rdocx.Document(), _fx241_legacy_text() + _fx241_legacy_text() + '<w:sectPr/>')
    before = document.to_bytes()
    with pytest.raises(ValueError, match="ambiguous"):
        document.set_legacy_form_field_value("legacy", "bad")
    assert document.to_bytes() == before
    selected = document.set_legacy_form_field_value("legacy", "second", source_part="/word/document.xml", ordinal=1)
    assert selected.ordinal == 1
    assert [x.value for x in document.legacy_form_fields()] == ["old", "second"]


def test_fx241_bind_contract_priority_nested_noop_owners_and_late_failure():
    import rdocx
    first = _fx241_control("tag", '<w:text/>', alias="alias")
    document = _replace_document_body(rdocx.Document(), first + '<w:p><w:r><w:t>held</w:t></w:r></w:p><w:sectPr/>')
    paragraph, section = document.paragraphs[-1], document.section(0)
    table = document.add_table(1, 1)
    cell = table.cell(0, 0)
    cell.text = "cell"
    before = document.to_bytes()
    for values in ({}, {"missing": "ignored"}, {"tag": "old"}):
        count = document.bind_content_controls(values)
        assert count == (1 if values.get("tag") else 0)
        assert document.to_bytes() == before
    for invalid in ({"tag": 12}, {"tag": True}, {1: "text"}, {"tag": "valid", "alias": None}):
        with pytest.raises(TypeError):
            document.bind_content_controls(invalid)
        assert document.to_bytes() == before
    assert document.bind_content_controls({"tag": "wins", "alias": "loses", "missing": "ignored"}) == 1
    assert document.content_controls()[0].text == "wins"
    assert paragraph.text == "held" and cell.text == "cell" and table.rows[0].cells[0].text == "cell" and section.xml
    nested = '<w:sdt><w:sdtPr><w:tag w:val="outer"/><w:group/></w:sdtPr><w:sdtContent>' + _fx241_control("inner", '<w:text/>') + '</w:sdtContent></w:sdt>'
    document = _replace_document_body(rdocx.Document(), first + nested + '<w:sectPr/>')
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        document.bind_content_controls({"tag": "early", "outer": "unsupported"})
    assert document.to_bytes() == before
    assert document.bind_content_controls({"inner": "nested"}) == 1
    assert any(x.tag == "inner" and x.text == "nested" for x in document.content_controls())


def _fx241_bound_document(missing=False):
    import rdocx
    binding = '<w:dataBinding w:storeItemID="{11111111-1111-1111-1111-111111111111}" w:xpath="/c:root/c:name" w:prefixMappings="xmlns:c=\'urn:customer\'"/>'
    document = _replace_document_body(rdocx.Document(), _fx241_control("bound", binding+'<w:text/>') + _fx241_legacy_text() + '<w:p><w:r><w:t>held</w:t></w:r></w:p><w:sectPr/>')
    source, result = io.BytesIO(document.to_bytes()), io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(result, 'w') as zout:
        for entry in zin.infolist():
            raw = zin.read(entry.filename)
            if entry.filename == 'word/_rels/document.xml.rels' and not missing:
                raw = raw.replace(b'</Relationships>', b'<Relationship Id="rIdBound241" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" Target="../customXml/item1.xml"/></Relationships>')
            if entry.filename == '[Content_Types].xml' and not missing:
                raw = raw.replace(b'</Types>', b'<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/></Types>')
            zout.writestr(entry, raw)
        if not missing:
            zout.writestr('customXml/item1.xml', '<c:root xmlns:c="urn:customer" opaque="keep"><c:name>old</c:name><c:other>untouched</c:other></c:root>')
            zout.writestr('customXml/itemProps1.xml', '<ds:datastoreItem xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml" ds:itemID="{11111111-1111-1111-1111-111111111111}"><ds:schemaRefs/></ds:datastoreItem>')
            zout.writestr('customXml/_rels/item1.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="props" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps1.xml"/></Relationships>')
    return rdocx.Document.from_bytes(result.getvalue())


def test_fx241_bound_xml_values_source_preservation_and_late_refusal():
    import rdocx
    import xml.etree.ElementTree as ET
    document = _fx241_bound_document()
    before = document.to_bytes()
    assert document.bind_content_controls({'bound': 'Ada & Co'}) == 1
    after = document.to_bytes()
    with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(io.BytesIO(after)) as new:
        assert set(old.namelist()) == set(new.namelist())
        for name in old.namelist():
            if name not in ('word/document.xml', 'customXml/item1.xml'):
                assert old.read(name) == new.read(name), name
        root = ET.fromstring(new.read('customXml/item1.xml'))
        assert root.find('{urn:customer}name').text == 'Ada & Co'
        assert root.find('{urn:customer}other').text == 'untouched' and root.get('opaque') == 'keep'
    assert rdocx.Document.from_bytes(after).content_controls()[0].text == 'Ada & Co'
    document = _fx241_bound_document(missing=True)
    paragraph = document.paragraphs[-1]
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        document.bind_content_controls({'bound': 'cannot resolve'})
    assert document.to_bytes() == before and paragraph.text == 'held'


def test_fx241_related_legacy_physical_part_and_noop_preserve_held_owners():
    import rdocx
    seed = rdocx.Document()
    seed.add_paragraph('body held')
    seed.section(0).header.add_paragraph('header held')
    source, result = io.BytesIO(seed.to_bytes()), io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(result, 'w') as zout:
        header = next(name for name in zin.namelist() if re.fullmatch(r'word/header\d+\.xml', name))
        for entry in zin.infolist():
            raw = zin.read(entry.filename)
            if entry.filename == header:
                raw = raw.replace(b'</w:hdr>', (_fx241_legacy_text('headerField')+'</w:hdr>').encode())
            zout.writestr(entry, raw)
    document = rdocx.Document.from_bytes(result.getvalue())
    paragraph, section = document.paragraphs[0], document.section(0)
    header_paragraph = next(p for p in section.header.paragraphs if p.text == 'header held')
    selected_story = next(x for x in document.stories if x.kind == 'header')
    # Flush a prior live related-story edit through the existing provider first.
    header_paragraph.text = 'header held edited'
    assert header_paragraph.text == 'header held edited'
    field = document.legacy_form_fields()[0]
    assert field.source_part == '/'+header and field.ordinal == 0
    with pytest.raises(rdocx.RdocxError):
        document.replace_image_for_story(selected_story, 'missing', b'')
    updated = document.set_legacy_form_field_value('headerField', 'changed', source_part=field.source_part, ordinal=0)
    assert updated.value == 'changed' and field.value == 'old'
    # Frozen Story metadata remains readable, but its old mutation epoch retires.
    with pytest.raises(rdocx.StaleElementError):
        document.replace_image_for_story(selected_story, 'missing', b'')
    assert paragraph.text == 'body held' and header_paragraph.text == 'header held edited' and selected_story.kind == 'header' and section.xml
    before = document.to_bytes()
    assert document.set_legacy_form_field_value('headerField', 'changed').value == 'changed'
    assert document.to_bytes() == before and header_paragraph.text == 'header held edited'
    reopened = rdocx.Document.from_bytes(before)
    assert reopened.legacy_form_fields()[0].source_part == field.source_part
    assert reopened.legacy_form_fields()[0].value == 'changed'


@pytest.mark.parametrize('broken', ['checkbox-state', 'legacy-default', 'legacy-cache'])
def test_fx241_source_observer_rejects_semantic_state_disagreement(broken):
    import xml.etree.ElementTree as ET
    document = _fx241_five_kind_document()
    document.bind_content_controls({'plain': 'Ada', 'check': 'true', 'choice': 'two', 'date': '2026-10-16'})
    document.set_legacy_form_field_value('legacy', 'filled')
    raw = document.to_bytes()
    _fx241_assert_five_kind_source(raw)
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    w14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'
    mutated = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(mutated, 'w') as zout:
        for entry in zin.infolist():
            data = zin.read(entry.filename)
            if entry.filename == 'word/document.xml':
                root = ET.fromstring(data)
                if broken == 'checkbox-state':
                    root.find('.//'+w14+'checked').set(w14+'val', '0')
                elif broken == 'legacy-default':
                    root.find('.//'+w+'ffData/'+w+'textInput/'+w+'default').set(w+'val', 'wrong')
                else:
                    field = next(p for p in root.iter(w+'p') if p.find('.//'+w+'ffData') is not None)
                    next(field.iter(w+'t')).text = 'wrong'
                data = ET.tostring(root)
            zout.writestr(entry, data)
    with pytest.raises(AssertionError):
        _fx241_assert_five_kind_source(mutated.getvalue())


def test_fx241_form_publication_retires_changed_owners_not_equal_text_routes():
    import rdocx
    document = _replace_document_body(rdocx.Document(), _fx241_legacy_text() + _fx241_legacy_text() + '<w:sectPr/>')
    unchanged, changed = document.paragraphs[0], document.paragraphs[1]
    old_run = unchanged.runs[0]
    section = document.section(0)
    section_xml = section.xml
    document.set_legacy_form_field_value('legacy', 'second', source_part='/word/document.xml', ordinal=1)
    assert unchanged.text == 'old' and section.xml == section_xml
    with pytest.raises(rdocx.StaleElementError):
        _ = changed.text
    with pytest.raises(rdocx.StaleElementError):
        _ = old_run.text
    assert [x.value for x in document.legacy_form_fields()] == ['old', 'second']


def test_fx241_changed_table_retires_cell_and_paragraph_descendants():
    import rdocx
    table_xml = '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid><w:tr><w:tc><w:tcPr/>' + _fx241_legacy_text() + '</w:tc></w:tr></w:tbl>'
    document = _replace_document_body(rdocx.Document(), table_xml + '<w:p><w:r><w:t>untouched</w:t></w:r></w:p><w:sectPr/>')
    table = document.tables[0]
    cell = table.cell(0, 0)
    paragraph = cell.paragraphs[0]
    outside = document.paragraphs[-1]
    document.set_legacy_form_field_value('legacy', 'changed')
    for access in (lambda: table.rows, lambda: cell.text, lambda: paragraph.text):
        with pytest.raises(rdocx.StaleElementError):
            access()
    assert outside.text == 'untouched'
    assert document.tables[0].cell(0, 0).text == 'changed'


def test_fx241_retired_recreated_owner_is_never_revived_by_equal_source():
    import rdocx
    document = _replace_document_body(rdocx.Document(), '<w:p><w:r><w:t>equal</w:t></w:r></w:p>' + _fx241_control('plain', '<w:text/>') + '<w:sectPr/>')
    retired = document.paragraphs[0]
    assert document.remove_content(0)
    current = document.add_paragraph('equal')
    document.bind_content_controls({'plain': 'changed'})
    with pytest.raises(rdocx.StaleElementError):
        _ = retired.text
    assert current.text == 'equal'


def _fx240_material():
    # Authentic existing native test identity, not a production or trusted key.
    import base64
    from pathlib import Path
    source = (Path(__file__).resolve().parents[2] / "oxml-opc" / "src" / "signature.rs").read_text(encoding="utf-8")
    def constant(name):
        return base64.b64decode(re.search(r'const ' + name + r': &str = "([^"]+)";', source)[1], validate=True)
    return constant("TEST_PRIVATE_KEY_PKCS8_BASE64"), constant("TEST_CERTIFICATE_DER_BASE64")


def _fx240_zip_edit(source, part=None, old=None, new=None, extra=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(output, "w") as result:
        for info in archive.infolist():
            data = archive.read(info.filename)
            if info.filename == part:
                assert old in data
                data = data.replace(old, new, 1)
            result.writestr(info, data)
        if extra:
            result.writestr("word/uncovered.xml" if "rdocx" == "rdocx" else "ppt/uncovered.xml", b"<uncovered/>")
    return output.getvalue()


def test_fx240_sign_verify_owned_reports_pem_and_multisignatures():
    import base64
    import rdocx as product
    owner = product.Document()
    assert owner.verify_signatures() == ()
    opaque = b'<unmodelled xmlns="urn:caller-owned"><payload a="keep">opaque</payload></unmodelled>'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(owner.to_bytes())) as source, zipfile.ZipFile(output, "w") as result:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]
        for name, data in entries:
            result.writestr(name, data)
        result.writestr("custom/opaque.xml", opaque)
    owner = product.Document.from_bytes(output.getvalue())
    key, cert = _fx240_material()
    def pem(label, value):
        encoded = base64.b64encode(value)
        return b" \r\n-----BEGIN " + label + b"-----\r\n" + b"\r\n".join(encoded[i:i+64] for i in range(0, len(encoded), 64)) + b"\r\n-----END " + label + b"-----\t\r\n"
    report = owner.sign(pem(b"PRIVATE KEY", key), pem(b"CERTIFICATE", cert))
    assert type(report) is dict and report["cryptographically_valid"] and report["coverage_complete"]
    assert set(report) == {"signature_part", "signer", "cryptographically_valid", "coverage_complete", "covered_parts", "covered_relationships", "issues"}
    assert set(report["signer"]) == {"subject", "issuer", "serial_number"}
    assert "F-172 Test Signer" in report["signer"]["subject"]
    assert report["issues"] == [] and report["covered_parts"] and report["covered_relationships"]
    assert set(report["covered_relationships"][0]) == {"source_part", "relationship_id"}
    report["covered_parts"].clear()
    report["signer"]["subject"] = "mutated"
    assert owner.verify_signatures()[0]["covered_parts"]
    assert owner.verify_signatures()[0]["signer"]["subject"] != "mutated"
    before = owner.to_bytes()
    with zipfile.ZipFile(io.BytesIO(before)) as preserved:
        assert preserved.read("custom/opaque.xml") == opaque
    with pytest.raises(product.PackageError):
        owner.sign(key, cert)  # The native provider refuses re-signing a signed graph.
    assert owner.to_bytes() == before
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(before)) as archive, zipfile.ZipFile(output, "w") as result:
        signature_bytes = archive.read("_xmlsignatures/sig1.xml")
        for info in archive.infolist():
            result.writestr(info, archive.read(info.filename))
        result.writestr("_xmlsignatures/sig2.xml", signature_bytes)
    reports = product.Document.from_bytes(output.getvalue()).verify_signatures()
    assert type(reports) is tuple and len(reports) == 2
    assert reports[1]["signer"] is None and not reports[1]["cryptographically_valid"]
    assert reports[1]["issues"] == [{"kind": "orphan_signature_part", "details": {"part_name": "/_xmlsignatures/sig2.xml"}}]
    assert not any(hasattr(product, name) for name in ("SignatureReport", "_SignatureReport", "_SignatureIssue"))


def test_fx240_current_tamper_partial_unsigned_and_malformed():
    import rdocx as product
    owner = product.Document()
    owner.add_paragraph("signed text")
    key, cert = _fx240_material()
    owner.sign(key, cert)
    original = owner.to_bytes()
    owner.paragraphs[0].text = "changed text"
    reports = owner.verify_signatures()
    assert len(reports) == 1 and not reports[0]["cryptographically_valid"]
    issues = reports[0]["issues"]
    changed = [issue for issue in issues if issue["kind"] == "changed_reference"]
    assert any(issue["details"]["uri"].startswith("/word/document.xml?") for issue in changed)
    assert all(set(issue["details"]) == {"uri", "expected_digest", "actual_digest"} for issue in changed)
    partial = product.Document.from_bytes(_fx240_zip_edit(original, extra=True))
    report = partial.verify_signatures()[0]
    assert report["cryptographically_valid"] and not report["coverage_complete"]
    assert {"kind": "uncovered_part", "details": {"part_name": "/word/uncovered.xml"}} in report["issues"]
    malformed = _fx240_zip_edit(original, "_xmlsignatures/sig1.xml", b"http://www.w3.org/2001/04/xmldsig-more#rsa-sha256", b"urn:unsupported-signature")
    with pytest.raises(product.PackageError):
        product.Document.from_bytes(malformed).verify_signatures()
    assert product.Document().verify_signatures() == ()


def test_fx240_refusal_preserves_current_bytes_and_held_authority():
    import rdocx as product
    owner = product.Document()
    owner.add_paragraph("signed text")
    paragraph = owner.paragraphs[0]; run = paragraph.runs[0]
    before = owner.to_bytes()
    key, cert = _fx240_material()
    for bad in (b"", b"x" * (1024 * 1024 + 1), b"-----BEGIN RSA PRIVATE KEY-----Zg==-----END RSA PRIVATE KEY-----", key + b" "):
        with pytest.raises(product.PackageError):
            owner.sign(bad, cert)
        assert owner.to_bytes() == before and owner.verify_signatures() == ()
        assert paragraph.text == "signed text" and run.text == "signed text"
    for value in ("not bytes", 123, None):
        with pytest.raises(TypeError):
            owner.sign(value, cert)
    report = owner.sign(key, cert)
    assert report["cryptographically_valid"]
    assert paragraph.text == "signed text" and run.text == "signed text"
    saved = owner.to_bytes()
    assert owner.verify_signatures()[0]["cryptographically_valid"]
    assert owner.to_bytes() == saved
    assert paragraph.text == "signed text" and run.text == "signed text"


def test_fx240_fill_then_sign_sign_then_fill_and_unchanged_story_owners():
    import rdocx
    key, cert = _fx240_material()
    document = _fx241_five_kind_document()
    document.bind_content_controls({"plain": "Ada", "check": "true", "choice": "two", "date": "2026-10-16"})
    document.set_legacy_form_field_value("legacy", "filled")
    assert document.sign(key, cert)["coverage_complete"]
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.verify_signatures()[0]["cryptographically_valid"]
    assert reopened.legacy_form_fields()[0].value == "filled"
    reopened.bind_content_controls({"plain": "after signing"})
    report = reopened.verify_signatures()[0]
    assert not report["cryptographically_valid"]
    assert any(issue["kind"] == "changed_reference" and issue["details"]["uri"].startswith("/word/document.xml?") for issue in report["issues"])

    unchanged = _replace_document_body(rdocx.Document(), '<w:p><w:bookmarkStart w:id="0" w:name="anchor"/><w:r><w:t>held</w:t></w:r><w:bookmarkEnd w:id="0"/><w:hyperlink w:anchor="anchor"><w:r><w:t>link</w:t></w:r></w:hyperlink></w:p><w:sectPr/>')
    paragraph = unchanged.paragraphs[0]
    run = paragraph.runs[0]
    link = unchanged.hyperlinks[0]
    story = next(item for item in unchanged.stories if item.kind == "body")
    snapshot = unchanged.story_items[0]
    unchanged.sign(key, cert)
    assert paragraph.text == "heldlink" and run.text == "held" and link.anchor == "anchor"
    assert unchanged.verify_signatures()[0]["cryptographically_valid"]
    assert snapshot.xml
    unchanged.add_hyperlink_to_story(story, "still current", "https://example.com/after")
    assert any(item.text == "still current" for item in unchanged.hyperlinks)
    assert not unchanged.verify_signatures()[0]["cryptographically_valid"]


def test_fx240_late_graph_refusal_preserves_link_run_story_and_raw_parts():
    import rdocx
    owner = rdocx.Document()
    owner.add_paragraph("held").add_hyperlink("external", "https://example.com/")
    paragraph = owner.paragraphs[0]
    run = paragraph.runs[0]
    link = owner.hyperlinks[0]
    story = next(item for item in owner.stories if item.kind == "body")
    before = owner.to_bytes()
    key, cert = _fx240_material()
    with pytest.raises(rdocx.PackageError, match="relationship"):
        owner.sign(key, cert)
    assert owner.to_bytes() == before and owner.verify_signatures() == ()
    assert paragraph.text == "heldexternal" and run.text == "held" and link.url == "https://example.com/"
    owner.add_hyperlink_to_story(story, "current", "https://example.com/after")
    assert any(item.text == "current" for item in owner.hyperlinks)
    changed = rdocx.Document()
    changed.add_paragraph("same")
    old_run = changed.paragraphs[0].runs[0]
    paragraph = changed.paragraphs[0]
    paragraph.replace_xml(paragraph.xml.replace("<w:r>", "<w:r><w:rPr><w:b/></w:rPr>", 1))
    with pytest.raises(rdocx.StaleElementError):
        _ = old_run.text
    changed.sign(key, cert)
    with pytest.raises(rdocx.StaleElementError):
        _ = old_run.text


def test_fx250_chart_snapshots_follow_physical_occurrences():
    import rdocx
    d = rdocx.Document()
    for kind, name in [(51, 'first'), (65, 'second')]:
        data = rdocx.CategoryChartData()
        data.categories = ['A', 'B']
        data.add_series(name, [1, 2])
        d.add_chart(kind, data, rdocx.Inches(4), rdocx.Inches(2))
    d = rdocx.Document.from_bytes(d.to_bytes())
    held = d.paragraphs[0]
    before = d.to_bytes()
    charts = d.charts
    assert [(c.chart_type, c.categories, [(s.name, s.values) for s in c.series]) for c in charts] == [
        (51, ('A', 'B'), [('first', (1.0, 2.0))]),
        (65, ('A', 'B'), [('second', (1.0, 2.0))])]
    assert d.to_bytes() == before and held.xml
    with pytest.raises(AttributeError):
        charts[0].chart_type = 5
    for kind in (1, -4120, 5, 81, 74):
        one = rdocx.Document()
        if kind == 74:
            data = rdocx.XyChartData()
            series = data.add_series('xy')
            series.add_data_point(1, 7)
            series.add_data_point(2, 8)
        else:
            data = rdocx.CategoryChartData()
            data.categories = ['A', 'B']
            data.add_series('series', [7, 8])
        one.add_chart(kind, data, rdocx.Inches(4), rdocx.Inches(2))
        assert rdocx.Document.from_bytes(one.to_bytes()).charts[0].chart_type == kind


def test_fx250_chart_snapshot_scope_and_refusal():
    import copy
    import xml.etree.ElementTree as ET
    import rdocx
    d = rdocx.Document()
    data = rdocx.CategoryChartData()
    data.categories = ['A']
    data.add_series('one', [7])
    d.add_chart(51, data, rdocx.Inches(4), rdocx.Inches(2))
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    c = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
    root = ET.fromstring(parts['word/document.xml'])
    run = root.find('.//' + w + 'r')
    drawing = run.find(w + 'drawing')
    run.append(copy.deepcopy(drawing))
    opaque = ET.SubElement(run, '{urn:fx250:opaque}holder')
    opaque.append(copy.deepcopy(drawing))
    parts['word/document.xml'] = ET.tostring(root)
    def package():
        b = io.BytesIO()
        with zipfile.ZipFile(b, 'w', zipfile.ZIP_DEFLATED) as z:
            for n, value in parts.items():
                z.writestr(n, value)
        return rdocx.Document.from_bytes(b.getvalue())
    d = package()
    before = d.to_bytes()
    assert len(d.charts) == 2 and d.to_bytes() == before
    relname = 'word/_rels/document.xml.rels'
    rels = ET.fromstring(parts[relname])
    relation = next(n for n in rels if n.get('Type', '').endswith('/chart'))
    original_target = posixpath.normpath(posixpath.join('word', relation.get('Target'))).lstrip('/')
    if relation.get('Target').startswith('/'):
        original_target = relation.get('Target').lstrip('/')
    renamed = 'word/renamed/physical-chart.xml'
    parts[renamed] = parts.pop(original_target)
    relation.set('Target', 'renamed/physical-chart.xml')
    parts[relname] = ET.tostring(rels)
    content_types = ET.fromstring(parts['[Content_Types].xml'])
    for item in content_types:
        if item.get('PartName') == '/' + original_target:
            item.set('PartName', '/' + renamed)
    parts['[Content_Types].xml'] = ET.tostring(content_types)
    assert len(package().charts) == 2
    chart = renamed
    tree = ET.fromstring(parts[chart])
    plot = tree.find('.//' + c + 'plotArea')
    for item in list(plot):
        if item.tag.endswith('barChart'):
            item.tag = c + 'bubbleChart'
    parts[chart] = ET.tostring(tree)
    d = package()
    before = d.to_bytes()
    with pytest.raises(Exception, match='chart|Chart|plot'):
        _ = d.charts
    assert d.to_bytes() == before


def test_fx250_display_equation_preserves_inline_and_diagnostics():
    import rdocx
    d = rdocx.Document()
    p = d.add_paragraph('math')
    result = p.add_equation(r'\frac{a}{b}', format='latex', display=True)
    assert result.equation.is_display and p.equations()[0].is_display
    p.add_equation('x', format='latex')
    assert not p.equations()[1].is_display
    reopened = rdocx.Document.from_bytes(d.to_bytes())
    assert [e.is_display for e in reopened.paragraphs[0].equations()] == [True, False]
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        xml = z.read('word/document.xml')
    assert b'<m:oMathPara' in xml and b'<m:oMath>' in xml
    before = d.to_bytes()
    with pytest.raises(TypeError):
        p.add_equation('x', format='latex', display=1)
    assert d.to_bytes() == before


def test_fx250_numbered_restart_and_font_roundtrip():
    import xml.etree.ElementTree as ET
    import rdocx
    d = rdocx.Document()
    for text, restart in [('one', False), ('two', False), ('again one', True), ('again two', False)]:
        d.add_numbered_list_item(text, restart=restart)
    assert d.paragraphs[2].numbering != d.paragraphs[1].numbering
    assert d.paragraphs[2].numbering == d.paragraphs[3].numbering
    d = rdocx.Document.from_bytes(d.to_bytes())
    d.add_numbered_list_item('continued')
    assert d.paragraphs[4].numbering == d.paragraphs[3].numbering
    new_id = d.restart_numbering(d.paragraphs[4], start=5)
    assert d.paragraphs[4].numbering == (new_id, 0)
    definition = d.add_numbering_definition([rdocx.ListLevel(format='bullet', text='-', font='Arial')])
    d.add_paragraph('bullet').numbering = (d.add_numbering_instance(definition), 0)
    d = rdocx.Document.from_bytes(d.to_bytes())
    assert d.paragraphs[4].numbering == (new_id, 0)
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        tree = ET.fromstring(z.read('word/numbering.xml'))
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    override = tree.find(f"{w}num[@{w}numId='{new_id}']/{w}lvlOverride/{w}startOverride")
    assert override.get(w + 'val') == '5'
    assert any(n.get(w + 'ascii') == n.get(w + 'hAnsi') == 'Arial' for n in tree.iter(w + 'rFonts'))
    legacy = rdocx.Document()
    definition = legacy.add_numbering_definition([rdocx.ListLevel(), rdocx.ListLevel(format='upperRoman')])
    instance = legacy.add_numbering_instance(definition)
    legacy.add_numbered_list_item('nested without restart', level=1)
    assert legacy.paragraphs[0].numbering == (instance, 1)


def test_fx250_restart_numbering_refuses_before_publication():
    import rdocx
    d = rdocx.Document()
    d.add_numbered_list_item('one')
    p = d.paragraphs[0]
    class Index:
        calls = 0
        def __index__(self):
            self.calls += 1
            return 5
    callback = Index()
    for value in (True, callback, -1, 2**32):
        before = d.to_bytes()
        with pytest.raises((TypeError, ValueError, OverflowError)):
            d.restart_numbering(p, start=value)
        assert d.to_bytes() == before and p.text == 'one'
    assert callback.calls == 0
    plain = d.add_paragraph('plain')
    other = rdocx.Document().add_paragraph('foreign')
    for owner in (plain, other):
        before = d.to_bytes()
        with pytest.raises(Exception):
            d.restart_numbering(owner, start=5)
        assert d.to_bytes() == before
    current = d.paragraphs[0]
    new_id = d.restart_numbering(current, start=0)
    assert d.paragraphs[0].numbering == (new_id, 0)
    definition = d.add_numbering_definition([rdocx.ListLevel(), rdocx.ListLevel()])
    instance = d.add_numbering_instance(definition)
    d.add_paragraph('nested').numbering = (instance, 1)
    nested = d.paragraphs[-1]
    restarted = d.restart_numbering(nested, start=2**32 - 1)
    assert d.paragraphs[-1].numbering == (restarted, 1)
    for font in ('', 'bad\x01font'):
        before = d.to_bytes()
        with pytest.raises(ValueError):
            d.add_numbering_definition([rdocx.ListLevel(font=font)])
        assert d.to_bytes() == before


def test_fx250_note_reference_styles_and_remove_by_id():
    import xml.etree.ElementTree as ET
    import rdocx
    d = rdocx.Document()
    d.add_paragraph('first')
    first = d.add_footnote(d.paragraphs[0], 'remove me').id
    d.add_paragraph('second')
    keep = d.add_footnote(d.paragraphs[1], 'keep me').id
    end = d.add_endnote(d.paragraphs[1], 'ending').id
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        main = z.read('word/document.xml')
        tree = ET.fromstring(main)
        w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        for kind, style in [('footnoteReference', 'FootnoteReference'), ('endnoteReference', 'EndnoteReference')]:
            runs = [run for run in tree.iter(w + 'r') if run.find(w + kind) is not None]
            assert runs and all(run.find(w + 'rPr/' + w + 'rStyle').get(w + 'val') == style for run in runs)
        for part, marker, style in [('word/footnotes.xml', 'footnoteRef', 'FootnoteReference'), ('word/endnotes.xml', 'endnoteRef', 'EndnoteReference')]:
            tree = ET.fromstring(z.read(part))
            runs = [run for run in tree.iter(w + 'r') if run.find(w + marker) is not None]
            assert runs and all(run.find(w + 'rPr/' + w + 'rStyle').get(w + 'val') == style for run in runs)
    d.remove_footnote(first)
    d = rdocx.Document.from_bytes(d.to_bytes())
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        notes = z.read('word/footnotes.xml')
        assert b'keep me' in notes and b'remove me' not in notes
        assert str(keep).encode() in notes
    before = d.to_bytes()
    for bad in (True, 2**31, first):
        with pytest.raises(Exception):
            d.remove_footnote(bad)
        assert d.to_bytes() == before
    d.remove_endnote(end)
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        assert b'ending' not in z.read('word/endnotes.xml')


@pytest.mark.parametrize('abstract_format,replacement_format', [('decimal', 'bullet'), ('bullet', 'decimal')])
def test_fx250_effective_numbering_override_restart_and_continuation(abstract_format, replacement_format):
    import copy
    import xml.etree.ElementTree as ET
    import rdocx
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    d = rdocx.Document()
    definition = d.add_numbering_definition([rdocx.ListLevel(format=abstract_format)])
    instance = d.add_numbering_instance(definition, start=1)
    d.add_paragraph('selected').numbering = (instance, 0)
    with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    numbering = ET.fromstring(parts['word/numbering.xml'])
    abstract = numbering.find(f"{w}abstractNum[@{w}abstractNumId='{definition}']/{w}lvl")
    concrete = numbering.find(f"{w}num[@{w}numId='{instance}']/{w}lvlOverride")
    level = copy.deepcopy(abstract)
    level.find(w + 'numFmt').set(w + 'val', replacement_format)
    level.find(w + 'lvlText').set(w + 'val', '-' if replacement_format == 'bullet' else '%1.')
    concrete.append(level)
    parts['word/numbering.xml'] = ET.tostring(numbering)
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, value in parts.items():
            z.writestr(name, value)
    d = rdocx.Document.from_bytes(b.getvalue())
    held = d.paragraphs[0]
    before = d.to_bytes()
    if replacement_format == 'bullet':
        with pytest.raises(Exception, match='numbered|bullet'):
            d.restart_numbering(held, start=5)
        assert d.to_bytes() == before and held.text == 'selected'
        d.add_numbered_list_item('decimal continuation')
        assert d.paragraphs[-1].numbering[0] != instance
    else:
        new_id = d.restart_numbering(held, start=5)
        assert d.paragraphs[0].numbering == (new_id, 0)
        d = rdocx.Document.from_bytes(d.to_bytes())
        d.add_numbered_list_item('continued after reopen')
        assert d.paragraphs[-1].numbering == (new_id, 0)
        with zipfile.ZipFile(io.BytesIO(d.to_bytes())) as z:
            tree = ET.fromstring(z.read('word/numbering.xml'))
        concrete = tree.find(f"{w}num[@{w}numId='{new_id}']/{w}lvlOverride")
        assert concrete.find(w + 'startOverride').get(w + 'val') == '5'
        assert concrete.find(w + 'lvl/' + w + 'numFmt').get(w + 'val') == 'decimal'



def test_f284_versioned_paths_filters_escapes_and_legacy_selection():
    import rdocx
    legacy = rdocx.Document()
    legacy.add_paragraph("{{a|upper}}")
    assert legacy.render_template({"a|upper": "legacy", "a": "versioned"}) == 1
    assert legacy.paragraphs[0].text == "legacy"
    versioned = rdocx.Document()
    versioned.add_paragraph('{{ person["full.name"]|trim|upper }} / {{values|join(";")}} / {{missing|default("guest")}} / {{{{name}}}}')
    assert versioned.render_template({"person": {"full.name": " Ada "}, "values": [0, False, None]}, grammar_version=1) == 3
    assert versioned.paragraphs[0].text == "ADA / 0;false; / guest / {{name}}"
    assert rdocx.Document.from_bytes(versioned.to_bytes()).paragraphs[0].text == versioned.paragraphs[0].text


def test_f284_container_expansion_uses_existing_native_transaction():
    import rdocx
    document = rdocx.Document()
    document.add_paragraph("{% for item in items %}")
    document.add_paragraph("BODY {{item.name|upper}}")
    document.add_table(1, 1).cell(0, 0).text = "CELL {{item.name|lower}}"
    document.add_paragraph("{% endfor %}")
    document.set_header("HEADER {{title|trim}}")
    document.set_footer("FOOTER {{title|trim}}")
    assert document.render_template({"items": [{"name": "One"}, {"name": "Two"}], "title": " Ada "}, grammar_version=1) == 6
    assert [p.text for p in document.paragraphs] == ["BODY ONE", "BODY TWO"]
    assert [t.cell(0, 0).text for t in document.tables] == ["CELL one", "CELL two"]
    assert document.section(0).header.paragraphs[0].text == "HEADER Ada"
    assert document.section(0).footer.paragraphs[0].text == "FOOTER Ada"
    assert rdocx.Document.from_bytes(document.to_bytes()).text() == document.text()


def test_f284_version_and_data_refusals_preserve_bytes_and_held_owners():
    import rdocx
    document = rdocx.Document()
    held = document.add_paragraph("{{name}}")
    for version, error in [(0, ValueError), (2, ValueError), (-1, OverflowError), ("1", TypeError), (1.0, TypeError)]:
        before = document.to_bytes()
        with pytest.raises(error):
            document.render_template({"name": "Ada"}, grammar_version=version)
        assert document.to_bytes() == before
        assert held.text == "{{name}}"
    for payload in [{"name": float("nan")}, {"name": 2**64}, {1: "bad"}]:
        before = document.to_bytes()
        with pytest.raises((ValueError, TypeError)):
            document.render_template(payload, grammar_version=1)
        assert document.to_bytes() == before
        assert held.text == "{{name}}"


def test_f284_wrong_intermediate_default_and_inserted_text_are_literal():
    import rdocx
    document = rdocx.Document()
    held = document.add_paragraph('{{wrong.child|default("fallback")}}')
    before = document.to_bytes()
    with pytest.raises(ValueError):
        document.render_template({"wrong": 3}, grammar_version=1)
    assert document.to_bytes() == before
    assert held.text == '{{wrong.child|default("fallback")}}'
    literal = rdocx.Document()
    literal.add_paragraph("{{name}} / {{other}}")
    assert literal.render_template({"name": "{{other}}\ue000rdocx-template-0\ue001", "other": "DATA"}, grammar_version=1) == 2
    assert literal.paragraphs[0].text == "{{other}}\ue000rdocx-template-0\ue001 / DATA"


W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c63f8cfc0f01f00050001ff89993d1d0000000049454e44ae426082"
)


def _fx261_replace_part(document, name, transform):
    import rdocx

    source = io.BytesIO(document.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == name:
                data = transform(data)
            result_zip.writestr(info, data)
    return rdocx.Document.from_bytes(output.getvalue())


def _fx261_add_part(document, name, payload, content_type):
    import rdocx

    source = io.BytesIO(document.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "[Content_Types].xml":
                override = (
                    f'<Override PartName="/{name}" ContentType="{content_type}"/>'
                ).encode()
                data = data.replace(b"</Types>", override + b"</Types>")
            result_zip.writestr(info, data)
        result_zip.writestr(name, payload)
    return rdocx.Document.from_bytes(output.getvalue())


def _fx261_set_doc_defaults(document, defaults):
    block = f"<w:docDefaults>{defaults}</w:docDefaults>"

    def transform(data):
        text = data.decode()
        text, count = re.subn(r"<w:docDefaults\b[^>]*/>", block, text, count=1)
        if count == 0:
            text, count = re.subn(
                r"<w:docDefaults\b.*?</w:docDefaults>", block, text, count=1, flags=re.S
            )
        if count == 0:
            root = text.find("<w:styles")
            root_end = text.find(">", root) + 1
            text = text[:root_end] + block + text[root_end:]
        return text.encode()

    return _fx261_replace_part(document, "word/styles.xml", transform)


def _fx261_defaults(xml):
    start = xml.find("<w:docDefaults")
    end = xml.find("</w:docDefaults>")
    return xml[start:end]


def _fx261_declared(xml, declaration_first):
    end = xml.find(">")
    declaration = f' xmlns:w14="{W14}"'
    attribute = ' w14:paraId="388B6825"'
    if declaration_first:
        return xml[:end] + declaration + attribute + xml[end:]
    return xml[:end] + attribute + declaration + xml[end:]


def _fx261_without_comments_extended(document):
    import rdocx

    source = io.BytesIO(document.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            if info.filename == "word/commentsExtended.xml":
                continue
            data = source_zip.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = re.sub(rb"<Override[^>]*commentsExtended\.xml[^>]*/>", b"", data)
            if info.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb"<Relationship[^>]*commentsExtended\.xml[^>]*/>", b"", data)
            result_zip.writestr(info, data)
    return rdocx.Document.from_bytes(output.getvalue())


def test_issue387_custom_properties_are_enumerable():
    import rdocx

    document = rdocx.Document()
    properties = document.custom_properties
    assert len(properties) == 0
    assert list(properties) == []
    assert properties.items() == ()
    with pytest.raises(KeyError):
        properties["Zebra"]
    properties.set("Zebra", "z")
    properties.set("Alpha", 2)
    assert [item.name for item in document.custom_properties] == ["Zebra", "Alpha"]
    assert [name for name, _item in document.custom_properties.items()] == ["Zebra", "Alpha"]
    assert document.custom_properties["Alpha"].value == 2
    assert document.custom_properties["Zebra"].value == "z"
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert [item.name for item in reopened.custom_properties] == ["Zebra", "Alpha"]


def test_issue390_comment_resolution_preserves_unknown_state():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("Comment target")
    comment_id = document.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=0, run_index=0),
            end=rdocx.RunPosition(body_index=0, run_index=1),
        ),
        author="Alice",
        text="Review this",
        initials="AL",
    )
    assert document.comments[0].resolved is False
    assert document.resolve_comment(comment_id, resolved=True) is True
    assert document.comments[0].resolved is True
    unknown = _fx261_without_comments_extended(document)
    assert unknown.comments[0].resolved is None
    assert unknown.comments[0].text == "Review this"
    explicit = rdocx.Comment(
        id=comment_id,
        author="Alice",
        initials="AL",
        date=None,
        text="Review this",
        parent_id=None,
        resolved=None,
    )
    assert explicit.resolved is None
    assert rdocx.Document.from_bytes(document.to_bytes()).comments[0].resolved is True


def test_issue391_part_inspection_reflects_current_package():
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("inspect")
    document = _fx261_add_part(
        document,
        "word/a-unreferenced.xml",
        b"<?xml version=\"1.0\"?><orphan/>",
        "application/xml",
    )
    names = document.part_names()
    assert names == tuple(sorted(names))
    assert all(name.startswith("/") for name in names)
    assert "/word/a-unreferenced.xml" in names
    assert document.part_xml("WORD/A-UNREFERENCED.XML") == '<?xml version="1.0"?><orphan/>'
    document.add_style("Inspected", "paragraph")
    assert "Inspected" in document.part_xml("/word/styles.xml")
    after_style = document.part_names()
    assert tuple(name for name in after_style if name in names) == names
    document.add_paragraph("rewritten")
    assert document.part_names() == after_style
    with pytest.raises(rdocx.RdocxError, match="missing"):
        document.part_xml("/word/missing.xml")
    pictured = rdocx.Document()
    pictured.add_picture(PNG, "pixel.png", rdocx.Pt(1), rdocx.Pt(1))
    media = next(name for name in pictured.part_names() if "/media/" in name)
    with pytest.raises(rdocx.RdocxError, match="not UTF-8"):
        pictured.part_xml(media)
    utf16 = _fx261_add_part(
        document,
        "word/utf16.xml",
        "<?xml version=\"1.0\"?><orphan/>".encode("utf-16"),
        "application/xml",
    )
    with pytest.raises(rdocx.RdocxError, match="not UTF-8"):
        utf16.part_xml("/word/utf16.xml")
    bom = _fx261_add_part(
        rdocx.Document(),
        "word/bom.xml",
        b"\xef\xbb\xbf<bom/>",
        "application/xml",
    )
    assert bom.part_xml("/word/bom.xml") == "<bom/>"


def test_issue391_default_paragraph_format_is_owned_and_atomic():
    import rdocx

    document = _fx261_set_doc_defaults(
        rdocx.Document(),
        '<w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri"/></w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr><w:spacing w:before="120" w:after="240" '
        'w:line="360" w:lineRule="auto"/></w:pPr></w:pPrDefault>'
        "<!--opaque-default-->",
    )
    current = document.default_paragraph_format
    assert current.space_before == rdocx.Pt(6)
    assert current.space_after == rdocx.Pt(12)
    assert current.line_spacing == 1.5
    before = document.to_bytes()
    document.set_default_paragraph_format(space_before=rdocx.Pt(6))
    assert document.to_bytes() == before
    document.set_default_paragraph_format(space_before=None)
    defaults = _fx261_defaults(document.part_xml("/word/styles.xml"))
    assert 'w:ascii="Calibri"' in defaults
    assert "<!--opaque-default-->" in defaults
    assert 'w:before="120"' not in defaults
    assert 'w:after="240"' in defaults
    assert 'w:line="360"' in defaults
    before = document.to_bytes()
    for value in (1.5, 12, True, "12"):
        with pytest.raises(TypeError):
            document.set_default_paragraph_format(line_spacing=value)
        assert document.to_bytes() == before
    document.set_default_paragraph_format(line_spacing=rdocx.Pt(12))
    defaults = _fx261_defaults(document.part_xml("/word/styles.xml"))
    assert 'w:line="240"' in defaults
    assert 'w:lineRule="exact"' in defaults
    assert 'w:ascii="Calibri"' in defaults


def test_issue391_default_paragraph_line_rules_keep_imported_states():
    import rdocx

    def opened(spacing):
        return _fx261_set_doc_defaults(
            rdocx.Document(),
            f"<w:pPrDefault><w:pPr><w:spacing {spacing}/></w:pPr></w:pPrDefault>",
        )

    assert opened('w:line="360" w:lineRule="auto"').default_paragraph_format.line_spacing == 1.5
    assert opened('w:line="259" w:lineRule="auto"').default_paragraph_format.line_spacing == 259 / 240
    assert opened('w:line="240" w:lineRule="exact"').default_paragraph_format.line_spacing == rdocx.Pt(12)
    assert opened('w:line="240" w:lineRule="atLeast"').default_paragraph_format.line_spacing == rdocx.Pt(12)
    assert opened('w:line="480"').default_paragraph_format.line_spacing == 2.0
    assert opened('w:before="120"').default_paragraph_format.line_spacing is None
    assert opened('w:line="240" w:lineRule="precise"').default_paragraph_format.line_spacing is None

    at_least = opened('w:line="240" w:lineRule="atLeast"')
    before = at_least.to_bytes()
    at_least.set_default_paragraph_format(space_before=rdocx.Pt(6))
    defaults = _fx261_defaults(at_least.part_xml("/word/styles.xml"))
    assert 'w:lineRule="atLeast"' in defaults
    assert 'w:line="240"' in defaults
    missing_rule = opened('w:line="480"')
    missing_rule.set_default_paragraph_format(space_after=rdocx.Pt(6))
    defaults = _fx261_defaults(missing_rule.part_xml("/word/styles.xml"))
    assert 'w:line="480"' in defaults
    assert "w:lineRule" not in defaults
    precise = opened('w:line="240" w:lineRule="precise"')
    precise.set_default_paragraph_format(space_before=rdocx.Pt(6))
    defaults = _fx261_defaults(precise.part_xml("/word/styles.xml"))
    assert 'w:lineRule="precise"' in defaults

    assigned = opened('w:line="360" w:lineRule="auto"')
    assigned.set_default_paragraph_format(line_spacing=rdocx.Pt(12))
    defaults = _fx261_defaults(assigned.part_xml("/word/styles.xml"))
    assert 'w:line="240"' in defaults and 'w:lineRule="exact"' in defaults
    assigned.set_default_paragraph_format(line_spacing=None)
    defaults = _fx261_defaults(assigned.part_xml("/word/styles.xml"))
    assert "w:line" not in defaults and "w:lineRule" not in defaults
    refused = opened('w:line="360" w:lineRule="auto"')
    before = refused.to_bytes()
    with pytest.raises(TypeError):
        refused.set_default_paragraph_format(line_spacing=1.5)
    assert refused.to_bytes() == before


def test_issue391_root_attributes_use_declared_namespace_scope():
    import xml.etree.ElementTree as ET

    import rdocx

    def assert_bound(owner, declaration_first):
        document = rdocx.Document()
        document.add_paragraph("held")
        document.add_table(1, 1).cell(0, 0).text = "cell"
        target = {
            "paragraph": document.paragraphs[0],
            "table": document.tables[0],
            "cell": document.tables[0].cell(0, 0),
        }[owner]
        target.replace_xml(_fx261_declared(target.xml, declaration_first))
        root = ET.fromstring(_document_xml(document))
        found = [
            element
            for element in root.iter()
            if element.get(f"{{{W14}}}paraId") == "388B6825"
        ]
        assert found, owner

    for owner in ("paragraph", "table", "cell"):
        assert_bound(owner, True)
        assert_bound(owner, False)
        document = rdocx.Document()
        document.add_paragraph("held")
        document.add_table(1, 1).cell(0, 0).text = "cell"
        target = {
            "paragraph": document.paragraphs[0],
            "table": document.tables[0],
            "cell": document.tables[0].cell(0, 0),
        }[owner]
        end = target.xml.find(">")
        unbound = target.xml[:end] + ' w14:paraId="388B6825"' + target.xml[end:]
        before = document.to_bytes()
        with pytest.raises(rdocx.RdocxError, match="unbound"):
            target.replace_xml(unbound)
        assert document.to_bytes() == before


def test_issue391_story_item_snapshots_are_namespace_complete():
    import xml.etree.ElementTree as ET

    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body paragraph")
    document.add_table(1, 1).cell(0, 0).text = "cell paragraph"
    document.set_header("header paragraph")
    content = "<w:txbxContent><w:p><w:r><w:t>box paragraph</w:t></w:r></w:p></w:txbxContent>"
    extra = (
        '<w:p><w:r><w:pict xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<v:shape id="only"><v:textbox>{content}</v:textbox></v:shape>'
        "</w:pict></w:r></w:p>"
    )

    def insert(data):
        text = data.decode()
        marker = text.rfind("<w:sectPr")
        if marker < 0:
            marker = text.rfind("</w:body>")
        return (text[:marker] + extra + text[marker:]).encode()

    document = _fx261_replace_part(document, "word/document.xml", insert)
    kinds = {(item.story.kind, item.kind) for item in document.story_items}
    assert ("body", "paragraph") in kinds
    assert ("body", "table") in kinds
    assert ("table_cell", "paragraph") in kinds
    assert ("header", "paragraph") in kinds
    assert ("text_box", "paragraph") in kinds
    for item in document.story_items:
        ET.fromstring(item.xml)
        if item.text:
            assert item.text.encode() in item.xml or item.text in item.xml.decode()


def _fx261_text_box(kind):
    import rdocx

    content = "<w:txbxContent><w:p><w:r><w:t>box NEEDLE</w:t></w:r></w:p></w:txbxContent>"
    drawing = (
        '<w:drawing xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wp:inline><wp:extent cx="914400" cy="457200"/>'
        '<wp:docPr id="1" name="box"/><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f"<wps:wsp><wps:txbx>{content}</wps:txbx></wps:wsp>"
        "</a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    vml = (
        '<w:pict xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<v:shape id="fallback"><v:textbox>{content}</v:textbox></v:shape></w:pict>'
    )
    if kind == "drawing":
        inner = drawing
    elif kind == "vml":
        inner = vml
    else:
        inner = (
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            f"<mc:Choice Requires=\"wps\">{drawing}</mc:Choice>"
            f"<mc:Fallback>{vml}<!--preserve alternate producer--></mc:Fallback></mc:AlternateContent>"
        )
    body = f'<w:p><w:r><w:t>held outside</w:t></w:r></w:p><w:p><w:r>{inner}</w:r></w:p>'
    return _replace_document_body(rdocx.Document(), body)


def test_issue391_textbox_paragraph_handles_preserve_physical_custody():
    import rdocx

    def box_paragraph(document):
        return next(
            item.paragraph
            for item in document.story_items
            if item.story.kind == "text_box" and item.kind == "paragraph"
        )

    for kind, copies in (("vml", 1), ("drawing", 1), ("pair", 2)):
        document = _fx261_text_box(kind)
        held = box_paragraph(document)
        held.paragraph_format.space_before = rdocx.Pt(18)
        xml = _document_xml(document)
        assert xml.count(b'w:before="360"') == copies
        assert b"box NEEDLE" in xml
        with pytest.raises(rdocx.StaleElementError):
            _ = held.text
        reopened = rdocx.Document.from_bytes(document.to_bytes())
        assert _document_xml(reopened).count(b'w:before="360"') == copies

    paired = _fx261_text_box("pair")
    outside = paired.paragraphs[0]
    outside.paragraph_format.space_before = rdocx.Pt(6)
    xml = _document_xml(paired)
    assert xml.count(b'w:before="120"') == 1
    assert b"box NEEDLE" in xml

    source = _document_xml(_fx261_text_box("pair")).decode()
    changed = source.replace(
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>box NEEDLE</w:t></w:r></w:p></w:txbxContent></v:textbox>",
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>different text</w:t></w:r></w:p></w:txbxContent></v:textbox>",
        1,
    )
    assert changed != source
    refused = _fx261_replace_part(
        rdocx.Document.from_bytes(_fx261_text_box("pair").to_bytes()),
        "word/document.xml",
        lambda _data: changed.encode(),
    )
    selected = box_paragraph(refused)
    before = refused.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="asymmetric"):
        selected.paragraph_format.space_before = rdocx.Pt(18)
    assert refused.to_bytes() == before

    table_document = rdocx.Document()
    table_document.add_table(1, 1)
    table_item = next(item for item in table_document.story_items if item.kind == "table")
    assert table_item.paragraph is None
    story = rdocx.Story(kind="body", part_name="/word/document.xml", owner_index=0)
    built = rdocx.StoryItem(story=story, kind="paragraph", index_path=(0,), text="built")
    assert built.paragraph is None


def test_f287_custom_xml_store_binding_and_refusal():
    import rdocx
    import inspect

    assert inspect.signature(rdocx.Document.add_custom_xml_store).parameters["schema_uris"].default == ()
    document = rdocx.Document()
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="empty"):
        document.add_custom_xml_store(b"")
    with pytest.raises(rdocx.RdocxError, match="fragment"):
        document.add_custom_xml_store(b"<a/><b/>", ["urn:schema"])
    with pytest.raises(rdocx.RdocxError, match="empty"):
        document.add_custom_xml_store(b"<root/>", [""])
    assert document.to_bytes() == before

    item = b'<ns:root xmlns:ns="urn:example"><ns:child>old</ns:child></ns:root>'
    item_id = document.add_custom_xml_store(item, ["urn:schema?a=1&b=2"])
    assert item_id == "{00000000-0000-4000-8000-000000000000}"
    second = document.add_custom_xml_store(b"<again/>", ["urn:repeat", "urn:repeat"])
    assert second != item_id
    body = next(story for story in document.stories if story.kind == "body")
    document.insert_content_control(
        rdocx.ContentLocation.end(body), "block", "rich_text", tag="bound"
    )
    location = _content_location(document, lambda item: item.kind == "content_control")
    before = document.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="XPath"):
        document.set_content_control_binding(
            location, item_id, "//ns:root/ns:child", "xmlns:ns='urn:example'"
        )
    assert document.to_bytes() == before
    document.set_content_control_binding(
        location, item_id, "/ns:root/ns:child", "xmlns:ns='urn:example'"
    )
    document.set_content_control_value_by_tag("bound", "filled")
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert reopened.content_controls()[0].text == "filled"
    import io
    import zipfile
    with zipfile.ZipFile(io.BytesIO(reopened.to_bytes())) as archive:
        stored = archive.read("customXml/item1.xml")
        properties = archive.read("customXml/itemProps1.xml")
    assert b"<ns:child>filled</ns:child>" in stored
    assert b"urn:schema?a=1&amp;b=2" in properties
    assert item_id.encode() in properties


def test_issue389_tail_removal_retires_held_owners_without_resurrection():
    import rdocx

    doc = rdocx.Document()
    for text in ("first", "middle", "last"):
        doc.add_paragraph(text)
    first, last = doc.paragraphs[0], doc.paragraphs[2]
    held_run = first.runs[0]
    held_format = first.paragraph_format
    before = doc.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        first.replace_xml("<broken")
    assert doc.to_bytes() == before
    first.replace_xml(first.xml.replace(">first<", ">changed<"))
    assert first.text == "changed"
    assert last.text == "last"
    held_run = first.runs[0]
    held_format = first.paragraph_format
    assert doc.remove_content(2)
    for operation in (lambda: first.text, lambda: held_run.text,
                      lambda: held_format.space_before, lambda: last.text,
                      lambda: last.replace_xml("<broken")):
        with pytest.raises(rdocx.StaleElementError, match="stale"):
            operation()
    fresh = doc.paragraphs[0]
    assert fresh.text == "changed"
    fresh.replace_xml(fresh.xml.replace(">changed<", ">again<"))
    assert fresh.text == "again"
    doc.add_paragraph("replacement")
    assert doc.paragraphs[2].text == "replacement"
    with pytest.raises(rdocx.StaleElementError, match="stale"):
        _ = last.text
    assert doc.remove_content(2)
    assert doc.remove_content(1)
    reopened = rdocx.Document.from_bytes(doc.to_bytes())
    assert [p.text for p in reopened.paragraphs] == ["again"]


@pytest.mark.parametrize("edit", ["body", "cell", "nested", "header", "xml", "property", "row", "tail", "noop_refusal"])
def test_issue398_read_indexes_follow_current_owners_after_edits(edit):
    import rdocx

    document = rdocx.Document()
    document.add_paragraph("body before")
    document.add_paragraph("second body")
    document.add_table(1, 1).cell(0, 0).text = "outer before"
    document.tables[0].cell(0, 0).add_table(1, 1).cell(0, 0).text = "nested before"
    document.section(0).header.add_paragraph("header before")

    def observe(doc):
        outer = doc.tables[0].cell(0, 0)
        return {
            "body": tuple(paragraph.text for paragraph in doc.paragraphs),
            "outer": tuple(paragraph.text for paragraph in outer.paragraphs),
            "nested": tuple(paragraph.text for paragraph in outer.tables[0].cell(0, 0).paragraphs),
            "header": tuple(paragraph.text for paragraph in doc.section(0).header.paragraphs),
            "rows": tuple(len(row.cells) for row in doc.tables[0].rows),
        }

    before = document.to_bytes()
    initial = observe(document)
    assert observe(document) == initial
    assert document.to_bytes() == before
    expected = dict(initial)
    if edit == "body":
        document.paragraphs[0].text = "body after"
        expected["body"] = ("body after", "second body")
    elif edit == "cell":
        document.tables[0].cell(0, 0).paragraphs[0].text = "outer after"
        expected["outer"] = ("outer after",) + initial["outer"][1:]
    elif edit == "nested":
        document.tables[0].cell(0, 0).tables[0].cell(0, 0).paragraphs[0].text = "nested after"
        expected["nested"] = ("nested after",)
    elif edit == "header":
        document.section(0).header.paragraphs[-1].text = "header after"
        expected["header"] = initial["header"][:-1] + ("header after",)
    elif edit == "xml":
        paragraph = document.paragraphs[0]
        paragraph.replace_xml(paragraph.xml.replace("body before", "xml after"))
        expected["body"] = ("xml after", "second body")
    elif edit == "property":
        document.tables[0].cell(0, 0).width = 200_000
    elif edit == "row":
        document.tables[0].add_row()
        expected["rows"] = (1, 1)
    elif edit == "tail":
        document.add_paragraph("removed")
        held = document.paragraphs[0]
        assert document.remove_content(document.find_content_index("removed"))
        with pytest.raises(rdocx.StaleElementError):
            _ = held.text
        assert observe(document) == expected
        document.add_paragraph("replacement")
        expected["body"] = initial["body"] + ("replacement",)
        with pytest.raises(rdocx.StaleElementError):
            _ = held.text
    else:
        held = document.paragraphs[0]
        held.replace_xml(held.xml)
        assert held.text == "body before"
        stable = document.to_bytes()
        with pytest.raises(rdocx.RdocxError):
            held.replace_xml(held.xml.replace("</w:p>", "<w:notAChild/></w:p>"))
        assert document.to_bytes() == stable
        assert held.text == "body before"
    assert observe(document) == expected
    assert observe(rdocx.Document.from_bytes(document.to_bytes())) == expected


@pytest.mark.parametrize("instruction, expected", [
    (r'TOC \o "1-1" \h', ("Chapter",)),
    (r'TOC \o "1-1" \h \z \u', ("Chapter", "Outline one")),
    (r'TOC \o "1-1" \u', ("Chapter", "Outline one")),
    (r'TOC \o "2-2" \u', ("Section", "Outline two")),
    (r'TOC \u', ("Chapter", "Section", "Outline two", "Outline one")),
])
def test_issue400_python_toc_rebuild_bounds_outline_entries_after_reopen(instruction, expected):
    import rdocx

    document = rdocx.Document()
    document.add_heading("Chapter", level=1)
    document.add_heading("Section", level=2)
    document.add_paragraph("Outline two").paragraph_format.outline_level = 1
    document.add_paragraph("Outline one").paragraph_format.outline_level = 0
    document.insert_toc(0, max_level=1)
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(document.to_bytes())) as source, zipfile.ZipFile(result, "w") as out:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "word/document.xml":
                data, count = re.subn(rb'(<w:instrText\b[^>]*>).*?(</w:instrText>)',
                                      lambda match: match[1] + instruction.encode() + match[2], data, flags=re.S)
                assert count == 1
            out.writestr(member, data)
    document = rdocx.Document.from_bytes(result.getvalue())

    def entries(doc):
        return tuple(paragraph.text.split("\t")[0] for paragraph in doc.paragraphs
                     if (paragraph.style or "").upper().startswith("TOC"))

    assert document.rebuild_toc().entry_count == len(expected)
    assert entries(document) == expected
    reopened = rdocx.Document.from_bytes(document.to_bytes())
    assert entries(reopened) == expected
    assert reopened.rebuild_toc().entry_count == len(expected)
    assert entries(reopened) == expected


@pytest.mark.parametrize("host", ["body", "cell", "nested", "header", "paired"])
@pytest.mark.parametrize("operation", ["tab", "line", "page", "column"])
def test_issue398_run_tabs_and_breaks_refresh_live_paragraph_text(host, operation):
    import rdocx

    if host == "paired":
        document = _fx261_text_box("pair")
    else:
        document = rdocx.Document()
        document.add_paragraph("before")
        document.add_table(1, 1).cell(0, 0).text = "before"
        document.tables[0].cell(0, 0).add_table(1, 1).cell(0, 0).text = "before"
        document.section(0).header.add_paragraph("before")

    def select(doc):
        if host == "body":
            return doc.paragraphs[0]
        if host == "cell":
            return doc.tables[0].cell(0, 0).paragraphs[0]
        if host == "nested":
            return doc.tables[0].cell(0, 0).tables[0].cell(0, 0).paragraphs[0]
        if host == "header":
            return doc.section(0).header.paragraphs[-1]
        return next(item.paragraph for item in doc.story_items
                    if item.story.kind == "text_box" and item.kind == "paragraph")

    paragraph = select(document)
    initial = "box NEEDLE" if host == "paired" else "before"
    assert paragraph.text == initial
    run = paragraph.runs[0]
    font = run.font
    bold = font.bold
    before = document.to_bytes()
    if host == "header" and operation == "page":
        with pytest.raises(ValueError, match="page break.*header"):
            run.add_break("page")
        assert document.to_bytes() == before
        expected = initial
    else:
        if operation == "tab":
            run.add_tab()
        elif operation == "line":
            run.add_break()
        else:
            run.add_break(operation)
        expected = initial + ("\t" if operation == "tab" else "\n")
    assert run.text == expected
    assert paragraph.text == expected
    assert select(document).text == expected
    assert font.bold == bold
    if host == "paired":
        tag = b"tab" if operation == "tab" else b"br"
        assert len(re.findall(rb"<w:" + tag + rb"(?:\s|/|>)", _document_xml(document))) == 2
    saved = document.to_bytes()
    assert paragraph.text == select(document).text == expected
    assert select(rdocx.Document.from_bytes(saved)).text == expected


@pytest.mark.parametrize("target", [0, 1])
@pytest.mark.parametrize("operation", ["tab", "line", "run_text", "paragraph_text", "append"])
@pytest.mark.parametrize("leading_direct", [False, True])
def test_issue398_controlled_cell_edit_selects_read_owner(target, operation, leading_direct):
    import rdocx
    content = ('<w:tbl><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid>'
               '<w:tr><w:tc><w:sdt><w:sdtContent><w:p><w:r><w:t>controlled</w:t>'
               '</w:r></w:p></w:sdtContent></w:sdt><w:p><w:r><w:t>after</w:t>'
               '</w:r></w:p></w:tc></w:tr></w:tbl>')
    if leading_direct:
        content = content.replace('<w:tc>', '<w:tc><w:p><w:r><w:t>before</w:t></w:r></w:p>')
        target += 1
    doc = _replace_document_body(rdocx.Document(), content)
    held = list(doc.tables[0].cell(0, 0).paragraphs)
    paragraph = held[target]
    run = paragraph.runs[0]
    font = run.font
    expected = (["before"] if leading_direct else []) + ["controlled", "after"]
    if operation == "tab":
        run.add_tab()
        expected[target] += "\t"
    elif operation == "line":
        run.add_break()
        expected[target] += "\n"
    elif operation == "run_text":
        run.text = "changed"
        expected[target] = "changed"
    elif operation == "paragraph_text":
        paragraph.text = "changed"
        expected[target] = "changed"
    else:
        paragraph.add_run("Z")
        expected[target] += "Z"
    assert [p.text for p in held] == expected
    assert [p.text for p in doc.tables[0].cell(0, 0).paragraphs] == expected
    if operation in ("tab", "line"):
        assert run.text == expected[target]
        assert font.bold is None
    out = io.BytesIO()
    doc.save(out)
    reopened = rdocx.Document.from_bytes(out.getvalue())
    assert [p.text for p in reopened.tables[0].cell(0, 0).paragraphs] == expected


@pytest.mark.parametrize("run_count", [1, 2])
def test_issue389_first_run_local_publication_keeps_exact_owner_lifetimes(run_count):
    import rdocx
    doc = rdocx.Document()
    doc.add_paragraph("before")
    doc.add_paragraph("tail")
    paragraph = doc.insert_paragraph(1, "")
    old_runs = paragraph.runs
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           + ''.join(f'<w:r><w:t>r{i}</w:t></w:r>' for i in range(run_count))
           + '</w:p>')
    paragraph.replace_xml(xml)
    expected = ''.join(f'r{i}' for i in range(run_count))
    assert paragraph.text == expected
    with pytest.raises(rdocx.StaleElementError):
        len(old_runs)
    run = paragraph.runs[0]
    font = run.font
    run.add_tab()
    font.bold = True
    assert font.bold is True
    expected = 'r0\t' + ''.join(f'r{i}' for i in range(1, run_count))
    assert paragraph.text == expected
    current_xml = paragraph.xml
    paragraph.replace_xml(current_xml)
    assert run.text == 'r0\t'
    before = doc.to_bytes()
    with pytest.raises(rdocx.RdocxError):
        paragraph.replace_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:invalid/></w:p>')
    assert doc.to_bytes() == before
    assert run.text == 'r0\t'
    paragraph.replace_xml(current_xml.replace('>r0<', '>changed<'))
    with pytest.raises(rdocx.StaleElementError):
        _ = run.text
    assert paragraph.runs[0].text == 'changed\t'
    # A later tail removal must update the appended run indices, then a full
    # publication must reconstruct the same physical owners without history.
    assert doc.remove_content(2)
    paragraph = doc.paragraphs[1]
    assert paragraph.runs[0].text == 'changed\t'
    paragraph.add_run('end')
    expected = 'changed\t' + ''.join(f'r{i}' for i in range(1, run_count)) + 'end'
    assert doc.paragraphs[1].text == expected
    stream = io.BytesIO()
    doc.save(stream)
    assert rdocx.Document.from_bytes(stream.getvalue()).paragraphs[1].text == expected


@pytest.mark.parametrize("path", ["first_run", "same_runs", "changed_runs"])
@pytest.mark.parametrize("section", ["<w:sectPr/>", "<w:sectPr></w:sectPr>", '<w:sectPr><w:pgSz w:w="12345"/></w:sectPr>'])
def test_issue389_paragraph_xml_refuses_section_introduction_on_every_path(path, section):
    import rdocx
    doc = rdocx.Document()
    paragraph = doc.add_paragraph("" if path == "first_run" else "old")
    runs = '<w:r><w:t>new</w:t></w:r>' * (2 if path == "changed_runs" else 1)
    xml = '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:pPr>' + section + '</w:pPr>' + runs + '</w:p>'
    before = doc.to_bytes()
    with pytest.raises(rdocx.RdocxError, match="section"):
        paragraph.replace_xml(xml)
    assert doc.to_bytes() == before
    assert paragraph.text == ("" if path == "first_run" else "old")


@pytest.mark.parametrize("prefix", ["w", "q"])
@pytest.mark.parametrize("section", ["<{p}:sectPr/>", "<{p}:sectPr></{p}:sectPr>", '<{p}:sectPr><{p}:pgSz {p}:w="12345"/></{p}:sectPr>', '<{p}:sectPr xmlns:x="urn:x" x:flag="keep"/>'])
def test_issue389_paragraph_xml_preserves_imported_section_and_refuses_removal(prefix, section):
    import rdocx
    section = section.format(p=prefix)
    if prefix != "w":
        section = section.replace(f'<{prefix}:sectPr', f'<{prefix}:sectPr xmlns:{prefix}="http://schemas.openxmlformats.org/wordprocessingml/2006/main"', 1)
    body = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:pPr>{section}</w:pPr><w:r><w:t>old</w:t></w:r></w:p>')
    doc = _replace_document_body(rdocx.Document(), body)
    paragraph = doc.paragraphs[0]
    original = paragraph.xml
    paragraph.replace_xml(original.replace('>old<', '>new<'))
    assert paragraph.xml == original.replace('>old<', '>new<')
    assert paragraph.text == 'new'
    assert rdocx.Document.from_bytes(doc.to_bytes()).paragraphs[0].text == 'new'
    if 'x:flag="keep"' in paragraph.xml:
        before_attribute = doc.to_bytes()
        with pytest.raises(rdocx.RdocxError, match='section'):
            paragraph.replace_xml(paragraph.xml.replace('x:flag="keep"', 'x:flag="changed"'))
        assert doc.to_bytes() == before_attribute
    # Removing the entire pPr avoids making assumptions about retained prefixes.
    without_section = re.sub(r'<w:pPr\b[^>]*>.*?</w:pPr>', '', paragraph.xml, flags=re.S)
    assert without_section != paragraph.xml
    before = doc.to_bytes()
    with pytest.raises(rdocx.RdocxError, match='section'):
        paragraph.replace_xml(without_section)
    assert doc.to_bytes() == before
    assert paragraph.text == 'new'


def test_issue389_retained_paragraph_alias_never_publishes_unsavable_text():
    import rdocx
    doc = _replace_document_body(rdocx.Document(), '<q:p xmlns:q="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><q:pPr><q:sectPr></q:sectPr></q:pPr><q:r><q:t>old</q:t></q:r></q:p>')
    paragraph = doc.paragraphs[0]
    before = doc.to_bytes()
    try:
        paragraph.replace_xml(paragraph.xml.replace('>old<', '>new<'))
    except rdocx.RdocxError as error:
        assert 'namespace owner' in str(error)
        assert doc.to_bytes() == before
        assert paragraph.text == 'old'
    else:
        assert paragraph.text == 'new'
        assert rdocx.Document.from_bytes(doc.to_bytes()).paragraphs[0].text == 'new'


def _issue407_document(root_w14):
    import rdocx

    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    body = "".join([
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>heading</w:t></w:r></w:p>',
        '<w:p><w:r><w:t xml:space="preserve">see </w:t></w:r><w:hyperlink w:anchor="target"><w:r><w:t>target</w:t></w:r></w:hyperlink></w:p>',
        '<w:p><w:bookmarkStart w:id="1" w:name="target"/><w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple><w:bookmarkEnd w:id="1"/></w:p>',
        '<w:tbl><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>cell</w:t></w:r></w:p>'
        '<w:tbl><w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>nested</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        '<w:p/></w:tc></w:tr></w:tbl>',
        '<w:sdt><w:sdtPr><w:alias w:val="block"/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>block one</w:t></w:r></w:p><w:p><w:r><w:t>block two</w:t></w:r></w:p></w:sdtContent></w:sdt>',
        *(f'<w:p><w:r><w:rPr><w:i/></w:rPr><w:t>{text}</w:t></w:r><w:r><w:tab/><w:t>tail</w:t></w:r></w:p>' for text in "abc"),
        f'<w:p xmlns:w14="{w14}" w14:paraId="388B6825"><w:r><w:t>identified</w:t></w:r></w:p>',
        '<w:p><w:r><w:t>last</w:t></w:r></w:p>',
    ])
    document = rdocx.Document()
    document.add_paragraph("first")
    document.section(0).header.add_paragraph("header text")
    source = io.BytesIO(document.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "word/document.xml":
                text = data.decode()
                at = text.index("<w:sectPr")
                text = text[:at] + body + text[at:]
                if root_w14:
                    text = text.replace("<w:document ", f'<w:document xmlns:w14="{w14}" ', 1)
                data = text.encode()
            result_zip.writestr(info, data)
    return rdocx.Document.from_bytes(output.getvalue())


@pytest.mark.parametrize("root_w14", [False, True])
def test_issue407_plain_body_splices_keep_owner_lifetimes(root_w14):
    import rdocx

    doc = _issue407_document(root_w14)

    def texts():
        return [paragraph.text for paragraph in doc.paragraphs]

    def read_all():
        # Debug builds compare every direct read with the whole-source read.
        for paragraph in doc.paragraphs:
            assert paragraph.xml.startswith("<w:p")
            for run in paragraph.runs:
                assert run.xml.startswith("<w:r")
        for table in doc.tables:
            assert table.xml.startswith("<w:tbl")
            for row in table.rows:
                for cell in row.cells:
                    assert cell.xml.startswith("<w:tc")

    expected = ["first", "heading", "see target", "1", "block one", "block two",
                "a\ttail", "b\ttail", "c\ttail", "identified", "last"]
    assert texts() == expected
    read_all()
    # A read declares w14 when its root or its own bytes use it, never only
    # because another paragraph does.
    assert ("xmlns:w14" in doc.paragraphs[0].xml) == root_w14
    assert "xmlns:w14" in doc.paragraphs[9].xml
    edits = [
        ("Document.insert_paragraph", lambda: doc.insert_paragraph(2, "inserted"), 2, "inserted"),
        ("Document.insert_paragraph", lambda: doc.insert_paragraph(0, "zero"), 0, "zero"),
        ("Document.remove_content", lambda: doc.remove_content(0), 0, None),
        ("Document.insert_paragraph", lambda: doc.insert_paragraph(5, "before table", "Heading1"), 5, "before table"),
        ("Document.insert_paragraph", lambda: doc.insert_paragraph(7, "before control"), 6, "before control"),
        ("Document.remove_content", lambda: doc.remove_content(9), 9, None),
        ("Document.insert_paragraph", lambda: doc.paragraphs[9].insert_paragraph_before("before b"), 9, "before b"),
        ("Document.remove_content_range", lambda: doc.remove_content_range(10, 12), 10, None),
        ("Document.insert_paragraph", lambda: doc.insert_paragraph(12, "end"), 12, "end"),
    ]
    for operation, edit, at, text in edits:
        held = doc.paragraphs[0]
        held_run = held.runs[0]
        held_table = doc.tables[0]
        before = len(expected)
        result = edit()
        if text is None:
            removed = 2 if operation.endswith("range") else 1
            del expected[at:at + removed]
            assert result == (2 if removed == 2 else True)
        else:
            expected.insert(at, text)
            assert result.text == text
            assert result.xml.startswith("<w:p")
        assert len(expected) != before
        for stale in (lambda: held.text, lambda: held_run.text, lambda: held_table.xml):
            with pytest.raises(rdocx.StaleElementError, match=operation):
                stale()
        assert texts() == expected
        read_all()
    held = doc.paragraphs[3]
    held_table = doc.tables[0]
    held_cell = held_table.cell(0, 0)
    table = doc.insert_table(4, 2, 3)
    assert len(table.rows) == 2 and table.cell(1, 2).text == ""
    for stale in (lambda: held.text, lambda: held_table.xml, lambda: held_cell.text):
        with pytest.raises(rdocx.StaleElementError, match="Document.insert_table"):
            stale()
    table.cell(1, 2).text = "new cell"
    assert doc.tables[0].cell(1, 2).text == "new cell"
    assert doc.tables[1].cell(0, 0).paragraphs[0].text == "cell"
    assert texts() == expected
    read_all()
    held = doc.paragraphs[3]
    held_table = doc.tables[1]
    appended_table = doc.add_table(1, 2)
    assert len(doc.tables) == 3 and appended_table.cell(0, 1).text == ""
    assert held.text == expected[3] and held_table.xml.startswith("<w:tbl")
    read_all()
    held = doc.paragraphs[3]
    held_run = held.runs[0]
    appended = doc.add_paragraph("appended")
    expected.append("appended")
    # An append keeps paragraph owners, while runs keep the document epoch.
    assert appended.text == "appended" and held.text == expected[3]
    with pytest.raises(rdocx.StaleElementError, match="Document.add_paragraph"):
        _ = held_run.xml
    assert texts() == expected
    read_all()
    reopened = rdocx.Document.from_bytes(doc.to_bytes())
    assert [paragraph.text for paragraph in reopened.paragraphs] == expected


def _issue407_work():
    from rdocx import _rdocx

    work = getattr(_rdocx, "_word_whole_body_work", None)
    if work is None:
        pytest.skip("the whole-body work counters exist in debug builds only")
    return work


def _issue407_review_document(paragraphs):
    """A body shaped like the python-docx review documents of the rdocx-skills scaling test: paragraphs of several
    runs with a picture and a table after every 10th, two sections with a header, comments, then the snippets
    python-docx appends, which declare `w` again: a tracked insertion, a block control and a repeating section whose
    item paragraph no story route reaches."""
    import rdocx

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    w15 = 'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"'
    document = rdocx.Document()
    for i in range(paragraphs):
        paragraph = document.add_paragraph()
        for k in range(3):
            paragraph.add_run(f"run {k} of paragraph {i} ").bold = False
        if i % 10 == 9:
            paragraph.add_run("").add_picture(_one_pixel_png(), "pixel.png")
            table = document.add_table(2, 2)
            table.cell(0, 0).text = f"cell of table {i // 10}"
    document.section(0).header.add_paragraph("header")
    for k in range(2):
        at = document.find_content_index(document.add_paragraph(f"commented {k}"))
        document.add_comment(
            rdocx.RunRange(
                start=rdocx.RunPosition(body_index=at, run_index=0),
                end=rdocx.RunPosition(body_index=at, run_index=0),
            ),
            author="Reviewer",
            text=f"comment {k}",
        )
    tail = "".join([
        f'<w:p><w:r><w:t xml:space="preserve">tracked </w:t></w:r><w:ins {w} w:id="901" w:author="Editor" '
        'w:date="2026-01-01T00:00:00Z"><w:r><w:t>inserted</w:t></w:r></w:ins></w:p>',
        f'<w:sdt {w}><w:sdtPr><w:tag w:val="probe0"/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>control</w:t></w:r>'
        '</w:p></w:sdtContent></w:sdt>',
        '<w:p><w:r><w:t>spot</w:t></w:r></w:p>',
        f'<w:sdt {w} {w15}><w:sdtPr><w:tag w:val="rows"/><w15:repeatingSection/></w:sdtPr><w:sdtContent><w:sdt>'
        '<w:sdtPr><w15:repeatingSectionItem/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>row</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt></w:sdtContent></w:sdt>',
        '<w:p><w:r><w:t>after the rows</w:t></w:r></w:p>',
        '<w:p><w:r><w:t>end</w:t></w:r></w:p>',
    ])
    source = io.BytesIO(document.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "word/document.xml":
                text = data.decode()
                at = text.rindex("<w:sectPr")
                # The first section ends at the comments and keeps the header.
                section = text[at:text.index("</w:sectPr>", at) + len("</w:sectPr>")]
                text = text[:at] + f"<w:p><w:pPr>{section}</w:pPr></w:p>" + tail + text[at:]
                data = text.encode()
            result_zip.writestr(info, data)
    return rdocx.Document.from_bytes(output.getvalue())


@pytest.mark.parametrize("paragraphs", [20, 80])
def test_issue407_review_shaped_edits_do_no_whole_body_work(paragraphs):
    import rdocx

    work = _issue407_work()
    doc = _issue407_review_document(paragraphs)
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts[-6:] == ["tracked inserted", "control", "spot", "row", "after the rows", "end"]
    assert len(doc.sections) == 2

    def direct(call):
        before = work()
        result = call()
        assert work() == before, (call, before, work())
        return result

    # Paragraphs inside content controls are not direct body children: they read the whole source.
    for paragraph in doc.paragraphs[-10:] + doc.paragraphs[8:11]:
        if paragraph.text in ("control", "row"):
            continue
        direct(lambda: paragraph.xml)
        for run in paragraph.runs:
            direct(lambda: run.xml)
    for table in (doc.tables[0], doc.tables[-1]):
        direct(lambda: table.xml)
        direct(lambda: table.cell(0, 0).xml)
    spot = doc.find_content_index(doc.paragraphs[-4])
    tables = len(doc.tables)
    direct(lambda: doc.insert_paragraph(spot, "inserted"))
    assert direct(lambda: doc.remove_content(spot))
    direct(lambda: doc.paragraphs[-2].insert_paragraph_before("before"))
    direct(lambda: doc.insert_paragraph(spot, "x"))
    assert direct(lambda: doc.remove_content_range(spot, spot + 2)) == 2
    direct(lambda: doc.add_paragraph("appended"))
    direct(lambda: doc.insert_table(spot, 2, 2))
    direct(lambda: doc.add_table(1, 2))
    assert [paragraph.text for paragraph in doc.paragraphs] == texts[:-4] + [
        "row", "before", "after the rows", "end", "appended"]
    assert len(doc.tables) == tables + 2
    reopened = rdocx.Document.from_bytes(doc.to_bytes())
    assert [paragraph.text for paragraph in reopened.paragraphs] == [paragraph.text for paragraph in doc.paragraphs]
    xml = reopened.part_xml("/word/document.xml")
    assert "<w15:repeatingSectionItem/>" in xml and 'w:val="probe0"' in xml and 'w:id="901"' in xml
    assert len(reopened.sections) == 2 and len(reopened.tables) == len(doc.tables)


_ISSUE407_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_ISSUE407_ROWS = (
    f'<w:sdt xmlns:w="{_ISSUE407_W}" xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"><w:sdtPr>'
    '<w:tag w:val="rows"/><w15:repeatingSection/></w:sdtPr><w:sdtContent><w:sdt><w:sdtPr><w15:repeatingSectionItem/>'
    '</w:sdtPr><w:sdtContent><w:p><w:r><w:t>row</w:t></w:r></w:p></w:sdtContent></w:sdt></w:sdtContent></w:sdt>'
)
_ISSUE407_TWIN = ('<w:sdt xmlns:x="urn:x"><w:sdtPr><x:meta/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>twin</w:t></w:r>'
                  '</w:p></w:sdtContent></w:sdt>')
_ISSUE407_ALIAS = (f'<q:sdt xmlns:q="{_ISSUE407_W}"><q:sdtPr><q:tag q:val="t"/></q:sdtPr><q:sdtContent><q:p><q:r>'
                   '<q:t>in q</q:t></q:r></q:p></q:sdtContent></q:sdt>')
_ISSUE407_EDITS = {
    "remove_content": lambda doc: doc.remove_content(1),
    "remove_content_range": lambda doc: doc.remove_content_range(0, 2),
    "replace_xml": lambda doc: doc.paragraphs[1].replace_xml(doc.paragraphs[1].xml.replace(">drop<", ">kept<")),
    "insert_paragraph": lambda doc: doc.insert_paragraph(1, "inserted"),
    "insert_paragraph_before": lambda doc: doc.paragraphs[1].insert_paragraph_before("before"),
    "add_paragraph": lambda doc: doc.add_paragraph("appended"),
    "insert_table": lambda doc: doc.insert_table(1, 1, 2),
    "add_table": lambda doc: doc.add_table(1, 2),
}


@pytest.mark.parametrize("owners", ["twin_sdt", "alias_sdt", "twin_repeating_sections", "repeating_section"])
@pytest.mark.parametrize("edit", sorted(_ISSUE407_EDITS))
def test_issue407_nested_owner_body_edits_never_publish_unsavable_bytes(owners, edit):
    import rdocx

    # A save replays nested declarations by owner identity: twin owners, or an
    # owner under an alias prefix, cannot be told apart once serialized.
    tail = {"twin_sdt": _ISSUE407_TWIN + '<w:p><w:r><w:t>mid</w:t></w:r></w:p>' + _ISSUE407_TWIN,
            "alias_sdt": _ISSUE407_ALIAS,
            "twin_repeating_sections": _ISSUE407_ROWS + _ISSUE407_ROWS,
            "repeating_section": _ISSUE407_ROWS}[owners]
    doc = _replace_document_body(
        rdocx.Document(),
        '<w:p><w:r><w:t>keep</w:t></w:r></w:p><w:p><w:r><w:t>drop</w:t></w:r></w:p>' + tail,
    )
    before = doc.to_bytes()
    texts = [paragraph.text for paragraph in doc.paragraphs]
    # A replacement splices the staged source, so it saves without a replay.
    refused = owners != "repeating_section" and edit != "replace_xml"
    try:
        _ISSUE407_EDITS[edit](doc)
    except rdocx.RdocxError as error:
        assert refused, error
        assert "namespace owner" in str(error)
        assert doc.to_bytes() == before
        assert [paragraph.text for paragraph in doc.paragraphs] == texts
    else:
        assert not refused
        saved = [paragraph.text for paragraph in doc.paragraphs]
        assert saved != texts or edit.endswith("table")
        reopened = rdocx.Document.from_bytes(doc.to_bytes())
        assert [paragraph.text for paragraph in reopened.paragraphs] == saved
        marker = {"twin_sdt": "<x:meta/>", "alias_sdt": "in q"}.get(owners, "<w15:repeatingSectionItem/>")
        assert marker in reopened.part_xml("/word/document.xml")


@pytest.mark.parametrize("root_w14", [False, True])
def test_issue407_direct_routes_do_no_whole_body_work(root_w14):
    import rdocx

    work = _issue407_work()
    doc = _issue407_document(root_w14)
    # A comments part does not keep plain paragraph removal off its route.
    doc.add_comment(
        rdocx.RunRange(
            start=rdocx.RunPosition(body_index=1, run_index=0),
            end=rdocx.RunPosition(body_index=1, run_index=1),
        ),
        author="Ada",
        text="review",
    )

    def direct(call):
        before = work()
        result = call()
        assert work() == before, (call, before, work())
        return result

    for index in (0, 1, 2, 3, 6, 7, 8, 9, 10):
        paragraph = doc.paragraphs[index]
        direct(lambda: paragraph.xml)
        for run in paragraph.runs:
            direct(lambda: run.xml)
    table = doc.tables[0]
    direct(lambda: table.xml)
    for row in table.rows:
        for cell in row.cells:
            direct(lambda: cell.xml)
    direct(lambda: doc.insert_paragraph(2, "inserted"))
    direct(lambda: doc.remove_content(2))
    direct(lambda: doc.paragraphs[7].insert_paragraph_before("before"))
    direct(lambda: doc.insert_paragraph(7, "x"))
    assert direct(lambda: doc.remove_content_range(7, 9)) == 2
    direct(lambda: doc.add_paragraph("appended"))
    direct(lambda: doc.insert_table(4, 2, 2))
    direct(lambda: doc.add_table(1, 2))
    direct(lambda: doc.tables[0].cell(1, 1).xml)
    assert [p.text for p in doc.paragraphs][-2:] == ["last", "appended"]
    # The counters do see the whole-body routes. A staged edit reads the
    # owner inventory by body block and story part, and a form publication
    # rebuilds it from the whole document.
    before = work()
    _ = doc.paragraphs[4].xml
    assert work()[1] > before[1]
    item = next(it for it in doc.story_items if it.story.kind == "body" and it.kind == "paragraph")
    doc.insert_content_control(rdocx.ContentLocation(item.story, item.kind, item.index_path), "block",
                               "plain_text", tag="t")
    before = work()
    assert doc.bind_content_controls({"t": "value"}) == 1
    assert work()[0] > before[0]


def _word_piece_work():
    from rdocx import _rdocx

    work = getattr(_rdocx, "_word_piece_work", None)
    if work is None:
        pytest.skip("the owner piece counter exists in debug builds only")
    return work


def _word_piece_edits():
    """Edits whose next owner inventory is read by body block and story part, each given the document and returning
    the call that makes the edit."""
    import rdocx

    def paragraph(doc):
        return next(p for p in reversed(list(doc.paragraphs)) if p.text.startswith("run 0 of paragraph"))

    def table(doc):
        return doc.tables[-1]

    def runs(doc):
        at = doc.find_content_index(paragraph(doc))
        return rdocx.RunRange(start=rdocx.RunPosition(body_index=at, run_index=0),
                              end=rdocx.RunPosition(body_index=at, run_index=1))

    def comment(doc):
        return next(c.id for c in doc.comments if c.text == "comment 0")

    def chart(doc):
        data = rdocx.CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("s", (1, 2))
        after = paragraph(doc)
        return lambda: doc.add_chart(rdocx.XL_CHART_TYPE.COLUMN_CLUSTERED, data, rdocx.Inches(4),
                                     rdocx.Inches(3), after=after)

    def settings(doc):
        settings = doc.settings
        return lambda: setattr(settings, "odd_and_even_pages_header_footer", True)

    def setter(target, name, value):
        def make(doc):
            obj = target(doc)
            return lambda: setattr(obj, name, value)
        return make

    def method(target, name, *args, **kwargs):
        def make(doc):
            obj = target(doc)
            values = [value(doc) if callable(value) else value for value in args]
            return lambda: getattr(obj, name)(*values, **kwargs)
        return make

    def document(doc):
        return doc

    return {
        "Table.indent": setter(table, "indent", rdocx.Inches(0.2)),
        "Table.set_borders": method(table, "set_borders", "single", size=8, color="000000"),
        "Row.height": setter(lambda doc: table(doc).rows[0], "height", rdocx.Inches(0.5)),
        "Cell.shading": setter(lambda doc: table(doc).cell(0, 0), "shading", rdocx.RGBColor(0x7B, 0x1E, 0x3A)),
        "Cell.add_paragraph": method(lambda doc: table(doc).cell(0, 0), "add_paragraph", "x"),
        "Paragraph.text": setter(paragraph, "text", "x"),
        "Paragraph.add_run": method(paragraph, "add_run", "x"),
        "Paragraph.add_hyperlink": method(paragraph, "add_hyperlink", "x", "https://example.com/a"),
        "Paragraph.add_equation": method(paragraph, "add_equation", "x^2", format="latex"),
        "Run.remove": method(lambda doc: paragraph(doc).runs[1], "remove"),
        "Run.add_field": method(lambda doc: paragraph(doc).runs[0], "add_field", "PAGE", "1"),
        "Document.split_run": method(document, "split_run", paragraph, 0, 3),
        "Document.add_comment": method(document, "add_comment", runs, author="Reviewer", text="c",
                                       date="2026-10-01T12:00:00Z"),
        "Document.reply_to": method(document, "reply_to", comment, author="Author", text="r",
                                    date="2026-10-01T12:00:00Z"),
        "Document.resolve_comment": method(document, "resolve_comment", comment),
        "Document.add_chart": chart,
        "Document.set_text_watermark": method(document, "set_text_watermark", "DRAFT"),
        "Document.set_note_policy": method(document, "set_note_policy", "footnote", start=2),
        "Settings.odd_and_even_pages_header_footer": settings,
    }


_WORD_PIECE_DOCUMENTS = {}


def _word_piece_document(paragraphs):
    import rdocx

    if paragraphs not in _WORD_PIECE_DOCUMENTS:
        _WORD_PIECE_DOCUMENTS[paragraphs] = _issue407_review_document(paragraphs).to_bytes()
    return rdocx.Document.from_bytes(_WORD_PIECE_DOCUMENTS[paragraphs])


@pytest.mark.parametrize("edit", sorted(_word_piece_edits()))
def test_word_owner_inventory_reads_only_what_an_edit_changes(edit):
    import rdocx

    # The owner inventory of an edit reads again only the body blocks and story parts the edit changed: as many in
    # a body four times as large, and never the whole document. Debug builds check each inventory read by block
    # and part against the rebuild.
    work = _issue407_work()
    pieces = _word_piece_work()
    reads = []
    for paragraphs in (20, 80):
        doc = _word_piece_document(paragraphs)
        call = _word_piece_edits()[edit](doc)
        before = (work(), pieces())
        call()
        assert work() == before[0], (edit, paragraphs)
        reads.append(pieces() - before[1])
        reopened = rdocx.Document.from_bytes(doc.to_bytes())
        assert [p.text for p in reopened.paragraphs] == [p.text for p in doc.paragraphs]
    assert reads[0] == reads[1] <= 2, (edit, reads)


_WORD_PIECE_BOX_DRAWING = (
    '<w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" '
    'allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset>'
    '</wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="914400" cy="914400"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
    '<wp:docPr id="7" name="Box"/><wp:cNvGraphicFramePr/>'
    '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData '
    'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:cNvSpPr txBox="1"/>'
    '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm><a:prstGeom prst="rect">'
    '<a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
    '<w:tbl><w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>boxed cell</w:t></w:r></w:p>'
    '</w:tc></w:tr></w:tbl><w:p/></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>'
    '</wp:anchor></w:drawing>'
)
_WORD_PIECE_BOXED_TABLE = ('<w:tbl><w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid><w:tr><w:tc><w:p><w:r>'
                           '<w:t>boxed cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p/>')
# A bare DrawingML box, the box Word writes (a DrawingML choice and its VML fallback, both stories), and a table
# whose cell holds a box. A table inside a box inside a table cell does not open, on main either.
_WORD_PIECE_TEXT_BOXES = (
    '<w:p><w:r>' + _WORD_PIECE_BOX_DRAWING + '</w:r></w:p>',
    '<w:p><w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><mc:Choice Requires="wps">'
    + _WORD_PIECE_BOX_DRAWING + '</mc:Choice><mc:Fallback><w:pict><v:rect xmlns:v="urn:schemas-microsoft-com:vml" '
    'style="width:72pt;height:72pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>{text} fallback</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox></v:rect></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>',
    '<w:tbl><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid><w:tr><w:tc><w:p><w:r>'
    + _WORD_PIECE_BOX_DRAWING.replace(_WORD_PIECE_BOXED_TABLE, "") + '</w:r></w:p></w:tc></w:tr></w:tbl>',
)


def _word_piece_text_box_document(paragraphs):
    import rdocx

    table = ('<w:tbl><w:tblGrid><w:gridCol w:w="1000"/><w:gridCol w:w="1000"/></w:tblGrid>'
             + '<w:tr><w:tc><w:p><w:r><w:t>c</w:t></w:r></w:p></w:tc><w:tc><w:p/></w:tc></w:tr>' * 2 + '</w:tbl>')
    body = [table]
    for i in range(paragraphs):
        body.append(f'<w:p><w:r><w:t>paragraph {i}</w:t></w:r></w:p>')
        if i % 10 in (2, 4, 6):
            body.append(_WORD_PIECE_TEXT_BOXES[i % 10 // 2 - 1].format(text=f"box {i}"))
        if i % 10 == 9:
            body.append(table)
    return _replace_document_body(rdocx.Document(), "".join(body) + "<w:sectPr/>")


def test_word_owner_inventory_reads_one_block_beside_text_boxes():
    # Text box stories, and the table cells inside them, number across the main part as table cells do, so the
    # inventory places them by offsets too and an edit reads only the block it changed.
    work = _issue407_work()
    pieces = _word_piece_work()
    reads = []
    for paragraphs in (20, 80):
        doc = _word_piece_text_box_document(paragraphs)
        edits = [
            lambda: setattr(doc.paragraphs[-1], "text", "edited"),
            lambda: setattr(doc.tables[-1].cell(0, 0), "width", 914400),
            # Cell paragraphs route up to the first table cell of a text box, as on main.
            lambda: doc.tables[0].cell(0, 1).add_paragraph("added"),
        ]
        for edit in edits:
            before = (work(), pieces())
            edit()
            assert work() == before[0], paragraphs
            reads.append((paragraphs, pieces() - before[1]))
        xml = doc.part_xml("/word/document.xml")
        assert all(text in xml for text in ("box 12", "box 14 fallback", "box 16", "boxed cell", "edited"))
    assert [count for _, count in reads] == [1] * 6, reads


def test_word_owner_inventory_reads_no_piece_for_a_part_read_twice():
    import rdocx

    # One part referenced as a header and as a footer is two story sources: the rebuild reads that inventory and
    # the pieces are not read at all.
    doc = rdocx.Document()
    for i in range(12):
        doc.add_paragraph(f"paragraph {i}")
    doc.add_table(1, 2)
    doc.section(0).header.add_paragraph("shared")
    source = io.BytesIO(doc.to_bytes())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(output, "w") as result_zip:
        rels = source_zip.read("word/_rels/document.xml.rels").decode()
        target = re.search(r'Type="[^"]*/header" Target="([^"]+)"', rels) or re.search(
            r'Target="([^"]+)" Type="[^"]*/header"', rels)
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                footer = ('<Relationship Id="rIdSharedFooter" Type="http://schemas.openxmlformats.org/officeDocument/'
                          f'2006/relationships/footer" Target="{target.group(1)}"/>')
                data = data.decode().replace("</Relationships>", footer + "</Relationships>").encode()
            if info.filename == "word/document.xml":
                text = data.decode()
                at = text.index("<w:headerReference")
                text = text[:at] + '<w:footerReference w:type="default" r:id="rIdSharedFooter"/>' + text[at:]
                data = text.encode()
            result_zip.writestr(info, data)
    work = _issue407_work()
    pieces = _word_piece_work()
    before = (work(), pieces())
    doc = rdocx.Document.from_bytes(output.getvalue())
    doc.paragraphs[3].text = "edited"
    doc.tables[0].cell(0, 0).width = 914400
    assert pieces() == before[1]
    assert work()[0] >= before[0][0] + 3
    assert rdocx.Document.from_bytes(doc.to_bytes()).paragraphs[3].text == "edited"


_WORD_PIECE_SPLICES = {
    "remove_content": lambda doc: doc.remove_content(0),
    "remove_content_range": lambda doc: doc.remove_content_range(0, 2),
    "insert_paragraph": lambda doc: doc.insert_paragraph(5, "inserted"),
    "insert_paragraph_before": lambda doc: doc.paragraphs[3].insert_paragraph_before("before"),
    "add_paragraph": lambda doc: doc.add_paragraph("appended"),
    "insert_table": lambda doc: doc.insert_table(5, 2, 2),
    "add_table": lambda doc: doc.add_table(1, 2),
}


@pytest.mark.parametrize("splice", sorted(_WORD_PIECE_SPLICES))
def test_word_owner_pieces_follow_a_body_splice(splice):
    import rdocx

    # A body splice reads no piece and keeps the pieces in step with the body, so an edit of the last paragraph
    # after it reads that paragraph's block alone. Debug builds check the inventory of that edit, which reuses the
    # pieces the splice placed, against the rebuild.
    work = _issue407_work()
    pieces = _word_piece_work()
    reads = []
    for paragraphs in (20, 80):
        doc = _word_piece_document(paragraphs)
        before = (work(), pieces())
        _WORD_PIECE_SPLICES[splice](doc)
        assert (work(), pieces()) == before, splice
        list(doc.paragraphs)[-1].text = "edited"
        assert work() == before[0], splice
        reads.append(pieces() - before[1])
        reopened = rdocx.Document.from_bytes(doc.to_bytes())
        assert [p.text for p in reopened.paragraphs] == [p.text for p in doc.paragraphs]
    assert reads == [1, 1], (splice, reads)


def test_word_owner_pieces_follow_a_tail_removal():
    import rdocx

    work = _issue407_work()
    pieces = _word_piece_work()
    for paragraphs in (20, 80):
        doc = rdocx.Document()
        for i in range(paragraphs):
            doc.add_paragraph(f"paragraph {i}")
            if i % 10 == 9:
                doc.add_table(2, 2)
        doc.add_paragraph("tail")
        doc = rdocx.Document.from_bytes(doc.to_bytes())
        before = (work(), pieces())
        # The last body child is a plain paragraph, which the tail removal takes in place.
        assert doc.remove_content(doc.find_content_index(list(doc.paragraphs)[-1]))
        assert (work(), pieces()) == before
        list(doc.paragraphs)[-1].text = "edited"
        assert work() == before[0] and pieces() - before[1] == 1, paragraphs


def _staging_work():
    from rdocx import _rdocx

    work = getattr(_rdocx, "_staging_work", None)
    if work is None:
        pytest.skip("the staging work counter exists in debug builds only")
    return work


_STAGING_DOCUMENTS = {}


def _staging_document(paragraphs):
    """The review-shaped document of the owner piece tests, with a font table and a numbering part, opened from
    bytes."""
    import rdocx

    if paragraphs not in _STAGING_DOCUMENTS:
        document = _issue407_review_document(paragraphs)
        document.set_font(rdocx.FontDefinition("Seed"))
        document.add_numbering_definition([rdocx.ListLevel(format="decimal", text="%1.")])
        _STAGING_DOCUMENTS[paragraphs] = document.to_bytes()
    return rdocx.Document.from_bytes(_STAGING_DOCUMENTS[paragraphs])


def _staging_edits():
    """Edits that stage a document or compare its saves and leave its main part as it was opened, or change only its
    model, each given the document and a fresh k and returning the call that makes the edit."""
    from pathlib import Path

    import rdocx

    font = (Path(__file__).resolve().parents[2] / "oxml-layout/fonts/Selawik-Regular.ttf").read_bytes()
    license = rdocx.FontEmbeddingLicense(authorized=True, identity="SIL Open Font License 1.1")

    def levels():
        return [rdocx.ListLevel(format="decimal", text="%1.")]

    def remove_font(doc, k):
        doc.set_font(rdocx.FontDefinition(f"Gone {k}"))
        return lambda: doc.remove_font(f"Gone {k}")

    def add_numbering_instance(doc, k):
        definition = doc.add_numbering_definition(levels())
        return lambda: doc.add_numbering_instance(definition)

    def link_style_to_numbering(doc, k):
        style = doc.add_style(f"Probe style {k}", "paragraph").style_id
        instance = doc.add_numbering_instance(doc.add_numbering_definition(levels()))
        return lambda: doc.link_style_to_numbering(style, instance, 0)

    return {
        "Document.set_font": lambda doc, k: lambda: doc.set_font(rdocx.FontDefinition(f"Probe {k}")),
        "Document.embed_font": lambda doc, k: lambda: doc.embed_font(
            "Selawik", font, ("regular", "bold", "italic")[k], license=license),
        "Document.remove_font": remove_font,
        "Document.embed_true_type_fonts": lambda doc, k: lambda: setattr(doc, "embed_true_type_fonts", k != 1),
        "Document.add_numbering_definition": lambda doc, k: lambda: doc.add_numbering_definition(levels()),
        "Document.add_numbering_instance": add_numbering_instance,
        "Document.link_style_to_numbering": link_style_to_numbering,
        "Document.add_bullet_list_item": lambda doc, k: lambda: doc.add_bullet_list_item(f"item {k}"),
        "Document.add_numbered_list_item": lambda doc, k: lambda: doc.add_numbered_list_item(f"item {k}"),
    }


@pytest.mark.parametrize("edit", sorted(_staging_edits()))
def test_staging_reads_the_main_part_once_while_it_keeps_its_opened_bytes(edit):
    # A document parses, scans and checks the main part it was opened with once for staging: the first edit reads
    # it. Later edits that leave that part as it was, or change only its model, read none of it again, in a body of 20
    # paragraphs as in one four times as large, and compare saves without writing a ZIP. Debug builds check every reuse
    # against the whole work.
    work = _staging_work()
    make = _staging_edits()[edit]
    counts = []
    for paragraphs in (20, 80):
        doc = _staging_document(paragraphs)
        work()
        make(doc, 0)()
        first = list(work())
        assert any(first), (paragraphs, first)
        sample = []
        for k in (1, 2):
            call = make(doc, k)
            work()
            call()
            sample.append(list(work()))
        counts.append(sample)
    assert counts[0] == counts[1], counts
    if edit in ("Document.add_bullet_list_item", "Document.add_numbered_list_item"):
        # A list item changes the model, which serializes and reads only the blocks it adds.
        assert all(sample[:4] == [0] * 4 for sample in counts[0]), counts
    else:
        assert counts[0] == [[0] * 6, [0] * 6], counts


def _reopen_work():
    from rdocx import _rdocx

    work = getattr(_rdocx, "_reopen_work", None)
    if work is None:
        pytest.skip("the reopen work counter exists in debug builds only")
    return work


def _main_part_keeping_edits():
    """Edits whose commit leaves the main document part as the last commit left it, each given the document and
    returning the call that makes the edit, or a document read again and that call: the calls of the rdocx-skills
    scaling test that keep the main part. A local edit of the body leaves the main part to the next commit, so the
    targets are made by commits, or read again."""
    import rdocx

    def paragraph(doc):
        return next(p for p in reversed(list(doc.paragraphs)) if p.text.startswith("run 0 of paragraph"))

    def footnote(doc):
        return doc.add_footnote(paragraph(doc), "note")

    def chart_legend(doc):
        data = rdocx.CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("s", (1, 2))
        doc.add_chart(rdocx.XL_CHART_TYPE.COLUMN_CLUSTERED, data, rdocx.Inches(4), rdocx.Inches(3))
        doc = rdocx.Document.from_bytes(doc.to_bytes())
        chart = doc.paragraphs[-1]
        return doc, lambda: doc.set_chart_legend(chart, visible=True, position="b")

    def watermark(doc):
        section = doc.section(0)
        section.set_text_watermark("DRAFT", width=rdocx.Inches(4), height=rdocx.Inches(1), font_size=rdocx.Pt(36))
        return section

    image = _one_pixel_png(b"\x00\x00\xff")
    return {
        "HeaderFooter.add_paragraph": lambda doc: (lambda header: lambda: header.add_paragraph("x"))(
            doc.section(0).header),
        "HeaderFooter.add_table": lambda doc: (lambda header: lambda: header.add_table(1, 1))(doc.section(0).header),
        "FootnoteEndnote.add_paragraph": lambda doc: (lambda note: lambda: note.add_paragraph("x"))(footnote(doc)),
        "FootnoteEndnote.add_table": lambda doc: (lambda note: lambda: note.add_table(1, 1))(footnote(doc)),
        "SectionHandle.set_text_watermark": lambda doc: lambda: doc.section(0).set_text_watermark(
            "COPY", width=rdocx.Inches(4), height=rdocx.Inches(1), font_size=rdocx.Pt(36)),
        "SectionHandle.set_image_watermark": lambda doc: lambda: doc.section(0).set_image_watermark(
            image, "mark.png", width=rdocx.Inches(2), height=rdocx.Inches(2)),
        "SectionHandle.remove_watermark": lambda doc: watermark(doc).remove_watermark,
        "Document.replace_image": lambda doc: (lambda rid: lambda: doc.replace_image(rid, image))(
            doc.pictures()[-1].relationship_id),
        "Document.replace_image_for_story": lambda doc: (lambda rid: lambda: doc.replace_image_for_story(
            next(s for s in doc.stories if s.kind == "body"), rid, image))(doc.pictures()[-1].relationship_id),
        "Document.add_custom_xml_store": lambda doc: lambda: doc.add_custom_xml_store(b"<data><v>1</v></data>"),
        "Document.set_chart_legend": chart_legend,
    }


@pytest.mark.parametrize("edit", sorted(_main_part_keeping_edits()))
def test_a_commit_that_keeps_the_main_part_reopens_without_reading_it(edit):
    import rdocx

    # A commit of a document read from bytes reopens its package in memory, as a write and a read of the ZIP would.
    # With the main document part left as the last commit left it, the reopen parses none of it and scans none of it
    # for identifiers: it scans as many bytes in a body four times as large. The staging of the first commit after an
    # opening still parses the main part once, which this reopen counter does not count. Debug builds also reopen
    # through the ZIP and assert the same document.
    work = _reopen_work()
    counts = []
    for paragraphs in (20, 80):
        doc = _word_piece_document(paragraphs)
        made = _main_part_keeping_edits()[edit](doc)
        doc, call = made if isinstance(made, tuple) else (doc, made)
        work()
        call()
        counts.append(work())
        reopened = rdocx.Document.from_bytes(doc.to_bytes())
        assert [p.text for p in reopened.paragraphs] == [p.text for p in doc.paragraphs]
    (zip_writes, main_parses, scanned), large = counts
    assert zip_writes == main_parses == large[0] == large[1] == 0, (edit, counts)
    # An identifier written into a changed part may take one more digit in the larger document. Scanning the main
    # part again would add the tens of kilobytes its 60 more paragraphs hold.
    assert abs(large[2] - scanned) <= 16, (edit, counts)


def _main_body_block_scans():
    from rdocx import _rdocx

    scans = getattr(_rdocx, "_main_body_block_scans", None)
    if scans is None:
        pytest.skip("the body child scan counter exists in debug builds only")
    return scans


def _story_lookup_calls():
    """Calls that address one item of the main part's stories, each given the document and returning the call."""
    import rdocx

    def paragraph(doc):
        return next(p for p in reversed(list(doc.paragraphs)) if p.text.startswith("run 0 of paragraph"))

    def item(doc, pick):
        found = next(it for it in doc.story_items if it.story.kind == "body" and pick(it))
        return rdocx.ContentLocation(found.story, found.kind, found.index_path)

    def cell(doc):
        return doc.tables[-1].cell(0, 0)

    def move_comment(doc):
        at = doc.find_content_index(paragraph(doc))
        found = next(it for it in doc.story_items if it.story.kind == "body" and it.kind == "paragraph"
                     and it.direct_body_index == at)
        target = rdocx.StoryRunRange(start=rdocx.StoryRunPosition(item=found, run_index=0),
                                     end=rdocx.StoryRunPosition(item=found, run_index=1))
        comment = next(c.id for c in doc.comments if c.text == "comment 0")
        return lambda: doc.move_comment(comment, target)

    def picture_size(doc):
        picture = doc.pictures()[-1]
        return lambda: doc.set_picture_size(picture, rdocx.Inches(2), rdocx.Inches(2))

    def control(doc):
        location = item(doc, lambda it: it.kind == "content_control" and b'w:val="probe0"' in it.xml)
        return lambda: doc.set_content_control_tag(location, "renamed")

    def made(target, name, *args):
        def make(doc):
            obj = target(doc)
            return lambda: getattr(obj, name)(*args)
        return make

    return {
        "Cell.replace_text": made(cell, "replace_text", "cell", "omega"),
        "Cell.add_table": made(cell, "add_table", 1, 1),
        "Cell.add_picture": made(cell, "add_picture", _one_pixel_png(), "pixel.png"),
        "Document.add_picture": lambda doc: lambda: doc.add_picture(_one_pixel_png(), "pixel.png"),
        "Document.set_picture_size": picture_size,
        "Document.move_comment": move_comment,
        "Document.set_content_control_tag": control,
        "Document.add_footnote": lambda doc: (lambda p: lambda: doc.add_footnote(p, "note"))(paragraph(doc)),
    }


@pytest.mark.parametrize("call", sorted(_story_lookup_calls()))
def test_story_lookups_scan_only_the_body_children_an_edit_changed(call):
    # Each story lookup of the main part reads its owners and items from the scans of the body children it keeps,
    # and scans again only a child whose bytes changed: as many in a body four times as large. Debug builds check
    # each lookup against the scans of the whole part.
    scans = _main_body_block_scans()
    counts = []
    for paragraphs in (20, 80):
        doc = _word_piece_document(paragraphs)
        # The first lookups read every child once, for its owners and then its items.
        assert doc.stories and doc.pictures()
        make = _story_lookup_calls()[call](doc)
        before = scans()
        make()
        counts.append(scans() - before)
    # Each call scans again the children it changed, at least one, as many at both sizes.
    assert 1 <= counts[0] == counts[1] <= 4, (call, counts)


def _producer_document(paragraphs, root_prefix=None, run_sibling=None):
    """The staging document with its main part as another producer may write it, which the model serializes
    otherwise: a line break after the body start, the root named with `root_prefix` bound to the Word namespace when
    given, and `run_sibling` after the first run of the body when given."""
    import io
    import zipfile

    import rdocx

    _staging_document(paragraphs)
    written = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_STAGING_DOCUMENTS[paragraphs])) as source, zipfile.ZipFile(written, "w") as copy:
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == "word/document.xml":
                data = data.replace(b"<w:body>", b"<w:body>\n", 1)
                if run_sibling:
                    data = data.replace(b"</w:r>", b"</w:r>" + run_sibling, 1)
                if root_prefix:
                    namespace = b'="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
                    data = data.replace(b"<w:document ", b"<%s:document xmlns:%s%s " % (
                        root_prefix.encode(), root_prefix.encode(), namespace), 1)
                    data = data.replace(b"</w:document>", b"</%s:document>" % root_prefix.encode(), 1)
            copy.writestr(entry, data)
    return rdocx.Document.from_bytes(written.getvalue())


def _block_edits():
    """Edits that change the main part of a document and commit it, each given the document and a fresh k and
    returning the call that makes the edit. The fragment imports, the scoped text replacements and the pictures edit
    the main part as the package holds it, and are made on a document whose main part another producer wrote."""
    import rdocx

    def run_text(doc, k):
        run = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("end")][-1].runs[0]
        return lambda: setattr(run, "text", f"end {k}")

    def footnote(doc, k):
        paragraph = [paragraph for paragraph in doc.paragraphs if paragraph.text == "spot"][-1]
        return lambda: doc.add_footnote(paragraph, f"note {k}")

    def import_fragment(doc, k):
        source = rdocx.Document()
        source.add_paragraph(f"imported {k}")
        source.add_paragraph("second")
        story = next(story for story in source.stories if story.kind == "body")
        fragment = rdocx.DocumentFragment.from_range(
            source, rdocx.ContentLocation(story, "paragraph", [0]), rdocx.ContentLocation(story, "paragraph", [1]))
        body = next(story for story in doc.stories if story.kind == "body")
        return lambda: doc.import_fragment(rdocx.ContentLocation.end(body), fragment)

    def replace_text_at(doc, k):
        index = doc.find_content_index(doc.add_paragraph(f"alpha {k}"))
        item = next(item for item in doc.story_items if item.story.kind == "body" and item.kind == "paragraph"
                    and item.direct_body_index == index)
        return lambda: doc.replace_text_at(item, "alpha", "omega")

    def cell_replace_text(doc, k):
        cell = doc.tables[-1].cell(0, 0)
        cell.text = f"alpha {k}"
        return lambda: cell.replace_text("alpha", "omega")

    return {
        "Table.add_row": lambda doc, k: doc.tables[-1].add_row,
        "Cell.text": lambda doc, k: lambda: setattr(doc.tables[-1].cell(0, 1), "text", f"cell {k}"),
        "Run.text": run_text,
        "Document.add_bullet_list_item": lambda doc, k: lambda: doc.add_bullet_list_item(f"item {k}"),
        "Document.insert_section": lambda doc, k: lambda: doc.insert_section(len(doc.sections)),
        "Document.insert_html": lambda doc, k: lambda: doc.insert_html(
            f"<p>inserted {k}</p>", after=[paragraph for paragraph in doc.paragraphs if paragraph.text == "spot"][-1]),
        "Document.add_footnote": footnote,
        "Document.import_fragment": import_fragment,
        "Paragraph.replace_text": lambda doc, k: (lambda paragraph: lambda: paragraph.replace_text("alpha", "omega"))(
            doc.add_paragraph(f"alpha {k}")),
        "Cell.replace_text": cell_replace_text,
        "Document.replace_text_at": replace_text_at,
        "Document.add_picture": lambda doc, k: lambda: doc.add_picture(
            _one_pixel_png(bytes([k, k, k])), f"p{k}.png", 12700, 12700),
        "Cell.add_picture": lambda doc, k: (lambda cell: lambda: cell.add_picture(
            _one_pixel_png(bytes([k, k, k])), f"c{k}.png", 12700, 12700))(doc.tables[-1].cell(0, 0)),
        "SectionHandle.set_page_border": lambda doc, k: lambda: doc.section(0).set_page_border(
            ("top", "left", "bottom", "right")[k]),
    }


@pytest.mark.parametrize("edit", sorted(_block_edits()))
def test_staging_writes_and_reads_only_the_main_part_blocks_an_edit_changes(edit):
    # Staging serializes the main part from the blocks kept for the body children an edit leaves equal, and reads in
    # their bytes what it read once: after two edits, a later one does none of the whole main part work and serializes
    # and reads as many blocks in a body of 20 paragraphs as in one four times as large. The reopen of its commit
    # parses none of the main part whole and scans for identifiers only the blocks no earlier reading parsed. Debug
    # builds check every block read against the whole work, and every reopen against a read of the ZIP.
    work = _staging_work()
    reopen = _reopen_work()
    make = _block_edits()[edit]
    counts = []
    reopens = []
    for paragraphs in (20, 80):
        producer = edit in ("Document.import_fragment", "Paragraph.replace_text", "Cell.replace_text",
                            "Document.replace_text_at", "Document.add_picture", "Cell.add_picture")
        doc = (_producer_document if producer else _staging_document)(paragraphs)
        for k in (0, 1):
            make(doc, k)()
        sample = []
        for k in (2, 3):
            call = make(doc, k)
            work()
            reopen()
            call()
            sample.append(list(work()))
            reopens.append(list(reopen()))
        counts.append(sample)
    # An import also reads the small document of its fragment whole and reopens its package, alike at both sizes.
    imported = edit == "Document.import_fragment"
    assert counts[0] == counts[1], counts
    assert all(sample[:4] == [0] * 4 or imported for sample in counts[0]), counts
    assert all(sample[:2] == [0, int(imported)] for sample in reopens), reopens
    # An identifier written into a changed block may take one more digit in the larger document.
    assert all(abs(large[2] - small[2]) <= 16 for small, large in zip(reopens[:2], reopens[2:])), reopens


def test_a_fragment_imports_into_a_main_part_whose_root_has_another_word_prefix():
    # A producer may name the root with another prefix bound to the Word namespace. A body child of such a main part
    # read alone would be closed with w:document, so a fragment spliced into it parses the part whole, as before.
    # Debug builds check every block read against the whole one.
    import rdocx

    doc = _producer_document(20, "x")
    texts = [paragraph.text for paragraph in doc.paragraphs]
    make = _block_edits()["Document.import_fragment"]
    for k in range(3):
        make(doc, k)()
    assert [paragraph.text for paragraph in doc.paragraphs] == texts + [f"imported {k}" for k in range(3)]
    reopened = rdocx.Document.from_bytes(doc.to_bytes())
    assert [paragraph.text for paragraph in reopened.paragraphs] == [paragraph.text for paragraph in doc.paragraphs]


def test_story_pictures_raise_where_the_relationship_rename_cannot_decode_the_main_part():
    # A foreign attribute value holding an undefined entity, in one body child of a main part another producer wrote.
    # Renaming the relationships of the part decodes every attribute value and raises on that one, so every story
    # picture raises, at every call, though the children the rename reads hold no such value.
    import rdocx

    for add in (lambda doc, k: doc.add_picture(_one_pixel_png(bytes([k, k, k])), f"p{k}.png", 12700, 12700),
                lambda doc, k: doc.tables[-1].cell(0, 0).add_picture(
                    _one_pixel_png(bytes([k, k, k])), f"c{k}.png", 12700, 12700)):
        doc = _producer_document(20, run_sibling=b'<x:foo xmlns:x="urn:x" x:a="&bogus;"/>')
        for k in range(3):
            with pytest.raises(rdocx.RdocxError, match="invalid story relationship id.*bogus"):
                add(doc, k)
